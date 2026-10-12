
from pathlib import Path
import csv
import os
import json
import re
import subprocess
import threading
import time
import traceback
import uuid
from datetime import datetime, timedelta
import random
from urllib.parse import quote
import phonenumbers
import pyautogui
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_from_directory
from werkzeug.utils import secure_filename
from playwright.sync_api import sync_playwright, TimeoutError
import openpyxl
import os as _os
import sys as _sys

# Debug copy of the app kept under dev/; make the project root importable.
_sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))
from core.main_scrapper import (
    open_group_and_members,
    extract_all_groups,
)
from core.group_db_exporter import save_group_contacts_to_postgres
from core.recovery_worker import RecoveryManager

try:
    import psycopg2
    from psycopg2 import sql
except Exception:
    psycopg2 = None
    sql = None
 

def load_env_file():
    env_path = Path(__file__).parent / ".env"
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


load_env_file()


def _env_bool(name, default=False):
    raw = os.getenv(name)
    if raw is None:
        return bool(default)
    return str(raw).strip().lower() in ("1", "true", "yes", "on")


def get_playwright_launch_kwargs(default_headless=False):
    headless = _env_bool("PLAYWRIGHT_HEADLESS", default_headless)
    start_maximized = _env_bool("PLAYWRIGHT_START_MAXIMIZED", True)
    no_viewport = _env_bool("PLAYWRIGHT_NO_VIEWPORT", True)
    slow_mo_raw = str(os.getenv("PLAYWRIGHT_SLOW_MO_MS", "0") or "0").strip()
    try:
        slow_mo = max(0, int(slow_mo_raw))
    except Exception:
        slow_mo = 0

    kwargs = {"headless": headless}
    if start_maximized:
        kwargs["args"] = ["--start-maximized"]
    if no_viewport:
        kwargs["no_viewport"] = True
    if slow_mo > 0:
        kwargs["slow_mo"] = slow_mo
    return kwargs


def get_playwright_user_data_dir():
    base = Path(__file__).parent
    multi_account = _env_bool("PLAYWRIGHT_MULTI_ACCOUNT", False)
    if not multi_account:
        raw = str(os.getenv("PLAYWRIGHT_USER_DATA_DIR", "./playwright_profile") or "").strip()
        path = Path(raw)
        if not path.is_absolute():
            path = (base / path).resolve()
        path.mkdir(parents=True, exist_ok=True)
        return path

    root_raw = str(os.getenv("PLAYWRIGHT_PROFILES_ROOT", "./playwright_profiles") or "").strip()
    root = Path(root_raw)
    if not root.is_absolute():
        root = (base / root).resolve()
    root.mkdir(parents=True, exist_ok=True)

    default_profile = str(os.getenv("PLAYWRIGHT_PROFILE_NAME", "default") or "default").strip() or "default"
    auto_select = _env_bool("PLAYWRIGHT_PROFILE_AUTO_SELECT", True)
    current_marker = root / ".current_profile"
    selected_name = default_profile

    if auto_select:
        marker_name = ""
        if current_marker.exists():
            try:
                marker_name = current_marker.read_text(encoding="utf-8").strip()
            except Exception:
                marker_name = ""
        if marker_name and (root / marker_name).is_dir():
            selected_name = marker_name
        else:
            profile_dirs = [p for p in root.iterdir() if p.is_dir() and not p.name.startswith(".")]
            if profile_dirs:
                profile_dirs.sort(key=lambda p: p.stat().st_mtime, reverse=True)
                selected_name = profile_dirs[0].name

    safe_name = "".join(ch for ch in selected_name if ch.isalnum() or ch in ("-", "_")).strip() or "default"
    path = root / safe_name
    path.mkdir(parents=True, exist_ok=True)
    try:
        current_marker.write_text(safe_name, encoding="utf-8")
    except Exception:
        pass
    return path


def _wait_for_any_selector(page, selectors, timeout_ms=120000):
    for selector in selectors:
        try:
            page.wait_for_selector(selector, timeout=timeout_ms)
            return selector
        except Exception:
            continue
    return ""


def _first_text_from_selectors(page, selectors, timeout_ms=1200):
    for selector in selectors:
        try:
            node = page.locator(selector).first
            if node.count() == 0:
                continue
            node.wait_for(state="visible", timeout=timeout_ms)
            value = (node.inner_text() or "").strip()
            if value:
                return value
        except Exception:
            continue
    return ""


def _click_first_selector(page, selectors, timeout_ms=1500):
    for selector in selectors:
        try:
            node = page.locator(selector).first
            if node.count() == 0:
                continue
            node.wait_for(state="visible", timeout=timeout_ms)
            try:
                node.scroll_into_view_if_needed(timeout=timeout_ms)
            except Exception:
                pass
            try:
                node.click(timeout=timeout_ms)
            except Exception:
                try:
                    node.click(timeout=timeout_ms, force=True)
                except Exception:
                    node.evaluate("(el) => el.click()")
            return True
        except Exception:
            continue
    return False


def _is_profile_panel_open(page):
    opened = _wait_for_any_selector(
        page,
        (
            "[data-testid='drawer-right']",
            "span[data-testid='pushname-input-read-only selectable-text']",
            "span[data-testid='pushname-input-read-only']",
            "[data-testid='phone']",
            "button[aria-label='Edit Name']",
            "button[aria-label='Edit profile photo']",
        ),
        timeout_ms=1400,
    )
    return bool(opened)


def _force_click_profile_navbar(page):
    try:
        clicked = bool(
            page.evaluate(
                """() => {
                  const btn = document.querySelector(
                    "header[data-testid='chatlist-header'] button[data-testid='navbar-item-profile'][aria-label='Profile']"
                  ) || document.querySelector(
                    "button[data-testid='navbar-item-profile'][aria-label='Profile']"
                  ) || document.querySelector(
                    "header[data-testid='chatlist-header'] button[data-testid='navbar-item-profile']"
                  ) || document.querySelector(
                    "button[data-testid='navbar-item-profile']"
                  ) || document.querySelector(
                    "[data-testid='navbar-item-profile-photo']"
                  );
                  if (!btn) return false;
                  btn.scrollIntoView({ block: "center", inline: "center" });
                  btn.dispatchEvent(new MouseEvent("mousedown", { bubbles: true, cancelable: true, view: window }));
                  btn.dispatchEvent(new MouseEvent("mouseup", { bubbles: true, cancelable: true, view: window }));
                  btn.dispatchEvent(new MouseEvent("click", { bubbles: true, cancelable: true, view: window }));
                  return true;
                }"""
            )
        )
        if not clicked:
            return False
        return _is_profile_panel_open(page)
    except Exception:
        return False


def _open_whatsapp_profile_panel(page):
    # 0) Preferred: direct navbar profile button (fast + stable)
    navbar_profile = _force_click_profile_navbar(page) or _click_first_selector(
        page,
        (
            "header[data-testid='chatlist-header'] button[data-testid='navbar-item-profile']",
            "button[data-testid='navbar-item-profile']",
            "[data-testid='navbar-item-profile-photo']",
            "button[aria-label='Profile'][data-navbar-item='true']",
        ),
        timeout_ms=1700,
    )
    if navbar_profile and _is_profile_panel_open(page):
        return True, "navbar_item_profile"

    # 1) Fallback: overflow menu -> Profile / Your profile
    menu_opened = _click_first_selector(
        page,
        (
            "button[aria-label='Menu']",
            "button[aria-label='Main menu']",
            "[data-testid='menu']",
            "header [data-testid='menu']",
        ),
        timeout_ms=1400,
    )
    if menu_opened:
        profile_from_menu = _click_first_selector(
            page,
            (
                "div[role='button']:has-text('Profile')",
                "div[role='button']:has-text('Your profile')",
                "li:has-text('Profile')",
                "li:has-text('Your profile')",
                "[role='menuitem']:has-text('Profile')",
                "[role='menuitem']:has-text('Your profile')",
            ),
            timeout_ms=1400,
        )
        if profile_from_menu and _is_profile_panel_open(page):
            return True, "menu_profile"

    return False, "none"


def _collect_whatsapp_account_snapshot(page):
    profile_name = ""
    profile_status = ""
    phone_number = ""
    opened_profile = False
    profile_open_method = "none"
    try:
        _wait_for_any_selector(
            page,
            (
                "[data-testid='chat-list']",
                "div[aria-label='Chat list']",
                "#pane-side",
            ),
            timeout_ms=120000,
        )
        page.wait_for_timeout(300)

        opened_profile, profile_open_method = _open_whatsapp_profile_panel(page)
        if opened_profile:
            page.wait_for_timeout(700)

            profile_name = _first_text_from_selectors(
                page,
                (
                    "span[data-testid='pushname-input-read-only selectable-text']",
                    "span[data-testid='pushname-input-read-only selectable-text'] span",
                    "[data-testid='drawer-right'] h2",
                    "[data-testid='drawer-right'] h1",
                    "header h2",
                    "header h1",
                    "h2",
                ),
                timeout_ms=1800,
            )
            if str(profile_name or "").strip().lower() in (
                "profile",
                "your profile",
                "you",
            ):
                profile_name = ""
            profile_status = _first_text_from_selectors(
                page,
                (
                    "[data-testid='drawer-right'] [data-testid='status']",
                    "[data-testid='drawer-right'] div:has-text('About') + div",
                    "span[data-testid='about']",
                    "[data-testid='drawer-right'] span[dir='auto']",
                    "div:has-text('About') + div",
                ),
                timeout_ms=1200,
            )
            phone_number = _first_text_from_selectors(
                page,
                (
                    "xpath=//*[@data-testid='phone']/following::span[contains(normalize-space(.),'+')][1]",
                    "div:has([data-testid='phone']) span",
                    "span:has-text('+92')",
                ),
                timeout_ms=1200,
            )
            panel_text = ""
            try:
                panel_text = (
                    page.locator("[data-testid='drawer-right']").first.inner_text(timeout=1200) or ""
                )
            except Exception:
                try:
                    panel_text = page.locator("#app").inner_text(timeout=1200) or ""
                except Exception:
                    panel_text = ""
            if panel_text:
                cleaned_panel_text = (
                    panel_text.replace("\u00a0", " ")
                    .replace("\u200e", "")
                    .replace("\u200f", "")
                    .replace("\u202a", "")
                    .replace("\u202b", "")
                    .replace("\u202c", "")
                    .replace("\u2066", "")
                    .replace("\u2067", "")
                    .replace("\u2068", "")
                    .replace("\u2069", "")
                )
                phone_match = re.search(r"\+\s*\d[\d\s\-\(\)]{6,}\d", cleaned_panel_text)
                if not phone_match:
                    phone_match = re.search(r"\b\d[\d\s\-\(\)]{6,}\d\b", cleaned_panel_text)
                if phone_match and not phone_number:
                    phone_number = re.sub(r"\s+", " ", phone_match.group(0)).strip()
                lines = [ln.strip() for ln in cleaned_panel_text.splitlines() if ln and ln.strip()]
                ignored_exact = {
                    "profile",
                    "your profile",
                    "about",
                    "phone",
                    "privacy",
                    "notifications",
                    "security",
                }
                if not profile_name:
                    for ln in lines:
                        low = ln.lower()
                        if (
                            low in ignored_exact
                            or "end-to-end encrypted" in low
                            or low.startswith("phone")
                            or re.search(r"\+\d", ln)
                            or re.search(r"\b\d{7,15}\b", ln)
                        ):
                            continue
                        if len(ln) > 1 and len(ln) <= 80 and not re.search(r"\+\d", ln):
                            profile_name = ln
                            break
                if not profile_status:
                    for idx, ln in enumerate(lines):
                        if ln.strip().lower() == "about" and idx + 1 < len(lines):
                            nxt = lines[idx + 1].strip()
                            if nxt:
                                profile_status = nxt
                                break
                if not profile_status:
                    for ln in lines:
                        low = ln.lower()
                        if (
                            not ln
                            or low in ignored_exact
                            or ln == profile_name
                            or "end-to-end encrypted" in low
                            or re.search(r"\+\d", ln)
                            or re.search(r"\b\d{7,15}\b", ln)
                        ):
                            continue
                        if 2 <= len(ln) <= 120:
                            profile_status = ln
                            break
    finally:
        try:
            page.keyboard.press("Escape")
            page.wait_for_timeout(120)
            page.keyboard.press("Escape")
        except Exception:
            pass

    return {
        "profile_name": str(profile_name or "").strip(),
        "profile_status": str(profile_status or "").strip(),
        "phone_number": str(phone_number or "").strip(),
        "profile_opened": bool(opened_profile),
        "profile_open_method": profile_open_method,
    }


def save_whatsapp_session_snapshot(page, flow_name, user_data_dir, login_state=""):
    snapshot = {
        "profile_name": "",
        "profile_status": "",
        "phone_number": "",
        "profile_opened": False,
    }
    snapshot_error = ""
    try:
        snapshot = _collect_whatsapp_account_snapshot(page)
    except Exception as exc:
        snapshot_error = str(exc)
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = (
        f"{ts} | flow={flow_name} | profile_name={snapshot.get('profile_name') or '-'} | "
        f"status={snapshot.get('profile_status') or '-'} | "
        f"number={snapshot.get('phone_number') or '-'} | "
        f"profile_opened={snapshot.get('profile_opened')} | "
        f"profile_open_method={snapshot.get('profile_open_method') or '-'} | "
        f"login_state={str(login_state or '').strip() or '-'} | "
        f"snapshot_error={snapshot_error or '-'} | "
        f"user_data_dir={user_data_dir}"
    )
    session_file = Path(__file__).parent / "session.txt"
    with session_file.open("a", encoding="utf-8") as fh:
        fh.write(line + "\n")
    return snapshot, line


CONTACT_COLUMN_ALIASES = (
    "phone",
    "phone number",
    "contacts",
    "contact",
    "phones",
)


def normalize_column_name(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    return "".join(ch for ch in text if ch.isalnum())


CONTACT_COLUMN_ALIASES_NORMALIZED = {
    normalize_column_name(v) for v in CONTACT_COLUMN_ALIASES
}

GROUP_COLUMN_CANDIDATES = ("whatspp_group", "whatsapp_group")
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".3gp", ".m4v"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".ogg", ".m4a", ".aac", ".flac", ".opus"}
DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".csv",
    ".txt",
    ".zip",
    ".ppt",
    ".pptx",
}
RECENT_MEDIA_EXTENSIONS = IMAGE_EXTENSIONS | VIDEO_EXTENSIONS | AUDIO_EXTENSIONS | DOCUMENT_EXTENSIONS


def get_db_connection():
    if psycopg2 is None:
        raise RuntimeError(
            "psycopg2 is not installed. Install it with: pip install psycopg2-binary"
        )

    conn = psycopg2.connect(
        host=os.getenv("PGHOST", "localhost"),
        port=int(os.getenv("PGPORT", "5432")),
        dbname=os.getenv("PGDATABASE", "whatsapp_db"),
        user=os.getenv("PGUSER", "postgres"),
        password=os.getenv("PGPASSWORD", ""),
    )
    conn.autocommit = True
    return conn


def get_db_schema():
    return os.getenv("PGSCHEMA", "public")


def ensure_message_log_tables(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL("DROP TABLE IF EXISTS {schema_name}.sent_sms CASCADE").format(
                schema_name=sql.Identifier(schema)
            )
        )
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.sms_sent (
                    id BIGSERIAL PRIMARY KEY,
                    campaign_id VARCHAR(6) NOT NULL,
                    campaign_name TEXT,
                    phone VARCHAR(20),
                    message TEXT,
                    status VARCHAR(20),
                    sent_at TIMESTAMP,
                    content_kind VARCHAR(50),
                    attachment_names TEXT,
                    attachment_count INTEGER DEFAULT 0,
                    payload_signature TEXT,
                    content_summary TEXT,
                    attachments_payload JSONB,
                    dedupe_signature TEXT
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.sms_not_sent (
                    id BIGSERIAL PRIMARY KEY,
                    campaign_id VARCHAR(6) NOT NULL,
                    campaign_name TEXT,
                    phone VARCHAR(20),
                    reason TEXT,
                    failed_at TIMESTAMP
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS content_kind VARCHAR(50)"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS attachment_names TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS attachment_count INTEGER DEFAULT 0'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS payload_signature TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS dedupe_signature TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS content_summary TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS attachments_payload JSONB"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS campaign_id VARCHAR(6)"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS campaign_name TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                ALTER TABLE {schema_name}.sms_sent
                ALTER COLUMN campaign_id TYPE VARCHAR(6)
                USING LPAD(COALESCE(campaign_id::text, '0'), 6, '0')
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_not_sent ADD COLUMN IF NOT EXISTS campaign_id VARCHAR(6)"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_not_sent ADD COLUMN IF NOT EXISTS campaign_name TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                ALTER TABLE {schema_name}.sms_sent
                ALTER COLUMN phone TYPE VARCHAR(255)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                ALTER TABLE {schema_name}.sms_not_sent
                ALTER COLUMN phone TYPE VARCHAR(255)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                ALTER TABLE {schema_name}.sms_not_sent
                ALTER COLUMN campaign_id TYPE VARCHAR(6)
                USING LPAD(COALESCE(campaign_id::text, '0'), 6, '0')
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL("DROP INDEX IF EXISTS {schema_name}.ux_sms_sent_phone_message").format(
                schema_name=sql.Identifier(schema)
            )
        )
        cur.execute(
            sql.SQL("DROP INDEX IF EXISTS {schema_name}.ux_sms_sent_phone_payload").format(
                schema_name=sql.Identifier(schema)
            )
        )
        cur.execute(
            sql.SQL(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS ux_sms_sent_phone_payload
                ON {schema_name}.sms_sent (phone, payload_signature, campaign_id)
                WHERE payload_signature IS NOT NULL
                  AND campaign_id IS NOT NULL
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_phone_dedupe
                ON {schema_name}.sms_sent (phone, dedupe_signature)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_phone
                ON {schema_name}.sms_sent (phone)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_sent_at
                ON {schema_name}.sms_sent (sent_at DESC)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_campaign_id
                ON {schema_name}.sms_sent (campaign_id)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_campaign_name
                ON {schema_name}.sms_sent (campaign_name)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_not_sent_phone
                ON {schema_name}.sms_not_sent (phone)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_not_sent_failed_at
                ON {schema_name}.sms_not_sent (failed_at DESC)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_not_sent_campaign_id
                ON {schema_name}.sms_not_sent (campaign_id)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_not_sent_campaign_name
                ON {schema_name}.sms_not_sent (campaign_name)
                """
            ).format(schema_name=sql.Identifier(schema))
        )


def _normalize_group_name(value):
    return str(value or "").strip()


def _normalize_group_key(value):
    text = _normalize_group_name(value)
    if not text:
        return ""
    text = text.replace("\u200b", "").replace("\u200c", "").replace("\u200d", "").replace("\ufeff", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text.casefold()


def _safe_int(value, default=0):
    try:
        return int(value)
    except Exception:
        return default


def _parse_optional_limit(raw_value, default=None):
    try:
        value = int(raw_value)
    except Exception:
        return default
    if value <= 0:
        return None
    return value


def generate_campaign_id():
    schema = get_db_schema()
    db_conn = get_db_connection()
    try:
        ensure_message_log_tables(db_conn)
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    WITH all_ids AS (
                        SELECT campaign_id::text AS campaign_id
                        FROM {schema_name}.sms_sent
                        UNION ALL
                        SELECT campaign_id::text AS campaign_id
                        FROM {schema_name}.sms_not_sent
                    )
                    SELECT COALESCE(
                        MAX(
                            CASE
                                WHEN campaign_id ~ '^[0-9]+$' THEN campaign_id::BIGINT
                                ELSE NULL
                            END
                        ),
                        0
                    )
                    FROM all_ids
                    """
                ).format(schema_name=sql.Identifier(schema))
            )
            row = cur.fetchone()
            last_value = int(row[0]) if row and row[0] is not None else 0
            next_value = last_value + 1
            return f"{next_value:06d}"
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def _normalize_campaign_id(value):
    text = str(value or "").strip()
    digits = "".join(ch for ch in text if ch.isdigit())
    if not digits:
        return ""
    return digits.zfill(6)


def _json_datetime(value):
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def _looks_like_phone_target(value):
    text = str(value or "").strip()
    if not text:
        return False
    compact = re.sub(r"[\s\-\(\)\.]", "", text)
    if compact.startswith("+"):
        compact = compact[1:]
    if not compact.isdigit():
        return False
    return 7 <= len(compact) <= 16


def list_campaign_summaries(limit=None, query_text=""):
    schema = get_db_schema()
    db_conn = get_db_connection()
    try:
        ensure_message_log_tables(db_conn)
        q = str(query_text or "").strip()
        q_like = f"%{q}%"
        with db_conn.cursor() as cur:
            query = """
                    WITH sent AS (
                        SELECT
                            LPAD(COALESCE(campaign_id::text, ''), 6, '0') AS campaign_id,
                            MAX(NULLIF(BTRIM(campaign_name), '')) AS campaign_name,
                            COUNT(*) AS sent_rows,
                            COUNT(DISTINCT phone) AS sent_total,
                            MIN(sent_at) AS campaign_start,
                            MAX(sent_at) AS max_sent_at
                        FROM {schema_name}.sms_sent
                        WHERE COALESCE(campaign_id::text, '') <> ''
                        GROUP BY 1
                    ),
                    failed AS (
                        SELECT
                            LPAD(COALESCE(campaign_id::text, ''), 6, '0') AS campaign_id,
                            MAX(NULLIF(BTRIM(campaign_name), '')) AS campaign_name,
                            COUNT(*) AS failed_rows,
                            COUNT(DISTINCT phone) AS failed_total,
                            MIN(failed_at) AS failed_start,
                            MAX(failed_at) AS max_failed_at
                        FROM {schema_name}.sms_not_sent
                        WHERE COALESCE(campaign_id::text, '') <> ''
                        GROUP BY 1
                    )
                    SELECT
                        COALESCE(sent.campaign_id, failed.campaign_id) AS campaign_id,
                        COALESCE(sent.campaign_name, failed.campaign_name, '') AS campaign_name,
                        COALESCE(sent.sent_total, 0) AS sent_total,
                        COALESCE(failed.failed_total, 0) AS failed_total,
                        COALESCE(sent.sent_rows, 0) AS sent_rows,
                        COALESCE(failed.failed_rows, 0) AS failed_rows,
                        COALESCE(sent.sent_total, 0) + COALESCE(failed.failed_total, 0) AS contacts_total,
                        sent.campaign_start,
                        failed.failed_start,
                        COALESCE(
                            GREATEST(sent.max_sent_at, failed.max_failed_at),
                            sent.max_sent_at,
                            failed.max_failed_at
                        ) AS campaign_end
                    FROM sent
                    FULL OUTER JOIN failed
                      ON sent.campaign_id = failed.campaign_id
                    WHERE (
                        %s = ''
                        OR COALESCE(sent.campaign_id, failed.campaign_id) ILIKE %s
                        OR COALESCE(sent.campaign_name, failed.campaign_name, '') ILIKE %s
                        OR (
                            COALESCE(sent.campaign_id, failed.campaign_id)
                            || '_'
                            || COALESCE(sent.campaign_name, failed.campaign_name, '')
                        ) ILIKE %s
                    )
                    ORDER BY campaign_end DESC NULLS LAST, campaign_id DESC
                    """
            params = [q, q_like, q_like, q_like]
            if limit is not None:
                query += "\nLIMIT %s"
                params.append(int(limit))
            cur.execute(
                sql.SQL(query).format(schema_name=sql.Identifier(schema)),
                tuple(params),
            )
            rows = cur.fetchall() or []

        campaigns = []
        for row in rows:
            sent_start = row[7]
            failed_start = row[8]
            if sent_start and failed_start:
                campaign_start = min(sent_start, failed_start)
            else:
                campaign_start = sent_start or failed_start
            campaigns.append(
                {
                    "campaign_id": str(row[0] or "").zfill(6),
                    "campaign_name": str(row[1] or "").strip(),
                    "sent_total": _safe_int(row[2]),
                    "failed_total": _safe_int(row[3]),
                    "sent_rows": _safe_int(row[4]),
                    "failed_rows": _safe_int(row[5]),
                    "contacts_total": _safe_int(row[6]),
                    "campaign_start": _json_datetime(campaign_start),
                    "campaign_end": _json_datetime(row[9]),
                }
            )
        return campaigns
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def get_campaign_detail(campaign_id, list_limit=None):
    normalized = _normalize_campaign_id(campaign_id)
    if not normalized:
        return None

    schema = get_db_schema()
    db_conn = get_db_connection()
    try:
        ensure_message_log_tables(db_conn)
        with db_conn.cursor() as cur:
            sent_query = """
                    SELECT phone, sent_at, message, content_kind, attachment_names, attachment_count, content_summary, campaign_name
                    FROM {schema_name}.sms_sent
                    WHERE LPAD(COALESCE(campaign_id::text, ''), 6, '0') = %s
                    ORDER BY sent_at ASC, id ASC
                    """
            sent_params = [normalized]
            if list_limit is not None:
                sent_query += "\nLIMIT %s"
                sent_params.append(int(list_limit))
            cur.execute(
                sql.SQL(sent_query).format(schema_name=sql.Identifier(schema)),
                tuple(sent_params),
            )
            sent_rows = cur.fetchall() or []

            failed_query = """
                    SELECT phone, failed_at, reason, campaign_name
                    FROM {schema_name}.sms_not_sent
                    WHERE LPAD(COALESCE(campaign_id::text, ''), 6, '0') = %s
                    ORDER BY failed_at ASC, id ASC
                    """
            failed_params = [normalized]
            if list_limit is not None:
                failed_query += "\nLIMIT %s"
                failed_params.append(int(list_limit))
            cur.execute(
                sql.SQL(failed_query).format(schema_name=sql.Identifier(schema)),
                tuple(failed_params),
            )
            failed_rows = cur.fetchall() or []

        sent_numbers = []
        sent_seen = set()
        sent_times = []
        kinds_seen = set()
        attachments_seen = set()
        message_samples = []
        campaign_name = ""
        for phone, sent_at, message, content_kind, attachment_names_raw, attachment_count, content_summary, row_campaign_name in sent_rows:
            if not campaign_name:
                campaign_name = str(row_campaign_name or "").strip()
            number = str(phone or "").strip()
            if number and number not in sent_seen:
                sent_numbers.append(number)
                sent_seen.add(number)
            if sent_at:
                sent_times.append(sent_at)
            ck = str(content_kind or "").strip().lower()
            if ck:
                kinds_seen.add(ck)
            if message:
                msg = str(message).strip()
                if msg and msg not in message_samples:
                    message_samples.append(msg)
            elif content_summary:
                msg = str(content_summary).strip()
                if msg and msg not in message_samples:
                    message_samples.append(msg)
            if attachment_names_raw:
                parsed_names = []
                try:
                    loaded = json.loads(attachment_names_raw) if isinstance(attachment_names_raw, str) else attachment_names_raw
                    if isinstance(loaded, list):
                        parsed_names = loaded
                except Exception:
                    parsed_names = []
                for item in parsed_names:
                    name = str(item or "").strip()
                    if name:
                        attachments_seen.add(name)
                if not parsed_names and int(attachment_count or 0) > 0:
                    attachments_seen.add(f"{int(attachment_count)} file(s)")

        failed_numbers = []
        failed_items = []
        failed_seen = set()
        failed_times = []
        for phone, failed_at, reason, row_campaign_name in failed_rows:
            if not campaign_name:
                campaign_name = str(row_campaign_name or "").strip()
            number = str(phone or "").strip()
            if number and number not in failed_seen:
                failed_numbers.append(number)
                failed_seen.add(number)
            if number:
                failed_items.append(
                    {
                        "phone": number,
                        "failed_at": _json_datetime(failed_at),
                        "reason": str(reason or ""),
                    }
                )
            if failed_at:
                failed_times.append(failed_at)

        all_numbers = set(sent_numbers) | set(failed_numbers)
        targets = []
        seen_targets = set()
        for number in sent_numbers + failed_numbers:
            if number and number not in seen_targets:
                targets.append(number)
                seen_targets.add(number)

        phone_like = 0
        non_phone_like = 0
        for target in targets:
            if _looks_like_phone_target(target):
                phone_like += 1
            else:
                non_phone_like += 1
        if phone_like and not non_phone_like:
            send_target_hint = "contacts"
        elif non_phone_like and not phone_like:
            send_target_hint = "groups"
        elif non_phone_like >= phone_like and non_phone_like > 0:
            send_target_hint = "groups"
        else:
            send_target_hint = "contacts"

        campaign_start = min(sent_times + failed_times) if (sent_times or failed_times) else None
        combined_times = sent_times + failed_times
        campaign_end = max(combined_times) if combined_times else None
        attachments_list = sorted(attachments_seen)
        kinds_list = sorted(kinds_seen)
        if attachments_list:
            content_mode = "files"
        elif kinds_list and all(k in ("text", "sms") for k in kinds_list):
            content_mode = "sms"
        elif not sent_rows:
            content_mode = "unknown"
        else:
            content_mode = "mixed"

        return {
            "campaign_id": normalized,
            "campaign_name": campaign_name,
            "contacts_total": len(all_numbers),
            "sent_total": len(sent_numbers),
            "failed_total": len(failed_numbers),
            "sent_rows": len(sent_rows),
            "failed_rows": len(failed_rows),
            "campaign_start": _json_datetime(campaign_start),
            "campaign_end": _json_datetime(campaign_end),
            "sent_numbers": sent_numbers,
            "failed_numbers": failed_numbers,
            "failed_items": failed_items,
            "targets": targets,
            "send_target_hint": send_target_hint,
            "campaign_content": {
                "mode": content_mode,
                "content_kinds": kinds_list,
                "messages": message_samples[:5],
                "attachments": attachments_list,
            },
        }
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def ensure_group_metadata_table(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.group_metadata (
                    group_name TEXT PRIMARY KEY,
                    description TEXT,
                    members_label TEXT,
                    members_count INTEGER,
                    created_by TEXT,
                    created_on TEXT,
                    invite_link TEXT,
                    source_scrap_id TEXT,
                    last_scraped_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS description TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS members_label TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS members_count INTEGER'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS created_by TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS created_on TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS invite_link TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS source_scrap_id TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS last_scraped_at TIMESTAMPTZ NOT NULL DEFAULT NOW()'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()'
            ).format(schema_name=sql.Identifier(schema))
        )


def get_all_contacts_group_column(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = %s
              AND table_name = 'all_contacts'
            """,
            (schema,),
        )
        columns = {row[0] for row in cur.fetchall() if row and row[0]}
    for name in GROUP_COLUMN_CANDIDATES:
        if name in columns:
            return name
    return None


def load_existing_group_names_from_all_contacts(db_conn):
    schema = get_db_schema()

    # Fast path: group_metadata has one row per group (much smaller than all_contacts).
    try:
        ensure_group_metadata_table(db_conn)
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT group_name
                    FROM {schema_name}.group_metadata
                    WHERE group_name IS NOT NULL
                      AND TRIM(group_name) <> ''
                    """
                ).format(schema_name=sql.Identifier(schema))
            )
            metadata_groups = {
                _normalize_group_key(row[0])
                for row in cur.fetchall()
                if row and _normalize_group_key(row[0])
            }
            if metadata_groups:
                return metadata_groups
    except Exception:
        pass

    with db_conn.cursor() as cur:
        cur.execute(
            """
            SELECT EXISTS (
                SELECT 1
                FROM information_schema.tables
                WHERE table_schema = %s
                  AND table_name = 'all_contacts'
            )
            """,
            (schema,),
        )
        table_exists = bool(cur.fetchone()[0])
        if not table_exists:
            return set()

    group_col = get_all_contacts_group_column(db_conn)
    if not group_col:
        return set()

    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                SELECT DISTINCT TRIM({group_col}) AS group_name
                FROM {schema_name}.all_contacts
                WHERE {group_col} IS NOT NULL
                  AND TRIM({group_col}) <> ''
                """
            ).format(
                schema_name=sql.Identifier(schema),
                group_col=sql.Identifier(group_col),
            )
        )
        return {
            _normalize_group_key(row[0])
            for row in cur.fetchall()
            if row and _normalize_group_key(row[0])
        }


def list_unique_groups_from_all_contacts(limit=None, query_text=""):
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        ensure_group_metadata_table(db_conn)
        group_col = get_all_contacts_group_column(db_conn)
        if not group_col:
            return []
        q = _normalize_group_name(query_text).lower()
        with db_conn.cursor() as cur:
            query = """
                    SELECT
                        src.group_name,
                        src.contacts_count,
                        gm.members_label,
                        gm.members_count,
                        gm.description,
                        gm.created_by,
                        gm.created_on,
                        gm.invite_link,
                        gm.source_scrap_id,
                        gm.last_scraped_at
                    FROM (
                        SELECT
                            MIN(TRIM({group_col})) AS group_name,
                            COUNT(DISTINCT "Phone") AS contacts_count
                        FROM {schema_name}.all_contacts
                        WHERE {group_col} IS NOT NULL
                          AND TRIM({group_col}) <> ''
                          AND LOWER(TRIM({group_col})) LIKE %s
                        GROUP BY LOWER(TRIM({group_col}))
                    ) src
                    LEFT JOIN {schema_name}.group_metadata gm
                      ON LOWER(TRIM(gm.group_name)) = LOWER(TRIM(src.group_name))
                    ORDER BY src.group_name ASC
                    """
            params = [f"%{q}%"]
            if limit is not None:
                query += "\nLIMIT %s"
                params.append(int(limit))
            cur.execute(
                sql.SQL(query).format(
                    schema_name=sql.Identifier(schema),
                    group_col=sql.Identifier(group_col),
                ),
                tuple(params),
            )
            rows = cur.fetchall()
            results = []
            for row in rows:
                if not row or not row[0]:
                    continue
                results.append(
                    {
                        "group_name": row[0],
                        "contacts_count": _safe_int(row[1]),
                        "members_label": row[2] or "",
                        "members_count": _safe_int(row[3], default=0) if row[3] is not None else None,
                        "description": row[4] or "",
                        "created_by": row[5] or "",
                        "created_on": row[6] or "",
                        "invite_link": row[7] or "",
                        "source_scrap_id": row[8] or "",
                        "last_scraped_at": row[9].isoformat() if row[9] else "",
                    }
                )
            return results
    finally:
        db_conn.close()


def load_group_contacts_from_all_contacts(group_name, limit=None):
    normalized = _normalize_group_name(group_name)
    if not normalized:
        return []
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        group_col = get_all_contacts_group_column(db_conn)
        if not group_col:
            return []
        with db_conn.cursor() as cur:
            query = """
                    SELECT DISTINCT "Phone"
                    FROM {schema_name}.all_contacts
                    WHERE LOWER(TRIM({group_col})) = LOWER(TRIM(%s))
                      AND "Phone" IS NOT NULL
                      AND TRIM("Phone") <> ''
                    ORDER BY "Phone" ASC
                    """
            params = [normalized]
            if limit is not None:
                query += "\nLIMIT %s"
                params.append(int(limit))
            cur.execute(
                sql.SQL(query).format(
                    schema_name=sql.Identifier(schema),
                    group_col=sql.Identifier(group_col),
                ),
                tuple(params),
            )
            return [str(row[0]).strip() for row in cur.fetchall() if row and row[0]]
    finally:
        db_conn.close()


def load_contacts_from_multiple_groups(group_names, limit=None):
    cleaned_groups = []
    seen = set()
    for name in group_names or []:
        value = _normalize_group_name(name)
        key = value.lower()
        if not value or key in seen:
            continue
        seen.add(key)
        cleaned_groups.append(value)
    if not cleaned_groups:
        return {"groups": [], "contacts": [], "contact_items": []}

    summaries = list_unique_groups_from_all_contacts(limit=None, query_text="")
    summary_map = {
        _normalize_group_name(item.get("group_name")).lower(): item
        for item in summaries
        if _normalize_group_name(item.get("group_name"))
    }

    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        group_col = get_all_contacts_group_column(db_conn)
        if not group_col:
            return {"groups": [], "contacts": [], "contact_items": []}
        with db_conn.cursor() as cur:
            query = """
                    SELECT
                        MIN(TRIM({group_col})) AS group_name,
                        "Phone"
                    FROM {schema_name}.all_contacts
                    WHERE LOWER(TRIM({group_col})) = ANY(%s)
                      AND "Phone" IS NOT NULL
                      AND TRIM("Phone") <> ''
                    GROUP BY LOWER(TRIM({group_col})), "Phone"
                    ORDER BY MIN(TRIM({group_col})) ASC, "Phone" ASC
                    """
            params = [[g.lower() for g in cleaned_groups]]
            if limit is not None:
                query += "\nLIMIT %s"
                params.append(int(limit))
            cur.execute(
                sql.SQL(query).format(
                    schema_name=sql.Identifier(schema),
                    group_col=sql.Identifier(group_col),
                ),
                tuple(params),
            )
            rows = cur.fetchall()
    finally:
        db_conn.close()

    grouped = {}
    merged_contacts = []
    merged_seen = set()
    for row in rows:
        if not row:
            continue
        group_name = _normalize_group_name(row[0])
        phone = str(row[1] or "").strip()
        if not group_name or not phone:
            continue
        grouped.setdefault(group_name, [])
        grouped[group_name].append(phone)
        if phone not in merged_seen:
            merged_seen.add(phone)
            merged_contacts.append(phone)

    selected_groups = []
    for group_name in cleaned_groups:
        key = group_name.lower()
        group_contacts = grouped.get(group_name, [])
        summary = summary_map.get(key) or {"group_name": group_name, "contacts_count": len(group_contacts)}
        selected_groups.append(
            {
                **summary,
                "group_name": summary.get("group_name") or group_name,
                "selected_contacts_count": len(group_contacts),
            }
        )

    merged_contacts = sorted(merged_contacts)
    return {
        "groups": selected_groups,
        "contacts": merged_contacts,
        "contact_items": build_contact_items(merged_contacts, {}),
    }


def get_group_detail_from_all_contacts(group_name, contact_limit=None):
    normalized = _normalize_group_name(group_name)
    if not normalized:
        return None
    summaries = list_unique_groups_from_all_contacts(limit=None, query_text=normalized)
    summary = None
    for item in summaries:
        if _normalize_group_name(item.get("group_name")).lower() == normalized.lower():
            summary = item
            break
    if summary is None:
        return None
    contacts = load_group_contacts_from_all_contacts(normalized, limit=contact_limit)
    return {
        **summary,
        "contacts": contacts,
        "contact_items": build_contact_items(contacts, {}),
    }


def list_group_tables():
    print("[DB] Loading tables: Connecting to database...")
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            print("[DB] Loading tables: Accessing information_schema...")
            cur.execute(
                """
                SELECT table_name
                FROM information_schema.tables
                WHERE table_schema = %s
                  AND table_type = 'BASE TABLE'
                ORDER BY table_name
                """,
                (schema,),
            )
            table_rows = cur.fetchall()
            tables = [
                row[0] for row in table_rows if row and len(row) > 0 and row[0]
            ]
            print(f"[DB] Loading tables: Loaded {len(tables)} tables from schema '{schema}'.")
            return tables
    except Exception as exc:
        print(f"[DB ERROR] list_group_tables failed: {exc}")
        traceback.print_exc()
        raise
    finally:
        db_conn.close()


def parse_selected_tables(raw_selected):
    if not raw_selected:
        return []
    if isinstance(raw_selected, list):
        values = raw_selected
        cleaned = []
        for item in values:
            text = str(item).strip()
            if text:
                cleaned.append(text)
        return list(dict.fromkeys(cleaned))
    try:
        parsed = json.loads(raw_selected)
        if isinstance(parsed, list):
            values = parsed
        else:
            values = []
    except Exception:
        values = str(raw_selected).split(",")
    cleaned = []
    for item in values:
        text = str(item).strip()
        if not text:
            continue
        cleaned.append(text)
    return list(dict.fromkeys(cleaned))


def load_numbers_from_tables(selected_tables):
    if not selected_tables:
        return []
    if psycopg2 is None or sql is None:
        raise RuntimeError(
            "psycopg2 is not installed. Install it with: pip install psycopg2-binary"
        )
    allowed_tables = set(list_group_tables())
    chosen = [t for t in selected_tables if t in allowed_tables]
    if not chosen:
        raise RuntimeError("No valid group_ tables selected.")
    print(f"[DB] Sending flow: Using {len(chosen)} selected tables.")

    db_conn = get_db_connection()
    schema = get_db_schema()
    numbers = []
    try:
        with db_conn.cursor() as cur:
            for table_name in chosen:
                print(f"[DB] Reading table: {table_name}")
                cur.execute(
                    """
                    SELECT column_name
                    FROM information_schema.columns
                    WHERE table_schema = %s
                      AND table_name = %s
                    """,
                    (schema, table_name),
                )
                phone_col = None
                for col_row in cur.fetchall():
                    if not col_row or len(col_row) == 0:
                        continue
                    col_name = col_row[0]
                    if normalize_column_name(col_name) in CONTACT_COLUMN_ALIASES_NORMALIZED:
                        phone_col = col_name
                        break
                if not phone_col:
                    continue
                cur.execute(
                    sql.SQL(
                        """
                        SELECT {phone_col}
                        FROM {schema_name}.{table_name}
                        WHERE {phone_col} IS NOT NULL
                        """
                    ).format(
                        phone_col=sql.Identifier(phone_col),
                        schema_name=sql.Identifier(schema),
                        table_name=sql.Identifier(table_name),
                    )
                )
                for row in cur.fetchall():
                    raw = row[0] if row and len(row) > 0 else None
                    if raw is None:
                        continue
                    value = str(raw).strip()
                    if not value:
                        continue
                    numbers.append(value)
    finally:
        db_conn.close()

    return list(dict.fromkeys(numbers))


def get_tables_info(selected_tables):
    if not selected_tables:
        return []
    if psycopg2 is None or sql is None:
        raise RuntimeError(
            "psycopg2 is not installed. Install it with: pip install psycopg2-binary"
        )

    allowed_tables = set(list_group_tables())
    chosen = [t for t in selected_tables if t in allowed_tables]
    if not chosen:
        return []

    print(f"[DB] Table info: Loading metadata for {len(chosen)} table(s)...")
    db_conn = get_db_connection()
    schema = get_db_schema()
    summaries = []
    try:
        with db_conn.cursor() as cur:
            for table_name in chosen:
                cur.execute(
                    """
                    SELECT column_name
                    FROM information_schema.columns
                    WHERE table_schema = %s
                      AND table_name = %s
                    ORDER BY ordinal_position
                    """,
                    (schema, table_name),
                )
                column_rows = cur.fetchall()
                columns = [row[0] for row in column_rows if row and len(row) > 0 and row[0]]

                phone_col = None
                for c in columns:
                    if normalize_column_name(c) in CONTACT_COLUMN_ALIASES_NORMALIZED:
                        phone_col = c
                        break

                phone_count = 0
                if phone_col:
                    cur.execute(
                        sql.SQL(
                            """
                            SELECT COUNT(DISTINCT {phone_col})
                            FROM {schema_name}.{table_name}
                            WHERE {phone_col} IS NOT NULL
                            """
                        ).format(
                            schema_name=sql.Identifier(schema),
                            table_name=sql.Identifier(table_name),
                            phone_col=sql.Identifier(phone_col),
                        )
                    )
                    count_row = cur.fetchone()
                    phone_count = int(count_row[0]) if count_row and len(count_row) > 0 else 0

                summaries.append(
                    {
                        "table": table_name,
                        "columns": columns,
                        "phone_column": phone_col,
                        "phone_count": phone_count,
                    }
                )
    finally:
        db_conn.close()

    print("[DB] Table info: Loaded metadata.")
    return summaries


def detect_attachment_kind(filename, mimetype=""):
    mime = str(mimetype or "").strip().lower()
    if mime.startswith("image/"):
        return "image"
    if mime.startswith("video/"):
        return "video"
    if mime.startswith("audio/"):
        return "audio"
    ext = Path(str(filename or "")).suffix.lower()
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in VIDEO_EXTENSIONS:
        return "video"
    if ext in AUDIO_EXTENSIONS:
        return "audio"
    return "file"


def is_supported_attachment_filename(filename):
    ext = Path(str(filename or "")).suffix.lower()
    if not ext:
        return False
    return (
        ext in IMAGE_EXTENSIONS
        or ext in VIDEO_EXTENSIONS
        or ext in AUDIO_EXTENSIONS
        or ext in DOCUMENT_EXTENSIONS
    )


def save_uploaded_attachments(file_storages):
    uploads_dir = Path(__file__).parent / "send_uploads"
    uploads_dir.mkdir(parents=True, exist_ok=True)
    saved = []
    for file_storage in file_storages or []:
        if not file_storage or not file_storage.filename:
            continue
        original_name = secure_filename(file_storage.filename or "")
        if not original_name:
            continue
        if not is_supported_attachment_filename(original_name):
            print(f"[WARN] Skipping unsupported attachment extension: {original_name}")
            continue
        suffix = Path(original_name).suffix
        stored_name = f"{Path(original_name).stem}_{uuid.uuid4().hex[:10]}{suffix}"
        target_path = uploads_dir / stored_name
        file_storage.save(target_path)
        saved.append(
            {
                "name": original_name,
                "path": str(target_path),
                "kind": detect_attachment_kind(original_name, getattr(file_storage, "mimetype", "")),
            }
        )
    return saved


def build_payload(message, attachments, message_variants=None):
    attachments = attachments or []
    normalized_variants = []
    for item in message_variants or []:
        if not isinstance(item, dict):
            continue
        text_value = str(item.get("text") or "").strip()
        if not text_value:
            continue
        try:
            slot = int(item.get("slot") or 0)
        except Exception:
            slot = 0
        if slot <= 0:
            slot = len(normalized_variants) + 1
        normalized_variants.append({"slot": slot, "text": text_value})

    normalized_attachments = []
    for item in attachments:
        normalized_attachments.append(
            {
                "name": item.get("name") or "",
                "kind": item.get("kind") or "file",
                "path": item.get("path") or "",
            }
        )
    kinds = {item.get("kind") or "file" for item in attachments}
    if not attachments:
        content_kind = "text"
    elif len(kinds) == 1:
        content_kind = next(iter(kinds))
    else:
        content_kind = "mixed"
    signature_payload = {
        "message": message or "",
        "attachments": normalized_attachments,
    }
    signature = json.dumps(signature_payload, ensure_ascii=False, sort_keys=True)
    dedupe_payload = {
        "message": message or "",
        "attachments": [
            {"name": item.get("name") or "", "kind": item.get("kind") or "file"}
            for item in normalized_attachments
        ],
    }
    dedupe_signature = json.dumps(dedupe_payload, ensure_ascii=False, sort_keys=True)
    content_summary = message or ""
    if attachments:
        names = ", ".join(item.get("name") or "" for item in attachments if item.get("name"))
        content_summary = (message + " | " if message else "") + names
    return {
        "message": message or "",
        "message_variants": normalized_variants,
        "attachments": normalized_attachments,
        "content_kind": content_kind,
        "attachment_count": len(attachments),
        "attachment_names": [item.get("name") or "" for item in normalized_attachments],
        "attachment_paths": [item.get("path") or "" for item in normalized_attachments],
        "payload_signature": signature,
        "dedupe_signature": dedupe_signature,
        "content_summary": content_summary[:1000],
    }


def _pick_random_message_for_target(payload):
    variants = list((payload or {}).get("message_variants") or [])
    if variants:
        picked = random.choice(variants)
        return int(picked.get("slot") or 1), str(picked.get("text") or ""), len(variants)
    return 1, str((payload or {}).get("message") or ""), 0


def _pick_random_attachment_for_target(payload):
    attachments = list((payload or {}).get("attachments") or [])
    total = len(attachments)
    if total == 0:
        return None, 0, 0
    if total == 1:
        item = attachments[0]
        return item, 1, 1
    picked_index = random.randint(1, total)
    item = attachments[picked_index - 1]
    return item, picked_index, total


def _init_live_rows(targets, batch_numbers=None):
    rows = []
    for idx, target in enumerate(targets or [], start=1):
        batch_no = 1
        if batch_numbers and (idx - 1) < len(batch_numbers):
            try:
                batch_no = int(batch_numbers[idx - 1])
            except Exception:
                batch_no = 1
        rows.append(
            {
                "count": idx,
                "contact": str(target or "").strip(),
                "batch": batch_no,
                "message": "-",
                "media": "-",
                "delay": "-",
                "status": "pending",
            }
        )
    return rows


def _update_live_row(state, row_index, **fields):
    try:
        idx = int(row_index)
    except Exception:
        return
    rows = list((state or {}).get("live_rows") or [])
    if idx < 0 or idx >= len(rows):
        return
    row = dict(rows[idx] or {})
    for k, v in fields.items():
        row[k] = v
    rows[idx] = row
    state["live_rows"] = rows


def _build_random_batch_sizes(total_items, min_size=20, max_size=30):
    try:
        total = max(0, int(total_items or 0))
    except Exception:
        total = 0
    if total <= 0:
        return []
    low = max(1, int(min_size or 1))
    high = max(low, int(max_size or low))
    sizes = []
    remaining = total
    while remaining > 0:
        if remaining <= high:
            size = remaining
        else:
            size = random.randint(low, high)
        sizes.append(size)
        remaining -= size
    return sizes


def _expand_batch_numbers(batch_sizes):
    numbers = []
    for batch_no, size in enumerate(batch_sizes or [], start=1):
        for _ in range(max(0, int(size or 0))):
            numbers.append(batch_no)
    return numbers


def is_already_sent(db_conn, phone, payload):
    schema = get_db_schema()
    payload_signature = str((payload or {}).get("payload_signature") or "")
    dedupe_signature = str((payload or {}).get("dedupe_signature") or "")
    message = str((payload or {}).get("message") or "")
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                SELECT 1
                FROM {schema_name}.sms_sent
                WHERE phone = %s
                  AND (
                    dedupe_signature = %s
                    OR
                    payload_signature = %s
                    OR (payload_signature IS NULL AND message = %s)
                  )
                LIMIT 1
                """
            ).format(schema_name=sql.Identifier(schema)),
            (phone, dedupe_signature, payload_signature, message),
        )
        return cur.fetchone() is not None


def save_sent_sms(db_conn, phone, payload, campaign_id=None, campaign_name=""):
    schema = get_db_schema()
    attachment_names = json.dumps((payload or {}).get("attachment_names") or [], ensure_ascii=False)
    attachments_payload = json.dumps((payload or {}).get("attachments") or [], ensure_ascii=False)
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                INSERT INTO {schema_name}.sms_sent (
                    phone,
                    message,
                    status,
                    sent_at,
                    campaign_id,
                    campaign_name,
                    content_kind,
                    attachment_names,
                    attachment_count,
                    payload_signature,
                    dedupe_signature,
                    content_summary,
                    attachments_payload
                )
                VALUES (%s, %s, 'sent', NOW(), %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
                ON CONFLICT DO NOTHING
                """
            ).format(schema_name=sql.Identifier(schema)),
            (
                phone,
                (payload or {}).get("message") or "",
                campaign_id,
                str(campaign_name or "").strip(),
                (payload or {}).get("content_kind") or "text",
                attachment_names,
                int((payload or {}).get("attachment_count") or 0),
                (payload or {}).get("payload_signature") or None,
                (payload or {}).get("dedupe_signature") or None,
                (payload or {}).get("content_summary") or "",
                attachments_payload,
            ),
        )


def save_not_sent_sms(db_conn, campaign_id, campaign_name, phone, reason):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                INSERT INTO {schema_name}.sms_not_sent (campaign_id, campaign_name, phone, reason, failed_at)
                VALUES (%s, %s, %s, %s, NOW())
                """
            ).format(schema_name=sql.Identifier(schema)),
            (str(campaign_id or "").zfill(6), str(campaign_name or "").strip(), phone, reason),
        )

def is_valid_phone(num):
    try:
        if not num.startswith("+"):
            num = "+" + num
        parsed = phonenumbers.parse(num, None)
        # Be permissive: accept numbers that are possible even if not strictly valid.
        return phonenumbers.is_possible_number(parsed) or phonenumbers.is_valid_number(parsed)
    except:
        # Fallback: basic check for + followed by 7-16 digits
        digits = "".join(ch for ch in str(num) if ch.isdigit())
        return 7 <= len(digits) <= 16


def normalize_phone(num):
    if num is None:
        return None
    num = str(num).strip()
    if not num:
        return None
    # Keep leading +, strip everything else to digits
    cleaned = "".join(ch for ch in num if ch.isdigit() or ch == "+")
    if cleaned.startswith("00"):
        cleaned = "+" + cleaned[2:]
    if cleaned.startswith("+"):
        return cleaned
    num = cleaned
    if num.startswith("+"):
        return num
    return "+" + num


def parse_country_code_filter(raw_filter):
    if not raw_filter:
        return []
    tokens = []
    for part in str(raw_filter).replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        digits = "".join(ch for ch in part if ch.isdigit())
        if not digits:
            continue
        tokens.append("+" + digits)
    # Preserve order, remove duplicates
    return list(dict.fromkeys(tokens))


def filter_numbers_by_country_codes(numbers, raw_filter):
    prefixes = parse_country_code_filter(raw_filter)
    if not prefixes:
        return numbers, 0, []
    filtered = [n for n in numbers if any(n.startswith(prefix) for prefix in prefixes)]
    skipped = max(0, len(numbers) - len(filtered))
    return filtered, skipped, prefixes


def _normalize_header(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    return "".join(ch for ch in text if ch.isalnum())


def load_numbers(file_storage):
    filename = secure_filename(file_storage.filename or "")
    ext = Path(filename).suffix.lower()
    numbers = []
 
    if ext == ".csv":
        text = file_storage.stream.read().decode("utf-8", errors="ignore")
        reader = csv.DictReader(text.splitlines())
        columns = [_normalize_header(c) for c in reader.fieldnames or []]
        key = None
        for candidate in CONTACT_COLUMN_ALIASES_NORMALIZED:
            if candidate in columns:
                key = candidate
                break
        for row in reader:
            if key:
                raw = row.get(key) or row.get(key.replace("_", ""))
                if raw is None:
                    # DictReader keeps original headers, so map by index if needed
                    try:
                        idx = columns.index(key)
                        raw = list(row.values())[idx]
                    except Exception:
                        raw = None
            else:
                values = list(row.values())
                raw = values[1] if len(values) > 1 else values[0] if values else None
            numbers.append(normalize_phone(raw))
        return [n for n in numbers if n]

    if ext == ".xlsx":
        try:
            import openpyxl
        except Exception:
            raise RuntimeError("openpyxl is required for .xlsx files")
        wb = openpyxl.load_workbook(file_storage, read_only=True)
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
        header = [_normalize_header(c) for c in header_row]
        col_index = None
        for candidate in CONTACT_COLUMN_ALIASES_NORMALIZED:
            if candidate in header:
                col_index = header.index(candidate)
                break
        for row in ws.iter_rows(min_row=2, values_only=True):
            if col_index is not None and col_index < len(row):
                raw = row[col_index]
            else:
                raw = row[1] if len(row) > 1 else row[0] if row else None
            numbers.append(normalize_phone(raw))
        return [n for n in numbers if n]

    raise RuntimeError("Unsupported file type. Use .csv or .xlsx")


def build_contact_items(contacts, status_map=None):
    status_map = status_map or {}
    items = []
    for number in contacts or []:
        info = status_map.get(number) or {}
        status = str(info.get("status") or "pending")
        reason = str(info.get("reason") or "")
        items.append(
            {
                "number": number,
                "status": status,
                "sent": status == "sent",
                "reason": reason,
            }
        )
    return items


def _click_first_available(locator):
    count = locator.count()
    for i in range(count):
        item = locator.nth(i)
        try:
            if item.is_visible():
                item.click()
                return True
        except Exception:
            continue
    return False


def _cursor_view_enabled():
    raw = str(os.getenv("CURSOR_VIEW", "0")).strip().lower()
    return raw in ("1", "true", "yes", "on")


def _debug_before_click(page, item, label="click"):
    if not _cursor_view_enabled():
        return
    try:
        page.bring_to_front()
    except Exception:
        pass
    try:
        box = item.bounding_box()
    except Exception:
        box = None
    if not box:
        print(f"[CURSOR] {label}")
        return
    x = float(box.get("x", 0)) + (float(box.get("width", 0)) / 2.0)
    y = float(box.get("y", 0)) + (float(box.get("height", 0)) / 2.0)
    try:
        page.mouse.move(x, y, steps=4)
    except Exception:
        pass
    try:
        page.evaluate(
            """
            ([x, y, label]) => {
              const old = document.getElementById('__cursor_debug_dot');
              if (old) old.remove();
              const dot = document.createElement('div');
              dot.id = '__cursor_debug_dot';
              dot.style.position = 'fixed';
              dot.style.left = (x - 10) + 'px';
              dot.style.top = (y - 10) + 'px';
              dot.style.width = '28px';
              dot.style.height = '28px';
              dot.style.border = '4px solid #ff2d2d';
              dot.style.borderRadius = '9999px';
              dot.style.background = 'rgba(255,45,45,0.28)';
              dot.style.boxShadow = '0 0 0 6px rgba(255,45,45,0.25)';
              dot.style.pointerEvents = 'none';
              dot.style.zIndex = '2147483647';
              document.body.appendChild(dot);

              const oldTag = document.getElementById('__cursor_debug_tag');
              if (oldTag) oldTag.remove();
              const tag = document.createElement('div');
              tag.id = '__cursor_debug_tag';
              tag.textContent = 'AUTO: ' + (label || 'click');
              tag.style.position = 'fixed';
              tag.style.left = '12px';
              tag.style.bottom = '12px';
              tag.style.padding = '6px 10px';
              tag.style.borderRadius = '8px';
              tag.style.background = '#ff2d2d';
              tag.style.color = '#fff';
              tag.style.font = 'bold 13px/1.2 Arial, sans-serif';
              tag.style.pointerEvents = 'none';
              tag.style.zIndex = '2147483647';
              document.body.appendChild(tag);
            }
            """,
            [x, y, label],
        )
        page.wait_for_timeout(70)
    except Exception:
        pass
    print(f"[CURSOR] {label} at ({int(x)}, {int(y)})")


def copy_file_to_windows_clipboard(file_path):
    resolved = str(Path(file_path).resolve())
    if os.name != "nt":
        raise RuntimeError("Clipboard file paste is currently supported only on Windows")
    attachment_kind = detect_attachment_kind(resolved)
    if attachment_kind == "image":
        ps_script = (
            "& { param([string]$path) "
            "Add-Type -AssemblyName System.Windows.Forms; "
            "Add-Type -AssemblyName System.Drawing; "
            "$img = [System.Drawing.Image]::FromFile($path); "
            "[System.Windows.Forms.Clipboard]::SetImage($img); "
            "$img.Dispose() }"
        )
    else:
        ps_script = (
            "& { param([string]$path) "
            "Add-Type -AssemblyName System.Windows.Forms; "
            "$files = New-Object System.Collections.Specialized.StringCollection; "
            "$null = $files.Add($path); "
            "[System.Windows.Forms.Clipboard]::SetFileDropList($files) }"
        )
    command = ["powershell", "-NoProfile", "-Command", ps_script, "-path", resolved]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        err = (result.stderr or result.stdout or "").strip()
        raise RuntimeError(f"Could not copy file to clipboard: {err or 'unknown error'}")


def _fill_caption_if_available(page, caption):
    if not caption:
        return
    caption_text = str(caption)
    candidates = [
        "div[role='textbox'][aria-label='Type a message'][contenteditable='true'][data-lexical-editor='true']",
        "div[role='textbox'][aria-label='Type a message'][contenteditable='true']",
        "div[aria-label='Type a message'][contenteditable='true']",
        "div[contenteditable='true'][data-tab='10']",
        "div[contenteditable='true'][data-tab='6']",
        "div[contenteditable='true'][role='textbox']",
    ]
    for selector in candidates:
        locator = page.locator(selector)
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count - 1, -1, -1):
            box = locator.nth(i)
            try:
                if box.is_visible():
                    _debug_before_click(page, box, "Type caption in preview")
                    box.click(force=True)
                    try:
                        box.fill(caption_text)
                    except Exception:
                        try:
                            page.keyboard.press("Control+A")
                            page.keyboard.press("Backspace")
                        except Exception:
                            pass
                        page.keyboard.type(caption_text, delay=0)
                    page.wait_for_timeout(60)
                    return
            except Exception:
                continue


def _composer_targets(page):
    targets = []
    selectors = [
        "footer div[contenteditable='true'][data-tab='10']",
        "footer div[contenteditable='true'][data-tab='6']",
        "footer div[contenteditable='true'][role='textbox']",
        "div[contenteditable='true'][data-tab='10']",
        "div[contenteditable='true'][data-tab='6']",
        "div[contenteditable='true'][role='textbox']",
    ]
    for selector in selectors:
        locator = page.locator(selector)
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count - 1, -1, -1):
            targets.append(locator.nth(i))

    text_nodes = [
        page.locator("footer p.copyable-text"),
        page.locator("p.copyable-text"),
    ]
    for locator in text_nodes:
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count - 1, -1, -1):
            try:
                targets.append(locator.nth(i).locator("xpath=ancestor::div[@contenteditable='true'][1]"))
            except Exception:
                continue
    return targets


def _focus_compose_box(page):
    page.bring_to_front()
    pyautogui.FAILSAFE = False
    for target in _composer_targets(page):
        try:
            if not target.is_visible():
                continue
            target.scroll_into_view_if_needed()
            target.click(force=True)
            target.evaluate("(el) => { el.focus(); }")
            page.wait_for_timeout(80)
            return True
        except Exception:
            continue
    return False


def _send_button_ready(page):
    send_buttons = page.locator(
        "button[aria-label='Send'], div[aria-label='Send'], span[data-icon='send'], span[data-icon='wds-ic-send-filled']"
    )
    try:
        for i in range(send_buttons.count()):
            if send_buttons.nth(i).is_visible():
                return True
    except Exception:
        return False
    return False


def _wait_random_send_delay(page, min_seconds=0, max_seconds=30, log_func=None, label="Send", on_delay=None):
    try:
        low = max(0, int(min_seconds))
    except Exception:
        low = 0
    try:
        high = max(low, int(max_seconds))
    except Exception:
        high = max(low, 30)

    delay_seconds = random.randint(low, high)
    planned_click_at = (datetime.now() + timedelta(seconds=delay_seconds)).strftime("%H:%M:%S")
    log_line = (
        f"{label} will click in {delay_seconds}s "
        f"(planned click time: {planned_click_at})."
    )
    print(f"[LOG] {log_line}")
    if callable(log_func):
        try:
            log_func(log_line)
        except Exception:
            pass
    if callable(on_delay):
        try:
            on_delay(delay_seconds, planned_click_at)
        except Exception:
            pass
    if delay_seconds > 0:
        page.wait_for_timeout(delay_seconds * 1000)
    return delay_seconds


def _open_attachment_menu(page):
    candidate_selectors = [
        "button[aria-label='Add file']",
        "button[title='Attach']",
        "div[title='Attach']",
        "span[data-icon='plus-rounded']",
    ]
    try:
        for selector in candidate_selectors:
            locator = page.locator(selector)
            for i in range(locator.count()):
                btn = locator.nth(i)
                if not btn.is_visible():
                    continue
                try:
                    _debug_before_click(page, btn, "Open + attachment menu")
                    btn.click(force=True)
                except Exception:
                    try:
                        parent_btn = btn.locator("xpath=ancestor::*[self::button or self::div][1]")
                        _debug_before_click(page, parent_btn, "Open + attachment menu (ancestor)")
                        parent_btn.click(force=True)
                    except Exception:
                        continue
                page.wait_for_timeout(180)
                return True
    except Exception:
        return False
    return False


def _open_document_picker(page):
    document_selectors = [
        "button[role='menuitem'][aria-label='Document']",
        "div[role='menuitem'][aria-label='Document']",
        "div[role='button'][aria-label='Document']",
        "button[aria-label='Document']",
        "div[role='button']:has-text('Document')",
        "button:has-text('Document')",
    ]
    for selector in document_selectors:
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                if not item.is_visible():
                    continue
                aria_label = (item.get_attribute("aria-label") or "").strip().lower()
                title_attr = (item.get_attribute("title") or "").strip().lower()
                # Avoid clicking preview thumbnail/remove controls that can trigger downloads.
                if "thumbnail" in aria_label or "remove attachment" in aria_label:
                    continue
                if "thumbnail" in title_attr or "remove attachment" in title_attr:
                    continue
                try:
                    _debug_before_click(page, item, "Pick Document option")
                    item.click(force=True)
                except Exception:
                    try:
                        parent_btn = item.locator("xpath=ancestor::*[self::button or self::div][1]")
                        _debug_before_click(page, parent_btn, "Pick Document option (ancestor)")
                        parent_btn.click(force=True)
                    except Exception:
                        continue
                page.wait_for_timeout(160)
                return True
        except Exception:
            continue
    return False


def _open_media_picker(page):
    media_selectors = [
        "button[role='menuitem'][type='button'][aria-label='Photos &amp; videos']",
        "button[role='menuitem'][type='button'][aria-label*='Photos'][aria-label*='videos']",
        "button[role='menuitem'][type='button'][aria-label='Photos & videos']",
        "button[role='menuitem'][aria-label='Photos & videos']",
        "div[role='menuitem'][aria-label='Photos & videos']",
        "div[role='button'][aria-label='Photos & videos']",
        "button[aria-label='Photos & videos']",
        "div[role='button']:has-text('Photos & videos')",
        "button:has-text('Photos & videos')",
        "div[role='button']:has-text('Photos and videos')",
        "button:has-text('Photos and videos')",
    ]
    for selector in media_selectors:
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                if not item.is_visible():
                    continue
                try:
                    _debug_before_click(page, item, "Pick Photos & videos option")
                    item.click(force=True)
                except Exception:
                    try:
                        parent_btn = item.locator("xpath=ancestor::*[self::button or self::div][1]")
                        _debug_before_click(page, parent_btn, "Pick Photos & videos option (ancestor)")
                        parent_btn.click(force=True)
                    except Exception:
                        continue
                page.wait_for_timeout(160)
                return True
        except Exception:
            continue
    return False


def _open_audio_picker(page):
    audio_selectors = [
        "button[role='menuitem'][aria-label='Audio']",
        "div[role='menuitem'][aria-label='Audio']",
        "div[role='button'][aria-label='Audio']",
        "button[aria-label='Audio']",
        "div[role='button']:has-text('Audio')",
        "button:has-text('Audio')",
    ]
    for selector in audio_selectors:
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                if not item.is_visible():
                    continue
                try:
                    _debug_before_click(page, item, "Pick Audio option")
                    item.click(force=True)
                except Exception:
                    try:
                        parent_btn = item.locator("xpath=ancestor::*[self::button or self::div][1]")
                        _debug_before_click(page, parent_btn, "Pick Audio option (ancestor)")
                        parent_btn.click(force=True)
                    except Exception:
                        continue
                page.wait_for_timeout(160)
                return True
        except Exception:
            continue
    return False


def _upload_via_menu_filechooser(page, selectors, attachment_path, label):
    for selector in selectors:
        locator = page.locator(selector)
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count):
            item = locator.nth(i)
            try:
                if not item.is_visible():
                    continue
                with page.expect_file_chooser(timeout=2200) as chooser_info:
                    _debug_before_click(page, item, label)
                    item.click(force=True)
                chooser = chooser_info.value
                chooser.set_files(attachment_path)
                print(f"[LOG] Uploaded via menu chooser: {label}")
                return _wait_for_attachment_ready(page, timeout_ms=2600, poll_ms=80)
            except Exception:
                continue
    return False


def _accept_matches_kind(accept_attr, attachment_kind, attachment_path):
    accept = str(accept_attr or "").lower().replace(" ", "")
    ext = Path(str(attachment_path or "")).suffix.lower()
    if not accept:
        return True

    if attachment_kind == "image":
        return ("image" in accept) or ("*" in accept) or (ext and ext in accept)
    if attachment_kind == "video":
        return ("video" in accept) or ("*" in accept) or (ext and ext in accept)
    if attachment_kind == "audio":
        # Audio uses dedicated audio option in WhatsApp attachment menu.
        return (
            ("audio" in accept)
            or ("*" in accept)
            or (ext and ext in accept)
        )

    # Document/other file mode: never allow media-only pickers like "image/*".
    if "image" in accept or "video" in accept:
        if "application" not in accept and "text" not in accept:
            if ext and ext in accept:
                return True
            # "image/*" or "video/*" should not be used for docs.
            return False
    if not accept:
        return True
    if ext and ext in accept:
        return True
    if "application" in accept or "text" in accept:
        return True
    if "." in accept:
        return True
    if accept == "*/*" or accept == "*":
        return True
    return False


def _attachment_preview_present(page):
    selectors = [
        "[aria-label='Remove attachment']",
        "div[role='tab'][aria-label*='thumbnail']",
        "button[aria-label='Add file']",
    ]
    for selector in selectors:
        try:
            locator = page.locator(selector)
            if locator.count() > 0 and locator.first.is_visible():
                return True
        except Exception:
            continue
    return False


def _wait_for_attachment_ready(page, timeout_ms=3000, poll_ms=80):
    started = time.time()
    timeout_s = max(0.3, float(timeout_ms) / 1000.0)
    poll = max(40, int(poll_ms))
    while (time.time() - started) < timeout_s:
        if _has_unsupported_attachment_notice(page):
            return False
        if _send_button_ready(page) or _attachment_preview_present(page):
            return True
        page.wait_for_timeout(poll)
    return _send_button_ready(page) or _attachment_preview_present(page)


def _click_attachment_send(page, on_delay=None):
    send_selectors = [
        "div[role='button'][aria-label='Send'][aria-disabled='false']",
        "button[aria-label='Send'][aria-disabled='false']",
        "div[aria-label='Send'][aria-disabled='false']",
        "span[data-icon='wds-ic-send-filled']",
        "span[data-icon='send']",
    ]
    delay_applied = False
    for _ in range(8):
        for selector in send_selectors:
            locator = page.locator(selector)
            try:
                count = locator.count()
            except Exception:
                count = 0
            for i in range(count):
                item = locator.nth(i)
                try:
                    if not item.is_visible():
                        continue
                    if not delay_applied:
                        _wait_random_send_delay(
                            page,
                            0,
                            30,
                            log_func=_send_log,
                            label="Attachment send button",
                            on_delay=on_delay,
                        )
                        delay_applied = True
                    _debug_before_click(page, item, "Click Send")
                    item.click(force=True)
                    # If click succeeded on a visible Send control, treat as sent.
                    # WhatsApp preview UI can linger briefly and previously caused false "failed" status.
                    page.wait_for_timeout(120)
                    return True
                except Exception:
                    continue
        page.wait_for_timeout(70)
    return False


def _find_attachment_input(page, attachment_kind, attachment_path):
    candidates = [
        "div[role='dialog'] input[type='file']",
        "input[type='file']",
    ]
    fallback = None
    for selector in candidates:
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                accept_attr = ""
                try:
                    accept_attr = item.get_attribute("accept") or ""
                except Exception:
                    accept_attr = ""
                if _accept_matches_kind(accept_attr, attachment_kind, attachment_path):
                    accept_norm = str(accept_attr or "").lower()
                    if attachment_kind == "file":
                        # For documents, prefer explicit document/file inputs and avoid media pickers.
                        if any(token in accept_norm for token in ["application", "text", ".", "*"]):
                            print(f"[LOG] Using file input accept='{accept_attr}' for document upload")
                            return item
                        if not accept_norm and fallback is None:
                            fallback = item
                            continue
                    else:
                        # For media, only accept media-style inputs. Avoid generic/document pickers.
                        if any(token in accept_norm for token in ["image", "video", "audio"]):
                            print(f"[LOG] Using file input accept='{accept_attr}' for media upload")
                            return item
                        if not accept_norm and fallback is None:
                            fallback = item
                            continue
                        # Skip wildcard/document-style input for media to force Photos & videos path.
                        continue
        except Exception:
            continue
    if attachment_kind == "file":
        return fallback
    return None


def _has_unsupported_attachment_notice(page):
    notices = [
        "file you tried adding is not supported",
        "files you tried adding are not supported",
        "not supported",
    ]
    for text in notices:
        try:
            if page.locator(f"text={text}").first.is_visible(timeout=300):
                return True
        except Exception:
            continue
    return False


def _attach_file_via_plus_button(page, attachment_path, attachment_kind):
    if not _open_attachment_menu(page):
        return False
    if attachment_kind == "file":
        print("[LOG] Choosing Document picker for file upload")
        if _upload_via_menu_filechooser(
            page,
            [
                "button[role='menuitem'][type='button'][aria-label='Document']",
                "button[role='menuitem'][aria-label='Document']",
                "div[role='menuitem'][aria-label='Document']",
                "button[aria-label='Document']",
            ],
            attachment_path,
            "Pick Document option",
        ):
            return _wait_for_attachment_ready(page, timeout_ms=2600, poll_ms=80)
        _open_document_picker(page)
    elif attachment_kind == "audio":
        print("[LOG] Choosing Audio picker for audio upload")
        if _upload_via_menu_filechooser(
            page,
            [
                "button[role='menuitem'][type='button'][aria-label='Audio']",
                "button[role='menuitem'][aria-label='Audio']",
                "div[role='menuitem'][aria-label='Audio']",
                "button[aria-label='Audio']",
            ],
            attachment_path,
            "Pick Audio option",
        ):
            return _wait_for_attachment_ready(page, timeout_ms=2600, poll_ms=80)
        if not _open_audio_picker(page):
            print("[WARN] Audio picker not found; refusing to use document flow for audio")
            return False
    else:
        print("[LOG] Choosing Photos & videos picker for image/video upload")
        if _upload_via_menu_filechooser(
            page,
            [
                "button[role='menuitem'][type='button'][aria-label='Photos &amp; videos']",
                "button[role='menuitem'][type='button'][aria-label*='Photos'][aria-label*='videos']",
                "button[role='menuitem'][type='button'][aria-label='Photos & videos']",
                "button[role='menuitem'][aria-label='Photos & videos']",
                "div[role='menuitem'][aria-label='Photos & videos']",
            ],
            attachment_path,
            "Pick Photos & videos option",
        ):
            return _wait_for_attachment_ready(page, timeout_ms=2600, poll_ms=80)
        if not _open_media_picker(page):
            print("[WARN] Photos & videos picker not found; refusing to use document flow for media")
            return False
    file_input = _find_attachment_input(page, attachment_kind, attachment_path)
    if file_input is None:
        if attachment_kind != "file":
            print("[WARN] No media file input found after opening Photos & videos")
        return False
    try:
        file_input.set_input_files(attachment_path)
    except Exception:
        return False
    return _wait_for_attachment_ready(page, timeout_ms=3400, poll_ms=80)


def _paste_attachment_into_composer(page):
    if not _focus_compose_box(page):
        return False
    for _ in range(2):
        try:
            pyautogui.hotkey("ctrl", "v")
        except Exception:
            page.keyboard.press("Control+V")
        page.wait_for_timeout(240)
        if _has_unsupported_attachment_notice(page):
            return False
        if _wait_for_attachment_ready(page, timeout_ms=2000, poll_ms=80):
            return True
    return False


def send_single_attachment(page, attachment, caption="", on_delay=None):
    attachment_path = str(attachment.get("path") or "").strip()
    if not attachment_path:
        raise RuntimeError("Attachment path is missing")
    attachment_kind = str(attachment.get("kind") or detect_attachment_kind(attachment_path))
    print(f"[LOG] Attaching '{Path(attachment_path).name}' as kind='{attachment_kind}'")

    attached = _attach_file_via_plus_button(page, attachment_path, attachment_kind)
    if not attached:
        if attachment_kind == "file":
            copy_file_to_windows_clipboard(attachment_path)
            attached = _paste_attachment_into_composer(page)
        else:
            # For media (image/video/audio), never fall back to generic paste,
            # because WhatsApp may treat it as document and skip media preview.
            raise RuntimeError(
                f"Could not attach {attachment_kind} via Photos & videos picker."
            )
    if not attached:
        if _has_unsupported_attachment_notice(page):
            raise RuntimeError(
                f"WhatsApp rejected attachment as unsupported: {Path(attachment_path).name}"
            )
        raise RuntimeError("Could not attach file into WhatsApp chat")

    # Wait for preview/send UI so we do not press send before upload is ready.
    preview_selectors = [
        "[aria-label='Remove attachment']",
        "button[aria-label='Add file']",
        "div[role='button'][aria-label='Send']",
        "button[aria-label='Send']",
        "span[data-icon='wds-ic-send-filled']",
        "span[data-icon='send']",
    ]
    preview_ready = False
    for selector in preview_selectors:
        try:
            page.wait_for_selector(selector, timeout=1000)
            preview_ready = True
            break
        except Exception:
            continue
    if not preview_ready:
        print("[WARN] File preview not detected before send. Trying to send anyway.")

    _fill_caption_if_available(page, caption)

    if not _click_attachment_send(page, on_delay=on_delay):
        raise RuntimeError("Send button not found for attachment after waiting for preview.")
    page.wait_for_timeout(120)


def _normalized_text(value):
    return " ".join(str(value or "").strip().lower().split())


def _fill_search_field(page, locator, text_value):
    text_value = str(text_value or "")
    try:
        locator.fill("")
        locator.fill(text_value)
        return True
    except Exception:
        pass
    try:
        locator.click(force=True)
        page.keyboard.press("Control+A")
        page.keyboard.press("Backspace")
        if text_value:
            page.keyboard.type(text_value, delay=8)
        return True
    except Exception:
        return False


def _send_plain_text_single_message(page, msg_box, text_value, on_delay=None):
    text_value = str(text_value or "")
    normalized = text_value.replace("\r\n", "\n").replace("\r", "\n")
    try:
        msg_box.first.click(force=True)
    except Exception:
        try:
            page.locator("div[contenteditable='true'][role='textbox']").first.click(force=True)
        except Exception:
            pass
    try:
        page.keyboard.press("Control+A")
        page.keyboard.press("Backspace")
    except Exception:
        pass

    lines = normalized.split("\n")
    for idx, line in enumerate(lines):
        if idx > 0:
            page.keyboard.press("Shift+Enter")
        if line:
            try:
                page.keyboard.insert_text(line)
            except Exception:
                page.keyboard.type(line, delay=0)
    _wait_random_send_delay(
        page,
        0,
        30,
        log_func=_send_log,
        label="Text message send",
        on_delay=on_delay,
    )
    page.keyboard.press("Enter")
    return True


def _click_groups_filter_for_send(page, timeout_ms=7000):
    candidates = [
        page.locator("#group-filter").first,
        page.get_by_role("tab", name="Groups").first,
        page.locator("button:has-text('Groups')").first,
    ]
    for item in candidates:
        try:
            if item.count() == 0:
                continue
            item.wait_for(state="visible", timeout=timeout_ms)
            item.click(timeout=timeout_ms)
            page.wait_for_timeout(120)
            return True
        except Exception:
            continue
    return False


def _open_group_chat_like_scrape_flow(page, group_name):
    target = str(group_name or "").strip()
    if not target:
        return False, "Group name is empty"

    try:
        page.wait_for_selector("#pane-side", timeout=12000)
    except Exception:
        return False, "Chat list pane not ready"

    if not _click_groups_filter_for_send(page):
        return False, "Could not open Groups filter"

    rows_selector = "div[aria-label='Chat list'][role='grid'] > div[role='row']"
    rows = page.locator(rows_selector)
    if rows.count() == 0:
        rows = page.locator("div[aria-label='Chat list'] div[role='row']")
    if rows.count() == 0:
        return False, "No group rows found in Groups list"

    target_norm = _normalized_text(target)
    matched_row = None
    max_rounds = 18
    no_move_rounds = 0
    last_tail = ""

    for _ in range(max_rounds):
        rows = page.locator(rows_selector)
        if rows.count() == 0:
            rows = page.locator("div[aria-label='Chat list'] div[role='row']")

        for i in range(rows.count()):
            row = rows.nth(i)
            try:
                if not row.is_visible():
                    continue
                title_node = row.locator("span[title]").first
                title_text = (title_node.get_attribute("title") or "").strip() if title_node.count() else ""
                row_text = (row.inner_text() or "").strip()
                title_norm = _normalized_text(title_text)
                row_norm = _normalized_text(row_text)
                if target_norm and (title_norm == target_norm or target_norm in title_norm or target_norm in row_norm):
                    matched_row = row
                    break
            except Exception:
                continue
        if matched_row is not None:
            break

        try:
            current_tail = ""
            if rows.count() > 0:
                current_tail = _normalized_text(rows.nth(rows.count() - 1).inner_text())
            if current_tail and current_tail == last_tail:
                no_move_rounds += 1
            else:
                no_move_rounds = 0
            last_tail = current_tail
        except Exception:
            pass

        if no_move_rounds >= 4:
            break

        try:
            rows.last.scroll_into_view_if_needed(timeout=1800)
            rows.last.hover()
            page.mouse.wheel(0, 1200)
            page.wait_for_timeout(220)
        except Exception:
            break

    if matched_row is None:
        return False, f"Group '{target}' not found in Groups list"

    try:
        matched_row.scroll_into_view_if_needed(timeout=3000)
    except Exception:
        pass
    try:
        matched_row.click()
    except Exception:
        try:
            matched_row.click(force=True)
        except Exception:
            return False, f"Could not click group '{target}' row"

    page.wait_for_timeout(180)

    try:
        page.locator("footer div[contenteditable='true'][role='textbox']").first.wait_for(
            state="visible", timeout=5000
        )
    except Exception:
        pass

    return True, ""


def send_messages(
    targets,
    payload,
    delay_seconds,
    allow_resend,
    state,
    send_target="contacts",
    campaign_id=None,
    campaign_name="",
    batch_min_size=20,
    batch_max_size=30,
):
    session_dir = get_playwright_user_data_dir()
    db_conn = get_db_connection()
    ensure_message_log_tables(db_conn)
    _send_log("Connected to PostgreSQL.")
    if campaign_id:
        _send_log(f"Campaign ID: {campaign_id}")
    if campaign_name:
        _send_log(f"Campaign Name: {campaign_name}")

    def set_contact_status(target, status, reason=""):
        if "contact_statuses" not in state:
            state["contact_statuses"] = {}
        state["contact_statuses"][target] = {"status": status, "reason": reason or ""}

    def record_failed(target, reason):
        current_status = (
            ((state.get("contact_statuses") or {}).get(target) or {}).get("status", "")
        )
        if current_status == "sent":
            _send_log(
                f"Post-send warning for {target}: {reason}. Keeping final status as sent."
            )
            return
        state["failed"].append({"number": target, "reason": reason})
        set_contact_status(target, "failed", reason)
        try:
            save_not_sent_sms(db_conn, campaign_id, campaign_name, target, reason)
        except Exception as db_exc:
            _send_log(f"DB warning: failed to save sms_not_sent for {target}: {db_exc}")

    try:
        with sync_playwright() as p:
            launch_kwargs = get_playwright_launch_kwargs(default_headless=False)
            _send_log(
                f"Playwright config -> headless={launch_kwargs.get('headless')} "
                f"start_maximized={'args' in launch_kwargs} "
                f"no_viewport={launch_kwargs.get('no_viewport', False)} "
                f"slow_mo={launch_kwargs.get('slow_mo', 0)}ms"
            )
            _send_log(f"Playwright user_data_dir: {session_dir}")
            browser = p.chromium.launch_persistent_context(
                user_data_dir=str(session_dir),
                **launch_kwargs,
            )

            page = browser.pages[0] if browser.pages else browser.new_page()
            _send_log("Opening WhatsApp Web...")
            page.goto("https://web.whatsapp.com/")

            _send_log("Scan QR if needed...")
            login_selector = _wait_for_any_selector(
                page,
                (
                    "[data-testid='chat-list']",
                    "div[aria-label='Chat list']",
                    "#pane-side",
                ),
                timeout_ms=5000,
            )
            login_state = "reused_logged_in" if login_selector else "logged_in_after_wait"
            if not login_selector:
                _wait_for_any_selector(
                    page,
                    (
                        "[data-testid='chat-list']",
                        "div[aria-label='Chat list']",
                        "#pane-side",
                    ),
                    timeout_ms=120000,
                )
            _send_log("WhatsApp loaded. Logged in.")
            try:
                _, session_line = save_whatsapp_session_snapshot(
                    page=page,
                    flow_name="send_messages",
                    user_data_dir=str(session_dir),
                    login_state=login_state,
                )
                _send_log(f"Session snapshot saved -> {session_line}")
            except Exception as session_exc:
                _send_log(f"Session snapshot warning: {session_exc}")
            _send_log(f"Cursor view: {'ON' if _cursor_view_enabled() else 'OFF'}")
            _send_log(f"Starting messaging loop for {len(targets)} target(s).")
            batch_sizes = _build_random_batch_sizes(
                len(targets),
                min_size=batch_min_size,
                max_size=batch_max_size,
            )
            batch_numbers = _expand_batch_numbers(batch_sizes)
            if len(batch_numbers) < len(targets):
                batch_numbers.extend([len(batch_sizes) or 1] * (len(targets) - len(batch_numbers)))
            state["total_batches"] = len(batch_sizes)
            state["current_batch_no"] = int(batch_numbers[0]) if batch_numbers else 0
            state["batch_pause_active"] = False
            state["batch_pause_remaining_sec"] = 0
            state["batch_next_start_at"] = ""
            state["live_rows"] = _init_live_rows(targets, batch_numbers=batch_numbers)
            if batch_sizes:
                _send_log(
                    f"Batch size range selected: {int(batch_min_size)}-{int(batch_max_size)}"
                )
                _send_log(
                    "Batch plan: "
                    + ", ".join(f"B{i+1}={size}" for i, size in enumerate(batch_sizes))
                )
            active_message_variants = list(payload.get("message_variants") or [])
            if active_message_variants:
                slots = ", ".join(str(int(item.get("slot") or 0)) for item in active_message_variants)
                _send_log(
                    f"Random message mode ON. Active message panels: {len(active_message_variants)} "
                    f"(Message {slots})."
                )
            total_attachments = len(payload.get("attachments") or [])
            if total_attachments > 0:
                _send_log(
                    f"Random attachment mode ON. Uploaded attachment(s): {total_attachments}. "
                    f"One attachment will be selected per target."
                )

            for target_index, target in enumerate(targets):
                target = str(target or "").strip()
                if not target:
                    continue
                batch_no = int(batch_numbers[target_index]) if target_index < len(batch_numbers) else 1
                prev_batch_no = int(batch_numbers[target_index - 1]) if target_index > 0 and (target_index - 1) < len(batch_numbers) else batch_no
                if target_index > 0 and batch_no != prev_batch_no:
                    pause_seconds = random.randint(180, 360)
                    state["batch_pause_active"] = True
                    state["batch_pause_remaining_sec"] = pause_seconds
                    state["current_batch_no"] = batch_no
                    _send_log(
                        f"Batch {prev_batch_no} complete. Next batch {batch_no} starts after "
                        f"{pause_seconds // 60}m {pause_seconds % 60:02d}s."
                    )
                    for remaining in range(pause_seconds, 0, -1):
                        state["batch_pause_remaining_sec"] = remaining
                        state["batch_next_start_at"] = (
                            datetime.now() + timedelta(seconds=remaining)
                        ).strftime("%H:%M:%S")
                        time.sleep(1)
                    state["batch_pause_active"] = False
                    state["batch_pause_remaining_sec"] = 0
                    state["batch_next_start_at"] = ""
                    _send_log(f"Batch {batch_no} started.")
                else:
                    state["current_batch_no"] = batch_no
                _send_log(f"Processing: {target}")

                selected_slot, selected_message, active_count = _pick_random_message_for_target(payload)
                if active_count > 0:
                    preview = " ".join(selected_message.split())
                    if len(preview) > 70:
                        preview = preview[:67] + "..."
                    _send_log(
                        f"Random message selected for {target}: Message {selected_slot} "
                        f"from {active_count} non-empty panel(s)"
                        + (f" | text='{preview}'" if preview else " | text=(empty)")
                    )
                selected_attachment, selected_attachment_index, selected_attachment_total = _pick_random_attachment_for_target(payload)
                selected_attachments = [selected_attachment] if selected_attachment else []
                message_label = f"txt{int(selected_slot)}" if active_count > 0 else "txt1"
                media_label = "-"
                if selected_attachment_total > 0:
                    attachment_name = str((selected_attachment or {}).get("name") or "").strip() or "unnamed"
                    media_label = attachment_name
                    _send_log(
                        f"Random attachment selected for {target}: "
                        f"{selected_attachment_index}/{selected_attachment_total} ({attachment_name})"
                    )
                _update_live_row(
                    state,
                    target_index,
                    batch=batch_no,
                    message=message_label,
                    media=media_label,
                    delay="Waiting...",
                    status="processing",
                )
                def _on_delay_selected(delay_seconds, planned_click_at):
                    _update_live_row(
                        state,
                        target_index,
                        delay=f"{int(delay_seconds)}s -> {planned_click_at}",
                    )
                target_payload = build_payload(
                    selected_message,
                    selected_attachments,
                    message_variants=payload.get("message_variants") or [],
                )

                if not allow_resend:
                    try:
                        if is_already_sent(db_conn, target, target_payload):
                            _send_log(f"Already exists in table. Skipping: {target}")
                            set_contact_status(target, "skipped", "Already sent")
                            _update_live_row(state, target_index, status="skipped")
                            continue
                    except Exception as db_exc:
                        _send_log(f"DB warning: duplicate check failed for {target}: {db_exc}")

                    if send_target != "groups":
                        if not is_valid_phone(target):
                            _send_log(f"Invalid format: {target}")
                            record_failed(target, "Invalid format")
                            _update_live_row(state, target_index, status="failed")
                            continue

                sent_committed = False
                try:
                    page.keyboard.press("Escape")
                    page.wait_for_timeout(60)

                    if send_target == "groups":
                        opened, open_err = _open_group_chat_like_scrape_flow(page, target)
                        if not opened:
                            _send_log(f"Group open failed: {target} ({open_err})")
                            record_failed(target, open_err or "Could not open group chat")
                            _update_live_row(state, target_index, status="failed")
                            continue
                        _send_log(f"Opened group chat via Groups tab: {target}")
                    else:
                        page.locator("button[aria-label*='New']").first.click()
                        search_box = page.locator("input[aria-label='Search name or number']")
                        search_box.wait_for(state="visible", timeout=5000)

                        search_box.fill("")
                        search_box.fill(target)
                        # Give WhatsApp search a little time to populate results before validation.
                        page.wait_for_timeout(220)

                        results = page.locator("div[role='button'][tabindex='0']")
                        invalid_text = page.locator("span:has-text('not on WhatsApp')")

                        try:
                            results.first.wait_for(state="visible", timeout=4200)
                            found = True
                        except TimeoutError:
                            found = False

                        if not found:
                            if invalid_text.count() > 0:
                                _send_log(f"Not on WhatsApp: {target}")
                                record_failed(target, "Not on WhatsApp")
                            else:
                                _send_log(f"No result found: {target}")
                                record_failed(target, "No result found")
                            _update_live_row(state, target_index, status="failed")
                            page.keyboard.press("Escape")
                            continue

                        _send_log(f"Exists on WhatsApp: {target}")

                        clicked = False
                        for i in range(results.count()):
                            item = results.nth(i)
                            try:
                                if not item.is_visible():
                                    continue
                                item.click()
                                clicked = True
                                break
                            except Exception:
                                continue

                        if not clicked:
                            _send_log(f"No clickable result: {target}")
                            record_failed(target, "No clickable result")
                            _update_live_row(state, target_index, status="failed")
                            page.keyboard.press("Escape")
                            continue

                    msg_box = page.locator("div[contenteditable='true'][role='textbox']")
                    msg_box.wait_for(state="visible", timeout=5000)

                    if target_payload.get("attachments"):
                        first_caption = target_payload.get("message") or ""
                        for index, attachment in enumerate(target_payload.get("attachments") or []):
                            send_single_attachment(
                                page,
                                attachment,
                                first_caption if index == 0 else "",
                                on_delay=_on_delay_selected,
                            )
                    else:
                        _send_plain_text_single_message(
                            page,
                            msg_box,
                            target_payload.get("message") or "",
                            on_delay=_on_delay_selected,
                        )

                    _send_log(f"Message sent to: {target}")
                    state["sent"] += 1
                    set_contact_status(target, "sent", "")
                    _update_live_row(state, target_index, status="sent")
                    sent_committed = True
                    try:
                        save_sent_sms(
                            db_conn,
                            target,
                            target_payload,
                            campaign_id=campaign_id,
                            campaign_name=campaign_name,
                        )
                    except Exception as db_exc:
                        _send_log(f"DB warning: failed to save sms_sent for {target}: {db_exc}")

                    page.keyboard.press("Escape")
                    page.wait_for_timeout(60)

                except Exception as e:
                    if sent_committed:
                        _send_log(
                            f"Post-send warning for {target}: {e}. Message already sent; status kept as sent."
                        )
                        try:
                            page.keyboard.press("Escape")
                            page.wait_for_timeout(120)
                        except Exception:
                            pass
                        continue
                    _send_log(f"Error for {target}: {e}")
                    record_failed(target, f"Error: {e}")
                    _update_live_row(state, target_index, status="failed")
                    try:
                        page.keyboard.press("Escape")
                        page.wait_for_timeout(120)
                    except Exception:
                        pass
                    continue

            _send_log("Done. You can close the browser or keep it open.")
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


app = Flask(__name__)
app.secret_key = "whatsapp-bot-secret"

send_state = {
    "running": False,
    "last_error": None,
    "last_result": None,
    "last_failed": [],
    "failed": [],
    "sent": 0,
    "total": 0,
    "last_message": "",
    "last_messages": [""] * 10,
    "last_delay": "0",
    "last_country_filter": "",
    "last_filename": "",
    "last_attachment_names": [],
    "last_resend": False,
    "last_source_type": "file",
    "last_send_target": "contacts",
    "last_selected_tables": [],
    "last_group_filter_name": "",
    "last_group_filter_names": [],
    "last_selected_contacts": [],
    "contacts": [],
    "last_contacts": [],
    "contact_statuses": {},
    "last_contact_statuses": {},
    "campaign_id": "",
    "last_campaign_id": "",
    "campaign_name": "",
    "last_campaign_name": "",
    "live_rows": [],
    "last_live_rows": [],
    "batch_pause_active": False,
    "batch_pause_remaining_sec": 0,
    "batch_next_start_at": "",
    "current_batch_no": 0,
    "total_batches": 0,
    "last_batch_min_size": 20,
    "last_batch_max_size": 30,
}

app_log_state = {
    "logs": [],
}
app_log_lock = threading.Lock()

scrape_state = {
    "running": False,
    "last_error": None,
    "last_result": None,
    "logs": [],
    "groups_total": 0,
    "groups_done": 0,
    "groups_skipped": 0,
    "groups": [],
    "started_at": None,
    "finished_at": None,
    "scrap_id": "",
}


def _append_app_log(message, source="APP"):
    ts = datetime.now().strftime("%H:%M:%S")
    line = f"[{ts}] [{str(source or 'APP').upper()}] {message}"
    with app_log_lock:
        app_log_state["logs"].append(line)
        if len(app_log_state["logs"]) > 1000:
            app_log_state["logs"] = app_log_state["logs"][-1000:]
    print(f"[{str(source or 'APP').upper()}] {message}")
    return line


def _send_log(message):
    return _append_app_log(message, "SEND")


def _scrape_log(message):
    scrape_state["logs"].append(_append_app_log(message, "SCRAPE"))
    # Keep memory bounded
    if len(scrape_state["logs"]) > 400:
        scrape_state["logs"] = scrape_state["logs"][-400:]


def _reset_scrape_state():
    scrape_state["last_error"] = None
    scrape_state["last_result"] = None
    scrape_state["logs"] = []
    scrape_state["groups_total"] = 0
    scrape_state["groups_done"] = 0
    scrape_state["groups_skipped"] = 0
    scrape_state["groups"] = []
    scrape_state["started_at"] = None
    scrape_state["finished_at"] = None
    scrape_state["scrap_id"] = ""


def ensure_scrape_runs_table(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.scrape_group_runs (
                    id BIGSERIAL PRIMARY KEY,
                    scrap_id TEXT NOT NULL,
                    group_index INTEGER NOT NULL,
                    groups_total INTEGER NOT NULL DEFAULT 0,
                    group_name TEXT NOT NULL,
                    scrapped_contacts_count INTEGER NOT NULL DEFAULT 0,
                    total_contacts_in_loop INTEGER NOT NULL DEFAULT 0,
                    imported_groups_in_loop INTEGER NOT NULL DEFAULT 0,
                    loop_started_at TIMESTAMPTZ NOT NULL,
                    event_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS idx_scrape_group_runs_scrap_id
                ON {schema_name}.scrape_group_runs (scrap_id, event_time)
                """
            ).format(schema_name=sql.Identifier(schema))
        )


def save_scrape_group_row(
    db_conn,
    scrap_id,
    group_index,
    groups_total,
    group_name,
    scrapped_contacts_count,
    total_contacts_in_loop,
    imported_groups_in_loop,
    loop_started_at,
):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                INSERT INTO {schema_name}.scrape_group_runs (
                    scrap_id,
                    group_index,
                    groups_total,
                    group_name,
                    scrapped_contacts_count,  
                    total_contacts_in_loop,
                    imported_groups_in_loop,
                    loop_started_at,
                    event_time 
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, NOW())
                """
            ).format(schema_name=sql.Identifier(schema)),
            (
                scrap_id,
                int(group_index),
                int(groups_total),
                group_name,
                int(scrapped_contacts_count),
                int(total_contacts_in_loop),
                int(imported_groups_in_loop),
                loop_started_at,
            ),
        )


def update_scrape_groups_total_for_run(db_conn, scrap_id, groups_total):
    if not scrap_id:
        return
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                UPDATE {schema_name}.scrape_group_runs
                SET groups_total = %s
                WHERE scrap_id = %s
                """
            ).format(schema_name=sql.Identifier(schema)),
            (int(groups_total), scrap_id),
        )


def load_scrape_rows(scrap_id):
    if not scrap_id:
        return []
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT
                        group_index,
                        groups_total,
                        group_name,
                        scrapped_contacts_count,
                        total_contacts_in_loop,
                        imported_groups_in_loop,
                        loop_started_at,
                        event_time
                    FROM {schema_name}.scrape_group_runs
                    WHERE scrap_id = %s
                    ORDER BY group_index ASC, event_time ASC
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (scrap_id,),
            )
            rows = cur.fetchall()
            return [
                {
                    "group_index": r[0],
                    "groups_total": r[1],
                    "group_name": r[2],
                    "scrapped_contacts_count": r[3],
                    "total_contacts_in_loop": r[4],
                    "imported_groups_in_loop": r[5],
                    "loop_started_at": r[6].isoformat() if r[6] else "",
                    "event_time": r[7].isoformat() if r[7] else "",
                }
                for r in rows
            ]
    except Exception as exc:
        print(f"[SCRAPE DB WARN] Could not load scrape rows: {exc}")
        return []
    finally:
        db_conn.close()


def list_scrap_ids(limit=100):
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT scrap_id, MAX(event_time) AS last_time
                    FROM {schema_name}.scrape_group_runs
                    GROUP BY scrap_id
                    ORDER BY last_time DESC
                    LIMIT %s
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (int(limit),),
            )
            return [row[0] for row in cur.fetchall() if row and row[0]]
    except Exception as exc:
        print(f"[SCRAPE DB WARN] Could not list scrap IDs: {exc}")
        return []
    finally:
        db_conn.close()


def generate_unique_numeric_scrap_id(db_conn, digits=6, max_attempts=300):
    schema = get_db_schema()
    d = max(6, int(digits))
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                SELECT COALESCE(
                    MAX(
                        CASE
                            WHEN scrap_id ~ '^[0-9]+$' THEN scrap_id::BIGINT
                            ELSE NULL
                        END
                    ),
                    0
                )
                FROM {schema_name}.scrape_group_runs
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        row = cur.fetchone()
        last_value = int(row[0]) if row and row[0] is not None else 0
        next_value = last_value + 1
        return str(next_value).zfill(d)


def search_groups_across_scrap_runs(query_text, limit=300):
    db_conn = get_db_connection()
    schema = get_db_schema()
    q = (query_text or "").strip().lower()
    try:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT
                        scrap_id,
                        group_name,
                        MAX(event_time) AS last_event_time
                    FROM {schema_name}.scrape_group_runs
                    WHERE group_name IS NOT NULL
                      AND group_name <> ''
                      AND scrap_id NOT LIKE '%-GROUP%'
                      AND scrap_id NOT LIKE '%-RECOVERY-%'
                      AND LOWER(group_name) LIKE %s
                    GROUP BY scrap_id, group_name
                    ORDER BY last_event_time DESC, scrap_id ASC, group_name ASC
                    LIMIT %s
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (f"%{q}%", int(limit)),
            )
            rows = cur.fetchall()
            results = []
            for row in rows:
                if not row:
                    continue
                scrap_id = row[0] or ""
                group_name = row[1] or ""
                last_event_time = row[2].isoformat() if row[2] else ""
                if not scrap_id or not group_name:
                    continue
                results.append(
                    {
                        "scrap_id": scrap_id,
                        "group_name": group_name,
                        "last_event_time": last_event_time,
                    }
                )
            return results
    finally:
        db_conn.close()


def list_unique_group_names(limit=2000):
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT MIN(group_name) AS group_name
                    FROM {schema_name}.scrape_group_runs
                    WHERE group_name IS NOT NULL
                      AND TRIM(group_name) <> ''
                    GROUP BY LOWER(TRIM(group_name))
                    ORDER BY group_name ASC
                    LIMIT %s
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (int(limit),),
            )
            return [row[0] for row in cur.fetchall() if row and row[0]]
    finally:
        db_conn.close()


def build_scrape_summary_from_rows(scrap_id, rows):
    if not rows:
        return {
            "scrap_id": scrap_id,
            "groups_total": 0,
            "groups_done": 0,
            "groups_skipped": 0,
            "total_contacts": 0,
            "last_result": None,
            "started_at": None,
            "finished_at": None,
        }
    groups_total = max(int(r.get("groups_total", 0)) for r in rows)
    groups_done = max(int(r.get("imported_groups_in_loop", 0)) for r in rows)
    groups_skipped = max(int(r.get("groups_skipped_in_loop", 0)) for r in rows)
    total_contacts = max(int(r.get("total_contacts_in_loop", 0)) for r in rows)
    started_at = rows[0].get("loop_started_at")
    finished_at = rows[-1].get("event_time")
    return {
        "scrap_id": scrap_id,
        "groups_total": groups_total,
        "groups_done": groups_done,
        "groups_skipped": groups_skipped,
        "total_contacts": total_contacts,
        "last_result": f"Scrapping complete. Groups imported: {groups_done}. Total contacts: {total_contacts}.",
        "started_at": started_at,
        "finished_at": finished_at,
    }


def run_scrape_job():
    scrape_state["running"] = True
    _reset_scrape_state()
    scrape_state["started_at"] = datetime.now().isoformat()
    _scrape_log("Starting WhatsApp contacts scrapping...")

    browser = None
    db_conn = None
    existing_group_names = set()
    logged_group_checks = set()
    try:
        db_conn = get_db_connection()
        ensure_scrape_runs_table(db_conn)
        existing_group_names = load_existing_group_names_from_all_contacts(db_conn)
        _scrape_log(
            f"Existing groups in DB: {len(existing_group_names)}. These groups will be skipped."
        )
        scrape_state["scrap_id"] = generate_unique_numeric_scrap_id(db_conn, digits=6)
        _scrape_log(f"Scrap ID: {scrape_state['scrap_id']}")

        user_data_dir = get_playwright_user_data_dir()

        with sync_playwright() as p:
            launch_kwargs = get_playwright_launch_kwargs(default_headless=False)
            _scrape_log(
                "Playwright config -> "
                f"headless={launch_kwargs.get('headless')} "
                f"start_maximized={'args' in launch_kwargs} "
                f"no_viewport={launch_kwargs.get('no_viewport', False)} "
                f"slow_mo={launch_kwargs.get('slow_mo', 0)}ms"
            )
            _scrape_log(f"Playwright user_data_dir: {user_data_dir}")
            browser = p.chromium.launch_persistent_context(
                user_data_dir=str(user_data_dir),
                **launch_kwargs,
            )
            page = browser.pages[0] if browser.pages else browser.new_page()
            _scrape_log("Opening WhatsApp Web...")
            page.goto("https://web.whatsapp.com")
            login_selector = _wait_for_any_selector(
                page,
                (
                    "[data-testid='chat-list']",
                    "div[aria-label='Chat list']",
                    "#pane-side",
                ),
                timeout_ms=5000,
            )
            login_state = "reused_logged_in" if login_selector else "logged_in_after_wait"
            if not login_selector:
                _wait_for_any_selector(
                    page,
                    (
                        "[data-testid='chat-list']",
                        "div[aria-label='Chat list']",
                        "#pane-side",
                    ),
                    timeout_ms=120000,
                )
            _scrape_log("WhatsApp loaded. Logged in.")
            try:
                _, session_line = save_whatsapp_session_snapshot(
                    page=page,
                    flow_name="run_scrape_job",
                    user_data_dir=str(user_data_dir),
                    login_state=login_state,
                )
                _scrape_log(f"Session snapshot saved -> {session_line}")
            except Exception as session_exc:
                _scrape_log(f"Session snapshot warning: {session_exc}")

            if not open_group_and_members(page):
                raise RuntimeError("Could not open Groups filter.")

            rows_selector = "div[aria-label='Chat list'][role='grid'] > div[role='row']"
            scrape_state["groups_total"] = page.locator(rows_selector).count()
            _scrape_log(f"Total groups found: {scrape_state['groups_total']}")

            def _on_group_complete(item):
                group_name = item.get("group", "group")
                contacts = item.get("contacts", []) or []
                metadata = dict(item.get("metadata") or {})
                metadata["group_name"] = group_name
                metadata["source_scrap_id"] = scrape_state["scrap_id"]
                metadata["last_scraped_at"] = scrape_state["started_at"]

                excel_files = []
                table_name = save_group_contacts_to_postgres(group_name, contacts, metadata=metadata)
                group_key = _normalize_group_key(group_name)
                if group_key:
                    existing_group_names.add(group_key)

                scrape_state["groups_done"] += 1
                loop_total_contacts = sum(
                    g.get("contacts_count", 0) for g in scrape_state["groups"]
                ) + len(contacts)
                scrape_state["groups"].append(
                    {
                        "name": group_name,
                        "contacts_count": len(contacts),
                        "excel_files": excel_files,
                        "db_table": table_name,
                        "group_index": scrape_state["groups_done"],
                        "groups_total": scrape_state["groups_total"],
                        "loop_total_contacts": loop_total_contacts,
                        "imported_groups_in_loop": scrape_state["groups_done"],
                        "metadata": metadata,
                    }
                )
                _scrape_log(
                    f"{group_name}: {len(contacts)} contacts imported "
                    f"(DB: {table_name}, Excel: {len(excel_files)} file)."
                )
                if db_conn:
                    try:
                        save_scrape_group_row(
                            db_conn=db_conn,
                            scrap_id=scrape_state["scrap_id"],
                            group_index=scrape_state["groups_done"],
                            groups_total=scrape_state["groups_total"],
                            group_name=group_name,
                            scrapped_contacts_count=len(contacts),
                            total_contacts_in_loop=loop_total_contacts,
                            imported_groups_in_loop=scrape_state["groups_done"],
                            loop_started_at=scrape_state["started_at"],
                        )
                    except Exception as exc:
                        _scrape_log(f"DB log warning for group {group_name}: {exc}")

            def _should_skip_group(group_name):
                group_key = _normalize_group_key(group_name)
                if not group_key:
                    return False
                if group_key in existing_group_names:
                    if group_key not in logged_group_checks:
                        scrape_state["groups_skipped"] += 1
                        _scrape_log(
                            f"Skipped:[ {group_name} ] already exists in list."
                        )
                        logged_group_checks.add(group_key)
                    return True
                return False

            def _on_group_skipped(event):
                reason = str((event or {}).get("reason") or "").strip().lower()
                group_name = str((event or {}).get("group") or "").strip()
                if reason == "community":
                    scrape_state["groups_total"] = max(
                        0, int(scrape_state.get("groups_total", 0)) - 1
                    )
                    if group_name:
                        _scrape_log(
                            f"Community skipped (not counted in groups): {group_name}"
                        )

            extract_all_groups(
                page,
                on_group_complete=_on_group_complete,
                should_skip_group=_should_skip_group,
                on_group_skipped=_on_group_skipped,
            )

            if db_conn and scrape_state.get("scrap_id"):
                try:
                    update_scrape_groups_total_for_run(
                        db_conn=db_conn,
                        scrap_id=scrape_state["scrap_id"],
                        groups_total=scrape_state["groups_total"],
                    )
                except Exception as exc:
                    _scrape_log(f"DB totals update warning: {exc}")

            total_contacts = sum(g["contacts_count"] for g in scrape_state["groups"])
            scrape_state["last_result"] = (
                f"Scrapping complete. Groups imported: {scrape_state['groups_done']}. "
                f"Total contacts: {total_contacts}."
            )
            _scrape_log(scrape_state["last_result"])

    except Exception as exc:
        scrape_state["last_error"] = str(exc)
        _scrape_log(f"Error: {exc}")
        traceback.print_exc()
    finally:
        try:
            if browser:
                browser.close()
        except Exception:
            pass
        try:
            if db_conn:
                db_conn.close()
        except Exception:
            pass
        scrape_state["running"] = False
        scrape_state["finished_at"] = datetime.now().isoformat()


def start_scrape_job():
    if scrape_state["running"]:
        _append_app_log("Scrapping job request ignored: another scrapping job is already running.", "SCRAPE")
        return False

    _append_app_log("Scrapping job started.", "SCRAPE")
    threading.Thread(target=run_scrape_job, daemon=True).start()
    return True


def _is_scrape_running():
    return scrape_state["running"]


recovery_manager = RecoveryManager(
    get_db_connection=get_db_connection,
    get_db_schema=get_db_schema,
    sql_module=sql,
    ensure_scrape_runs_table=ensure_scrape_runs_table,
    save_scrape_group_row=save_scrape_group_row,
    is_scrape_running=_is_scrape_running,
)


def start_send_job(
    targets,
    payload,
    delay_seconds,
    allow_resend,
    send_target="contacts",
    campaign_id=None,
    campaign_name="",
    batch_min_size=20,
    batch_max_size=30,
):
    if send_state["running"]:
        _send_log("Send job request ignored: another send job is already running.")
        return False

    def _job():
        active_campaign_id = str(campaign_id or generate_campaign_id()).zfill(6)
        active_campaign_name = str(campaign_name or "").strip()
        send_state["running"] = True
        send_state["last_error"] = None
        send_state["last_result"] = None
        send_state["campaign_id"] = active_campaign_id
        send_state["last_campaign_id"] = active_campaign_id
        send_state["campaign_name"] = active_campaign_name
        send_state["last_campaign_name"] = active_campaign_name
        send_state["last_live_rows"] = []
        send_state["batch_pause_active"] = False
        send_state["batch_pause_remaining_sec"] = 0
        send_state["batch_next_start_at"] = ""
        send_state["current_batch_no"] = 0
        send_state["total_batches"] = 0
        send_state["failed"] = []
        send_state["sent"] = 0
        send_state["total"] = len(targets)
        send_state["contacts"] = list(targets)
        send_state["live_rows"] = _init_live_rows(targets)
        send_state["contact_statuses"] = {
            phone: {"status": "pending", "reason": ""}
            for phone in targets
        }
        _send_log(
            f"Send job started. Targets={len(targets)} | "
            f"target_type={send_target} | campaign={active_campaign_id}_{active_campaign_name or '-'}"
        )
        try:
            send_messages(
                targets,
                payload,
                delay_seconds,
                allow_resend,
                send_state,
                send_target=send_target,
                campaign_id=active_campaign_id,
                campaign_name=active_campaign_name,
                batch_min_size=batch_min_size,
                batch_max_size=batch_max_size,
            )
            send_state["last_failed"] = list(send_state["failed"])
            send_state["last_contacts"] = list(send_state["contacts"])
            send_state["last_contact_statuses"] = {
                k: {"status": v.get("status", "pending"), "reason": v.get("reason", "")}
                for k, v in send_state["contact_statuses"].items()
            }
            send_state["last_live_rows"] = list(send_state.get("live_rows") or [])
            send_state["last_result"] = f"Sent to {send_state['sent']} of {send_state['total']} contacts."
            _send_log(send_state["last_result"])
        except Exception as exc:
            send_state["last_error"] = str(exc)
            _send_log(f"Send job error: {exc}")
        finally:
            if send_state["contacts"] and not send_state["last_contacts"]:
                send_state["last_contacts"] = list(send_state["contacts"])
            if send_state["contact_statuses"] and not send_state["last_contact_statuses"]:
                send_state["last_contact_statuses"] = {
                    k: {"status": v.get("status", "pending"), "reason": v.get("reason", "")}
                    for k, v in send_state["contact_statuses"].items()
                }
            if send_state["live_rows"] and not send_state["last_live_rows"]:
                send_state["last_live_rows"] = list(send_state.get("live_rows") or [])
            send_state["campaign_id"] = ""
            send_state["campaign_name"] = ""
            send_state["batch_pause_active"] = False
            send_state["batch_pause_remaining_sec"] = 0
            send_state["batch_next_start_at"] = ""
            send_state["running"] = False
            _send_log("Send job finished.")

    threading.Thread(target=_job, daemon=True).start()
    return True


@app.route("/", methods=["GET"])
def index():
    keep = request.args.get("keep") == "1"
    if not send_state["running"] and not keep:
        send_state["last_message"] = ""
        send_state["last_messages"] = [""] * 10
        send_state["last_delay"] = "0"
        send_state["last_country_filter"] = ""
        send_state["last_filename"] = ""
        send_state["last_attachment_names"] = []
        send_state["last_resend"] = False
        send_state["last_source_type"] = "file"
        send_state["last_send_target"] = "contacts"
        send_state["last_selected_tables"] = []
        send_state["last_group_filter_name"] = ""
        send_state["last_group_filter_names"] = []
        send_state["last_selected_contacts"] = []
        send_state["contacts"] = []
        send_state["last_contacts"] = []
        send_state["contact_statuses"] = {}
        send_state["last_contact_statuses"] = {}
        send_state["last_error"] = None
        send_state["last_result"] = None
        send_state["last_failed"] = []
        send_state["failed"] = []
        send_state["sent"] = 0
        send_state["total"] = 0
        send_state["campaign_id"] = ""
        send_state["last_campaign_id"] = ""
        send_state["campaign_name"] = ""
        send_state["last_campaign_name"] = ""
        send_state["live_rows"] = []
        send_state["last_live_rows"] = []
        send_state["batch_pause_active"] = False
        send_state["batch_pause_remaining_sec"] = 0
        send_state["batch_next_start_at"] = ""
        send_state["current_batch_no"] = 0
        send_state["total_batches"] = 0
        send_state["last_batch_min_size"] = 20
        send_state["last_batch_max_size"] = 30
        with app_log_lock:
            app_log_state["logs"] = []
    if not scrape_state["running"] and not keep:
        _reset_scrape_state()
    if not keep:
        recovery_manager.reset_if_idle()
    return render_template("index.html", state=send_state, keep=keep)


@app.route("/reset", methods=["GET"])
def reset():
    if send_state["running"]:
        flash("Cannot reset while sending is in progress.")
        return redirect(url_for("index"))
    send_state["last_message"] = ""
    send_state["last_messages"] = [""] * 10
    send_state["last_delay"] = "0"
    send_state["last_country_filter"] = ""
    send_state["last_filename"] = ""
    send_state["last_attachment_names"] = []
    send_state["last_resend"] = False
    send_state["last_source_type"] = "file"
    send_state["last_send_target"] = "contacts"
    send_state["last_selected_tables"] = []
    send_state["last_group_filter_name"] = ""
    send_state["last_group_filter_names"] = []
    send_state["last_selected_contacts"] = []
    send_state["contacts"] = []
    send_state["last_contacts"] = []
    send_state["contact_statuses"] = {}
    send_state["last_contact_statuses"] = {}
    send_state["last_error"] = None
    send_state["last_result"] = None
    send_state["last_failed"] = []
    send_state["failed"] = []
    send_state["sent"] = 0
    send_state["total"] = 0
    send_state["campaign_id"] = ""
    send_state["last_campaign_id"] = ""
    send_state["campaign_name"] = ""
    send_state["last_campaign_name"] = ""
    send_state["live_rows"] = []
    send_state["last_live_rows"] = []
    send_state["batch_pause_active"] = False
    send_state["batch_pause_remaining_sec"] = 0
    send_state["batch_next_start_at"] = ""
    send_state["current_batch_no"] = 0
    send_state["total_batches"] = 0
    send_state["last_batch_min_size"] = 20
    send_state["last_batch_max_size"] = 30
    with app_log_lock:
        app_log_state["logs"] = []
    if not scrape_state["running"]:
        _reset_scrape_state()
    recovery_manager.reset_if_idle()
    flash("")
    return redirect(url_for("index"))


@app.route("/send", methods=["POST"])
def send():
    is_ajax = (request.headers.get("X-Requested-With") or "").lower() == "xmlhttprequest"

    def send_error(message, status_code=400):
        if is_ajax:
            return jsonify({"ok": False, "error": message}), status_code
        flash(message)
        return redirect(url_for("index"))

    def send_ok(message, campaign_id="", campaign_name=""):
        if is_ajax:
            cid = str(campaign_id or "").zfill(6) if str(campaign_id or "").strip() else ""
            return jsonify({
                "ok": True,
                "message": message,
                "running": send_state["running"],
                "campaign_id": cid,
                "campaign_name": str(campaign_name or "").strip(),
            })
        flash(message)
        return redirect(url_for("index", keep=1))

    if send_state["running"]:
        return send_error("A send job is already running. Please wait.", 409)

    file = request.files.get("contacts")
    source_type = (request.form.get("source_type") or "file").strip().lower()
    send_target = (request.form.get("send_target") or "contacts").strip().lower()
    if send_target not in ("contacts", "groups"):
        send_target = "contacts"

    selected_tables_list = request.form.getlist("selected_tables")
    selected_tables_raw = request.form.get("selected_tables") or "[]"
    selected_tables = parse_selected_tables(
        selected_tables_list if selected_tables_list else selected_tables_raw
    )
    selected_contacts_raw = request.form.get("selected_contacts") or "[]"
    selected_contacts = parse_selected_tables(selected_contacts_raw)
    reuse_campaign_id = _normalize_campaign_id(request.form.get("reuse_campaign_id") or "")
    campaign_name = (request.form.get("campaign_name") or "").strip()
    group_filter_name = (request.form.get("group_filter_name") or "").strip()
    group_filter_names = parse_selected_tables(request.form.get("group_filter_names") or "[]")
    attachment_files = request.files.getlist("attachments")
    message_slots = []
    message_variants = []
    for i in range(1, 11):
        msg = str(request.form.get(f"message_{i}") or "").strip()
        message_slots.append(msg)
        if msg:
            message_variants.append({"slot": i, "text": msg})
    legacy_message = str(request.form.get("message") or "").strip()
    if not message_variants and legacy_message:
        message_slots[0] = legacy_message
        message_variants = [{"slot": 1, "text": legacy_message}]
    message = message_variants[0]["text"] if message_variants else legacy_message
    delay = request.form.get("delay") or "0"
    batch_min_raw = request.form.get("batch_min_size") or "20"
    batch_max_raw = request.form.get("batch_max_size") or "30"
    country_filter = (request.form.get("country") or "").strip()
    resend_sms = True

    unsupported_uploads = []
    for item in attachment_files or []:
        name = secure_filename((item.filename or "").strip())
        if not name:
            continue
        if not is_supported_attachment_filename(name):
            unsupported_uploads.append(name)
    if unsupported_uploads:
        return send_error(
            "Unsupported attachment type: "
            + ", ".join(unsupported_uploads)
            + ". Allowed: image/video/audio/pdf/doc/docx/xls/xlsx/csv/txt/zip/ppt/pptx"
        )

    saved_attachments = save_uploaded_attachments(attachment_files)
    if not message and not saved_attachments:
        return send_error("Please enter a message or attach at least one file.")
    payload = build_payload(message, saved_attachments, message_variants=message_variants)

    try:
        delay_seconds = max(0, int(delay))
    except Exception:
        delay_seconds = 0
    try:
        batch_min_size = max(1, int(batch_min_raw))
    except Exception:
        batch_min_size = 20
    try:
        batch_max_size = max(1, int(batch_max_raw))
    except Exception:
        batch_max_size = 30
    if batch_max_size < batch_min_size:
        batch_min_size, batch_max_size = batch_max_size, batch_min_size
    batch_min_size = min(500, batch_min_size)
    batch_max_size = min(500, batch_max_size)

    send_state["last_message"] = message
    send_state["last_messages"] = list(message_slots)
    send_state["last_delay"] = str(delay_seconds)
    send_state["last_country_filter"] = country_filter
    send_state["last_attachment_names"] = payload["attachment_names"]
    send_state["last_resend"] = resend_sms
    send_state["last_source_type"] = source_type
    send_state["last_send_target"] = send_target
    send_state["last_selected_tables"] = selected_tables
    send_state["last_campaign_name"] = campaign_name
    send_state["last_group_filter_name"] = group_filter_name
    send_state["last_group_filter_names"] = group_filter_names
    send_state["last_selected_contacts"] = list(selected_contacts)
    send_state["last_filename"] = file.filename if file else ""
    send_state["last_batch_min_size"] = batch_min_size
    send_state["last_batch_max_size"] = batch_max_size

    try:
        used_selected_contacts = False
        skipped_count = 0
        applied_prefixes = []
        targets = []
        if not campaign_name:
            return send_error("Campaign name is required.")

        if send_target == "groups":
            selected_groups = [g for g in (selected_contacts or []) if str(g).strip()]
            selected_groups = sorted(set(str(g).strip() for g in selected_groups if str(g).strip()))
            if not selected_groups:
                return send_error("Please add group(s) into Selected Contacts for group send mode.")
            targets = selected_groups
            send_state["last_selected_contacts"] = list(targets)
            used_selected_contacts = True
        else:
            # Safety rule: contacts send must always use Selected Contacts only.
            # If selection is empty, do not fall back to all contacts.
            selected_contacts = [c for c in (selected_contacts or []) if str(c).strip()]
            if not selected_contacts:
                return send_error("No contacts in Selected Contacts list. Please add contacts before sending.")
            if reuse_campaign_id:
                numbers = list(selected_contacts)
            elif source_type == "database":
                if group_filter_names:
                    multi_group_payload = load_contacts_from_multiple_groups(group_filter_names)
                    numbers = multi_group_payload.get("contacts", [])
                    if not group_filter_name and group_filter_names:
                        group_filter_name = ", ".join(group_filter_names)
                        send_state["last_group_filter_name"] = group_filter_name
                    if not selected_tables:
                        selected_tables = ["all_contacts"]
                        send_state["last_selected_tables"] = selected_tables
                elif group_filter_name:
                    numbers = load_group_contacts_from_all_contacts(group_filter_name)
                    if not selected_tables:
                        selected_tables = ["all_contacts"]
                        send_state["last_selected_tables"] = selected_tables
                elif not selected_tables:
                    return send_error("Please select at least one database table.")
                else:
                    numbers = load_numbers_from_tables(selected_tables)
            else:
                if not file or not file.filename:
                    return send_error("Please upload a contacts file.")
                numbers = load_numbers(file)

            numbers, skipped_count, applied_prefixes = filter_numbers_by_country_codes(numbers, country_filter)

            selected_set = set(selected_contacts)
            numbers = [n for n in numbers if n in selected_set]
            if not numbers:
                return send_error("No selected contacts matched the loaded source/filter.")
            used_selected_contacts = True

            if not numbers:
                if country_filter:
                    return send_error("No numbers matched the provided country code filter.")
                if source_type == "database":
                    if group_filter_name or group_filter_names:
                        return send_error("No phone numbers found in the selected group.")
                    return send_error("No phone numbers found in selected database tables.")
                return send_error("No phone numbers found in the file.")
            targets = numbers
    except Exception as exc:
        return send_error(str(exc))

    try:
        campaign_id = generate_campaign_id()
    except Exception as exc:
        return send_error(f"Could not generate campaign ID: {exc}")
    started = start_send_job(
        targets,
        payload,
        delay_seconds,
        resend_sms,
        send_target=send_target,
        campaign_id=campaign_id,
        campaign_name=campaign_name,
        batch_min_size=batch_min_size,
        batch_max_size=batch_max_size,
    )
    if not started:
        return send_error("A send job is already running.", 409)

    if reuse_campaign_id:
        source_text = f"Re-use Campaign {reuse_campaign_id}"
    else:
        source_text = f"DB tables ({len(selected_tables)})" if source_type == "database" else "file"
    source_text += f" - Target: {'Groups' if send_target == 'groups' else 'Contacts'}"
    source_text += f" - Batch range: {batch_min_size}-{batch_max_size}"
    if used_selected_contacts:
        source_text += " - Selected Contacts"
    if group_filter_names:
        source_text += f" - Groups: {', '.join(group_filter_names)}"
    elif group_filter_name:
        source_text += f" - Group: {group_filter_name}"

    if applied_prefixes:
        return send_ok(
            f"Started sending to {len(targets)} filtered contacts (skipped {skipped_count}). "
            f"Filter: {', '.join(applied_prefixes)}. Source: {source_text}. "
            f"Campaign: {str(campaign_id).zfill(6)}_{campaign_name}.",
            campaign_id=campaign_id,
            campaign_name=campaign_name,
        )
    return send_ok(
        f"Started sending to {len(targets)} contacts. Source: {source_text}. "
        f"Campaign: {str(campaign_id).zfill(6)}_{campaign_name}.",
        campaign_id=campaign_id,
        campaign_name=campaign_name,
    )


@app.route("/status", methods=["GET"])
def status():
    contacts = send_state["contacts"] if send_state["running"] else send_state["last_contacts"]
    statuses = (
        send_state["contact_statuses"] if send_state["running"] else send_state["last_contact_statuses"]
    )
    live_rows = send_state["live_rows"] if send_state["running"] else send_state["last_live_rows"]
    contact_items = build_contact_items(contacts, statuses)
    with app_log_lock:
        logs = list(app_log_state["logs"])
    return jsonify({
        "running": send_state["running"],
        "last_error": send_state["last_error"],
        "last_result": send_state["last_result"],
        "campaign_id": (
            str(send_state["campaign_id"]).zfill(6)
            if send_state["running"] and str(send_state["campaign_id"]).strip()
            else (
                str(send_state["last_campaign_id"]).zfill(6)
                if str(send_state["last_campaign_id"]).strip()
                else ""
            )
        ),
        "campaign_name": (
            str(send_state["campaign_name"]).strip()
            if send_state["running"] and str(send_state["campaign_name"]).strip()
            else str(send_state["last_campaign_name"]).strip()
        ),
        "contacts": contacts,
        "contact_items": contact_items,
        "sent": send_state["sent"],
        "total": send_state["total"],
        "live_rows": live_rows,
        "batch_pause_active": bool(send_state.get("batch_pause_active")),
        "batch_pause_remaining_sec": int(send_state.get("batch_pause_remaining_sec") or 0),
        "batch_next_start_at": str(send_state.get("batch_next_start_at") or ""),
        "current_batch_no": int(send_state.get("current_batch_no") or 0),
        "total_batches": int(send_state.get("total_batches") or 0),
        "batch_min_size": int(send_state.get("last_batch_min_size") or 20),
        "batch_max_size": int(send_state.get("last_batch_max_size") or 30),
        "logs": logs[-300:],
    })


@app.route("/logs/status", methods=["GET"])
def logs_status():
    with app_log_lock:
        logs = list(app_log_state["logs"])
    return jsonify({
        "logs": logs[-500:],
        "send_running": bool(send_state.get("running")),
        "scrape_running": bool(scrape_state.get("running")),
    })


@app.route("/assets/pic/<path:filename>", methods=["GET"])
def template_pic(filename):
    pic_dir = Path(__file__).parent / "templates" / "pic"
    return send_from_directory(str(pic_dir), filename)


@app.route("/assets/send_uploads/<path:filename>", methods=["GET"])
def send_upload_asset(filename):
    uploads_dir = Path(__file__).parent / "send_uploads"
    return send_from_directory(str(uploads_dir), filename)


@app.route("/api/recent-send-media", methods=["GET"])
def api_recent_send_media():
    try:
        uploads_dir = Path(__file__).parent / "send_uploads"
        uploads_dir.mkdir(parents=True, exist_ok=True)
        limit_raw = request.args.get("limit") or "200"
        try:
            limit = max(1, min(1000, int(limit_raw)))
        except Exception:
            limit = 200

        files = []
        for p in uploads_dir.iterdir():
            if not p.is_file():
                continue
            ext = p.suffix.lower()
            if ext and ext not in RECENT_MEDIA_EXTENSIONS:
                continue
            try:
                stat = p.stat()
            except Exception:
                continue
            kind = "other"
            if ext in IMAGE_EXTENSIONS:
                kind = "image"
            elif ext in VIDEO_EXTENSIONS:
                kind = "video"
            elif ext in AUDIO_EXTENSIONS:
                kind = "audio"
            elif ext in DOCUMENT_EXTENSIONS:
                kind = "document"
            files.append(
                {
                    "name": p.name,
                    "size": int(stat.st_size or 0),
                    "modified_at": datetime.fromtimestamp(stat.st_mtime).isoformat(),
                    "kind": kind,
                    "url": f"/assets/send_uploads/{quote(p.name)}",
                }
            )

        files.sort(key=lambda x: x.get("modified_at") or "", reverse=True)
        return jsonify({"ok": True, "files": files[:limit]})
    except Exception as exc:
        print(f"[MEDIA ERROR] api_recent_send_media failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "files": []}), 500


@app.route("/scrape/start", methods=["POST"])
def scrape_start():
    started = start_scrape_job()
    if not started:
        return jsonify({"ok": False, "error": "Scrapping is already running."}), 409
    return jsonify({"ok": True})


@app.route("/scrape/status", methods=["GET"])
def scrape_status():
    selected_scrap_id = (request.args.get("scrap_id") or "").strip()
    available_scrap_ids = list_scrap_ids()

    # Historical mode: return data for explicitly requested scrape ID.
    if selected_scrap_id and selected_scrap_id != scrape_state["scrap_id"]:
        rows = load_scrape_rows(selected_scrap_id)
        summary = build_scrape_summary_from_rows(selected_scrap_id, rows)
        return jsonify(
            {
                "running": False,
                "last_error": None,
                "last_result": summary["last_result"],
                "scrap_id": selected_scrap_id,
                "groups_total": summary["groups_total"],
                "groups_done": summary["groups_done"],
                "groups_skipped": summary.get("groups_skipped", 0),
                "total_contacts": summary["total_contacts"],
                "groups": [],
                "db_rows": rows,
                "logs": [],
                "started_at": summary["started_at"],
                "finished_at": summary["finished_at"],
                "available_scrap_ids": available_scrap_ids,
            }
        )

    total_contacts = sum(g.get("contacts_count", 0) for g in scrape_state["groups"])
    db_rows = load_scrape_rows(scrape_state["scrap_id"])
    return jsonify(
        {
            "running": scrape_state["running"],
            "last_error": scrape_state["last_error"],
            "last_result": scrape_state["last_result"],
            "scrap_id": scrape_state["scrap_id"],
            "groups_total": scrape_state["groups_total"],
            "groups_done": scrape_state["groups_done"],
            "groups_skipped": scrape_state["groups_skipped"],
            "total_contacts": total_contacts,
            "groups": scrape_state["groups"],
            "db_rows": db_rows,
            "logs": scrape_state["logs"],
            "started_at": scrape_state["started_at"],
            "finished_at": scrape_state["finished_at"],
            "available_scrap_ids": available_scrap_ids,
        }
    )


@app.route("/scrape/recover/start", methods=["POST"])
def scrape_recover_start():
    payload = request.get_json(silent=True) or {}
    mode = (payload.get("mode") or "by_scrap_id").strip().lower()
    base_scrap_id = (payload.get("scrap_id") or "").strip()
    group_name = (payload.get("group_name") or "").strip()
    ok, err = recovery_manager.start(base_scrap_id, group_name, mode=mode)
    if not ok:
        return jsonify({"ok": False, "error": err}), 409
    return jsonify({"ok": True})


@app.route("/scrape/recover/status", methods=["GET"])
def scrape_recover_status():
    return jsonify(recovery_manager.status())


@app.route("/api/group-tables", methods=["GET"])
def api_group_tables():
    try:
        tables = list_group_tables()
        return jsonify({"ok": True, "tables": tables})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc), "tables": []}), 500


@app.route("/api/table-info", methods=["POST"])
def api_table_info():
    try:
        payload = request.get_json(silent=True) or {}
        selected_tables = parse_selected_tables(payload.get("tables") or [])
        info = get_tables_info(selected_tables)
        return jsonify({"ok": True, "info": info})
    except Exception as exc:
        print(f"[DB ERROR] api_table_info failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "info": []}), 500


@app.route("/api/preview-table-contacts", methods=["POST"])
def api_preview_table_contacts():
    try:
        payload = request.get_json(silent=True) or {}
        selected_tables = parse_selected_tables(payload.get("tables") or [])
        contacts = load_numbers_from_tables(selected_tables)
        return jsonify(
            {
                "ok": True,
                "contacts": contacts,
                "contact_items": build_contact_items(contacts, {}),
            }
        )
    except Exception as exc:
        print(f"[DB ERROR] api_preview_table_contacts failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "contacts": []}), 500


@app.route("/api/preview-file-contacts", methods=["POST"])
def api_preview_file_contacts():
    try:
        file = request.files.get("contacts")
        if not file or not file.filename:
            return jsonify({"ok": False, "error": "No file uploaded.", "contacts": []}), 400
        contacts = load_numbers(file)
        return jsonify(
            {
                "ok": True,
                "contacts": contacts,
                "contact_items": build_contact_items(contacts, {}),
            }
        )
    except Exception as exc:
        print(f"[FILE ERROR] api_preview_file_contacts failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "contacts": []}), 500


@app.route("/api/groups-panel", methods=["GET"])
def api_groups_panel():
    try:
        q = (request.args.get("q") or "").strip()
        limit = _parse_optional_limit(request.args.get("limit"), default=None)
        groups = list_unique_groups_from_all_contacts(limit=limit, query_text=q)
        return jsonify({"ok": True, "groups": groups})
    except Exception as exc:
        print(f"[GROUPS ERROR] api_groups_panel failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "groups": []}), 500


@app.route("/api/campaigns-panel", methods=["GET"])
def api_campaigns_panel():
    try:
        q = (request.args.get("q") or "").strip()
        limit = _parse_optional_limit(request.args.get("limit"), default=None)
        campaigns = list_campaign_summaries(limit=limit, query_text=q)
        return jsonify({"ok": True, "campaigns": campaigns})
    except Exception as exc:
        print(f"[CAMPAIGNS ERROR] api_campaigns_panel failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "campaigns": []}), 500


@app.route("/api/campaign-details", methods=["GET"])
def api_campaign_details():
    try:
        campaign_id = (request.args.get("campaign_id") or "").strip()
        limit = _parse_optional_limit(request.args.get("limit"), default=None)
        detail = get_campaign_detail(campaign_id, list_limit=limit)
        if not detail:
            return jsonify({"ok": False, "error": "Campaign not found.", "campaign": None}), 404
        return jsonify({"ok": True, "campaign": detail})
    except Exception as exc:
        print(f"[CAMPAIGNS ERROR] api_campaign_details failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "campaign": None}), 500


@app.route("/api/campaign-reuse", methods=["GET"])
def api_campaign_reuse():
    try:
        campaign_id = (request.args.get("campaign_id") or "").strip()
        limit = _parse_optional_limit(request.args.get("limit"), default=None)
        detail = get_campaign_detail(campaign_id, list_limit=limit)
        if not detail:
            return jsonify({"ok": False, "error": "Campaign not found."}), 404
        targets = detail.get("targets") or []
        send_target_hint = detail.get("send_target_hint") or "contacts"
        return jsonify(
            {
                "ok": True,
                "campaign_id": detail.get("campaign_id") or "",
                "send_target": send_target_hint if send_target_hint in ("contacts", "groups") else "contacts",
                "contacts": targets,
                "contact_items": build_contact_items(targets, {}),
                "targets_total": len(targets),
            }
        )
    except Exception as exc:
        print(f"[CAMPAIGNS ERROR] api_campaign_reuse failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc)}), 500


@app.route("/api/group-details", methods=["GET"])
def api_group_details(): 
    try:
        group_name = (request.args.get("group_name") or "").strip()
        limit = _parse_optional_limit(request.args.get("limit"), default=None)
        detail = get_group_detail_from_all_contacts(group_name, contact_limit=limit)
        if not detail:
            return jsonify({"ok": False, "error": "Group not found.", "group": None}), 404
        return jsonify({"ok": True, "group": detail})
    except Exception as exc:
        print(f"[GROUPS ERROR] api_group_details failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "group": None}), 500


@app.route("/api/group-details-bulk", methods=["POST"])
def api_group_details_bulk():
    try:
        payload = request.get_json(silent=True) or {}
        group_names = parse_selected_tables(payload.get("group_names") or [])
        limit = _parse_optional_limit(payload.get("limit"), default=None)
        result = load_contacts_from_multiple_groups(group_names, limit=limit)
        return jsonify({"ok": True, **result})
    except Exception as exc:
        print(f"[GROUPS ERROR] api_group_details_bulk failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "groups": [], "contacts": [], "contact_items": []}), 500


@app.route("/api/search-scrape-groups", methods=["GET"])
def api_search_scrape_groups():
    try:
        q = (request.args.get("q") or "").strip()
        limit_raw = request.args.get("limit") or "300"
        try:
            limit = max(1, min(1000, int(limit_raw)))
        except Exception:
            limit = 300
        results = search_groups_across_scrap_runs(q, limit=limit)
        return jsonify({"ok": True, "results": results})
    except Exception as exc:
        print(f"[SCRAPE DB ERROR] api_search_scrape_groups failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "results": []}), 500


@app.route("/api/unique-scrape-groups", methods=["GET"])
def api_unique_scrape_groups():
    try:
        limit_raw = request.args.get("limit") or "2000"
        try:
            limit = max(1, min(5000, int(limit_raw)))
        except Exception:
            limit = 2000
        groups = list_unique_group_names(limit=limit)
        return jsonify({"ok": True, "groups": groups})
    except Exception as exc:
        print(f"[SCRAPE DB ERROR] api_unique_scrape_groups failed: {exc}")
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc), "groups": []}), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
