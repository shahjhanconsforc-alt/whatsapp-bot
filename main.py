
from pathlib import Path
import csv
import os
import json
import re
import shutil
import socket
import subprocess
import threading
import time
import traceback
from datetime import datetime, timedelta
import random
from urllib.parse import quote
import phonenumbers
import pyautogui
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_from_directory
from werkzeug.utils import secure_filename
from playwright.sync_api import sync_playwright, TimeoutError
import openpyxl
from core.main_scrapper import (
    open_group_and_members,
    extract_all_groups,
)
from core.recovery_worker import RecoveryManager
from panels import (
    messaging_panel,
    new_scrapping_panel,
    previous_scrapping_panel,
    groups_joining_panel,
    groups_panel,
    campaigns_panel,
    campaign_monitoring_panel,
    logs_panel,
)

try:
    import psycopg2
    from psycopg2 import sql
except Exception:
    psycopg2 = None
    sql = None
 

def load_env_file():
    env_path = Path(__file__).parent / ".env"
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


load_env_file()


def _env_bool(name, default=False):
    raw = os.getenv(name)
    if raw is None:
        return bool(default)
    return str(raw).strip().lower() in ("1", "true", "yes", "on")


def get_playwright_launch_kwargs(default_headless=False):
    headless = _env_bool("PLAYWRIGHT_HEADLESS", default_headless)
    start_maximized = _env_bool("PLAYWRIGHT_START_MAXIMIZED", True)
    no_viewport = _env_bool("PLAYWRIGHT_NO_VIEWPORT", True)
    slow_mo_raw = str(os.getenv("PLAYWRIGHT_SLOW_MO_MS", "0") or "0").strip()
    try:
        slow_mo = max(0, int(slow_mo_raw))
    except Exception:
        slow_mo = 0

    kwargs = {"headless": headless}
    if start_maximized:
        kwargs["args"] = ["--start-maximized"]
    if no_viewport:
        kwargs["no_viewport"] = True
    if slow_mo > 0:
        kwargs["slow_mo"] = slow_mo
    return kwargs


def _normalize_multi_account_base_profile_name(root, selected_name):
    name = str(selected_name or "").strip()
    if not name:
        return name
    try:
        profile_root = Path(root).resolve()
    except Exception:
        profile_root = Path(root)
    matches = list(re.finditer(r"_profile_\d+", name))
    if len(matches) <= 1:
        return name

    first_candidate = name[:matches[0].end()]
    if first_candidate and (profile_root / first_candidate).is_dir():
        return first_candidate

    candidate = name
    while True:
        stripped = re.sub(r"_profile_\d+$", "", candidate)
        if not stripped or stripped == candidate:
            break
        if (profile_root / stripped).is_dir():
            candidate = stripped
            if len(list(re.finditer(r"_profile_\d+", candidate))) <= 1:
                return candidate
            continue
        break
    return candidate


def get_playwright_user_data_dir():
    base = Path(__file__).parent
    multi_account = _env_bool("PLAYWRIGHT_MULTI_ACCOUNT", False)
    if not multi_account:
        raw = str(os.getenv("PLAYWRIGHT_USER_DATA_DIR", "./playwright_profile") or "").strip()
        path = Path(raw)
        if not path.is_absolute():
            path = (base / path).resolve()
        path.mkdir(parents=True, exist_ok=True)
        return path

    root_raw = str(os.getenv("PLAYWRIGHT_PROFILES_ROOT", "./playwright_profiles") or "").strip()
    root = Path(root_raw)
    if not root.is_absolute():
        root = (base / root).resolve()
    root.mkdir(parents=True, exist_ok=True)

    default_profile = str(os.getenv("PLAYWRIGHT_PROFILE_NAME", "default") or "default").strip() or "default"
    auto_select = _env_bool("PLAYWRIGHT_PROFILE_AUTO_SELECT", True)
    current_marker = root / ".current_profile"
    selected_name = default_profile

    if auto_select:
        marker_name = ""
        if current_marker.exists():
            try:
                marker_name = current_marker.read_text(encoding="utf-8").strip()
            except Exception:
                marker_name = ""
        if marker_name and (root / marker_name).is_dir():
            selected_name = marker_name
        else:
            profile_dirs = [p for p in root.iterdir() if p.is_dir() and not p.name.startswith(".")]
            if profile_dirs:
                profile_dirs.sort(key=lambda p: p.stat().st_mtime, reverse=True)
                selected_name = profile_dirs[0].name

    selected_name = _normalize_multi_account_base_profile_name(root, selected_name)

    safe_name = "".join(ch for ch in selected_name if ch.isalnum() or ch in ("-", "_")).strip() or "default"
    path = root / safe_name
    path.mkdir(parents=True, exist_ok=True)
    try:
        current_marker.write_text(safe_name, encoding="utf-8")
    except Exception:
        pass
    return path


_PLAYWRIGHT_PROFILE_OWNERS = {}
_PLAYWRIGHT_PROFILE_OWNERS_LOCK = threading.Lock()
_PLAYWRIGHT_SINGLETON_FILES = ("SingletonLock", "SingletonCookie", "SingletonSocket")
_PROFILE_SESSION_JOBS = {}
_PROFILE_SESSION_JOBS_LOCK = threading.Lock()
_OPEN_SESSION_PREVIEW_WINDOWS = {}
_OPEN_SESSION_PREVIEW_WINDOWS_LOCK = threading.Lock()


def _normalize_playwright_profile_key(profile_dir):
    try:
        key = str(Path(profile_dir).resolve())
    except Exception:
        key = str(profile_dir)
    return key.lower() if os.name == "nt" else key


def claim_playwright_profile(profile_dir, owner_label):
    profile_key = _normalize_playwright_profile_key(profile_dir)
    owner_text = str(owner_label or "another task").strip() or "another task"
    with _PLAYWRIGHT_PROFILE_OWNERS_LOCK:
        active_owner = _PLAYWRIGHT_PROFILE_OWNERS.get(profile_key)
        if active_owner:
            raise RuntimeError(
                f"Playwright profile is already in use by {active_owner}. "
                "Please wait for that browser task to finish."
            )
        _PLAYWRIGHT_PROFILE_OWNERS[profile_key] = owner_text


def release_playwright_profile(profile_dir):
    if not str(profile_dir or "").strip():
        return
    profile_key = _normalize_playwright_profile_key(profile_dir)
    with _PLAYWRIGHT_PROFILE_OWNERS_LOCK:
        _PLAYWRIGHT_PROFILE_OWNERS.pop(profile_key, None)


def _is_local_pid_running(pid_value):
    try:
        pid_num = int(pid_value)
    except Exception:
        return False
    if pid_num <= 0:
        return False
    if os.name != "nt":
        return Path("/proc").joinpath(str(pid_num)).exists()
    try:
        os.kill(pid_num, 0)
    except Exception:
        return False
    return True


def cleanup_stale_playwright_profile_locks(profile_dir):
    profile_path = Path(profile_dir)
    result = {
        "removed": [],
        "present": [],
        "lock_exists": False,
        "lock_parseable": False,
        "lock_host": "",
        "lock_pid": "",
        "lock_active": False,
    }
    current_hosts = {
        str(socket.gethostname() or "").strip().lower(),
        str(socket.gethostname() or "").strip().split(".")[0].lower(),
    }

    lock_link = profile_path / "SingletonLock"
    if os.path.lexists(lock_link):
        result["lock_exists"] = True
        try:
            lock_target = os.readlink(lock_link)
        except OSError:
            lock_target = ""
        match = re.match(r"^(?P<host>.+)-(?P<pid>\d+)$", Path(lock_target).name) if lock_target else None
        if match:
            result["lock_parseable"] = True
            result["lock_host"] = str(match.group("host") or "").strip()
            result["lock_pid"] = str(match.group("pid") or "").strip()
            if result["lock_host"].lower() in current_hosts and _is_local_pid_running(result["lock_pid"]):
                result["lock_active"] = True

    for name in _PLAYWRIGHT_SINGLETON_FILES:
        candidate = profile_path / name
        if not os.path.lexists(candidate):
            continue
        result["present"].append(name)
        if result["lock_active"]:
            continue
        if result["lock_exists"] and not result["lock_parseable"]:
            continue
        if not candidate.is_symlink():
            continue

        extra_targets = []
        try:
            link_target = os.readlink(candidate)
        except OSError:
            link_target = ""
        if link_target:
            resolved_target = Path(link_target)
            if not resolved_target.is_absolute():
                resolved_target = (candidate.parent / resolved_target).resolve()
            extra_targets.append(resolved_target)

        try:
            candidate.unlink()
            result["removed"].append(name)
        except FileNotFoundError:
            pass
        except Exception:
            continue

        for extra_target in extra_targets:
            try:
                if extra_target.parent == profile_path and extra_target.exists():
                    extra_target.unlink()
            except Exception:
                continue

    return result


def _is_profile_in_use_launch_error(exc):
    message = str(exc or "").lower()
    return (
        "profile appears to be in use" in message
        or "chromium has locked the profile" in message
        or "process_singleton_posix" in message
    )


def launch_persistent_context_safe(browser_type, profile_dir, launch_kwargs, owner_label="", log_fn=None):
    owner_text = str(owner_label or "browser task").strip() or "browser task"
    log = log_fn if callable(log_fn) else None
    claim_playwright_profile(profile_dir, owner_text)
    try:
        stale = cleanup_stale_playwright_profile_locks(profile_dir)
        if stale["removed"] and log:
            log(
                f"Removed stale Chromium profile locks from {profile_dir}: "
                + ", ".join(stale["removed"])
            )
        try:
            return browser_type.launch_persistent_context(
                user_data_dir=str(profile_dir),
                **launch_kwargs,
            )
        except Exception as exc:
            if not _is_profile_in_use_launch_error(exc):
                raise

            stale = cleanup_stale_playwright_profile_locks(profile_dir)
            if stale["removed"]:
                if log:
                    log(
                        f"Retrying browser launch after clearing stale Chromium locks: "
                        + ", ".join(stale["removed"])
                    )
                return browser_type.launch_persistent_context(
                    user_data_dir=str(profile_dir),
                    **launch_kwargs,
                )

            details = ""
            if stale["lock_active"] and stale["lock_pid"]:
                details = (
                    f" Active Chromium PID {stale['lock_pid']} is still holding the profile"
                    f" on host {stale['lock_host'] or 'current host'}."
                )
            raise RuntimeError(
                f"Playwright profile '{Path(profile_dir).name}' is locked by another Chromium process."
                f"{details}"
            ) from exc
    except Exception:
        release_playwright_profile(profile_dir)
        raise


def _profile_number_file_path(profile_dir):
    return Path(profile_dir) / "profile_number.txt"


def _profile_session_ready_file_path(profile_dir):
    return Path(profile_dir) / ".session_ready.json"


def _store_profile_number(profile_dir, profile_number):
    value = str(profile_number or "").strip()
    path = _profile_number_file_path(profile_dir)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value, encoding="utf-8")
    return str(path)


def _write_profile_number(profile_dir, profile_number):
    value = str(profile_number or "").strip()
    if not value:
        raise ValueError("Profile number is required.")
    return _store_profile_number(profile_dir, value)


def _read_profile_number(profile_dir):
    path = _profile_number_file_path(profile_dir)
    if not path.exists():
        return ""
    try:
        return str(path.read_text(encoding="utf-8") or "").strip()
    except Exception:
        return ""


def _ensure_and_read_profile_number(profile_dir, seed_value=""):
    seed = str(seed_value or "").strip()
    path = _profile_number_file_path(profile_dir)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        _store_profile_number(profile_dir, seed)
        return seed
    try:
        return str(path.read_text(encoding="utf-8") or "").strip()
    except Exception:
        return ""


def _read_profile_session_ready(profile_dir):
    path = _profile_session_ready_file_path(profile_dir)
    if not path.exists():
        return {}
    try:
        raw = json.loads(path.read_text(encoding="utf-8") or "{}")
    except Exception:
        return {}
    if not isinstance(raw, dict):
        return {}
    return raw


def _mark_profile_session_ready(profile_dir, ready=True, **extra):
    path = _profile_session_ready_file_path(profile_dir)
    payload = {
        "ready": bool(ready),
        "updated_at": time.time(),
    }
    for key, value in dict(extra or {}).items():
        payload[str(key)] = value
    if not payload["ready"]:
        try:
            if path.exists():
                path.unlink()
        except Exception:
            pass
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=True), encoding="utf-8")


def _is_profile_session_ready(profile_dir):
    data = _read_profile_session_ready(profile_dir)
    return bool(data.get("ready"))


def _get_send_profile_dir(base_session_dir, profile_idx, create=False):
    base_path = Path(base_session_dir).resolve()
    try:
        idx = int(profile_idx)
    except Exception:
        idx = 1
    # Slot folders must always be numbered off the ROOT profile name. The
    # caller usually passes the currently-selected profile, which may itself
    # already be a slot folder (e.g. "default_profile_6"); composing on top of
    # that produced nested names like "default_profile_6_profile_2" that the
    # slot discovery/removal code cannot see, leaving orphaned folders behind.
    root_name = re.sub(r"(?:_profile_\d+)+$", "", base_path.name) or base_path.name
    if idx <= 1:
        profile_dir = base_path.parent / root_name
    else:
        profile_dir = base_path.parent / f"{root_name}_profile_{idx}"
    if create:
        profile_dir.mkdir(parents=True, exist_ok=True)
    return profile_dir


def _profile_numbers_sequence_file_path():
    data_dir = Path(__file__).parent / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    return data_dir / "profile_numbers_sequence.txt"


def _load_profile_numbers_sequence():
    path = _profile_numbers_sequence_file_path()
    data = {}
    if not path.exists():
        return data
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except Exception:
        return data
    for idx, raw in enumerate(lines, start=1):
        value = str(raw or "").strip()
        if value:
            data[idx] = value
    return data


def _save_profile_numbers_sequence(mapping, min_size=0):
    cleaned = {}
    for k, v in dict(mapping or {}).items():
        try:
            idx = int(k)
        except Exception:
            continue
        if idx < 1:
            continue
        val = str(v or "").strip()
        if val:
            cleaned[idx] = val
    max_idx = max([min_size] + list(cleaned.keys())) if cleaned or min_size > 0 else 0
    lines = []
    for idx in range(1, max_idx + 1):
        lines.append(str(cleaned.get(idx, "")).strip())
    path = _profile_numbers_sequence_file_path()
    path.write_text("\n".join(lines), encoding="utf-8")


def _set_profile_number_in_sequence(sequence_no, profile_number, min_size=0):
    numbers = _load_profile_numbers_sequence()
    try:
        idx = int(sequence_no)
    except Exception:
        return
    if idx < 1:
        return
    value = str(profile_number or "").strip()
    if value:
        numbers[idx] = value
    else:
        numbers.pop(idx, None)
    _save_profile_numbers_sequence(numbers, min_size=max(min_size, idx))


def _profile_dir_is_claimed(profile_dir):
    profile_key = _normalize_playwright_profile_key(profile_dir)
    with _PLAYWRIGHT_PROFILE_OWNERS_LOCK:
        return bool(_PLAYWRIGHT_PROFILE_OWNERS.get(profile_key))


def _shift_profile_sequence_after_removal(removed_idx):
    removed_idx = _parse_profile_index(removed_idx)
    if removed_idx < 1:
        return {}
    numbers = _load_profile_numbers_sequence()
    shifted = {}
    for key, value in dict(numbers or {}).items():
        current_idx = _parse_profile_index(key)
        if current_idx < 1 or current_idx == removed_idx:
            continue
        target_idx = current_idx if current_idx < removed_idx else (current_idx - 1)
        shifted[target_idx] = str(value or "").strip()
    max_idx = max(shifted.keys()) if shifted else 0
    _save_profile_numbers_sequence(shifted, min_size=max_idx)
    return shifted


def _shift_profile_session_jobs_after_removal(removed_idx):
    removed_idx = _parse_profile_index(removed_idx)
    if removed_idx < 1:
        return
    with _PROFILE_SESSION_JOBS_LOCK:
        existing = dict(_PROFILE_SESSION_JOBS or {})
        _PROFILE_SESSION_JOBS.clear()
        for key, value in existing.items():
            current_idx = _parse_profile_index(key)
            if current_idx < 1 or current_idx == removed_idx:
                continue
            target_idx = current_idx if current_idx < removed_idx else (current_idx - 1)
            record = dict(value or {})
            record["profile_idx"] = target_idx
            _PROFILE_SESSION_JOBS[target_idx] = record


def _shift_send_state_profiles_after_removal(send_state, removed_idx):
    removed_idx = _parse_profile_index(removed_idx)
    if removed_idx < 1 or not isinstance(send_state, dict):
        return
    for key in ("profile_states", "last_profile_states"):
        rows = []
        for raw in list(send_state.get(key) or []):
            row = dict(raw or {})
            current_idx = _parse_profile_index(row.get("profile_idx") or row.get("name"))
            if current_idx < 1 or current_idx == removed_idx:
                continue
            if current_idx > removed_idx:
                current_idx -= 1
                row["profile_idx"] = current_idx
                row["name"] = f"Profile {current_idx}"
            rows.append(row)
        send_state[key] = rows


def _force_rmtree(target_dir):
    """Delete a profile directory robustly on Windows.

    Chromium profile folders frequently contain read-only files and files that
    are momentarily locked, which makes a plain shutil.rmtree raise
    PermissionError. We clear the read-only attribute on error and retry a few
    times so a transient lock does not abort the whole removal.
    """
    path = Path(target_dir)
    if not path.exists():
        return

    def _on_error(func, problem_path, exc_info):
        try:
            os.chmod(problem_path, 0o700)
        except Exception:
            pass
        try:
            func(problem_path)
        except Exception:
            pass

    last_exc = None
    for attempt in range(5):
        try:
            shutil.rmtree(path, onerror=_on_error)
        except Exception as exc:
            last_exc = exc
        if not path.exists():
            return
        time.sleep(0.3)
    if path.exists():
        raise RuntimeError(
            f"Could not fully delete profile folder '{path.name}'. "
            "Close any open WhatsApp/Chromium window using it and try again."
            + (f" ({last_exc})" if last_exc else "")
        )


def _cleanup_profile_compact_staging(base_dir):
    """Remove any leftover staging folders from an interrupted compaction."""
    try:
        base_path = Path(base_dir)
        parent = base_path.parent
        prefix = f"{base_path.name}__compact_"
        if not parent.exists():
            return
        for child in parent.iterdir():
            try:
                if child.is_dir() and str(child.name or "").startswith(prefix):
                    _force_rmtree(child)
            except Exception:
                continue
    except Exception:
        pass


def _remove_profile_session_slot(profile_idx, send_state=None, limit=100):
    removed_idx = _parse_profile_index(profile_idx)
    if removed_idx < 1:
        raise ValueError("Invalid profile session.")
    if isinstance(send_state, dict) and (
        bool(send_state.get("running"))
        or bool(send_state.get("session_opening"))
        or bool(send_state.get("awaiting_session_start"))
    ):
        raise ValueError("Cannot remove profile sessions while WhatsApp sending/sessions are active.")

    base_dir = get_playwright_user_data_dir()
    _cleanup_profile_compact_staging(base_dir)

    # Snapshot every profile folder physically on disk (real sessions + any
    # empty/garbage leftovers), so we can clean and renumber all of them.
    existing = {
        idx: _get_send_profile_dir(base_dir, idx, create=False)
        for idx in _discover_all_profile_dir_indexes(limit=limit)
    }

    # Refuse to remove while a session in range is still open in a browser, so
    # we never yank a folder out from under a live Chromium process.
    for idx, path in existing.items():
        if idx >= removed_idx and path.exists() and _profile_dir_is_claimed(path):
            raise ValueError(f"Profile session #{idx} is currently open. Close it before removing.")

    removed_path = _get_send_profile_dir(base_dir, removed_idx, create=False)
    removed_exists = removed_path.exists()

    # Decide what to KEEP (real sessions, excluding the one being removed) and
    # capture each kept session's phone number before anything moves on disk.
    old_sequence = _load_profile_numbers_sequence()
    keep = []  # list of (path, phone_number)
    for idx in sorted(existing.keys()):
        path = existing[idx]
        if idx == removed_idx:
            continue
        if _profile_dir_has_session_data(path):
            number = _read_profile_number(path) or str(old_sequence.get(idx) or "").strip()
            keep.append((path, str(number or "").strip()))
        else:
            # Empty / garbage folder — delete it outright so it can never come
            # back as a phantom logged-out session.
            _force_rmtree(path)

    # Strictly delete the removed session's whole folder (all files + subdirs).
    if removed_exists:
        _force_rmtree(removed_path)

    # Compact the kept sessions into contiguous slots 1..N. Stage every kept
    # folder under a temporary name first so renumbering can never clobber a
    # folder we still need (the base folder doubles as slot #1).
    staged = []  # list of (stage_path, phone_number)
    for n, (src, number) in enumerate(keep):
        stage = base_dir.parent / f"{base_dir.name}__compact_{n}"
        if stage.exists():
            _force_rmtree(stage)
        shutil.move(str(src), str(stage))
        staged.append((stage, number))
    new_sequence = {}
    for n, (stage, number) in enumerate(staged, start=1):
        dst = _get_send_profile_dir(base_dir, n, create=False)
        if dst.exists():
            _force_rmtree(dst)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(stage), str(dst))
        if number:
            new_sequence[n] = number

    # Rebuild the stored number sequence and transient state for the new order.
    _save_profile_numbers_sequence(new_sequence, min_size=len(staged))
    _shift_profile_session_jobs_after_removal(removed_idx)
    _shift_send_state_profiles_after_removal(send_state, removed_idx)

    remaining_count = len(staged)
    if isinstance(send_state, dict):
        send_state["last_profiles_numbers"] = max(1, remaining_count) if remaining_count > 0 else 1
    return {
        "removed_index": removed_idx,
        "removed_path": str(removed_path),
        "removed_exists": bool(removed_exists),
        "remaining_count": int(remaining_count),
    }


def _load_profile_numbers_from_session_dirs(count, create_missing=False):
    try:
        total = int(count or 0)
    except Exception:
        total = 0
    total = max(0, min(100, total))
    if total < 1:
        return {}

    stored = _load_profile_numbers_sequence()
    base_session_dir = get_playwright_user_data_dir()
    resolved = {}
    updated_sequence = {}
    for idx in range(1, total + 1):
        # Read the profile dir WITHOUT creating it. Fabricating a fresh
        # directory here is what made removed ("garbage") sessions reappear:
        # a stale `count` would recreate the deleted `..._profile_N` folders,
        # which the next session scan then discovered as logged-out sessions.
        # Only ensure the number TXT inside directories that already exist.
        profile_dir = _get_send_profile_dir(base_session_dir, idx, create=False)
        try:
            dir_exists = profile_dir.exists()
        except Exception:
            dir_exists = False
        seed_value = str(stored.get(idx) or "").strip()
        profile_txt_path = _profile_number_file_path(profile_dir)
        if create_missing and dir_exists:
            value = _ensure_and_read_profile_number(profile_dir, seed_value=seed_value)
        elif dir_exists and profile_txt_path.exists():
            value = _read_profile_number(profile_dir)
        else:
            value = seed_value
        value = str(value or "").strip()
        if value:
            updated_sequence[idx] = value
            resolved[idx] = value
    _save_profile_numbers_sequence(updated_sequence, min_size=total)
    return resolved


def _parse_profile_index(value):
    if isinstance(value, int):
        return value if value >= 1 else 0
    try:
        parsed = int(value)
        return parsed if parsed >= 1 else 0
    except Exception:
        pass
    match = re.search(r"(\d+)$", str(value or ""))
    if not match:
        return 0
    try:
        parsed = int(match.group(1))
        return parsed if parsed >= 1 else 0
    except Exception:
        return 0


def _profile_state_row_defaults(profile_idx, profile_number=""):
    idx = max(1, int(profile_idx or 1))
    return {
        "profile_idx": idx,
        "name": f"Profile {idx}",
        "profile_number": str(profile_number or "").strip(),
        "status": "unknown",
        "reason": "",
        "targets_assigned": 0,
        "sent": 0,
        "failed": 0,
        "last_target": "",
    }


def _sync_profile_number_into_send_state(send_state, profile_idx, profile_number):
    idx = _parse_profile_index(profile_idx)
    if idx < 1:
        return
    value = str(profile_number or "").strip()
    for key in ("profile_states", "last_profile_states"):
        rows = list(send_state.get(key) or [])
        if not rows:
            continue
        changed = False
        updated_rows = []
        for row in rows:
            current = dict(row or {})
            row_idx = _parse_profile_index(current.get("profile_idx") or current.get("name"))
            if row_idx == idx:
                current["profile_idx"] = idx
                current["profile_number"] = value
                changed = True
            updated_rows.append(current)
        if changed:
            send_state[key] = updated_rows


# Files/folders that only exist inside a real WhatsApp profile session. An
# empty, freshly auto-created, or half-deleted folder has none of these, so we
# can tell a genuine session apart from "garbage" leftovers on disk.
_PROFILE_SESSION_DATA_MARKERS = (
    "profile_number.txt",
    ".session_ready.json",
    "Default",
    "Local State",
    "Cookies",
)


def _profile_dir_has_session_data(profile_dir):
    """True only if the folder actually holds a real profile session.

    Prevents empty/auto-created/half-removed folders from being surfaced as
    phantom ("garbage") logged-out sessions.
    """
    path = Path(profile_dir)
    try:
        if not path.is_dir():
            return False
    except Exception:
        return False
    for marker in _PROFILE_SESSION_DATA_MARKERS:
        try:
            if (path / marker).exists():
                return True
        except Exception:
            continue
    return False


def _iter_existing_profile_dirs(limit=100):
    """Yield (idx, Path) for every physically-present profile folder on disk.

    Includes empty/garbage folders; callers decide whether to keep them.
    """
    try:
        base_session_dir = get_playwright_user_data_dir().resolve()
    except Exception:
        return
    if not base_session_dir:
        return
    # The current-profile marker may point at a slot folder (for example
    # "default_profile_12"). Slot discovery must always start from the ROOT
    # profile ("default"), exactly like _get_send_profile_dir does, otherwise
    # only that one folder is found and every other account disappears.
    root_name = re.sub(r"(?:_profile_\d+)+$", "", base_session_dir.name) or base_session_dir.name
    base_session_dir = base_session_dir.parent / root_name
    cap = max(1, min(int(limit or 100), 100))
    try:
        if base_session_dir.exists():
            yield 1, base_session_dir
    except Exception:
        pass
    try:
        base_parent = base_session_dir.parent
        prefix = f"{base_session_dir.name}_profile_"
        if base_parent.exists():
            for child in base_parent.iterdir():
                try:
                    if not child.is_dir():
                        continue
                except Exception:
                    continue
                name = str(child.name or "")
                if not name.startswith(prefix):
                    continue
                suffix = name[len(prefix):]
                try:
                    idx = int(suffix)
                except Exception:
                    continue
                if 2 <= idx <= cap:
                    yield idx, child
    except Exception:
        pass


def _discover_all_profile_dir_indexes(limit=100):
    """Every physically-present profile folder index (real + garbage).

    Used by the remover to physically clean up and renumber folders on disk.
    """
    cap = max(1, min(int(limit or 100), 100))
    return sorted({idx for idx, _ in _iter_existing_profile_dirs(limit=limit) if 1 <= idx <= cap})


def _discover_profile_session_indexes(limit=100):
    # Real sessions only: folders that actually contain session data. Empty,
    # auto-created, or half-deleted "garbage" folders are intentionally ignored
    # so they never resurface as phantom logged-out sessions.
    cap = max(1, min(int(limit or 100), 100))
    indexes = set()
    for idx, path in _iter_existing_profile_dirs(limit=limit):
        if 1 <= idx <= cap and _profile_dir_has_session_data(path):
            indexes.add(idx)
    return sorted(indexes)


def _discover_profile_session_count(limit=100):
    session_indexes = _discover_profile_session_indexes(limit=limit)
    if session_indexes:
        return max(session_indexes)
    try:
        stored = _load_profile_numbers_sequence()
        if stored:
            return max(1, min(int(limit or 100), max(int(k) for k in stored.keys())))
    except Exception:
        pass
    return 1


def _filter_logged_in_profile_rows(rows):
    filtered = []
    for row in list(rows or []):
        current = dict(row or {})
        status = str(current.get("status") or "").strip().lower()
        if status not in ("logged_in", "in_use"):
            continue
        current["profile_idx"] = _parse_profile_index(
            current.get("profile_idx") or current.get("name")
        ) or 1
        filtered.append(current)
    filtered.sort(key=lambda item: int(item.get("profile_idx") or 1))
    return filtered


def _save_profile_number_for_session(profile_idx, profile_number, count):
    idx = _parse_profile_index(profile_idx)
    total = max(1, min(100, int(count or idx or 1)))
    if idx < 1 or idx > total:
        raise ValueError("Invalid profile index.")
    profile_dir = _get_send_profile_dir(get_playwright_user_data_dir(), idx, create=True)
    value = str(profile_number or "").strip()
    _store_profile_number(profile_dir, value)
    _set_profile_number_in_sequence(idx, value, min_size=total)
    return value, profile_dir


def _set_profile_session_job(profile_idx, status, message="", **extra):
    idx = _parse_profile_index(profile_idx)
    if idx < 1:
        return
    record = {
        "profile_idx": idx,
        "status": str(status or "").strip().lower() or "unknown",
        "message": str(message or "").strip(),
        "updated_at": time.time(),
    }
    for key, value in dict(extra or {}).items():
        record[str(key)] = value
    with _PROFILE_SESSION_JOBS_LOCK:
        _PROFILE_SESSION_JOBS[idx] = record


def _get_profile_session_jobs(max_age_seconds=1800):
    now = time.time()
    keep = {}
    rows = []
    with _PROFILE_SESSION_JOBS_LOCK:
        for idx, raw in list(_PROFILE_SESSION_JOBS.items()):
            record = dict(raw or {})
            status = str(record.get("status") or "").strip().lower()
            updated_at = float(record.get("updated_at") or 0.0)
            if status in ("opening", "waiting_login") or (now - updated_at) <= float(max_age_seconds):
                keep[idx] = record
        _PROFILE_SESSION_JOBS.clear()
        _PROFILE_SESSION_JOBS.update(keep)
        for idx in sorted(keep.keys()):
            rows.append(dict(keep[idx]))
    return rows


def _set_open_session_preview_window(profile_idx, **fields):
    idx = _parse_profile_index(profile_idx)
    if idx < 1:
        return
    with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
        current = dict(_OPEN_SESSION_PREVIEW_WINDOWS.get(idx) or {})
        current["profile_idx"] = idx
        for key, value in dict(fields or {}).items():
            current[str(key)] = value
        current["updated_at"] = time.time()
        _OPEN_SESSION_PREVIEW_WINDOWS[idx] = current


def _clear_open_session_preview_window(profile_idx):
    idx = _parse_profile_index(profile_idx)
    if idx < 1:
        return
    with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
        _OPEN_SESSION_PREVIEW_WINDOWS.pop(idx, None)


def _get_open_session_preview_rows():
    rows = []
    with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
        for idx in sorted(_OPEN_SESSION_PREVIEW_WINDOWS.keys()):
            current = dict(_OPEN_SESSION_PREVIEW_WINDOWS.get(idx) or {})
            current["profile_idx"] = idx
            current.pop("thread", None)
            current.pop("stop_event", None)
            current.pop("context", None)
            current.pop("page", None)
            current.pop("handoff_ready_event", None)
            current.pop("handoff_claimed_event", None)
            current.pop("handoff_requested", None)
            rows.append(current)
    return rows


def _get_open_session_preview_indexes():
    return [
        int((row or {}).get("profile_idx") or 0)
        for row in _get_open_session_preview_rows()
        if int((row or {}).get("profile_idx") or 0) >= 1
    ]


def _is_open_session_preview_active(profile_idx):
    idx = _parse_profile_index(profile_idx)
    if idx < 1:
        return False
    with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
        current = dict(_OPEN_SESSION_PREVIEW_WINDOWS.get(idx) or {})
    if not current:
        return False
    thread = current.get("thread")
    if thread is None:
        return True
    try:
        return bool(thread.is_alive())
    except Exception:
        return True


def _take_open_session_preview_window_for_send(profile_idx, wait_timeout_seconds=10.0):
    idx = _parse_profile_index(profile_idx)
    if idx < 1:
        return None

    with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
        current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
        if not current:
            return None
        ready_event = current.get("handoff_ready_event")
        if ready_event is None:
            ready_event = threading.Event()
            current["handoff_ready_event"] = ready_event
        claimed_event = current.get("handoff_claimed_event")
        if claimed_event is None:
            claimed_event = threading.Event()
            current["handoff_claimed_event"] = claimed_event
        try:
            ready_event.clear()
        except Exception:
            pass
        try:
            claimed_event.clear()
        except Exception:
            pass
        current["handoff_requested"] = True
        current["status"] = "handoff"
        current["message"] = (
            f"Reusing already opened WhatsApp session window for sending on slot #{idx}..."
        )
        current["updated_at"] = time.time()

    wait_seconds = max(0.5, float(wait_timeout_seconds or 0.0))
    try:
        ready = bool(ready_event.wait(timeout=wait_seconds))
    except Exception:
        ready = False
    if not ready:
        with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
            current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
            if current:
                current["handoff_requested"] = False
                if str(current.get("status") or "").strip().lower() == "handoff":
                    current["status"] = "open"
                    current["message"] = (
                        f"WhatsApp preview window is open for session slot #{idx}."
                    )
                current["updated_at"] = time.time()
        return None

    with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
        current = _OPEN_SESSION_PREVIEW_WINDOWS.pop(idx, None)
    try:
        claimed_event.set()
    except Exception:
        pass
    if not current:
        return None

    context = current.get("context")
    page = current.get("page")
    profile_dir = current.get("profile_dir")
    if context is None or page is None or not str(profile_dir or "").strip():
        if context is not None:
            try:
                context.close()
            except Exception:
                pass
        if str(profile_dir or "").strip():
            try:
                release_playwright_profile(profile_dir)
            except Exception:
                pass
        return None
    return {
        "profile_idx": idx,
        "context": context,
        "page": page,
        "profile_dir": Path(profile_dir),
        "profile_number": str(current.get("profile_number") or "").strip(),
    }


def _run_in_open_session_preview_window(profile_idx, fn, wait_for_start_seconds=15.0):
    """Hand `fn(context, page)` to the thread that owns the open preview window
    for this slot. Playwright objects are bound to their creating thread, so the
    window's own thread must run the work. Returns the request dict (wait on
    request["done_event"], then read request["error"]) or None when no window
    picked the request up in time."""
    idx = _parse_profile_index(profile_idx)
    if idx < 1 or not callable(fn):
        return None
    request = {
        "fn": fn,
        "started_event": threading.Event(),
        "done_event": threading.Event(),
        "result": None,
        "error": None,
        "requested_at": time.time(),
    }
    with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
        current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
        if not current or current.get("run_request") is not None:
            return None
        if str(current.get("status") or "").strip().lower() in ("closing",):
            return None
        current["run_request"] = request
        current["updated_at"] = time.time()
    started = request["started_event"].wait(timeout=max(0.5, float(wait_for_start_seconds or 0.0)))
    if not started:
        with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
            current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
            if current and current.get("run_request") is request:
                current.pop("run_request", None)
        return None
    return request


def _apply_open_session_preview_overlays(rows):
    preview_rows = _get_open_session_preview_rows()
    if not preview_rows:
        return list(rows or [])
    preview_by_idx = {
        int((row or {}).get("profile_idx") or 0): dict(row or {})
        for row in preview_rows
        if int((row or {}).get("profile_idx") or 0) >= 1
    }
    merged = []
    for raw in list(rows or []):
        current = dict(raw or {})
        idx = _parse_profile_index(current.get("profile_idx") or current.get("name"))
        preview = preview_by_idx.get(idx)
        if preview:
            preview_status = str(preview.get("status") or "").strip().lower()
            preview_message = str(preview.get("message") or "").strip()
            preview_number = str(preview.get("profile_number") or "").strip()
            if preview_number:
                current["profile_number"] = preview_number
            if preview_status == "opening":
                current["status"] = "opening"
                current["reason"] = preview_message or (
                    f"Opening WhatsApp preview window for session slot #{idx}."
                )
            elif preview_status == "closing":
                current["status"] = "opening"
                current["reason"] = preview_message or (
                    f"Closing WhatsApp preview window for session slot #{idx}."
                )
            elif preview_status == "login_required":
                current["status"] = "logged_out"
                current["reason"] = preview_message or (
                    f"WhatsApp window is open for session slot #{idx}; scan the QR code to log in."
                )
            else:
                current["status"] = "in_use"
                current["reason"] = preview_message or (
                    "Logged-in session is already open in a preview browser window."
                )
        merged.append(current)
    return merged


def _start_open_session_preview_window(profile_idx, profile_number=""):
    idx = _parse_profile_index(profile_idx)
    if idx < 1:
        raise ValueError("Invalid profile session.")
    if _is_open_session_preview_active(idx):
        return False

    profile_dir = _get_send_profile_dir(get_playwright_user_data_dir(), idx, create=False)
    if not profile_dir.exists():
        raise ValueError(f"Profile session #{idx} was not found.")

    preview_number = _ensure_and_read_profile_number(profile_dir, seed_value=profile_number)
    stop_event = threading.Event()
    handoff_ready_event = threading.Event()
    handoff_claimed_event = threading.Event()
    _set_open_session_preview_window(
        idx,
        status="opening",
        message=f"Opening WhatsApp preview window for session slot #{idx}...",
        profile_dir=str(profile_dir),
        profile_number=preview_number,
        stop_event=stop_event,
        handoff_ready_event=handoff_ready_event,
        handoff_claimed_event=handoff_claimed_event,
        handoff_requested=False,
    )

    def _job():
        nonlocal preview_number
        launch_kwargs = dict(get_playwright_launch_kwargs(default_headless=False))
        launch_kwargs["headless"] = False
        launch_kwargs.setdefault("slow_mo", 0)
        context = None
        locked_profile_dir = None
        handed_off_to_send = False

        def _maybe_run_request(restore_message):
            """Execute a campaign work request inside this window's own thread."""
            with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
                current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
                request = current.get("run_request") if current else None
                if not request:
                    return False
                current["status"] = "sending"
                current["message"] = f"Session slot #{idx} is being used by the running campaign."
                current["updated_at"] = time.time()
            request["started_event"].set()
            try:
                request["result"] = request["fn"](context, page)
            except BaseException as exc:  # noqa: BLE001 - propagate to the requester
                request["error"] = exc
            finally:
                with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
                    current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
                    if current is not None:
                        if current.get("run_request") is request:
                            current.pop("run_request", None)
                        if str(current.get("status") or "") == "sending":
                            current["status"] = "open"
                            current["message"] = str(restore_message or "").strip() or (
                                f"WhatsApp preview window is open for session slot #{idx}."
                            )
                        current["updated_at"] = time.time()
                request["done_event"].set()
            return True

        def _store_logged_in_session(reason="preview_login"):
            try:
                snapshot, _ = save_whatsapp_session_snapshot(
                    page=page,
                    flow_name=reason,
                    user_data_dir=str(profile_dir),
                    login_state="logged_in",
                )
            except Exception:
                snapshot = {}
            snapshot_number = str((snapshot or {}).get("phone_number") or "").strip()
            number_value = str(preview_number or "").strip()
            if snapshot_number and not number_value:
                try:
                    _write_profile_number(profile_dir, snapshot_number)
                    _set_profile_number_in_sequence(idx, snapshot_number, min_size=idx)
                except Exception:
                    pass
                number_value = snapshot_number
            try:
                _mark_profile_session_ready(
                    profile_dir,
                    ready=True,
                    profile_number=number_value,
                    profile_name=str((snapshot or {}).get("profile_name") or "").strip(),
                    phone_number=snapshot_number,
                )
            except Exception:
                pass
            return number_value

        def _maybe_handoff_to_send(restore_message):
            try:
                with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
                    current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
                    if not current or not bool(current.get("handoff_requested")):
                        return False
                    current["context"] = context
                    current["page"] = page
                    current["status"] = "handoff"
                    current["message"] = (
                        f"Reusing already opened WhatsApp session window for sending on slot #{idx}..."
                    )
                    current["updated_at"] = time.time()
                    ready_event = current.get("handoff_ready_event")
                    claimed_event = current.get("handoff_claimed_event")
                try:
                    if ready_event is not None:
                        ready_event.set()
                except Exception:
                    pass
                deadline = time.time() + 12.0
                while time.time() < deadline and not stop_event.is_set():
                    try:
                        if claimed_event is not None and claimed_event.wait(timeout=0.1):
                            return True
                    except Exception:
                        time.sleep(0.1)
                with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
                    current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
                    if current:
                        current["handoff_requested"] = False
                        current["status"] = "open"
                        current["message"] = str(restore_message or "").strip() or (
                            f"WhatsApp preview window is open for session slot #{idx}."
                        )
                        current["updated_at"] = time.time()
                try:
                    if ready_event is not None:
                        ready_event.clear()
                except Exception:
                    pass
                try:
                    if claimed_event is not None:
                        claimed_event.clear()
                except Exception:
                    pass
            except Exception:
                return False
            return False

        try:
            _append_app_log(
                f"Opening logged-in WhatsApp preview window for session slot #{idx}.",
                "SESSIONS",
            )
            with sync_playwright() as p:
                context = launch_persistent_context_safe(
                    browser_type=p.chromium,
                    profile_dir=profile_dir,
                    launch_kwargs=launch_kwargs,
                    owner_label=f"session preview (Profile {idx})",
                    log_fn=lambda msg: _append_app_log(
                        f"[Profile {idx}] {msg}",
                        "SESSIONS",
                    ),
                )
                locked_profile_dir = profile_dir
                page = context.pages[0] if context.pages else context.new_page()
                page.goto(
                    "https://web.whatsapp.com/",
                    wait_until="commit",
                    timeout=15000,
                )
                try:
                    page.bring_to_front()
                except Exception:
                    pass
                _set_open_session_preview_window(
                    idx,
                    context=context,
                    page=page,
                    profile_number=preview_number,
                )
                if _maybe_handoff_to_send(
                    f"WhatsApp preview window is open for session slot #{idx}."
                ):
                    handed_off_to_send = True
                    return
                logged_in = False
                login_required = False
                started = time.time()
                while (time.time() - started) < 6.0 and not stop_event.is_set():
                    if _maybe_run_request(
                        f"WhatsApp preview window is open for session slot #{idx}."
                    ):
                        continue
                    if _maybe_handoff_to_send(
                        f"WhatsApp preview window is open for session slot #{idx}."
                    ):
                        handed_off_to_send = True
                        return
                    try:
                        if page.is_closed():
                            return
                    except Exception:
                        return
                    try:
                        if _is_whatsapp_profile_logged_in(page):
                            logged_in = True
                            break
                    except Exception:
                        pass
                    try:
                        if _is_whatsapp_login_required(page):
                            login_required = True
                            break
                    except Exception:
                        pass
                    try:
                        page.wait_for_timeout(250)
                    except Exception:
                        time.sleep(0.25)
                open_message = f"WhatsApp preview window is open for session slot #{idx}."
                session_marked_ready = False
                if logged_in:
                    open_message = f"WhatsApp session slot #{idx} is logged in and ready for sending."
                    stored_number = _store_logged_in_session("preview_login")
                    if stored_number:
                        preview_number = stored_number
                    session_marked_ready = True
                elif login_required:
                    open_message = (
                        f"WhatsApp window is open for session slot #{idx}; scan the QR code to log in."
                    )
                _set_open_session_preview_window(
                    idx,
                    status=("login_required" if (login_required and not logged_in) else "open"),
                    message=open_message,
                    profile_number=preview_number,
                    context=context,
                    page=page,
                )
                if _maybe_handoff_to_send(open_message):
                    handed_off_to_send = True
                    return

                last_login_check = 0.0
                while not stop_event.is_set():
                    if _maybe_run_request(open_message):
                        continue
                    if _maybe_handoff_to_send(open_message):
                        handed_off_to_send = True
                        return
                    if (time.time() - last_login_check) >= 4.0:
                        last_login_check = time.time()
                        now_logged_in = False
                        try:
                            now_logged_in = _is_whatsapp_profile_logged_in(page)
                        except Exception:
                            now_logged_in = False
                        if now_logged_in and not session_marked_ready:
                            session_marked_ready = True
                            stored_number = _store_logged_in_session("preview_login")
                            if stored_number:
                                preview_number = stored_number
                            open_message = f"WhatsApp session slot #{idx} is logged in and ready for sending."
                            _set_open_session_preview_window(
                                idx,
                                status="open",
                                message=open_message,
                                profile_number=preview_number,
                            )
                            _append_app_log(
                                f"WhatsApp session slot #{idx} logged in and stored from the open window.",
                                "SESSIONS",
                            )
                        elif not now_logged_in and session_marked_ready:
                            session_marked_ready = False
                            try:
                                _mark_profile_session_ready(profile_dir, ready=False)
                            except Exception:
                                pass
                            open_message = (
                                f"WhatsApp window is open for session slot #{idx}; scan the QR code to log in."
                            )
                            _set_open_session_preview_window(
                                idx,
                                status="login_required",
                                message=open_message,
                                profile_number=preview_number,
                            )
                    try:
                        if page.is_closed():
                            _append_app_log(
                                f"WhatsApp preview window closed for session slot #{idx}.",
                                "SESSIONS",
                            )
                            return
                    except Exception:
                        return
                    try:
                        page.wait_for_timeout(500)
                    except Exception:
                        time.sleep(0.5)
        except Exception as exc:
            _append_app_log(
                f"Could not open WhatsApp preview window for session slot #{idx}: {exc}",
                "SESSIONS",
            )
        finally:
            if handed_off_to_send:
                _append_app_log(
                    f"Reusing the already opened WhatsApp preview window for sending on session slot #{idx}.",
                    "SESSIONS",
                )
                return
            _set_open_session_preview_window(
                idx,
                status="closing",
                message=f"Closing WhatsApp preview window for session slot #{idx}...",
                profile_number=preview_number,
            )
            if context is not None:
                try:
                    context.close()
                except Exception:
                    pass
            if str(locked_profile_dir or "").strip():
                release_playwright_profile(locked_profile_dir)
            _clear_open_session_preview_window(idx)

    thread = threading.Thread(target=_job, daemon=True)
    _set_open_session_preview_window(idx, thread=thread)
    thread.start()
    return True


def _request_close_open_session_preview_windows(profile_indexes=None, wait_timeout_seconds=12.0):
    target_indexes = set()
    if profile_indexes is None:
        target_indexes = set(_get_open_session_preview_indexes())
    else:
        for raw in list(profile_indexes or []):
            idx = _parse_profile_index(raw)
            if idx >= 1:
                target_indexes.add(idx)
    if not target_indexes:
        return {"requested": 0, "closed": 0, "remaining": []}

    with _OPEN_SESSION_PREVIEW_WINDOWS_LOCK:
        targeted = []
        for idx in sorted(target_indexes):
            current = _OPEN_SESSION_PREVIEW_WINDOWS.get(idx)
            if not current:
                continue
            current["status"] = "closing"
            current["message"] = f"Closing WhatsApp preview window for session slot #{idx}..."
            current["updated_at"] = time.time()
            targeted.append((idx, current.get("stop_event")))
    if not targeted:
        return {"requested": 0, "closed": 0, "remaining": []}

    for _, stop_event in targeted:
        try:
            if stop_event is not None:
                stop_event.set()
        except Exception:
            continue

    deadline = time.time() + max(0.5, float(wait_timeout_seconds or 0.0))
    remaining = {idx for idx, _ in targeted}
    while remaining and time.time() < deadline:
        remaining = {idx for idx in remaining if _is_open_session_preview_active(idx)}
        if remaining:
            time.sleep(0.1)

    return {
        "requested": len(targeted),
        "closed": len(targeted) - len(remaining),
        "remaining": sorted(remaining),
    }


def _run_profile_session_login_window(profile_idx):
    idx = _parse_profile_index(profile_idx)
    if idx < 1:
        return

    def _job():
        profile_dir = _get_send_profile_dir(get_playwright_user_data_dir(), idx, create=True)
        profile_number = _ensure_and_read_profile_number(profile_dir, seed_value="")
        launch_kwargs = dict(get_playwright_launch_kwargs(default_headless=False))
        launch_kwargs["headless"] = False
        launch_kwargs.setdefault("slow_mo", 0)
        context = None
        locked_profile_dir = None
        try:
            _mark_profile_session_ready(profile_dir, ready=False)
            _set_profile_session_job(
                idx,
                "opening",
                f"Opening WhatsApp Web for session slot #{idx}...",
                profile_number=profile_number,
            )
            with sync_playwright() as p:
                context = launch_persistent_context_safe(
                    browser_type=p.chromium,
                    profile_dir=profile_dir,
                    launch_kwargs=launch_kwargs,
                    owner_label=f"session setup (Profile {idx})",
                )
                locked_profile_dir = profile_dir
                page = context.pages[0] if context.pages else context.new_page()
                page.goto(
                    "https://web.whatsapp.com/",
                    wait_until="commit",
                    timeout=15000,
                )
                try:
                    page.bring_to_front()
                except Exception:
                    pass
                _set_profile_session_job(
                    idx,
                    "waiting_login",
                    f"WhatsApp Web is open for session slot #{idx}. Scan QR and complete login.",
                    profile_number=profile_number,
                )
                deadline = time.time() + 600
                while time.time() < deadline:
                    try:
                        if page.is_closed():
                            _set_profile_session_job(
                                idx,
                                "cancelled",
                                f"WhatsApp session slot #{idx} browser was closed before login completed.",
                                profile_number=profile_number,
                            )
                            return
                    except Exception:
                        pass
                    try:
                        if _is_whatsapp_profile_logged_in(page):
                            _set_profile_session_job(
                                idx,
                                "storing",
                                f"Login detected for session slot #{idx}. Saving WhatsApp session...",
                                profile_number=profile_number,
                            )
                            snapshot = {}
                            try:
                                snapshot, _ = save_whatsapp_session_snapshot(
                                    page=page,
                                    flow_name="profile_session_add",
                                    user_data_dir=str(profile_dir),
                                    login_state="logged_in",
                                )
                            except Exception:
                                snapshot = {}
                            snapshot_number = str((snapshot or {}).get("phone_number") or "").strip()
                            if snapshot_number and not str(profile_number or "").strip():
                                try:
                                    _write_profile_number(profile_dir, snapshot_number)
                                    _set_profile_number_in_sequence(idx, snapshot_number, min_size=idx)
                                    profile_number = snapshot_number
                                except Exception:
                                    profile_number = snapshot_number
                            profile_number = _ensure_and_read_profile_number(
                                profile_dir,
                                seed_value=profile_number,
                            )
                            stabilization_until = time.time() + 8.0
                            while time.time() < stabilization_until:
                                try:
                                    if page.is_closed():
                                        break
                                except Exception:
                                    break
                                try:
                                    page.wait_for_timeout(500)
                                except Exception:
                                    time.sleep(0.5)
                            _mark_profile_session_ready(
                                profile_dir,
                                ready=True,
                                profile_number=profile_number,
                                profile_name=str((snapshot or {}).get("profile_name") or "").strip(),
                                phone_number=snapshot_number,
                            )
                            _set_profile_session_job(
                                idx,
                                "stored",
                                f"WhatsApp session slot #{idx} logged in and stored successfully.",
                                profile_number=profile_number,
                            )
                            try:
                                page.wait_for_timeout(2200)
                            except Exception:
                                time.sleep(2.2)
                            return
                    except Exception:
                        pass
                    try:
                        page.wait_for_timeout(700)
                    except Exception:
                        time.sleep(0.7)
                _set_profile_session_job(
                    idx,
                    "timeout",
                    f"WhatsApp session slot #{idx} did not finish login yet. You can try again.",
                    profile_number=profile_number,
                )
                _mark_profile_session_ready(profile_dir, ready=False)
        except Exception as exc:
            _set_profile_session_job(
                idx,
                "failed",
                f"Could not open WhatsApp session slot #{idx}: {exc}",
                profile_number=profile_number,
            )
            _mark_profile_session_ready(profile_dir, ready=False)
        finally:
            if context is not None:
                try:
                    context.close()
                except Exception:
                    pass
            if str(locked_profile_dir or "").strip():
                release_playwright_profile(locked_profile_dir)

    thread = threading.Thread(target=_job, daemon=True)
    thread.start()
    return thread


def _create_next_profile_session_slot(send_state=None, limit=100):
    existing_indexes = _discover_profile_session_indexes(limit=limit)
    stored_sequence = _load_profile_numbers_sequence()
    stored_max = 0
    if stored_sequence:
        try:
            stored_max = max(int(k) for k in stored_sequence.keys())
        except Exception:
            stored_max = 0
    live_max = 0
    if isinstance(send_state, dict):
        for key in ("profile_states", "last_profile_states"):
            for row in list(send_state.get(key) or []):
                live_max = max(
                    live_max,
                    _parse_profile_index((row or {}).get("profile_idx") or (row or {}).get("name")),
                )
    current_max = max([0] + list(existing_indexes) + [stored_max, live_max])
    next_idx = 1 if current_max < 1 else (current_max + 1)
    max_limit = max(1, min(int(limit or 100), 100))
    if next_idx > max_limit:
        raise ValueError(f"Maximum {max_limit} WhatsApp sessions reached.")
    profile_dir = _get_send_profile_dir(
        get_playwright_user_data_dir(),
        next_idx,
        create=True,
    )
    profile_number = _ensure_and_read_profile_number(profile_dir, seed_value="")
    _set_profile_number_in_sequence(next_idx, profile_number, min_size=next_idx)
    if isinstance(send_state, dict):
        send_state["last_profiles_numbers"] = max(
            int(send_state.get("last_profiles_numbers") or 1),
            next_idx,
        )
    return next_idx, profile_dir, profile_number


def _build_profile_rows_for_count(count, sequence_map=None, live_rows=None):
    total = max(1, min(100, int(count or 1)))
    seq = dict(sequence_map or {})
    live_by_idx = {}
    for row in list(live_rows or []):
        current = dict(row or {})
        row_idx = _parse_profile_index(current.get("profile_idx") or current.get("name"))
        if row_idx >= 1:
            current["profile_idx"] = row_idx
            live_by_idx[row_idx] = current
    rows = []
    for idx in range(1, total + 1):
        seed_number = str(seq.get(idx) or seq.get(str(idx)) or "").strip()
        row = _profile_state_row_defaults(idx, seed_number)
        if idx in live_by_idx:
            row.update(live_by_idx[idx])
            row["profile_idx"] = idx
            row["name"] = row.get("name") or f"Profile {idx}"
        row["profile_number"] = str(row.get("profile_number") or seed_number or "").strip()
        rows.append(row)
    return rows


def _scan_profile_sessions_status(count, send_state=None, validate=False):
    discovered_indexes = _discover_profile_session_indexes(limit=100)
    discovered_total = max(discovered_indexes) if discovered_indexes else 0
    current_rows = []
    base_session_dir = get_playwright_user_data_dir()
    if isinstance(send_state, dict) and bool(send_state.get("running")):
        current_rows = list(send_state.get("profile_states") or [])
    current_row_max = 0
    for row in current_rows:
        current_row_max = max(
            current_row_max,
            _parse_profile_index((row or {}).get("profile_idx") or (row or {}).get("name")),
        )
    total = max(1, min(100, max(int(count or 1), discovered_total, current_row_max)))
    sequence_map = _load_profile_numbers_from_session_dirs(
        total,
        create_missing=bool(current_rows),
    )
    if current_rows:
        rows = _build_profile_rows_for_count(total, sequence_map=sequence_map, live_rows=current_rows)
        active_rows = _filter_logged_in_profile_rows(rows)
        return {
            "count": total,
            "sequence": {str(k): v for k, v in sequence_map.items()},
            "rows": rows,
            "active_rows": active_rows,
            "active_profile_indexes": [int(row.get("profile_idx") or 1) for row in active_rows],
            "source": "live",
        }

    if not discovered_indexes:
        return {
            "count": 0,
            "sequence": {str(k): v for k, v in sequence_map.items()},
            "rows": [],
            "active_rows": [],
            "active_profile_indexes": [],
            "source": "empty",
        }

    if not bool(validate):
        rows = []
        for idx in discovered_indexes:
            profile_dir = _get_send_profile_dir(base_session_dir, idx, create=False)
            profile_number = _ensure_and_read_profile_number(
                profile_dir,
                seed_value=str(sequence_map.get(idx) or "").strip(),
            )
            session_ready = _is_profile_session_ready(profile_dir)
            row = _profile_state_row_defaults(idx, profile_number)
            if session_ready:
                row["status"] = "logged_in"
                row["reason"] = "Stored logged-in session found."
            else:
                row["status"] = "logged_out"
                row["reason"] = "QR/login required."
            rows.append(row)
        rows = _apply_open_session_preview_overlays(rows)
        active_rows = _filter_logged_in_profile_rows(rows)
        return {
            # Report the count actually present on disk, not the (possibly
            # stale) requested count, so removals shrink the count instead of
            # leaving a sticky value that keeps re-summoning deleted sessions.
            "count": discovered_total,
            "sequence": {str(k): v for k, v in sequence_map.items()},
            "rows": rows,
            "active_rows": active_rows,
            "active_profile_indexes": [int(row.get("profile_idx") or 1) for row in active_rows],
            "source": "cache",
        }

    launch_kwargs = dict(get_playwright_launch_kwargs(default_headless=True))
    launch_kwargs["headless"] = True
    launch_kwargs["slow_mo"] = 0
    launch_kwargs.pop("args", None)
    launch_kwargs.pop("no_viewport", None)
    rows = []
    with sync_playwright() as p:
        for idx in discovered_indexes:
            profile_dir = _get_send_profile_dir(base_session_dir, idx, create=False)
            profile_number = _ensure_and_read_profile_number(
                profile_dir,
                seed_value=str(sequence_map.get(idx) or "").strip(),
            )
            session_ready = _is_profile_session_ready(profile_dir)
            row = _profile_state_row_defaults(idx, profile_number)
            context = None
            locked_profile_dir = None
            try:
                context = launch_persistent_context_safe(
                    browser_type=p.chromium,
                    profile_dir=profile_dir,
                    launch_kwargs=launch_kwargs,
                    owner_label=f"profile session scan (Profile {idx})",
                )
                locked_profile_dir = profile_dir
                page = context.pages[0] if context.pages else context.new_page()
                page.goto("https://web.whatsapp.com/", wait_until="commit", timeout=12000)
                logged_in = False
                login_required = False
                started = time.time()
                while (time.time() - started) < 5.5:
                    try:
                        if _is_whatsapp_profile_logged_in(page):
                            logged_in = True
                            break
                    except Exception:
                        pass
                    try:
                        if _is_whatsapp_login_required(page):
                            login_required = True
                            break
                    except Exception:
                        pass
                    page.wait_for_timeout(250)
                if logged_in:
                    _mark_profile_session_ready(profile_dir, ready=True, profile_number=profile_number)
                    row["status"] = "logged_in"
                    row["reason"] = "Logged in and usable."
                elif session_ready and not login_required:
                    row["status"] = "logged_in"
                    row["reason"] = "Stored logged-in session found."
                else:
                    _mark_profile_session_ready(profile_dir, ready=False)
                    row["status"] = "logged_out"
                    row["reason"] = "QR/login required."
            except Exception as exc:
                exc_text = str(exc or "").strip()
                if "locked by another chromium process" in exc_text.lower():
                    row["status"] = "in_use" if session_ready else "logged_out"
                    row["reason"] = (
                        "Logged-in session is already open in another browser process."
                        if session_ready
                        else "Session is open in another browser process."
                    )
                elif session_ready:
                    row["status"] = "logged_in"
                    row["reason"] = "Stored session is available. Validation check had a temporary warning."
                else:
                    row["status"] = "error"
                    row["reason"] = ("Check failed: " + exc_text)[:240]
            finally:
                if context is not None:
                    try:
                        context.close()
                    except Exception:
                        pass
                if str(locked_profile_dir or "").strip():
                    release_playwright_profile(locked_profile_dir)
            rows.append(row)

    rows = _apply_open_session_preview_overlays(rows)
    active_rows = _filter_logged_in_profile_rows(rows)
    return {
        # See note above: report the on-disk session count, not the requested
        # one, so a removed session does not silently come back on next poll.
        "count": discovered_total,
        "sequence": {str(k): v for k, v in sequence_map.items()},
        "rows": rows,
        "active_rows": active_rows,
        "active_profile_indexes": [int(row.get("profile_idx") or 1) for row in active_rows],
        "source": "scan",
    }


def resolve_scrape_profile_number():
    stored_numbers = _load_profile_numbers_sequence()
    value = str(stored_numbers.get(1) or "").strip()
    if value:
        return value
    try:
        return _read_profile_number(get_playwright_user_data_dir())
    except Exception:
        return ""


def _wait_for_any_selector(page, selectors, timeout_ms=120000):
    for selector in selectors:
        try:
            page.wait_for_selector(selector, timeout=timeout_ms)
            return selector
        except Exception:
            continue
    return ""


def _first_text_from_selectors(page, selectors, timeout_ms=1200):
    for selector in selectors:
        try:
            node = page.locator(selector).first
            if node.count() == 0:
                continue
            node.wait_for(state="visible", timeout=timeout_ms)
            value = (node.inner_text() or "").strip()
            if value:
                return value
        except Exception:
            continue
    return ""


def _click_first_selector(page, selectors, timeout_ms=1500, label="click", force=False, state=None):
    for selector in selectors:
        try:
            node = page.locator(selector).first
            if node.count() == 0:
                continue
            node.wait_for(state="visible", timeout=timeout_ms)
            try:
                node.scroll_into_view_if_needed(timeout=timeout_ms)
            except Exception:
                pass
            if _human_click_item(
                page,
                node,
                label=label or selector,
                timeout_ms=timeout_ms,
                force=force,
                state=state,
            ):
                return True
            if _delayed_dom_click_fallback(page, node, label=label or selector, state=state):
                return True
        except Exception:
            continue
    return False


def _is_profile_panel_open(page):
    opened = _wait_for_any_selector(
        page,
        (
            "[data-testid='drawer-right']",
            "span[data-testid='pushname-input-read-only selectable-text']",
            "span[data-testid='pushname-input-read-only']",
            "[data-testid='phone']",
            "button[aria-label='Edit Name']",
            "button[aria-label='Edit profile photo']",
        ),
        timeout_ms=1400,
    )
    return bool(opened)


def _force_click_profile_navbar(page):
    # Profile opening is globally disabled by user request.
    return False


def _open_whatsapp_profile_panel(page):
    # Profile opening is globally disabled by user request.
    return False, "disabled_no_profile_click"


def _collect_whatsapp_account_snapshot(page):
    profile_name = ""
    profile_status = ""
    phone_number = ""
    opened_profile = False
    profile_open_method = "none"
    try:
        _wait_for_any_selector(
            page,
            (
                "[data-testid='chat-list']",
                "div[aria-label='Chat list']",
                "#pane-side",
            ),
            timeout_ms=120000,
        )
        page.wait_for_timeout(300)

        # User-requested behavior: do not click/open WhatsApp Profile during startup.
        opened_profile = False
        profile_open_method = "disabled_no_profile_click"
        if opened_profile:
            page.wait_for_timeout(700)

            profile_name = _first_text_from_selectors(
                page,
                (
                    "span[data-testid='pushname-input-read-only selectable-text']",
                    "span[data-testid='pushname-input-read-only selectable-text'] span",
                    "[data-testid='drawer-right'] h2",
                    "[data-testid='drawer-right'] h1",
                    "header h2",
                    "header h1",
                    "h2",
                ),
                timeout_ms=1800,
            )
            if str(profile_name or "").strip().lower() in (
                "profile",
                "your profile",
                "you",
            ):
                profile_name = ""
            profile_status = _first_text_from_selectors(
                page,
                (
                    "[data-testid='drawer-right'] [data-testid='status']",
                    "[data-testid='drawer-right'] div:has-text('About') + div",
                    "span[data-testid='about']",
                    "[data-testid='drawer-right'] span[dir='auto']",
                    "div:has-text('About') + div",
                ),
                timeout_ms=1200,
            )
            phone_number = _first_text_from_selectors(
                page,
                (
                    "xpath=//*[@data-testid='phone']/following::span[contains(normalize-space(.),'+')][1]",
                    "div:has([data-testid='phone']) span",
                    "span:has-text('+92')",
                ),
                timeout_ms=1200,
            )
            panel_text = ""
            try:
                panel_text = (
                    page.locator("[data-testid='drawer-right']").first.inner_text(timeout=1200) or ""
                )
            except Exception:
                try:
                    panel_text = page.locator("#app").inner_text(timeout=1200) or ""
                except Exception:
                    panel_text = ""
            if panel_text:
                cleaned_panel_text = (
                    panel_text.replace("\u00a0", " ")
                    .replace("\u200e", "")
                    .replace("\u200f", "")
                    .replace("\u202a", "")
                    .replace("\u202b", "")
                    .replace("\u202c", "")
                    .replace("\u2066", "")
                    .replace("\u2067", "")
                    .replace("\u2068", "")
                    .replace("\u2069", "")
                )
                phone_match = re.search(r"\+\s*\d[\d\s\-\(\)]{6,}\d", cleaned_panel_text)
                if not phone_match:
                    phone_match = re.search(r"\b\d[\d\s\-\(\)]{6,}\d\b", cleaned_panel_text)
                if phone_match and not phone_number:
                    phone_number = re.sub(r"\s+", " ", phone_match.group(0)).strip()
                lines = [ln.strip() for ln in cleaned_panel_text.splitlines() if ln and ln.strip()]
                ignored_exact = {
                    "profile",
                    "your profile",
                    "about",
                    "phone",
                    "privacy",
                    "notifications",
                    "security",
                }
                if not profile_name:
                    for ln in lines:
                        low = ln.lower()
                        if (
                            low in ignored_exact
                            or "end-to-end encrypted" in low
                            or low.startswith("phone")
                            or re.search(r"\+\d", ln)
                            or re.search(r"\b\d{7,15}\b", ln)
                        ):
                            continue
                        if len(ln) > 1 and len(ln) <= 80 and not re.search(r"\+\d", ln):
                            profile_name = ln
                            break
                if not profile_status:
                    for idx, ln in enumerate(lines):
                        if ln.strip().lower() == "about" and idx + 1 < len(lines):
                            nxt = lines[idx + 1].strip()
                            if nxt:
                                profile_status = nxt
                                break
                if not profile_status:
                    for ln in lines:
                        low = ln.lower()
                        if (
                            not ln
                            or low in ignored_exact
                            or ln == profile_name
                            or "end-to-end encrypted" in low
                            or re.search(r"\+\d", ln)
                            or re.search(r"\b\d{7,15}\b", ln)
                        ):
                            continue
                        if 2 <= len(ln) <= 120:
                            profile_status = ln
                            break
    finally:
        try:
            page.keyboard.press("Escape")
            page.wait_for_timeout(120)
            page.keyboard.press("Escape")
        except Exception:
            pass

    return {
        "profile_name": str(profile_name or "").strip(),
        "profile_status": str(profile_status or "").strip(),
        "phone_number": str(phone_number or "").strip(),
        "profile_opened": bool(opened_profile),
        "profile_open_method": profile_open_method,
    }


def save_whatsapp_session_snapshot(page, flow_name, user_data_dir, login_state=""):
    snapshot = {
        "profile_name": "",
        "profile_status": "",
        "phone_number": "",
        "profile_opened": False,
    }
    snapshot_error = ""
    try:
        snapshot = _collect_whatsapp_account_snapshot(page)
    except Exception as exc:
        snapshot_error = str(exc)
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = (
        f"{ts} | flow={flow_name} | profile_name={snapshot.get('profile_name') or '-'} | "
        f"status={snapshot.get('profile_status') or '-'} | "
        f"number={snapshot.get('phone_number') or '-'} | "
        f"profile_opened={snapshot.get('profile_opened')} | "
        f"profile_open_method={snapshot.get('profile_open_method') or '-'} | "
        f"login_state={str(login_state or '').strip() or '-'} | "
        f"snapshot_error={snapshot_error or '-'} | "
        f"user_data_dir={user_data_dir}"
    )
    return snapshot, line


CONTACT_COLUMN_ALIASES = (
    "phone",
    "phone number",
    "contacts",
    "contact",
    "phones",
)


def normalize_column_name(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    return "".join(ch for ch in text if ch.isalnum())


CONTACT_COLUMN_ALIASES_NORMALIZED = {
    normalize_column_name(v) for v in CONTACT_COLUMN_ALIASES
}

GROUP_COLUMN_CANDIDATES = ("whatspp_group", "whatsapp_group")
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".3gp", ".m4v"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".ogg", ".m4a", ".aac", ".flac", ".opus"}
DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".csv",
    ".txt",
    ".zip",
    ".ppt",
    ".pptx",
}
RECENT_MEDIA_EXTENSIONS = IMAGE_EXTENSIONS | VIDEO_EXTENSIONS | AUDIO_EXTENSIONS | DOCUMENT_EXTENSIONS


def get_db_connection():
    if psycopg2 is None:
        raise RuntimeError(
            "psycopg2 is not installed. Install it with: pip install psycopg2-binary"
        )

    conn = psycopg2.connect(
        host=os.getenv("PGHOST", "localhost"),
        port=int(os.getenv("PGPORT", "5432")),
        dbname=os.getenv("PGDATABASE", "whatsapp_db"),
        user=os.getenv("PGUSER", "postgres"),
        password=os.getenv("PGPASSWORD", ""),
    )
    conn.autocommit = True
    return conn


def get_db_schema():
    return os.getenv("PGSCHEMA", "public")


def ensure_message_log_tables(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL("DROP TABLE IF EXISTS {schema_name}.sent_sms CASCADE").format(
                schema_name=sql.Identifier(schema)
            )
        )
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.sms_sent (
                    id BIGSERIAL PRIMARY KEY,
                    campaign_id VARCHAR(6) NOT NULL,
                    campaign_name TEXT,
                    phone VARCHAR(20),
                    message TEXT,
                    status VARCHAR(20),
                    sent_at TIMESTAMP,
                    content_kind VARCHAR(50),
                    attachment_names TEXT,
                    attachment_count INTEGER DEFAULT 0,
                    payload_signature TEXT,
                    content_summary TEXT,
                    attachments_payload JSONB,
                    dedupe_signature TEXT
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.sms_not_sent (
                    id BIGSERIAL PRIMARY KEY,
                    campaign_id VARCHAR(6) NOT NULL,
                    campaign_name TEXT,
                    phone VARCHAR(20),
                    reason TEXT,
                    failed_at TIMESTAMP
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS content_kind VARCHAR(50)"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS attachment_names TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS attachment_count INTEGER DEFAULT 0'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS payload_signature TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS dedupe_signature TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS content_summary TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS attachments_payload JSONB"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS campaign_id VARCHAR(6)"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_sent ADD COLUMN IF NOT EXISTS campaign_name TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                ALTER TABLE {schema_name}.sms_sent
                ALTER COLUMN campaign_id TYPE VARCHAR(6)
                USING LPAD(COALESCE(campaign_id::text, '0'), 6, '0')
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_not_sent ADD COLUMN IF NOT EXISTS campaign_id VARCHAR(6)"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                "ALTER TABLE {schema_name}.sms_not_sent ADD COLUMN IF NOT EXISTS campaign_name TEXT"
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                ALTER TABLE {schema_name}.sms_sent
                ALTER COLUMN phone TYPE VARCHAR(255)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                ALTER TABLE {schema_name}.sms_not_sent
                ALTER COLUMN phone TYPE VARCHAR(255)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                ALTER TABLE {schema_name}.sms_not_sent
                ALTER COLUMN campaign_id TYPE VARCHAR(6)
                USING LPAD(COALESCE(campaign_id::text, '0'), 6, '0')
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL("DROP INDEX IF EXISTS {schema_name}.ux_sms_sent_phone_message").format(
                schema_name=sql.Identifier(schema)
            )
        )
        cur.execute(
            sql.SQL("DROP INDEX IF EXISTS {schema_name}.ux_sms_sent_phone_payload").format(
                schema_name=sql.Identifier(schema)
            )
        )
        cur.execute(
            sql.SQL(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS ux_sms_sent_phone_payload
                ON {schema_name}.sms_sent (phone, payload_signature, campaign_id)
                WHERE payload_signature IS NOT NULL
                  AND campaign_id IS NOT NULL
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_phone_dedupe
                ON {schema_name}.sms_sent (phone, dedupe_signature)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_phone
                ON {schema_name}.sms_sent (phone)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_sent_at
                ON {schema_name}.sms_sent (sent_at DESC)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_campaign_id
                ON {schema_name}.sms_sent (campaign_id)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_sent_campaign_name
                ON {schema_name}.sms_sent (campaign_name)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_not_sent_phone
                ON {schema_name}.sms_not_sent (phone)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_not_sent_failed_at
                ON {schema_name}.sms_not_sent (failed_at DESC)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_not_sent_campaign_id
                ON {schema_name}.sms_not_sent (campaign_id)
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS ix_sms_not_sent_campaign_name
                ON {schema_name}.sms_not_sent (campaign_name)
                """
            ).format(schema_name=sql.Identifier(schema))
        )


def _normalize_group_name(value):
    return str(value or "").strip()


def _normalize_group_key(value):
    text = _normalize_group_name(value)
    if not text:
        return ""
    text = text.replace("\u200b", "").replace("\u200c", "").replace("\u200d", "").replace("\ufeff", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text.casefold()


def _safe_int(value, default=0):
    try:
        return int(value)
    except Exception:
        return default


def _parse_optional_limit(raw_value, default=None):
    try:
        value = int(raw_value)
    except Exception:
        return default
    if value <= 0:
        return None
    return value


def generate_campaign_id():
    schema = get_db_schema()
    db_conn = get_db_connection()
    try:
        ensure_message_log_tables(db_conn)
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    WITH all_ids AS (
                        SELECT campaign_id::text AS campaign_id
                        FROM {schema_name}.sms_sent
                        UNION ALL
                        SELECT campaign_id::text AS campaign_id
                        FROM {schema_name}.sms_not_sent
                    )
                    SELECT COALESCE(
                        MAX(
                            CASE
                                WHEN campaign_id ~ '^[0-9]+$' THEN campaign_id::BIGINT
                                ELSE NULL
                            END
                        ),
                        0
                    )
                    FROM all_ids
                    """
                ).format(schema_name=sql.Identifier(schema))
            )
            row = cur.fetchone()
            last_value = int(row[0]) if row and row[0] is not None else 0
            next_value = last_value + 1
            return f"{next_value:06d}"
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def _normalize_campaign_name_base(value, default_value="campaign"):
    text = re.sub(r"\s+", " ", str(value or "").strip())
    if text:
        matched = re.match(r"^(.*?)(?:_(\d+))$", text)
        if matched and str(matched.group(1) or "").strip():
            text = str(matched.group(1) or "").strip()
    if not text:
        text = str(default_value or "campaign").strip() or "campaign"
    return text[:120]


def _split_campaign_name_base_and_suffix(value, default_value="campaign"):
    text = re.sub(r"\s+", " ", str(value or "").strip())
    if not text:
        return _normalize_campaign_name_base("", default_value=default_value), None
    matched = re.match(r"^(.*?)(?:_(\d+))$", text)
    if matched and str(matched.group(1) or "").strip():
        base_name = _normalize_campaign_name_base(
            matched.group(1),
            default_value=default_value,
        )
        try:
            suffix_value = int(matched.group(2))
        except Exception:
            suffix_value = None
        return base_name, suffix_value
    return _normalize_campaign_name_base(text, default_value=default_value), None


def _campaign_name_usage_state(base_name, current_name=""):
    base_text = _normalize_campaign_name_base(base_name)
    current_text = re.sub(r"\s+", " ", str(current_name or "").strip())
    schema = get_db_schema()
    db_conn = get_db_connection()
    try:
        ensure_message_log_tables(db_conn)
        pattern_sql = r"(?i)^" + re.escape(base_text) + r"_(\d+)$"
        pattern_py = re.compile(r"^" + re.escape(base_text) + r"_(\d+)$", re.IGNORECASE)
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    WITH all_names AS (
                        SELECT NULLIF(BTRIM(campaign_name), '') AS campaign_name
                        FROM {schema_name}.sms_sent
                        UNION ALL
                        SELECT NULLIF(BTRIM(campaign_name), '') AS campaign_name
                        FROM {schema_name}.sms_not_sent
                    )
                    SELECT campaign_name
                    FROM all_names
                    WHERE campaign_name IS NOT NULL
                      AND (
                          campaign_name ~* %s
                          OR (%s <> '' AND LOWER(campaign_name) = LOWER(%s))
                      )
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (pattern_sql, current_text, current_text),
            )
            rows = cur.fetchall() or []
        max_suffix = 0
        exact_exists = False
        for row in rows:
            name_text = str((row[0] if row else "") or "").strip()
            if not name_text:
                continue
            if current_text and name_text.casefold() == current_text.casefold():
                exact_exists = True
            matched = pattern_py.match(name_text)
            if not matched:
                continue
            try:
                suffix_value = int(matched.group(1))
            except Exception:
                continue
            if suffix_value > max_suffix:
                max_suffix = suffix_value
        return {"max_suffix": int(max_suffix), "exact_exists": bool(exact_exists)}
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def generate_autonomous_campaign_name(value="", default_base="campaign"):
    current_text = re.sub(r"\s+", " ", str(value or "").strip())
    base_name, explicit_suffix = _split_campaign_name_base_and_suffix(
        current_text,
        default_value=default_base,
    )
    state = _campaign_name_usage_state(base_name, current_name=current_text)
    if explicit_suffix is not None and explicit_suffix > 0 and not bool(state.get("exact_exists")):
        return f"{base_name}_{int(explicit_suffix):04d}"[:120]
    next_suffix = max(1, int(state.get("max_suffix") or 0) + 1)
    return f"{base_name}_{next_suffix:04d}"[:120]


def _normalize_campaign_id(value):
    text = str(value or "").strip()
    digits = "".join(ch for ch in text if ch.isdigit())
    if not digits:
        return ""
    return digits.zfill(6)


def _json_datetime(value):
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def _looks_like_phone_target(value):
    text = str(value or "").strip()
    if not text:
        return False
    compact = re.sub(r"[\s\-\(\)\.]", "", text)
    if compact.startswith("+"):
        compact = compact[1:]
    if not compact.isdigit():
        return False
    return 7 <= len(compact) <= 16


def list_campaign_summaries(limit=None, query_text=""):
    schema = get_db_schema()
    db_conn = get_db_connection()
    try:
        ensure_message_log_tables(db_conn)
        q = str(query_text or "").strip()
        q_like = f"%{q}%"
        with db_conn.cursor() as cur:
            query = """
                    WITH sent AS (
                        SELECT
                            LPAD(COALESCE(campaign_id::text, ''), 6, '0') AS campaign_id,
                            MAX(NULLIF(BTRIM(campaign_name), '')) AS campaign_name,
                            COUNT(*) AS sent_rows,
                            COUNT(DISTINCT phone) AS sent_total,
                            MIN(sent_at) AS campaign_start,
                            MAX(sent_at) AS max_sent_at
                        FROM {schema_name}.sms_sent
                        WHERE COALESCE(campaign_id::text, '') <> ''
                        GROUP BY 1
                    ),
                    failed AS (
                        SELECT
                            LPAD(COALESCE(campaign_id::text, ''), 6, '0') AS campaign_id,
                            MAX(NULLIF(BTRIM(campaign_name), '')) AS campaign_name,
                            COUNT(*) AS failed_rows,
                            COUNT(DISTINCT phone) AS failed_total,
                            MIN(failed_at) AS failed_start,
                            MAX(failed_at) AS max_failed_at
                        FROM {schema_name}.sms_not_sent
                        WHERE COALESCE(campaign_id::text, '') <> ''
                        GROUP BY 1
                    )
                    SELECT
                        COALESCE(sent.campaign_id, failed.campaign_id) AS campaign_id,
                        COALESCE(sent.campaign_name, failed.campaign_name, '') AS campaign_name,
                        COALESCE(sent.sent_total, 0) AS sent_total,
                        COALESCE(failed.failed_total, 0) AS failed_total,
                        COALESCE(sent.sent_rows, 0) AS sent_rows,
                        COALESCE(failed.failed_rows, 0) AS failed_rows,
                        COALESCE(sent.sent_total, 0) + COALESCE(failed.failed_total, 0) AS contacts_total,
                        sent.campaign_start,
                        failed.failed_start,
                        COALESCE(
                            GREATEST(sent.max_sent_at, failed.max_failed_at),
                            sent.max_sent_at,
                            failed.max_failed_at
                        ) AS campaign_end
                    FROM sent
                    FULL OUTER JOIN failed
                      ON sent.campaign_id = failed.campaign_id
                    WHERE (
                        %s = ''
                        OR COALESCE(sent.campaign_id, failed.campaign_id) ILIKE %s
                        OR COALESCE(sent.campaign_name, failed.campaign_name, '') ILIKE %s
                        OR (
                            COALESCE(sent.campaign_id, failed.campaign_id)
                            || '_'
                            || COALESCE(sent.campaign_name, failed.campaign_name, '')
                        ) ILIKE %s
                    )
                    ORDER BY campaign_end DESC NULLS LAST, campaign_id DESC
                    """
            params = [q, q_like, q_like, q_like]
            if limit is not None:
                query += "\nLIMIT %s"
                params.append(int(limit))
            cur.execute(
                sql.SQL(query).format(schema_name=sql.Identifier(schema)),
                tuple(params),
            )
            rows = cur.fetchall() or []

        campaigns = []
        for row in rows:
            sent_start = row[7]
            failed_start = row[8]
            if sent_start and failed_start:
                campaign_start = min(sent_start, failed_start)
            else:
                campaign_start = sent_start or failed_start
            campaigns.append(
                {
                    "campaign_id": str(row[0] or "").zfill(6),
                    "campaign_name": str(row[1] or "").strip(),
                    "sent_total": _safe_int(row[2]),
                    "failed_total": _safe_int(row[3]),
                    "sent_rows": _safe_int(row[4]),
                    "failed_rows": _safe_int(row[5]),
                    "contacts_total": _safe_int(row[6]),
                    "campaign_start": _json_datetime(campaign_start),
                    "campaign_end": _json_datetime(row[9]),
                }
            )
        return campaigns
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def get_campaign_detail(campaign_id, list_limit=None):
    normalized = _normalize_campaign_id(campaign_id)
    if not normalized:
        return None

    schema = get_db_schema()
    db_conn = get_db_connection()
    try:
        ensure_message_log_tables(db_conn)
        with db_conn.cursor() as cur:
            sent_query = """
                    SELECT phone, sent_at, message, content_kind, attachment_names, attachment_count, content_summary, campaign_name
                    FROM {schema_name}.sms_sent
                    WHERE LPAD(COALESCE(campaign_id::text, ''), 6, '0') = %s
                    ORDER BY sent_at ASC, id ASC
                    """
            sent_params = [normalized]
            if list_limit is not None:
                sent_query += "\nLIMIT %s"
                sent_params.append(int(list_limit))
            cur.execute(
                sql.SQL(sent_query).format(schema_name=sql.Identifier(schema)),
                tuple(sent_params),
            )
            sent_rows = cur.fetchall() or []

            failed_query = """
                    SELECT phone, failed_at, reason, campaign_name
                    FROM {schema_name}.sms_not_sent
                    WHERE LPAD(COALESCE(campaign_id::text, ''), 6, '0') = %s
                    ORDER BY failed_at ASC, id ASC
                    """
            failed_params = [normalized]
            if list_limit is not None:
                failed_query += "\nLIMIT %s"
                failed_params.append(int(list_limit))
            cur.execute(
                sql.SQL(failed_query).format(schema_name=sql.Identifier(schema)),
                tuple(failed_params),
            )
            failed_rows = cur.fetchall() or []

        sent_numbers = []
        sent_seen = set()
        sent_times = []
        kinds_seen = set()
        attachments_seen = set()
        message_samples = []
        campaign_name = ""
        for phone, sent_at, message, content_kind, attachment_names_raw, attachment_count, content_summary, row_campaign_name in sent_rows:
            if not campaign_name:
                campaign_name = str(row_campaign_name or "").strip()
            number = str(phone or "").strip()
            if number and number not in sent_seen:
                sent_numbers.append(number)
                sent_seen.add(number)
            if sent_at:
                sent_times.append(sent_at)
            ck = str(content_kind or "").strip().lower()
            if ck:
                kinds_seen.add(ck)
            if message:
                msg = str(message).strip()
                if msg and msg not in message_samples:
                    message_samples.append(msg)
            elif content_summary:
                msg = str(content_summary).strip()
                if msg and msg not in message_samples:
                    message_samples.append(msg)
            if attachment_names_raw:
                parsed_names = []
                try:
                    loaded = json.loads(attachment_names_raw) if isinstance(attachment_names_raw, str) else attachment_names_raw
                    if isinstance(loaded, list):
                        parsed_names = loaded
                except Exception:
                    parsed_names = []
                for item in parsed_names:
                    name = str(item or "").strip()
                    if name:
                        attachments_seen.add(name)
                if not parsed_names and int(attachment_count or 0) > 0:
                    attachments_seen.add(f"{int(attachment_count)} file(s)")

        failed_numbers = []
        failed_items = []
        failed_seen = set()
        failed_times = []
        for phone, failed_at, reason, row_campaign_name in failed_rows:
            if not campaign_name:
                campaign_name = str(row_campaign_name or "").strip()
            number = str(phone or "").strip()
            if number and number not in failed_seen:
                failed_numbers.append(number)
                failed_seen.add(number)
            if number:
                failed_items.append(
                    {
                        "phone": number,
                        "failed_at": _json_datetime(failed_at),
                        "reason": str(reason or ""),
                    }
                )
            if failed_at:
                failed_times.append(failed_at)

        all_numbers = set(sent_numbers) | set(failed_numbers)
        targets = []
        seen_targets = set()
        for number in sent_numbers + failed_numbers:
            if number and number not in seen_targets:
                targets.append(number)
                seen_targets.add(number)

        phone_like = 0
        non_phone_like = 0
        for target in targets:
            if _looks_like_phone_target(target):
                phone_like += 1
            else:
                non_phone_like += 1
        if phone_like and not non_phone_like:
            send_target_hint = "contacts"
        elif non_phone_like and not phone_like:
            send_target_hint = "groups"
        elif non_phone_like >= phone_like and non_phone_like > 0:
            send_target_hint = "groups"
        else:
            send_target_hint = "contacts"

        campaign_start = min(sent_times + failed_times) if (sent_times or failed_times) else None
        combined_times = sent_times + failed_times
        campaign_end = max(combined_times) if combined_times else None
        attachments_list = sorted(attachments_seen)
        kinds_list = sorted(kinds_seen)
        if attachments_list:
            content_mode = "files"
        elif kinds_list and all(k in ("text", "sms") for k in kinds_list):
            content_mode = "sms"
        elif not sent_rows:
            content_mode = "unknown"
        else:
            content_mode = "mixed"

        return {
            "campaign_id": normalized,
            "campaign_name": campaign_name,
            "contacts_total": len(all_numbers),
            "sent_total": len(sent_numbers),
            "failed_total": len(failed_numbers),
            "sent_rows": len(sent_rows),
            "failed_rows": len(failed_rows),
            "campaign_start": _json_datetime(campaign_start),
            "campaign_end": _json_datetime(campaign_end),
            "sent_numbers": sent_numbers,
            "failed_numbers": failed_numbers,
            "failed_items": failed_items,
            "targets": targets,
            "send_target_hint": send_target_hint,
            "campaign_content": {
                "mode": content_mode,
                "content_kinds": kinds_list,
                "messages": message_samples[:5],
                "attachments": attachments_list,
            },
        }
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def ensure_group_metadata_table(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.group_metadata (
                    group_name TEXT PRIMARY KEY,
                    description TEXT,
                    members_label TEXT,
                    members_count INTEGER,
                    created_by TEXT,
                    created_on TEXT,
                    invite_link TEXT,
                    source_scrap_id TEXT,
                    last_scraped_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS description TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS members_label TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS members_count INTEGER'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS created_by TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS created_on TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS invite_link TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS source_scrap_id TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS last_scraped_at TIMESTAMPTZ NOT NULL DEFAULT NOW()'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.group_metadata ADD COLUMN IF NOT EXISTS updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()'
            ).format(schema_name=sql.Identifier(schema))
        )


def _normalize_scrape_group_metadata(group_name, metadata=None):
    payload = dict(metadata or {})
    safe_group = _normalize_group_name(payload.get("group_name") or group_name or "group")
    members_count = payload.get("members_count")
    if str(members_count or "").strip():
        members_count = _safe_int(members_count, default=0)
    else:
        members_count = None
    return {
        "group_name": safe_group or "group",
        "description": str(payload.get("description") or "").strip(),
        "members_label": str(payload.get("members_label") or "").strip(),
        "members_count": members_count,
        "created_by": str(payload.get("created_by") or "").strip(),
        "created_on": str(payload.get("created_on") or "").strip(),
        "invite_link": str(payload.get("invite_link") or "").strip(),
        "source_scrap_id": str(payload.get("source_scrap_id") or "").strip(),
        "last_scraped_at": str(payload.get("last_scraped_at") or "").strip(),
    }


def _get_all_contacts_columns(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = %s
              AND table_name = 'all_contacts'
            """,
            (schema,),
        )
        return {row[0] for row in cur.fetchall() if row and row[0]}


def ensure_all_contacts_table(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.all_contacts (
                    "Phone" TEXT,
                    "whatspp_group" TEXT,
                    "time" TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.all_contacts ADD COLUMN IF NOT EXISTS "Phone" TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.all_contacts ADD COLUMN IF NOT EXISTS "whatspp_group" TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.all_contacts ADD COLUMN IF NOT EXISTS "time" TIMESTAMPTZ NOT NULL DEFAULT NOW()'
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                'ALTER TABLE {schema_name}.all_contacts ADD COLUMN IF NOT EXISTS "whatsapp_group" TEXT'
            ).format(schema_name=sql.Identifier(schema))
        )

        existing_cols = _get_all_contacts_columns(db_conn)
        has_legacy_whatsapp_group = "whatsapp_group" in existing_cols

        if has_legacy_whatsapp_group:
            cur.execute(
                sql.SQL(
                    """
                    UPDATE {schema_name}.all_contacts
                    SET "whatspp_group" = COALESCE(NULLIF(TRIM("whatspp_group"), ''), NULLIF(TRIM("whatsapp_group"), ''))
                    WHERE COALESCE(NULLIF(TRIM("whatspp_group"), ''), '') = ''
                      AND COALESCE(NULLIF(TRIM("whatsapp_group"), ''), '') <> ''
                    """
                ).format(schema_name=sql.Identifier(schema))
            )
            cur.execute(
                sql.SQL(
                    """
                    UPDATE {schema_name}.all_contacts
                    SET "whatsapp_group" = COALESCE(NULLIF(TRIM("whatsapp_group"), ''), NULLIF(TRIM("whatspp_group"), ''))
                    WHERE COALESCE(NULLIF(TRIM("whatsapp_group"), ''), '') = ''
                      AND COALESCE(NULLIF(TRIM("whatspp_group"), ''), '') <> ''
                    """
                ).format(schema_name=sql.Identifier(schema))
            )

        cur.execute(
            """
            SELECT tc.constraint_name
            FROM information_schema.table_constraints tc
            JOIN information_schema.key_column_usage kcu
              ON tc.constraint_name = kcu.constraint_name
             AND tc.table_schema = kcu.table_schema
            WHERE tc.table_schema = %s
              AND tc.table_name = 'all_contacts'
              AND tc.constraint_type = 'UNIQUE'
            GROUP BY tc.constraint_name
            HAVING COUNT(*) = 1
               AND MIN(kcu.column_name) = 'Phone'
               AND MAX(kcu.column_name) = 'Phone'
            """,
            (schema,),
        )
        for row in cur.fetchall():
            constraint_name = str(row[0] or "").strip()
            if not constraint_name:
                continue
            cur.execute(
                sql.SQL(
                    'ALTER TABLE {schema_name}.all_contacts DROP CONSTRAINT IF EXISTS {constraint_name}'
                ).format(
                    schema_name=sql.Identifier(schema),
                    constraint_name=sql.Identifier(constraint_name),
                )
            )

        cur.execute(
            """
            SELECT indexname, indexdef
            FROM pg_indexes
            WHERE schemaname = %s
              AND tablename = 'all_contacts'
            """,
            (schema,),
        )
        for index_name, index_def in cur.fetchall():
            normalized_def = " ".join(str(index_def or "").split())
            if "CREATE UNIQUE INDEX" not in normalized_def.upper():
                continue
            if '("Phone")' not in normalized_def:
                continue
            if '("Phone", ' in normalized_def or ', "Phone")' in normalized_def:
                continue
            if str(index_name or "").strip() == "all_contacts_pkey":
                continue
            cur.execute(
                sql.SQL(
                    "DROP INDEX IF EXISTS {schema_name}.{index_name}"
                ).format(
                    schema_name=sql.Identifier(schema),
                    index_name=sql.Identifier(index_name),
                )
            )

        cur.execute(
            sql.SQL(
                """
                CREATE UNIQUE INDEX IF NOT EXISTS ux_all_contacts_phone_group
                ON {schema_name}.all_contacts ("Phone", "whatspp_group")
                """
            ).format(schema_name=sql.Identifier(schema))
        )

    return has_legacy_whatsapp_group


def save_scraped_group_contacts(db_conn, group_name, contacts, metadata=None):
    schema = get_db_schema()
    normalized = _normalize_scrape_group_metadata(group_name, metadata)
    safe_group = normalized["group_name"]
    unique_contacts = []
    seen_contacts = set()
    for raw_contact in contacts or []:
        number = str(raw_contact or "").strip()
        if not number or number in seen_contacts:
            continue
        seen_contacts.add(number)
        unique_contacts.append(number)

    ensure_group_metadata_table(db_conn)
    has_legacy_whatsapp_group = ensure_all_contacts_table(db_conn)

    with db_conn.cursor() as cur:
        for number in unique_contacts:
            if has_legacy_whatsapp_group:
                cur.execute(
                    sql.SQL(
                        """
                        INSERT INTO {schema_name}.all_contacts ("Phone", "whatspp_group", "whatsapp_group", "time")
                        VALUES (%s, %s, %s, NOW())
                        ON CONFLICT ("Phone", "whatspp_group") DO NOTHING
                        """
                    ).format(schema_name=sql.Identifier(schema)),
                    (number, safe_group, safe_group),
                )
            else:
                cur.execute(
                    sql.SQL(
                        """
                        INSERT INTO {schema_name}.all_contacts ("Phone", "whatspp_group", "time")
                        VALUES (%s, %s, NOW())
                        ON CONFLICT ("Phone", "whatspp_group") DO NOTHING
                        """
                    ).format(schema_name=sql.Identifier(schema)),
                    (number, safe_group),
                )

        cur.execute(
            sql.SQL(
                """
                INSERT INTO {schema_name}.group_metadata (
                    group_name,
                    description,
                    members_label,
                    members_count,
                    created_by,
                    created_on,
                    invite_link,
                    source_scrap_id,
                    last_scraped_at,
                    updated_at
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, COALESCE(%s::timestamptz, NOW()), NOW())
                ON CONFLICT (group_name)
                DO UPDATE SET
                    description = EXCLUDED.description,
                    members_label = EXCLUDED.members_label,
                    members_count = EXCLUDED.members_count,
                    created_by = EXCLUDED.created_by,
                    created_on = EXCLUDED.created_on,
                    invite_link = EXCLUDED.invite_link,
                    source_scrap_id = EXCLUDED.source_scrap_id,
                    last_scraped_at = EXCLUDED.last_scraped_at,
                    updated_at = NOW()
                """
            ).format(schema_name=sql.Identifier(schema)),
            (
                safe_group,
                normalized["description"] or None,
                normalized["members_label"] or None,
                normalized["members_count"],
                normalized["created_by"] or None,
                normalized["created_on"] or None,
                normalized["invite_link"] or None,
                normalized["source_scrap_id"] or None,
                normalized["last_scraped_at"] or datetime.now().isoformat(),
            ),
        )

    return "all_contacts"


def get_all_contacts_group_column(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = %s
              AND table_name = 'all_contacts'
            """,
            (schema,),
        )
        columns = {row[0] for row in cur.fetchall() if row and row[0]}
    for name in GROUP_COLUMN_CANDIDATES:
        if name in columns:
            return name
    return None


def load_existing_group_names_from_all_contacts(db_conn):
    schema = get_db_schema()
    ensure_group_metadata_table(db_conn)
    ensure_all_contacts_table(db_conn)

    group_col = get_all_contacts_group_column(db_conn)
    if not group_col:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT group_name
                    FROM {schema_name}.group_metadata
                    WHERE group_name IS NOT NULL
                      AND TRIM(group_name) <> ''
                    """
                ).format(schema_name=sql.Identifier(schema))
            )
            return {
                _normalize_group_key(row[0])
                for row in cur.fetchall()
                if row and _normalize_group_key(row[0])
            }

    incomplete_groups = set()
    groups_to_skip = set()

    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                WITH contact_counts AS (
                    SELECT
                        LOWER(TRIM({group_col})) AS group_key,
                        COUNT(DISTINCT "Phone") AS contacts_count
                    FROM {schema_name}.all_contacts
                    WHERE {group_col} IS NOT NULL
                      AND TRIM({group_col}) <> ''
                      AND "Phone" IS NOT NULL
                      AND TRIM("Phone") <> ''
                    GROUP BY LOWER(TRIM({group_col}))
                )
                SELECT
                    gm.group_name,
                    gm.members_count,
                    COALESCE(cc.contacts_count, 0) AS contacts_count
                FROM {schema_name}.group_metadata gm
                LEFT JOIN contact_counts cc
                  ON cc.group_key = LOWER(TRIM(gm.group_name))
                WHERE gm.group_name IS NOT NULL
                  AND TRIM(gm.group_name) <> ''
                """
            ).format(
                schema_name=sql.Identifier(schema),
                group_col=sql.Identifier(group_col),
            )
        )
        for row in cur.fetchall():
            if not row:
                continue
            group_key = _normalize_group_key(row[0])
            if not group_key:
                continue
            expected_count = _safe_int(row[1], default=0) if row[1] is not None else 0
            stored_count = _safe_int(row[2], default=0)
            if expected_count > 0 and stored_count < expected_count:
                incomplete_groups.add(group_key)
                continue
            groups_to_skip.add(group_key)

    try:
        with db_conn.cursor() as cur:
            cur.execute(
                """
                SELECT EXISTS (
                    SELECT 1
                    FROM information_schema.tables
                    WHERE table_schema = %s
                      AND table_name = 'scrape_group_runs'
                )
                """,
                (schema,),
            )
            has_scrape_group_runs = bool(cur.fetchone()[0])
        if has_scrape_group_runs:
            with db_conn.cursor() as cur:
                cur.execute(
                    sql.SQL(
                        """
                        WITH db_counts AS (
                            SELECT
                                LOWER(TRIM({group_col})) AS group_key,
                                COUNT(DISTINCT "Phone") AS contacts_count
                            FROM {schema_name}.all_contacts
                            WHERE {group_col} IS NOT NULL
                              AND TRIM({group_col}) <> ''
                              AND "Phone" IS NOT NULL
                              AND TRIM("Phone") <> ''
                            GROUP BY LOWER(TRIM({group_col}))
                        ),
                        scrape_counts AS (
                            SELECT
                                LOWER(TRIM(group_name)) AS group_key,
                                MAX(scrapped_contacts_count) AS scrapped_contacts_count
                            FROM {schema_name}.scrape_group_runs
                            WHERE group_name IS NOT NULL
                              AND TRIM(group_name) <> ''
                            GROUP BY LOWER(TRIM(group_name))
                        )
                        SELECT
                            sc.group_key,
                            sc.scrapped_contacts_count,
                            COALESCE(dc.contacts_count, 0) AS contacts_count
                        FROM scrape_counts sc
                        LEFT JOIN db_counts dc
                          ON dc.group_key = sc.group_key
                        WHERE sc.scrapped_contacts_count > COALESCE(dc.contacts_count, 0)
                        """
                    ).format(
                        schema_name=sql.Identifier(schema),
                        group_col=sql.Identifier(group_col),
                    )
                )
                for row in cur.fetchall():
                    if not row:
                        continue
                    group_key = _normalize_group_key(row[0])
                    if not group_key:
                        continue
                    incomplete_groups.add(group_key)
                    groups_to_skip.discard(group_key)
    except Exception:
        pass

    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                SELECT DISTINCT TRIM({group_col}) AS group_name
                FROM {schema_name}.all_contacts
                WHERE {group_col} IS NOT NULL
                  AND TRIM({group_col}) <> ''
                """
            ).format(
                schema_name=sql.Identifier(schema),
                group_col=sql.Identifier(group_col),
            )
        )
        all_contacts_groups = {
            _normalize_group_key(row[0])
            for row in cur.fetchall()
            if row and _normalize_group_key(row[0]) and _normalize_group_key(row[0]) not in incomplete_groups
        }
    return groups_to_skip | all_contacts_groups


def list_unique_groups_from_all_contacts(limit=None, query_text=""):
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        ensure_group_metadata_table(db_conn)
        group_col = get_all_contacts_group_column(db_conn)
        if not group_col:
            return []
        q = _normalize_group_name(query_text).lower()
        with db_conn.cursor() as cur:
            query = """
                    SELECT
                        src.group_name,
                        src.contacts_count,
                        gm.members_label,
                        gm.members_count,
                        gm.description,
                        gm.created_by,
                        gm.created_on,
                        gm.invite_link,
                        gm.source_scrap_id,
                        gm.last_scraped_at
                    FROM (
                        SELECT
                            MIN(TRIM({group_col})) AS group_name,
                            COUNT(DISTINCT "Phone") AS contacts_count
                        FROM {schema_name}.all_contacts
                        WHERE {group_col} IS NOT NULL
                          AND TRIM({group_col}) <> ''
                          AND LOWER(TRIM({group_col})) LIKE %s
                        GROUP BY LOWER(TRIM({group_col}))
                    ) src
                    LEFT JOIN {schema_name}.group_metadata gm
                      ON LOWER(TRIM(gm.group_name)) = LOWER(TRIM(src.group_name))
                    ORDER BY src.group_name ASC
                    """
            params = [f"%{q}%"]
            if limit is not None:
                query += "\nLIMIT %s"
                params.append(int(limit))
            cur.execute(
                sql.SQL(query).format(
                    schema_name=sql.Identifier(schema),
                    group_col=sql.Identifier(group_col),
                ),
                tuple(params),
            )
            rows = cur.fetchall()
            results = []
            for row in rows:
                if not row or not row[0]:
                    continue
                results.append(
                    {
                        "group_name": row[0],
                        "contacts_count": _safe_int(row[1]),
                        "members_label": row[2] or "",
                        "members_count": _safe_int(row[3], default=0) if row[3] is not None else None,
                        "description": row[4] or "",
                        "created_by": row[5] or "",
                        "created_on": row[6] or "",
                        "invite_link": row[7] or "",
                        "source_scrap_id": row[8] or "",
                        "last_scraped_at": row[9].isoformat() if row[9] else "",
                    }
                )
            return results
    finally:
        db_conn.close()


def load_group_contacts_from_all_contacts(group_name, limit=None):
    normalized = _normalize_group_name(group_name)
    if not normalized:
        return []
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        group_col = get_all_contacts_group_column(db_conn)
        if not group_col:
            return []
        with db_conn.cursor() as cur:
            query = """
                    SELECT
                        "Phone",
                        MIN("time") AS first_seen
                    FROM {schema_name}.all_contacts
                    WHERE LOWER(TRIM({group_col})) = LOWER(TRIM(%s))
                      AND "Phone" IS NOT NULL
                      AND TRIM("Phone") <> ''
                    GROUP BY "Phone"
                    ORDER BY first_seen ASC, "Phone" ASC
                    """
            params = [normalized]
            if limit is not None:
                query += "\nLIMIT %s"
                params.append(int(limit))
            cur.execute(
                sql.SQL(query).format(
                    schema_name=sql.Identifier(schema),
                    group_col=sql.Identifier(group_col),
                ),
                tuple(params),
            )
            return [str(row[0]).strip() for row in cur.fetchall() if row and row[0]]
    finally:
        db_conn.close()


def load_contacts_from_multiple_groups(group_names, limit=None):
    cleaned_groups = []
    seen = set()
    for name in group_names or []:
        value = _normalize_group_name(name)
        key = value.lower()
        if not value or key in seen:
            continue
        seen.add(key)
        cleaned_groups.append(value)
    if not cleaned_groups:
        return {"groups": [], "contacts": [], "contact_items": []}

    summaries = list_unique_groups_from_all_contacts(limit=None, query_text="")
    summary_map = {
        _normalize_group_name(item.get("group_name")).lower(): item
        for item in summaries
        if _normalize_group_name(item.get("group_name"))
    }

    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        group_col = get_all_contacts_group_column(db_conn)
        if not group_col:
            return {"groups": [], "contacts": [], "contact_items": []}
        with db_conn.cursor() as cur:
            query = """
                    SELECT
                        MIN(TRIM({group_col})) AS group_name,
                        "Phone",
                        MIN("time") AS first_seen
                    FROM {schema_name}.all_contacts
                    WHERE LOWER(TRIM({group_col})) = ANY(%s)
                      AND "Phone" IS NOT NULL
                      AND TRIM("Phone") <> ''
                    GROUP BY LOWER(TRIM({group_col})), "Phone"
                    ORDER BY MIN(TRIM({group_col})) ASC, first_seen ASC, "Phone" ASC
                    """
            params = [[g.lower() for g in cleaned_groups]]
            if limit is not None:
                query += "\nLIMIT %s"
                params.append(int(limit))
            cur.execute(
                sql.SQL(query).format(
                    schema_name=sql.Identifier(schema),
                    group_col=sql.Identifier(group_col),
                ),
                tuple(params),
            )
            rows = cur.fetchall()
    finally:
        db_conn.close()

    grouped = {}
    merged_contacts = []
    merged_row_ids = {}
    merged_seen = set()
    for row in rows:
        if not row:
            continue
        group_name = _normalize_group_name(row[0])
        phone = str(row[1] or "").strip()
        if not group_name or not phone:
            continue
        grouped.setdefault(group_name, [])
        grouped[group_name].append(phone)
        if phone not in merged_seen:
            merged_seen.add(phone)
            merged_contacts.append(phone)
            merged_row_ids[phone] = len(merged_contacts)

    selected_groups = []
    for group_name in cleaned_groups:
        key = group_name.lower()
        group_contacts = grouped.get(group_name, [])
        summary = summary_map.get(key) or {"group_name": group_name, "contacts_count": len(group_contacts)}
        selected_groups.append(
            {
                **summary,
                "group_name": summary.get("group_name") or group_name,
                "selected_contacts_count": len(group_contacts),
            }
        )

    # Keep DB sequence as loaded from query; do not sort alphabetically.
    return {
        "groups": selected_groups,
        "contacts": merged_contacts,
        "contact_items": build_contact_items(merged_contacts, {}, row_id_map=merged_row_ids),
    }


def get_group_detail_from_all_contacts(group_name, contact_limit=None):
    normalized = _normalize_group_name(group_name)
    if not normalized:
        return None
    summaries = list_unique_groups_from_all_contacts(limit=None, query_text=normalized)
    summary = None
    for item in summaries:
        if _normalize_group_name(item.get("group_name")).lower() == normalized.lower():
            summary = item
            break
    if summary is None:
        return None
    contacts = load_group_contacts_from_all_contacts(normalized, limit=contact_limit)
    row_id_map = {number: idx + 1 for idx, number in enumerate(contacts)}
    return {
        **summary,
        "contacts": contacts,
        "contact_items": build_contact_items(contacts, {}, row_id_map=row_id_map),
    }


def list_group_tables():
    print("[DB] Loading tables: Connecting to database...")
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            print("[DB] Loading tables: Accessing information_schema...")
            cur.execute(
                """
                SELECT table_name
                FROM information_schema.tables
                WHERE table_schema = %s
                  AND table_type = 'BASE TABLE'
                ORDER BY table_name
                """,
                (schema,),
            )
            table_rows = cur.fetchall()
            tables = [
                row[0] for row in table_rows if row and len(row) > 0 and row[0]
            ]
            print(f"[DB] Loading tables: Loaded {len(tables)} tables from schema '{schema}'.")
            return tables
    except Exception as exc:
        print(f"[DB ERROR] list_group_tables failed: {exc}")
        traceback.print_exc()
        raise
    finally:
        db_conn.close()


def parse_selected_tables(raw_selected):
    if not raw_selected:
        return []
    if isinstance(raw_selected, list):
        values = raw_selected
        cleaned = []
        for item in values:
            text = str(item).strip()
            if text:
                cleaned.append(text)
        return list(dict.fromkeys(cleaned))
    try:
        parsed = json.loads(raw_selected)
        if isinstance(parsed, list):
            values = parsed
        else:
            values = []
    except Exception:
        values = str(raw_selected).split(",")
    cleaned = []
    for item in values:
        text = str(item).strip()
        if not text:
            continue
        cleaned.append(text)
    return list(dict.fromkeys(cleaned))


def _parse_row_bounds(row_start=None, row_end=None):
    try:
        start_idx = int(row_start) if row_start is not None else None
    except Exception:
        start_idx = None
    try:
        end_idx = int(row_end) if row_end is not None else None
    except Exception:
        end_idx = None
    if start_idx is not None and start_idx < 1:
        start_idx = 1
    if end_idx is not None and end_idx < 1:
        end_idx = 1
    if start_idx is not None and end_idx is not None and end_idx < start_idx:
        start_idx, end_idx = end_idx, start_idx
    return start_idx, end_idx


def _resolve_phone_column(cur, schema, table_name):
    cur.execute(
        """
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = %s
          AND table_name = %s
        """,
        (schema, table_name),
    )
    for col_row in cur.fetchall():
        if not col_row or len(col_row) == 0:
            continue
        col_name = col_row[0]
        if normalize_column_name(col_name) in CONTACT_COLUMN_ALIASES_NORMALIZED:
            return col_name
    return None


def get_max_row_number_for_tables(selected_tables):
    if not selected_tables:
        return 0
    if psycopg2 is None or sql is None:
        raise RuntimeError(
            "psycopg2 is not installed. Install it with: pip install psycopg2-binary"
        )
    allowed_tables = set(list_group_tables())
    chosen = [t for t in selected_tables if t in allowed_tables]
    if not chosen:
        return 0

    max_rows = 0
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            for table_name in chosen:
                phone_col = _resolve_phone_column(cur, schema, table_name)
                if not phone_col:
                    continue
                cur.execute(
                    sql.SQL(
                        """
                        SELECT COUNT(*)
                        FROM {schema_name}.{table_name}
                        WHERE {phone_col} IS NOT NULL
                          AND TRIM(CAST({phone_col} AS TEXT)) <> ''
                        """
                    ).format(
                        schema_name=sql.Identifier(schema),
                        table_name=sql.Identifier(table_name),
                        phone_col=sql.Identifier(phone_col),
                    )
                )
                row = cur.fetchone()
                count = int(row[0]) if row and row[0] is not None else 0
                if count > max_rows:
                    max_rows = count
    finally:
        db_conn.close()

    return max_rows


def load_numbers_from_tables(selected_tables, row_start=None, row_end=None, include_row_ids=False):
    if not selected_tables:
        return []
    if psycopg2 is None or sql is None:
        raise RuntimeError(
            "psycopg2 is not installed. Install it with: pip install psycopg2-binary"
        )
    allowed_tables = set(list_group_tables())
    chosen = [t for t in selected_tables if t in allowed_tables]
    if not chosen:
        raise RuntimeError("No valid group_ tables selected.")
    print(f"[DB] Sending flow: Using {len(chosen)} selected tables.")

    start_idx, end_idx = _parse_row_bounds(row_start=row_start, row_end=row_end)

    db_conn = get_db_connection()
    schema = get_db_schema()
    numbers = []
    seen_numbers = set()
    row_id_map = {}
    try:
        with db_conn.cursor() as cur:
            for table_name in chosen:
                print(f"[DB] Reading table: {table_name}")
                phone_col = _resolve_phone_column(cur, schema, table_name)
                if not phone_col:
                    continue
                if start_idx is not None or end_idx is not None:
                    cur.execute(
                        sql.SQL(
                            """
                            SELECT phone_val, rn
                            FROM (
                                SELECT
                                    {phone_col} AS phone_val,
                                    ROW_NUMBER() OVER (ORDER BY ctid) AS rn
                                FROM {schema_name}.{table_name}
                                WHERE {phone_col} IS NOT NULL
                            ) rows_q
                            WHERE (%s IS NULL OR rn >= %s)
                              AND (%s IS NULL OR rn <= %s)
                            ORDER BY rn ASC
                            """
                        ).format(
                            phone_col=sql.Identifier(phone_col),
                            schema_name=sql.Identifier(schema),
                            table_name=sql.Identifier(table_name),
                        ),
                        (start_idx, start_idx, end_idx, end_idx),
                    )
                else:
                    if include_row_ids:
                        cur.execute(
                            sql.SQL(
                                """
                                SELECT phone_val, rn
                                FROM (
                                    SELECT
                                        {phone_col} AS phone_val,
                                        ROW_NUMBER() OVER (ORDER BY ctid) AS rn
                                    FROM {schema_name}.{table_name}
                                    WHERE {phone_col} IS NOT NULL
                                ) rows_q
                                ORDER BY rn ASC
                                """
                            ).format(
                                phone_col=sql.Identifier(phone_col),
                                schema_name=sql.Identifier(schema),
                                table_name=sql.Identifier(table_name),
                            )
                        )
                    else:
                        cur.execute(
                            sql.SQL(
                                """
                                SELECT {phone_col}
                                FROM {schema_name}.{table_name}
                                WHERE {phone_col} IS NOT NULL
                                """
                            ).format(
                                phone_col=sql.Identifier(phone_col),
                                schema_name=sql.Identifier(schema),
                                table_name=sql.Identifier(table_name),
                            )
                        )
                for row in cur.fetchall():
                    raw = row[0] if row and len(row) > 0 else None
                    rn = int(row[1]) if row and len(row) > 1 and row[1] is not None else None
                    if raw is None:
                        continue
                    value = str(raw).strip()
                    if not value:
                        continue
                    if value in seen_numbers:
                        continue
                    seen_numbers.add(value)
                    numbers.append(value)
                    if include_row_ids and rn is not None:
                        row_id_map[value] = rn
    finally:
        db_conn.close()

    if include_row_ids:
        return numbers, row_id_map
    return numbers


def get_tables_info(selected_tables):
    if not selected_tables:
        return []
    if psycopg2 is None or sql is None:
        raise RuntimeError(
            "psycopg2 is not installed. Install it with: pip install psycopg2-binary"
        )

    allowed_tables = set(list_group_tables())
    chosen = [t for t in selected_tables if t in allowed_tables]
    if not chosen:
        return []

    print(f"[DB] Table info: Loading metadata for {len(chosen)} table(s)...")
    db_conn = get_db_connection()
    schema = get_db_schema()
    summaries = []
    try:
        with db_conn.cursor() as cur:
            for table_name in chosen:
                cur.execute(
                    """
                    SELECT column_name
                    FROM information_schema.columns
                    WHERE table_schema = %s
                      AND table_name = %s
                    ORDER BY ordinal_position
                    """,
                    (schema, table_name),
                )
                column_rows = cur.fetchall()
                columns = [row[0] for row in column_rows if row and len(row) > 0 and row[0]]

                phone_col = None
                for c in columns:
                    if normalize_column_name(c) in CONTACT_COLUMN_ALIASES_NORMALIZED:
                        phone_col = c
                        break

                cur.execute(
                    sql.SQL(
                        """
                        SELECT COUNT(*)
                        FROM {schema_name}.{table_name}
                        """
                    ).format(
                        schema_name=sql.Identifier(schema),
                        table_name=sql.Identifier(table_name),
                    )
                )
                row_count_row = cur.fetchone()
                row_count = int(row_count_row[0]) if row_count_row and len(row_count_row) > 0 else 0

                phone_count = 0
                if phone_col:
                    cur.execute(
                        sql.SQL(
                            """
                            SELECT COUNT(DISTINCT {phone_col})
                            FROM {schema_name}.{table_name}
                            WHERE {phone_col} IS NOT NULL
                              AND TRIM(CAST({phone_col} AS TEXT)) <> ''
                            """
                        ).format(
                            schema_name=sql.Identifier(schema),
                            table_name=sql.Identifier(table_name),
                            phone_col=sql.Identifier(phone_col),
                        )
                    )
                    count_row = cur.fetchone()
                    phone_count = int(count_row[0]) if count_row and len(count_row) > 0 else 0

                summaries.append(
                    {
                        "table": table_name,
                        "columns": columns,
                        "columns_count": len(columns),
                        "row_count": row_count,
                        "phone_column": phone_col,
                        "phone_count": phone_count,
                    }
                )
    finally:
        db_conn.close()

    print("[DB] Table info: Loaded metadata.")
    return summaries


def detect_attachment_kind(filename, mimetype=""):
    mime = str(mimetype or "").strip().lower()
    if mime.startswith("image/"):
        return "image"
    if mime.startswith("video/"):
        return "video"
    if mime.startswith("audio/"):
        return "audio"
    ext = Path(str(filename or "")).suffix.lower()
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in VIDEO_EXTENSIONS:
        return "video"
    if ext in AUDIO_EXTENSIONS:
        return "audio"
    return "file"


def is_supported_attachment_filename(filename):
    ext = Path(str(filename or "")).suffix.lower()
    if not ext:
        return False
    return (
        ext in IMAGE_EXTENSIONS
        or ext in VIDEO_EXTENSIONS
        or ext in AUDIO_EXTENSIONS
        or ext in DOCUMENT_EXTENSIONS
    )


def save_uploaded_attachments(file_storages):
    uploads_dir = Path(__file__).parent / "send_uploads"
    uploads_dir.mkdir(parents=True, exist_ok=True)
    saved = []
    for file_storage in file_storages or []:
        if not file_storage or not file_storage.filename:
            continue
        original_name = secure_filename(file_storage.filename or "")
        if not original_name:
            continue
        if not is_supported_attachment_filename(original_name):
            print(f"[WARN] Skipping unsupported attachment extension: {original_name}")
            continue
        target_path = uploads_dir / original_name
        file_storage.save(target_path)
        saved.append(
            {
                "name": original_name,
                "path": str(target_path),
                "kind": detect_attachment_kind(original_name, getattr(file_storage, "mimetype", "")),
            }
        )
    return saved


def build_payload(message, attachments, message_variants=None):
    attachments = attachments or []
    normalized_variants = []
    for item in message_variants or []:
        if not isinstance(item, dict):
            continue
        text_value = str(item.get("text") or "")
        if not text_value.strip():
            continue
        try:
            slot = int(item.get("slot") or 0)
        except Exception:
            slot = 0
        if slot <= 0:
            slot = len(normalized_variants) + 1
        normalized_variants.append({"slot": slot, "text": text_value})

    normalized_attachments = []
    for item in attachments:
        normalized_attachments.append(
            {
                "name": item.get("name") or "",
                "kind": item.get("kind") or "file",
                "path": item.get("path") or "",
            }
        )
    kinds = {item.get("kind") or "file" for item in attachments}
    if not attachments:
        content_kind = "text"
    elif len(kinds) == 1:
        content_kind = next(iter(kinds))
    else:
        content_kind = "mixed"
    signature_payload = {
        "message": message or "",
        "attachments": normalized_attachments,
    }
    signature = json.dumps(signature_payload, ensure_ascii=False, sort_keys=True)
    dedupe_payload = {
        "message": message or "",
        "attachments": [
            {"name": item.get("name") or "", "kind": item.get("kind") or "file"}
            for item in normalized_attachments
        ],
    }
    dedupe_signature = json.dumps(dedupe_payload, ensure_ascii=False, sort_keys=True)
    content_summary = message or ""
    if attachments:
        names = ", ".join(item.get("name") or "" for item in attachments if item.get("name"))
        content_summary = (message + " | " if message else "") + names
    return {
        "message": message or "",
        "message_variants": normalized_variants,
        "attachments": normalized_attachments,
        "content_kind": content_kind,
        "attachment_count": len(attachments),
        "attachment_names": [item.get("name") or "" for item in normalized_attachments],
        "attachment_paths": [item.get("path") or "" for item in normalized_attachments],
        "payload_signature": signature,
        "dedupe_signature": dedupe_signature,
        "content_summary": content_summary[:1000],
    }


def _pick_random_message_for_target(payload, random_enabled=True, slot_min=1, slot_max=100):
    variants = list((payload or {}).get("message_variants") or [])
    slot_low = max(1, int(slot_min or 1))
    slot_high = max(slot_low, int(slot_max or slot_low))
    scoped_variants = []
    for item in variants:
        try:
            slot = int(item.get("slot") or 0)
        except Exception:
            slot = 0
        if slot_low <= slot <= slot_high:
            scoped_variants.append(item)
    if not scoped_variants:
        scoped_variants = variants
    if scoped_variants:
        if random_enabled:
            picked = random.choice(scoped_variants)
        else:
            picked = sorted(scoped_variants, key=lambda x: int(x.get("slot") or 0))[0]
        return int(picked.get("slot") or 1), str(picked.get("text") or ""), len(scoped_variants)
    return 1, str((payload or {}).get("message") or ""), 0


def _pick_random_attachment_for_target(payload, random_enabled=True, min_count=1, max_count=1):
    attachments = list((payload or {}).get("attachments") or [])
    total = len(attachments)
    if total == 0:
        return [], 0, 0, []
    low = max(1, int(min_count or 1))
    high = max(low, int(max_count or low))
    high = min(high, total)
    low = min(low, high)
    if random_enabled:
        pick_count = random.randint(low, high)
        selected_indexes = sorted(random.sample(list(range(total)), pick_count))
    else:
        pick_count = low
        selected_indexes = list(range(0, pick_count))
    picked = [attachments[idx] for idx in selected_indexes]
    index_labels = [idx + 1 for idx in selected_indexes]
    return picked, pick_count, total, index_labels


def _init_live_rows(targets, batch_numbers=None):
    rows = []
    for idx, target in enumerate(targets or [], start=1):
        batch_no = 1
        if batch_numbers and (idx - 1) < len(batch_numbers):
            try:
                batch_no = int(batch_numbers[idx - 1])
            except Exception:
                batch_no = 1
        planned = _planned_automation_action_for_target(idx - 1)
        rows.append(
            {
                "count": idx,
                "contact": str(target or "").strip(),
                "profile": "-",
                "batch": batch_no,
                "message": "-",
                "media": "-",
                "media_delay": "-",
                "send_delay": "-",
                "status": "pending",
                "action": str(planned.get("label") or "").strip() or "-",
                "action_id": str(planned.get("id") or "").strip(),
                "action_phase": str(planned.get("phase") or "after_send").strip() or "after_send",
                "action_status": "assigned",
                "action_detail": "",
                "action_performed_at": "",
                "action_checked": False,
            }
        )
    return rows


# Guards read-modify-write updates of shared campaign state while several
# WhatsApp sessions send in parallel threads.
_SEND_STATE_LOCK = threading.RLock()


def _update_live_row(state, row_index, **fields):
    try:
        idx = int(row_index)
    except Exception:
        return
    with _SEND_STATE_LOCK:
        rows = list((state or {}).get("live_rows") or [])
        if idx < 0 or idx >= len(rows):
            return
        row = dict(rows[idx] or {})
        for k, v in fields.items():
            row[k] = v
        rows[idx] = row
        state["live_rows"] = rows


def _action_definition_by_id(action_id):
    action_key = str(action_id or "").strip()
    if not action_key:
        return {}
    for item in CAMPAIGN_OPTIONAL_ACTIONS:
        if str((item or {}).get("id") or "").strip() == action_key:
            return dict(item or {})
    return {}


def _planned_automation_action_for_target(target_index):
    catalog = list(CAMPAIGN_OPTIONAL_ACTIONS or [])
    if not catalog:
        return {"id": "", "label": "", "phase": "after_send"}
    try:
        idx = max(0, int(target_index or 0))
    except Exception:
        idx = 0
    item = dict(catalog[idx % len(catalog)] or {})
    return {
        "id": str(item.get("id") or "").strip(),
        "label": str(item.get("label") or "").strip(),
        "phase": str(item.get("phase") or "after_send").strip() or "after_send",
    }


def _init_automation_action_rows(targets):
    rows = []
    for idx, target in enumerate(targets or [], start=1):
        planned = _planned_automation_action_for_target(idx - 1)
        rows.append(
            {
                "count": idx,
                "sms_number": str(target or "").strip(),
                "action_id": str(planned.get("id") or "").strip(),
                "action": str(planned.get("label") or "").strip() or "-",
                "phase": str(planned.get("phase") or "after_send").strip() or "after_send",
                "checked": False,
                "status": "assigned",
                "detail": "",
                "performed_at": "",
            }
        )
    return rows


AUTOMATION_ACTION_TO_LIVE_ROW_FIELDS = {
    "action": "action",
    "action_id": "action_id",
    "phase": "action_phase",
    "status": "action_status",
    "detail": "action_detail",
    "performed_at": "action_performed_at",
    "checked": "action_checked",
}


def _update_automation_action_row(state, row_index, **fields):
    try:
        idx = int(row_index)
    except Exception:
        return
    with _SEND_STATE_LOCK:
        rows = list((state or {}).get("automation_action_rows") or [])
        if idx < 0 or idx >= len(rows):
            return
        row = dict(rows[idx] or {})
        for key, value in fields.items():
            row[key] = value
        rows[idx] = row
        state["automation_action_rows"] = rows
        # Mirror into the Campaign Monitoring live row for the same target.
        live_fields = {}
        for key, value in fields.items():
            live_key = AUTOMATION_ACTION_TO_LIVE_ROW_FIELDS.get(key)
            if live_key:
                live_fields[live_key] = value
        if live_fields:
            _update_live_row(state, idx, **live_fields)


def _automation_action_row_snapshot(state, row_index):
    try:
        idx = int(row_index)
    except Exception:
        return {}
    rows = list((state or {}).get("automation_action_rows") or [])
    if idx < 0 or idx >= len(rows):
        return {}
    return dict(rows[idx] or {})


def _build_random_batch_sizes(total_items, min_size=20, max_size=30):
    try:
        total = max(0, int(total_items or 0))
    except Exception:
        total = 0
    if total <= 0:
        return []
    low = max(1, int(min_size or 1))
    high = max(low, int(max_size or low))
    sizes = []
    remaining = total
    while remaining > 0:
        if remaining <= low:
            size = remaining
        else:
            size = random.randint(low, min(high, remaining))
        tail = remaining - size
        # Do not leave a final batch smaller than the minimum when it still
        # fits into this one.
        if 0 < tail < low and (size + tail) <= high:
            size += tail
        sizes.append(size)
        remaining -= size
    return sizes


def _build_fixed_batch_sizes(total_items, batch_size=20):
    try:
        total = max(0, int(total_items or 0))
    except Exception:
        total = 0
    if total <= 0:
        return []
    size = max(1, int(batch_size or 1))
    sizes = []
    remaining = total
    while remaining > 0:
        chunk = min(size, remaining)
        sizes.append(chunk)
        remaining -= chunk
    return sizes


def _expand_batch_numbers(batch_sizes):
    numbers = []
    for batch_no, size in enumerate(batch_sizes or [], start=1):
        for _ in range(max(0, int(size or 0))):
            numbers.append(batch_no)
    return numbers


def is_already_sent(db_conn, phone, payload):
    schema = get_db_schema()
    payload_signature = str((payload or {}).get("payload_signature") or "")
    dedupe_signature = str((payload or {}).get("dedupe_signature") or "")
    message = str((payload or {}).get("message") or "")
    with db_conn.cursor() as cur:
        if dedupe_signature:
            cur.execute(
                sql.SQL(
                    """
                    SELECT 1
                    FROM {schema_name}.sms_sent
                    WHERE phone = %s
                      AND dedupe_signature = %s
                    LIMIT 1
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (phone, dedupe_signature),
            )
            if cur.fetchone() is not None:
                return True

        if payload_signature:
            cur.execute(
                sql.SQL(
                    """
                    SELECT 1
                    FROM {schema_name}.sms_sent
                    WHERE phone = %s
                      AND payload_signature = %s
                    LIMIT 1
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (phone, payload_signature),
            )
            if cur.fetchone() is not None:
                return True

        # Legacy fallback for old rows that might not have signatures.
        cur.execute(
            sql.SQL(
                """
                SELECT 1
                FROM {schema_name}.sms_sent
                WHERE phone = %s
                  AND payload_signature IS NULL
                  AND message = %s
                LIMIT 1
                """
            ).format(schema_name=sql.Identifier(schema)),
            (phone, message),
        )
        return cur.fetchone() is not None


def save_sent_sms(db_conn, phone, payload, campaign_id=None, campaign_name=""):
    schema = get_db_schema()
    attachment_names = json.dumps((payload or {}).get("attachment_names") or [], ensure_ascii=False)
    attachments_payload = json.dumps((payload or {}).get("attachments") or [], ensure_ascii=False)
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                INSERT INTO {schema_name}.sms_sent (
                    phone,
                    message,
                    status,
                    sent_at,
                    campaign_id,
                    campaign_name,
                    content_kind,
                    attachment_names,
                    attachment_count,
                    payload_signature,
                    dedupe_signature,
                    content_summary,
                    attachments_payload
                )
                VALUES (%s, %s, 'sent', NOW(), %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
                ON CONFLICT DO NOTHING
                """
            ).format(schema_name=sql.Identifier(schema)),
            (
                phone,
                (payload or {}).get("message") or "",
                campaign_id,
                str(campaign_name or "").strip(),
                (payload or {}).get("content_kind") or "text",
                attachment_names,
                int((payload or {}).get("attachment_count") or 0),
                (payload or {}).get("payload_signature") or None,
                (payload or {}).get("dedupe_signature") or None,
                (payload or {}).get("content_summary") or "",
                attachments_payload,
            ),
        )


def save_not_sent_sms(db_conn, campaign_id, campaign_name, phone, reason):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                INSERT INTO {schema_name}.sms_not_sent (campaign_id, campaign_name, phone, reason, failed_at)
                VALUES (%s, %s, %s, %s, NOW())
                """
            ).format(schema_name=sql.Identifier(schema)),
            (str(campaign_id or "").zfill(6), str(campaign_name or "").strip(), phone, reason),
        )

def is_valid_phone(num):
    try:
        if not num.startswith("+"):
            num = "+" + num
        parsed = phonenumbers.parse(num, None)
        # Be permissive: accept numbers that are possible even if not strictly valid.
        return phonenumbers.is_possible_number(parsed) or phonenumbers.is_valid_number(parsed)
    except:
        # Fallback: basic check for + followed by 7-16 digits
        digits = "".join(ch for ch in str(num) if ch.isdigit())
        return 7 <= len(digits) <= 16


def normalize_phone(num):
    if num is None:
        return None
    num = str(num).strip()
    if not num:
        return None
    # Keep leading +, strip everything else to digits
    cleaned = "".join(ch for ch in num if ch.isdigit() or ch == "+")
    if cleaned.startswith("00"):
        cleaned = "+" + cleaned[2:]
    if cleaned.startswith("+"):
        return cleaned
    num = cleaned
    if num.startswith("+"):
        return num
    return "+" + num


def parse_country_code_filter(raw_filter):
    if not raw_filter:
        return []
    tokens = []
    for part in str(raw_filter).replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        digits = "".join(ch for ch in part if ch.isdigit())
        if not digits:
            continue
        tokens.append("+" + digits)
    # Preserve order, remove duplicates
    return list(dict.fromkeys(tokens))


def filter_numbers_by_country_codes(numbers, raw_filter):
    prefixes = parse_country_code_filter(raw_filter)
    if not prefixes:
        return numbers, 0, []
    filtered = [n for n in numbers if any(n.startswith(prefix) for prefix in prefixes)]
    skipped = max(0, len(numbers) - len(filtered))
    return filtered, skipped, prefixes


def _parse_yes_no(raw_value, default=True):
    if raw_value is None:
        return bool(default)
    value = str(raw_value).strip().lower()
    if value in {"yes", "true", "1", "on"}:
        return True
    if value in {"no", "false", "0", "off"}:
        return False
    return bool(default)


def _parse_int_with_bounds(raw_value, default_value, min_value=None, max_value=None):
    try:
        value = int(str(raw_value).strip())
    except Exception:
        value = int(default_value)
    if min_value is not None:
        value = max(int(min_value), value)
    if max_value is not None:
        value = min(int(max_value), value)
    return value


def _normalize_header(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    return "".join(ch for ch in text if ch.isalnum())


def load_numbers(file_storage):
    filename = secure_filename(file_storage.filename or "")
    ext = Path(filename).suffix.lower()
    numbers = []
 
    if ext == ".csv":
        text = file_storage.stream.read().decode("utf-8", errors="ignore")
        reader = csv.DictReader(text.splitlines())
        columns = [_normalize_header(c) for c in reader.fieldnames or []]
        key = None
        for candidate in CONTACT_COLUMN_ALIASES_NORMALIZED:
            if candidate in columns:
                key = candidate
                break
        for row in reader:
            if key:
                raw = row.get(key) or row.get(key.replace("_", ""))
                if raw is None:
                    # DictReader keeps original headers, so map by index if needed
                    try:
                        idx = columns.index(key)
                        raw = list(row.values())[idx]
                    except Exception:
                        raw = None
            else:
                values = list(row.values())
                raw = values[1] if len(values) > 1 else values[0] if values else None
            numbers.append(normalize_phone(raw))
        return [n for n in numbers if n]

    if ext == ".xlsx":
        try:
            import openpyxl
        except Exception:
            raise RuntimeError("openpyxl is required for .xlsx files")
        wb = openpyxl.load_workbook(file_storage, read_only=True)
        ws = wb.active
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
        header = [_normalize_header(c) for c in header_row]
        col_index = None
        for candidate in CONTACT_COLUMN_ALIASES_NORMALIZED:
            if candidate in header:
                col_index = header.index(candidate)
                break
        for row in ws.iter_rows(min_row=2, values_only=True):
            if col_index is not None and col_index < len(row):
                raw = row[col_index]
            else:
                raw = row[1] if len(row) > 1 else row[0] if row else None
            numbers.append(normalize_phone(raw))
        return [n for n in numbers if n]

    raise RuntimeError("Unsupported file type. Use .csv or .xlsx")


def _dedupe_numbers_preserve_order(numbers):
    cleaned = []
    seen = set()
    for raw in numbers or []:
        number = normalize_phone(raw)
        if not number or number in seen:
            continue
        seen.add(number)
        cleaned.append(number)
    return cleaned


def load_numbers_manual_flexible(file_storage):
    """
    Flexible parser for manual ALL Contacts upload:
    - accepts CSV/XLSX with or without headers
    - accepts 1-column files (phone only)
    - prefers alias header column when available
    """
    filename = secure_filename(file_storage.filename or "")
    ext = Path(filename).suffix.lower()

    if ext == ".csv":
        text = file_storage.stream.read().decode("utf-8", errors="ignore")
        rows = list(csv.reader(text.splitlines()))
        if not rows:
            return []
        first = rows[0] if rows else []
        first_norm = [_normalize_header(c) for c in first]
        alias_idx = None
        for candidate in CONTACT_COLUMN_ALIASES_NORMALIZED:
            if candidate in first_norm:
                alias_idx = first_norm.index(candidate)
                break
        has_header = alias_idx is not None
        start_idx = 1 if has_header else 0
        values = []
        for row in rows[start_idx:]:
            if not row:
                continue
            raw = None
            if has_header and alias_idx is not None and alias_idx < len(row):
                raw = row[alias_idx]
            if raw is None or str(raw).strip() == "":
                raw = next((cell for cell in row if str(cell or "").strip()), None)
            values.append(raw)
        return _dedupe_numbers_preserve_order(values)

    if ext == ".xlsx":
        try:
            import openpyxl
        except Exception:
            raise RuntimeError("openpyxl is required for .xlsx files")
        wb = openpyxl.load_workbook(file_storage, read_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(min_row=1, values_only=True)]
        if not rows:
            return []
        first = rows[0]
        first_norm = [_normalize_header(c) for c in first]
        alias_idx = None
        for candidate in CONTACT_COLUMN_ALIASES_NORMALIZED:
            if candidate in first_norm:
                alias_idx = first_norm.index(candidate)
                break
        has_header = alias_idx is not None
        start_idx = 1 if has_header else 0
        values = []
        for row in rows[start_idx:]:
            if not row:
                continue
            raw = None
            if has_header and alias_idx is not None and alias_idx < len(row):
                raw = row[alias_idx]
            if raw is None or str(raw).strip() == "":
                raw = next((cell for cell in row if str(cell or "").strip()), None)
            values.append(raw)
        return _dedupe_numbers_preserve_order(values)

    raise RuntimeError("Unsupported file type. Use .csv or .xlsx")


def import_manual_contacts_to_all_contacts(numbers, group_name):
    group_value = _normalize_group_name(group_name)
    if not group_value:
        raise RuntimeError("Group name is required.")

    cleaned = _dedupe_numbers_preserve_order(numbers)
    if not cleaned:
        raise RuntimeError("No valid phone numbers found in uploaded file.")

    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        has_legacy_whatsapp_group = ensure_all_contacts_table(db_conn)
        with db_conn.cursor() as cur:
            for number in cleaned:
                if has_legacy_whatsapp_group:
                    cur.execute(
                        sql.SQL(
                            """
                            INSERT INTO {schema_name}.all_contacts ("Phone", "whatspp_group", "whatsapp_group", "time")
                            VALUES (%s, %s, %s, NOW())
                            ON CONFLICT ("Phone", "whatspp_group") DO NOTHING
                            """
                        ).format(schema_name=sql.Identifier(schema)),
                        (number, group_value, group_value),
                    )
                else:
                    cur.execute(
                        sql.SQL(
                            """
                            INSERT INTO {schema_name}.all_contacts ("Phone", "whatspp_group", "time")
                            VALUES (%s, %s, NOW())
                            ON CONFLICT ("Phone", "whatspp_group") DO NOTHING
                            """
                        ).format(schema_name=sql.Identifier(schema)),
                        (number, group_value),
                    )
        return cleaned
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def group_exists_in_all_contacts(group_name):
    normalized = _normalize_group_name(group_name)
    if not normalized:
        return False
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        group_col = get_all_contacts_group_column(db_conn)
        if not group_col:
            return False
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT 1
                    FROM {schema_name}.all_contacts
                    WHERE LOWER(TRIM({group_col})) = LOWER(TRIM(%s))
                    LIMIT 1
                    """
                ).format(
                    schema_name=sql.Identifier(schema),
                    group_col=sql.Identifier(group_col),
                ),
                (normalized,),
            )
            row = cur.fetchone()
            return bool(row)
    finally:
        try:
            db_conn.close()
        except Exception:
            pass


def build_contact_items(contacts, status_map=None, row_id_map=None):
    status_map = status_map or {}
    row_id_map = row_id_map or {}
    items = []
    for number in contacts or []:
        info = status_map.get(number) or {}
        status = str(info.get("status") or "pending")
        reason = str(info.get("reason") or "")
        row_id = row_id_map.get(number)
        items.append(
            {
                "number": number,
                "status": status,
                "sent": status == "sent",
                "reason": reason,
                "row_id": row_id,
            }
        )
    return items


def _click_first_available(locator, page=None, label="click", timeout_ms=1800, force=False, state=None):
    count = locator.count()
    for i in range(count):
        item = locator.nth(i)
        try:
            if item.is_visible():
                if page is not None:
                    if _human_click_item(
                        page,
                        item,
                        label=label,
                        timeout_ms=timeout_ms,
                        force=force,
                        state=state,
                    ):
                        return True
                    if _delayed_click_fallback(
                        page,
                        item,
                        label=label,
                        timeout_ms=timeout_ms,
                        force=force,
                        state=state,
                    ):
                        return True
                click_kwargs = {}
                if timeout_ms is not None:
                    click_kwargs["timeout"] = timeout_ms
                if force:
                    click_kwargs["force"] = True
                if _humanize_actions_enabled():
                    click_kwargs["delay"] = random.randint(45, 135)
                item.click(**click_kwargs)
                return True
        except Exception:
            continue
    return False


def _cursor_view_enabled():
    # Default ON so automation movement is visible unless explicitly disabled.
    raw = str(os.getenv("CURSOR_VIEW", "1")).strip().lower()
    return raw in ("1", "true", "yes", "on")


def _humanize_actions_enabled():
    raw = str(os.getenv("PLAYWRIGHT_HUMANIZE", "1")).strip().lower()
    return raw in ("1", "true", "yes", "on")


def _resolve_interaction_state(state=None):
    if isinstance(state, dict):
        return state
    try:
        active = globals().get("send_state")
        if isinstance(active, dict):
            return active
    except Exception:
        pass
    return state


def _get_click_delay_settings(state=None):
    resolved_state = _resolve_interaction_state(state)
    defaults = {"enabled": True, "min_ms": 150, "max_ms": 900}
    if not isinstance(resolved_state, dict):
        return defaults
    config = dict(resolved_state.get("_click_delay_config") or {})
    if not config:
        config = {
            "enabled": resolved_state.get("last_random_click_delay_enabled", defaults["enabled"]),
            "min_ms": resolved_state.get("last_click_delay_min_ms", defaults["min_ms"]),
            "max_ms": resolved_state.get("last_click_delay_max_ms", defaults["max_ms"]),
        }
    try:
        enabled = bool(config.get("enabled", defaults["enabled"]))
    except Exception:
        enabled = defaults["enabled"]
    try:
        min_ms = max(0, int(config.get("min_ms", defaults["min_ms"]) or defaults["min_ms"]))
    except Exception:
        min_ms = defaults["min_ms"]
    try:
        max_ms = max(min_ms, int(config.get("max_ms", defaults["max_ms"]) or defaults["max_ms"]))
    except Exception:
        max_ms = max(min_ms, defaults["max_ms"])
    return {"enabled": enabled, "min_ms": min_ms, "max_ms": max_ms}


def _wait_before_click_action(page, label="click", state=None):
    settings = _get_click_delay_settings(state)
    if not settings.get("enabled"):
        return 0
    low = max(0, int(settings.get("min_ms") or 0))
    high = max(low, int(settings.get("max_ms") or low))
    wait_ms = random.randint(low, high) if high > 0 else 0
    if wait_ms > 0:
        _cooperative_wait(
            page,
            wait_ms,
            state=_resolve_interaction_state(state),
            step_ms=max(40, min(wait_ms, 120)),
        )
    return wait_ms


def _get_optional_side_action_settings(state=None):
    resolved_state = _resolve_interaction_state(state)
    defaults = {
        "enabled": True,
        "chance_percent": 20,
        "pause_min_ms": 700,
        "pause_max_ms": 2000,
    }
    if not isinstance(resolved_state, dict):
        return defaults
    config = dict(resolved_state.get("_optional_side_action_config") or {})
    try:
        enabled = bool(config.get("enabled", defaults["enabled"]))
    except Exception:
        enabled = defaults["enabled"]
    try:
        chance_percent = max(0, min(100, int(config.get("chance_percent", defaults["chance_percent"]) or defaults["chance_percent"])))
    except Exception:
        chance_percent = defaults["chance_percent"]
    try:
        pause_min_ms = max(0, int(config.get("pause_min_ms", defaults["pause_min_ms"]) or defaults["pause_min_ms"]))
    except Exception:
        pause_min_ms = defaults["pause_min_ms"]
    try:
        pause_max_ms = max(pause_min_ms, int(config.get("pause_max_ms", defaults["pause_max_ms"]) or defaults["pause_max_ms"]))
    except Exception:
        pause_max_ms = max(pause_min_ms, defaults["pause_max_ms"])
    return {
        "enabled": enabled,
        "chance_percent": chance_percent,
        "pause_min_ms": pause_min_ms,
        "pause_max_ms": pause_max_ms,
    }


def _optional_side_action_pause_ms(state=None):
    settings = _get_optional_side_action_settings(state)
    low = max(0, int(settings.get("pause_min_ms") or 0))
    high = max(low, int(settings.get("pause_max_ms") or low))
    return random.randint(low, high) if high > 0 else 0


CAMPAIGN_OPTIONAL_ACTIONS = (
    {
        "id": "groups_scroll",
        "label": "Open Groups and scroll",
        "phase": "after_send",
        "description": "After sending, open the Groups tab and scroll through the group list.",
    },
    {
        "id": "group_members_peek",
        "label": "Check group members",
        "phase": "after_send",
        "description": "After sending, open a random group and briefly inspect its members/info.",
    },
    {
        "id": "menu_peek",
        "label": "Open WhatsApp menu",
        "phase": "after_send",
        "description": "After sending, briefly open the WhatsApp menu and return.",
    },
    {
        "id": "filter_peek",
        "label": "Open chat filter",
        "phase": "after_send",
        "description": "After sending, open a safe chat filter like Unread, Favorites, or Archived.",
    },
    {
        "id": "profile_peek",
        "label": "Open WhatsApp profile",
        "phase": "after_send",
        "description": "After sending, briefly open the WhatsApp profile panel and return.",
    },
    {
        "id": "settings_peek",
        "label": "Open WhatsApp settings",
        "phase": "after_send",
        "description": "After sending, briefly open the WhatsApp settings panel and return.",
    },
    {
        "id": "chat_scroll",
        "label": "Scroll opened chat",
        "phase": "after_send",
        "description": "After sending, scroll the currently opened chat before moving to the next target.",
    },
    {
        "id": "contact_info_peek",
        "label": "Open number info",
        "phase": "after_send",
        "description": "After sending, open the currently active contact info and return before the next target.",
    },
    {
        "id": "attach_document_click",
        "label": "Click Document option",
        "phase": "after_send",
        "description": "After sending, open the attachment menu, click Document, then return safely.",
    },
    {
        "id": "attach_media_click",
        "label": "Click Photos & videos option",
        "phase": "after_send",
        "description": "After sending, open the attachment menu, click Photos & videos, then return safely.",
    },
    {
        "id": "attach_camera_click",
        "label": "Click Camera option",
        "phase": "after_send",
        "description": "After sending, open the attachment menu, click Camera, then return safely.",
    },
    {
        "id": "attach_audio_click",
        "label": "Click Audio option",
        "phase": "after_send",
        "description": "After sending, open the attachment menu, click Audio, then return safely.",
    },
    {
        "id": "attach_contact_click",
        "label": "Click Contact option",
        "phase": "after_send",
        "description": "After sending, open the attachment menu, click Contact, then return safely.",
    },
    {
        "id": "attach_poll_click",
        "label": "Click Poll option",
        "phase": "after_send",
        "description": "After sending, open the attachment menu, click Poll, then return safely.",
    },
    {
        "id": "attach_event_click",
        "label": "Click Event option",
        "phase": "after_send",
        "description": "After sending, open the attachment menu, click Event, then return safely.",
    },
    {
        "id": "attach_sticker_click",
        "label": "Click New sticker option",
        "phase": "after_send",
        "description": "After sending, open the attachment menu, click New sticker, then return safely.",
    },
)


def _build_campaign_action_state(existing=None):
    existing_map = {}
    for raw in list(existing or []):
        row = dict(raw or {})
        action_id = str(row.get("id") or "").strip()
        if action_id:
            existing_map[action_id] = row

    rows = []
    for position, item in enumerate(CAMPAIGN_OPTIONAL_ACTIONS, start=1):
        merged = {
            "id": str(item.get("id") or "").strip(),
            "label": str(item.get("label") or "").strip(),
            "phase": str(item.get("phase") or "after_send").strip() or "after_send",
            "description": str(item.get("description") or "").strip(),
            "position": position,
            "done": False,
            "count": 0,
            "last_at": "",
            "completed_at": "",
            "last_target": "",
            "last_detail": "",
        }
        previous = dict(existing_map.get(merged["id"]) or {})
        if previous:
            merged["done"] = bool(previous.get("done"))
            try:
                merged["count"] = max(0, int(previous.get("count") or 0))
            except Exception:
                merged["count"] = 0
            merged["last_at"] = str(previous.get("last_at") or "").strip()
            merged["completed_at"] = str(previous.get("completed_at") or "").strip()
            merged["last_target"] = str(previous.get("last_target") or "").strip()
            merged["last_detail"] = str(previous.get("last_detail") or "").strip()
        rows.append(merged)
    return rows


def _get_active_campaign_actions(state=None):
    resolved_state = _resolve_interaction_state(state)
    if not isinstance(resolved_state, dict):
        return _build_campaign_action_state()
    return _build_campaign_action_state(resolved_state.get("campaign_actions") or [])


def _get_campaign_actions_snapshot(state=None):
    resolved_state = _resolve_interaction_state(state)
    if not isinstance(resolved_state, dict):
        return _build_campaign_action_state()
    if bool(resolved_state.get("running")):
        source_rows = resolved_state.get("campaign_actions") or []
    else:
        source_rows = resolved_state.get("last_campaign_actions") or resolved_state.get("campaign_actions") or []
    return _build_campaign_action_state(source_rows)


def _pending_campaign_action_ids(state=None, phase=""):
    rows = _get_active_campaign_actions(state)
    phase_value = str(phase or "").strip().lower()
    pending = []
    for row in rows:
        if phase_value and str(row.get("phase") or "").strip().lower() != phase_value:
            continue
        if not bool(row.get("done")):
            pending.append(str(row.get("id") or "").strip())
    return pending


def _mark_campaign_action_done(state, action_id, target="", detail=""):
    with _SEND_STATE_LOCK:
        return _mark_campaign_action_done_unlocked(state, action_id, target=target, detail=detail)


def _mark_campaign_action_done_unlocked(state, action_id, target="", detail=""):
    resolved_state = _resolve_interaction_state(state)
    action_key = str(action_id or "").strip()
    if not isinstance(resolved_state, dict) or not action_key:
        return False

    rows = _build_campaign_action_state(resolved_state.get("campaign_actions") or [])
    now_text = datetime.now().strftime("%H:%M:%S")
    changed = False
    action_label = action_key
    for row in rows:
        if str(row.get("id") or "").strip() != action_key:
            continue
        action_label = str(row.get("label") or action_key).strip() or action_key
        row["done"] = True
        try:
            row["count"] = max(0, int(row.get("count") or 0)) + 1
        except Exception:
            row["count"] = 1
        row["last_at"] = now_text
        if not str(row.get("completed_at") or "").strip():
            row["completed_at"] = now_text
        if str(target or "").strip():
            row["last_target"] = str(target or "").strip()
        if str(detail or "").strip():
            row["last_detail"] = str(detail or "").strip()
        changed = True
        break
    if not changed:
        return False

    resolved_state["campaign_actions"] = rows
    event_text = action_label
    if str(target or "").strip():
        event_text += f" for {str(target or '').strip()}"
    if str(detail or "").strip():
        event_text += f" ({str(detail or '').strip()})"
    resolved_state["campaign_action_event_text"] = event_text.strip()
    return True


def _human_pause(page, min_ms=35, max_ms=120, state=None, step_ms=90):
    state = _resolve_interaction_state(state)
    if not _humanize_actions_enabled():
        return 0
    try:
        low = max(0, int(min_ms or 0))
    except Exception:
        low = 0
    try:
        high = max(low, int(max_ms or low))
    except Exception:
        high = low
    wait_ms = random.randint(low, high) if high > 0 else 0
    if wait_ms > 0:
        _cooperative_wait(page, wait_ms, state=state, step_ms=max(40, min(int(step_ms or 90), wait_ms)))
    return wait_ms


def _read_human_cursor(page):
    fallback = {
        "x": 72.0,
        "y": 96.0,
        "viewport_width": 1366.0,
        "viewport_height": 768.0,
    }
    try:
        payload = page.evaluate(
            """
            () => {
              const width = window.innerWidth || document.documentElement.clientWidth || 1366;
              const height = window.innerHeight || document.documentElement.clientHeight || 768;
              const saved = window.__codexHumanCursor || {};
              return {
                x: Number.isFinite(saved.x) ? saved.x : Math.max(42, Math.round(width * 0.08)),
                y: Number.isFinite(saved.y) ? saved.y : Math.max(72, Math.round(height * 0.18)),
                viewport_width: width,
                viewport_height: height,
              };
            }
            """
        )
        if isinstance(payload, dict):
            return payload
    except Exception:
        pass
    return fallback


def _remember_human_cursor(page, x, y):
    try:
        page.evaluate(
            """
            ([x, y]) => {
              window.__codexHumanCursor = { x, y };
            }
            """,
            [float(x), float(y)],
        )
    except Exception:
        pass


def _human_click_target_from_box(box):
    if not isinstance(box, dict):
        return None
    try:
        width = max(1.0, float(box.get("width", 0) or 0))
        height = max(1.0, float(box.get("height", 0) or 0))
        origin_x = float(box.get("x", 0) or 0)
        origin_y = float(box.get("y", 0) or 0)
    except Exception:
        return None

    pad_x = min(max(2.0, width * 0.18), max(2.0, (width / 2.0) - 1.0))
    pad_y = min(max(2.0, height * 0.22), max(2.0, (height / 2.0) - 1.0))
    if width <= 8.0:
        pad_x = 1.0
    if height <= 8.0:
        pad_y = 1.0

    local_x_min = max(1.0, pad_x)
    local_x_max = max(local_x_min, width - pad_x)
    local_y_min = max(1.0, pad_y)
    local_y_max = max(local_y_min, height - pad_y)
    local_x = random.uniform(local_x_min, local_x_max)
    local_y = random.uniform(local_y_min, local_y_max)
    abs_x = origin_x + local_x
    abs_y = origin_y + local_y
    return {
        "local_x": float(local_x),
        "local_y": float(local_y),
        "abs_x": float(abs_x),
        "abs_y": float(abs_y),
        "width": width,
        "height": height,
    }


def _human_move_mouse_to_box(page, box):
    point = _human_click_target_from_box(box)
    if not point:
        return None
    if not _humanize_actions_enabled():
        try:
            page.mouse.move(point["abs_x"], point["abs_y"], steps=18)
            _remember_human_cursor(page, point["abs_x"], point["abs_y"])
        except Exception:
            pass
        return point

    cursor_state = _read_human_cursor(page)
    viewport_width = max(200.0, float(cursor_state.get("viewport_width") or 1366.0))
    viewport_height = max(200.0, float(cursor_state.get("viewport_height") or 768.0))
    current_x = float(cursor_state.get("x") or 72.0)
    current_y = float(cursor_state.get("y") or 96.0)
    target_x = float(point["abs_x"])
    target_y = float(point["abs_y"])
    distance = ((target_x - current_x) ** 2 + (target_y - current_y) ** 2) ** 0.5
    overshoot_x = max(4.0, min(viewport_width - 4.0, target_x + random.uniform(-16.0, 16.0)))
    overshoot_y = max(4.0, min(viewport_height - 4.0, target_y + random.uniform(-12.0, 12.0)))
    try:
        if distance >= 90.0:
            page.mouse.move(overshoot_x, overshoot_y, steps=random.randint(12, 26))
            _remember_human_cursor(page, overshoot_x, overshoot_y)
        page.mouse.move(target_x, target_y, steps=random.randint(7, 18))
        _remember_human_cursor(page, target_x, target_y)
    except Exception:
        pass
    return point


def _human_click_item(page, item, label="click", timeout_ms=1800, force=False, state=None):
    state = _resolve_interaction_state(state)
    point = _debug_before_click(page, item, label=label, state=state)
    if not point:
        _wait_before_click_action(page, label=label, state=state)
    try:
        local_x = 1
        local_y = 1
        if point:
            max_x = max(1, int(point.get("width") or 1) - 1)
            max_y = max(1, int(point.get("height") or 1) - 1)
            local_x = min(max(1, int(round(point["local_x"]))), max_x)
            local_y = min(max(1, int(round(point["local_y"]))), max_y)
        click_kwargs = {"timeout": timeout_ms}
        if force:
            click_kwargs["force"] = True
        if point:
            click_kwargs["position"] = {
                "x": local_x,
                "y": local_y,
            }
        if _humanize_actions_enabled():
            click_kwargs["delay"] = random.randint(45, 135)
        item.click(**click_kwargs)
        _human_pause(page, min_ms=20, max_ms=85, state=state, step_ms=80)
        return True
    except Exception:
        pass

    try:
        fallback_kwargs = {"timeout": timeout_ms}
        if force:
            fallback_kwargs["force"] = True
        item.click(**fallback_kwargs)
        return True
    except Exception:
        return False


def _delayed_click_fallback(page, item, label="click", timeout_ms=1800, force=False, state=None):
    state = _resolve_interaction_state(state)
    point = _debug_before_click(page, item, label=label, state=state)
    if not point:
        _wait_before_click_action(page, label=label, state=state)
    click_kwargs = {}
    if timeout_ms is not None:
        click_kwargs["timeout"] = timeout_ms
    if force:
        click_kwargs["force"] = True
    if _humanize_actions_enabled():
        click_kwargs["delay"] = random.randint(45, 135)
    try:
        item.click(**click_kwargs)
        _human_pause(page, min_ms=20, max_ms=85, state=state, step_ms=80)
        return True
    except Exception:
        fallback_kwargs = dict(click_kwargs)
        fallback_kwargs.pop("delay", None)
        try:
            item.click(**fallback_kwargs)
            return True
        except Exception:
            return False


def _delayed_dom_click_fallback(page, item, label="click", state=None):
    state = _resolve_interaction_state(state)
    point = _debug_before_click(page, item, label=label, state=state)
    if not point:
        _wait_before_click_action(page, label=label, state=state)
    try:
        item.evaluate("(el) => el.click()")
        _human_pause(page, min_ms=20, max_ms=85, state=state, step_ms=80)
        return True
    except Exception:
        return False


def _human_type_text_basic(page, text_value, state=None, min_delay_ms=24, max_delay_ms=56):
    text_value = str(text_value or "")
    normalized = text_value.replace("\r\n", "\n").replace("\r", "\n")
    if not normalized:
        return False
    if not _humanize_actions_enabled():
        for idx, line in enumerate(normalized.split("\n")):
            if idx > 0:
                page.keyboard.press("Shift+Enter")
            if line:
                page.keyboard.type(line, delay=12)
        return True

    typed_since_pause = 0
    next_pause_after = random.randint(6, 14)
    index = 0
    while index < len(normalized):
        _wait_if_send_paused(state, page=page)
        char = normalized[index]
        if char == "\n":
            page.keyboard.press("Shift+Enter")
            _human_pause(page, min_ms=35, max_ms=110, state=state, step_ms=70)
            typed_since_pause = 0
            next_pause_after = random.randint(6, 14)
            index += 1
            continue

        chunk_size = random.randint(1, 3)
        chunk_chars = []
        while index < len(normalized) and normalized[index] != "\n" and len(chunk_chars) < chunk_size:
            chunk_chars.append(normalized[index])
            index += 1
        page.keyboard.type("".join(chunk_chars), delay=random.randint(min_delay_ms, max_delay_ms))
        typed_since_pause += len(chunk_chars)
        if typed_since_pause >= next_pause_after:
            _human_pause(page, min_ms=85, max_ms=240, state=state, step_ms=90)
            typed_since_pause = 0
            next_pause_after = random.randint(6, 14)
    return True


def _debug_before_click(page, item, label="click", state=None):
    state = _resolve_interaction_state(state)
    try:
        page.bring_to_front()
    except Exception:
        pass
    if _humanize_actions_enabled():
        _human_pause(page, min_ms=35, max_ms=120, state=state, step_ms=80)
    try:
        item.scroll_into_view_if_needed(timeout=1400)
    except Exception:
        pass
    try:
        box = item.bounding_box()
    except Exception:
        box = None
    if not box:
        print(f"[CURSOR] {label}")
        return None
    point = _human_move_mouse_to_box(page, box)
    if not point:
        print(f"[CURSOR] {label}")
        return None
    x = float(point["abs_x"])
    y = float(point["abs_y"])
    if not _cursor_view_enabled():
        _wait_before_click_action(page, label=label, state=state)
        return point
    try:
        page.evaluate(
            """
            ([x, y, label]) => {
              const old = document.getElementById('__cursor_debug_dot');
              if (old) old.remove();
              const dot = document.createElement('div');
              dot.id = '__cursor_debug_dot';
              dot.style.position = 'fixed';
              dot.style.left = (x - 10) + 'px';
              dot.style.top = (y - 10) + 'px';
              dot.style.width = '28px';
              dot.style.height = '28px';
              dot.style.border = '4px solid #ff2d2d';
              dot.style.borderRadius = '9999px';
              dot.style.background = 'rgba(255,45,45,0.28)';
              dot.style.boxShadow = '0 0 0 6px rgba(255,45,45,0.25)';
              dot.style.pointerEvents = 'none';
              dot.style.zIndex = '2147483647';
              document.body.appendChild(dot);

              const oldTag = document.getElementById('__cursor_debug_tag');
              if (oldTag) oldTag.remove();
              const tag = document.createElement('div');
              tag.id = '__cursor_debug_tag';
              tag.textContent = 'AUTO: ' + (label || 'click');
              tag.style.position = 'fixed';
              tag.style.left = '12px';
              // Lifted off the bottom-left corner so WhatsApp's profile picture
              // in the left rail stays visible under the caption.
              tag.style.bottom = '16%';
              tag.style.padding = '6px 10px';
              tag.style.borderRadius = '8px';
              tag.style.background = '#ff2d2d';
              tag.style.color = '#fff';
              tag.style.font = 'bold 13px/1.2 Arial, sans-serif';
              tag.style.pointerEvents = 'none';
              tag.style.zIndex = '2147483647';
              document.body.appendChild(tag);
            }
            """,
            [x, y, label],
        )
        page.wait_for_timeout(180)
    except Exception:
        pass
    if _humanize_actions_enabled():
        _human_pause(page, min_ms=45, max_ms=105, state=state, step_ms=80)
    _wait_before_click_action(page, label=label, state=state)
    print(f"[CURSOR] {label} at ({int(x)}, {int(y)})")
    return point


def copy_file_to_windows_clipboard(file_path):
    resolved = str(Path(file_path).resolve())
    if os.name != "nt":
        raise RuntimeError("Clipboard file paste is currently supported only on Windows")
    attachment_kind = detect_attachment_kind(resolved)
    if attachment_kind == "image":
        ps_script = (
            "& { param([string]$path) "
            "Add-Type -AssemblyName System.Windows.Forms; "
            "Add-Type -AssemblyName System.Drawing; "
            "$img = [System.Drawing.Image]::FromFile($path); "
            "[System.Windows.Forms.Clipboard]::SetImage($img); "
            "$img.Dispose() }"
        )
    else:
        ps_script = (
            "& { param([string]$path) "
            "Add-Type -AssemblyName System.Windows.Forms; "
            "$files = New-Object System.Collections.Specialized.StringCollection; "
            "$null = $files.Add($path); "
            "[System.Windows.Forms.Clipboard]::SetFileDropList($files) }"
        )
    command = ["powershell", "-NoProfile", "-Command", ps_script, "-path", resolved]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        err = (result.stderr or result.stdout or "").strip()
        raise RuntimeError(f"Could not copy file to clipboard: {err or 'unknown error'}")


def _type_text_with_random_pauses(
    page,
    text_value,
    typing_pause_chars_min=0,
    typing_pause_chars_max=0,
    typing_pause_delay_min_ms=0,
    typing_pause_delay_max_ms=0,
    remove_chars_during_typing_enabled=False,
    remove_chars_times_min=0,
    remove_chars_times_max=0,
    remove_chars_count_min=0,
    remove_chars_count_max=0,
    state=None,
):
    text_value = str(text_value or "")
    normalized = text_value.replace("\r\n", "\n").replace("\r", "\n")
    typing_pause_chars_min = max(0, int(typing_pause_chars_min or 0))
    typing_pause_chars_max = max(typing_pause_chars_min, int(typing_pause_chars_max or typing_pause_chars_min))
    typing_pause_delay_min_ms = max(0, int(typing_pause_delay_min_ms or 0))
    typing_pause_delay_max_ms = max(
        typing_pause_delay_min_ms,
        int(typing_pause_delay_max_ms or typing_pause_delay_min_ms),
    )
    remove_chars_during_typing_enabled = bool(remove_chars_during_typing_enabled)
    remove_chars_times_min = max(0, int(remove_chars_times_min or 0))
    remove_chars_times_max = max(
        remove_chars_times_min,
        int(remove_chars_times_max or remove_chars_times_min),
    )
    remove_chars_count_min = max(0, int(remove_chars_count_min or 0))
    remove_chars_count_max = max(
        remove_chars_count_min,
        int(remove_chars_count_max or remove_chars_count_min),
    )
    use_human_typing_pauses = bool(
        normalized
        and typing_pause_chars_max > 0
        and typing_pause_delay_max_ms > 0
    )
    remaining_typeable_chars = sum(1 for char in normalized if char != "\n")
    if remaining_typeable_chars <= 0:
        return False
    use_remove_chars_behavior = bool(
        normalized
        and remove_chars_during_typing_enabled
        and remove_chars_times_max > 0
        and remove_chars_count_max > 0
    )
    if not use_human_typing_pauses and not use_remove_chars_behavior:
        return False

    typing_pause_chars_min = max(1, typing_pause_chars_min) if use_human_typing_pauses else 0
    typing_pause_chars_max = max(typing_pause_chars_min, typing_pause_chars_max) if use_human_typing_pauses else 0
    remove_chars_count_floor = max(1, remove_chars_count_min) if use_remove_chars_behavior else 0
    if use_remove_chars_behavior:
        remove_chars_count_max = max(remove_chars_count_floor, remove_chars_count_max)

    current_line_chars = []
    typed_typeable_chars = 0
    typed_since_pause = 0
    next_pause_after = random.randint(typing_pause_chars_min, typing_pause_chars_max) if use_human_typing_pauses else 0
    remove_trigger_index = 0
    remove_triggers = []

    if use_remove_chars_behavior:
        desired_remove_events = random.randint(remove_chars_times_min, remove_chars_times_max)
        if desired_remove_events > 0:
            available_positions = list(range(max(1, remove_chars_count_floor), remaining_typeable_chars + 1))
            if available_positions:
                desired_remove_events = min(desired_remove_events, len(available_positions))
                remove_triggers = sorted(random.sample(available_positions, desired_remove_events))

    def _reset_pause_window():
        nonlocal typed_since_pause, next_pause_after
        typed_since_pause = 0
        if use_human_typing_pauses:
            next_pause_after = random.randint(typing_pause_chars_min, typing_pause_chars_max)

    def _maybe_pause():
        if not use_human_typing_pauses:
            return
        if typed_since_pause < next_pause_after:
            return
        pause_ms = random.randint(
            typing_pause_delay_min_ms,
            typing_pause_delay_max_ms,
        )
        if pause_ms > 0:
            _cooperative_wait(page, pause_ms, state=state, step_ms=200)
        _reset_pause_window()

    def _maybe_remove_and_rewrite():
        nonlocal remove_trigger_index
        if not use_remove_chars_behavior or remove_trigger_index >= len(remove_triggers):
            return False
        if typed_typeable_chars < remove_triggers[remove_trigger_index]:
            return False
        if len(current_line_chars) < remove_chars_count_floor:
            return False
        remove_count = min(
            len(current_line_chars),
            random.randint(remove_chars_count_floor, remove_chars_count_max),
        )
        if remove_count <= 0:
            remove_trigger_index += 1
            return False
        removed_text = "".join(current_line_chars[-remove_count:])
        _human_pause(page, min_ms=65, max_ms=180, state=state, step_ms=85)
        for _ in range(remove_count):
            _wait_if_send_paused(state, page=page)
            page.keyboard.press("Backspace")
            page.wait_for_timeout(random.randint(22, 48))
        _human_pause(page, min_ms=55, max_ms=150, state=state, step_ms=80)
        if removed_text:
            page.keyboard.type(removed_text, delay=random.randint(20, 46))
        _human_pause(page, min_ms=35, max_ms=95, state=state, step_ms=75)
        remove_trigger_index += 1
        _reset_pause_window()
        return True

    for char in normalized:
        _wait_if_send_paused(state, page=page)
        if char == "\n":
            page.keyboard.press("Shift+Enter")
            _human_pause(page, min_ms=35, max_ms=110, state=state, step_ms=70)
            current_line_chars = []
            _reset_pause_window()
            continue
        page.keyboard.type(char, delay=random.randint(18, 42))
        current_line_chars.append(char)
        typed_typeable_chars += 1
        typed_since_pause += 1
        if _maybe_remove_and_rewrite():
            continue
        _maybe_pause()
    return True


def _fill_caption_if_available(
    page,
    caption,
    typing_pause_chars_min=0,
    typing_pause_chars_max=0,
    typing_pause_delay_min_ms=0,
    typing_pause_delay_max_ms=0,
    remove_chars_during_typing_enabled=False,
    remove_chars_times_min=0,
    remove_chars_times_max=0,
    remove_chars_count_min=0,
    remove_chars_count_max=0,
    state=None,
):
    if not caption:
        return
    caption_text = str(caption)
    candidates = [
        "div[contenteditable='true'][aria-label*='caption' i]",
        "div[contenteditable='true'][aria-placeholder*='caption' i]",
        "div[role='dialog'] div[contenteditable='true'][role='textbox']",
        "div[role='textbox'][aria-label='Type a message'][contenteditable='true'][data-lexical-editor='true']",
        "div[role='textbox'][aria-label='Type a message'][contenteditable='true']",
        "div[aria-label='Type a message'][contenteditable='true']",
        "div[contenteditable='true'][data-tab='10']",
        "div[contenteditable='true'][data-tab='6']",
        "div[contenteditable='true'][role='textbox']",
    ]
    for selector in candidates:
        locator = page.locator(selector)
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count - 1, -1, -1):
            box = locator.nth(i)
            try:
                if box.is_visible():
                    if not _human_click_item(
                        page,
                        box,
                        label="Type caption in preview",
                        timeout_ms=1600,
                        force=True,
                        state=state,
                    ):
                        _delayed_click_fallback(
                            page,
                            box,
                            label="Type caption in preview",
                            timeout_ms=1600,
                            force=True,
                            state=state,
                        )
                    try:
                        page.keyboard.press("Control+A")
                        page.keyboard.press("Backspace")
                    except Exception:
                        pass
                    used_human_typing = _type_text_with_random_pauses(
                        page,
                        caption_text,
                        typing_pause_chars_min=typing_pause_chars_min,
                        typing_pause_chars_max=typing_pause_chars_max,
                        typing_pause_delay_min_ms=typing_pause_delay_min_ms,
                        typing_pause_delay_max_ms=typing_pause_delay_max_ms,
                        remove_chars_during_typing_enabled=remove_chars_during_typing_enabled,
                        remove_chars_times_min=remove_chars_times_min,
                        remove_chars_times_max=remove_chars_times_max,
                        remove_chars_count_min=remove_chars_count_min,
                        remove_chars_count_max=remove_chars_count_max,
                        state=state,
                    )
                    if not used_human_typing:
                        _human_type_text_basic(page, caption_text, state=state)
                    page.wait_for_timeout(60)
                    return
            except Exception:
                continue


def _composer_targets(page):
    targets = []
    selectors = [
        "footer div[contenteditable='true'][data-tab='10']",
        "footer div[contenteditable='true'][data-tab='6']",
        "footer div[contenteditable='true'][role='textbox']",
        "div[contenteditable='true'][data-tab='10']",
        "div[contenteditable='true'][data-tab='6']",
        "div[contenteditable='true'][role='textbox']",
    ]
    for selector in selectors:
        locator = page.locator(selector)
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count - 1, -1, -1):
            targets.append(locator.nth(i))

    text_nodes = [
        page.locator("footer p.copyable-text"),
        page.locator("p.copyable-text"),
    ]
    for locator in text_nodes:
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count - 1, -1, -1):
            try:
                targets.append(locator.nth(i).locator("xpath=ancestor::div[@contenteditable='true'][1]"))
            except Exception:
                continue
    return targets


def _focus_compose_box(page, state=None):
    page.bring_to_front()
    pyautogui.FAILSAFE = False
    for target in _composer_targets(page):
        try:
            if not target.is_visible():
                continue
            target.scroll_into_view_if_needed()
            if not _human_click_item(
                page,
                target,
                label="Focus attachment composer",
                timeout_ms=1600,
                force=True,
                state=state,
            ):
                if not _delayed_click_fallback(
                    page,
                    target,
                    label="Focus attachment composer",
                    timeout_ms=1600,
                    force=True,
                    state=state,
                ):
                    continue
            target.evaluate("(el) => { el.focus(); }")
            page.wait_for_timeout(80)
            return True
        except Exception:
            continue
    return False


ATTACHMENT_SEND_BUTTON_SELECTOR = (
    "button[aria-label='Send'], div[aria-label='Send'], div[role='button'][aria-label='Send'], "
    "button[aria-label='Send message'], span[data-icon='send'], span[data-icon='wds-ic-send-filled'], "
    "span[data-icon='send-filled'], button:has(svg title:has-text('ic-send')), "
    "div[role='button']:has(svg title:has-text('ic-send'))"
)


def _send_button_ready(page):
    send_buttons = page.locator(ATTACHMENT_SEND_BUTTON_SELECTOR)
    try:
        for i in range(send_buttons.count()):
            if send_buttons.nth(i).is_visible():
                return True
    except Exception:
        return False
    return False


def _is_send_pause_requested(state):
    try:
        return bool(state and state.get("pause_requested"))
    except Exception:
        return False


class SendJobCancelled(RuntimeError):
    pass


def _is_send_cancel_requested(state):
    try:
        return bool(state and state.get("cancel_requested"))
    except Exception:
        return False


def _raise_if_send_cancel_requested(state, message="Campaign canceled by user."):
    if _is_send_cancel_requested(state):
        raise SendJobCancelled(str(message or "Campaign canceled by user."))
    return False


def _wait_if_send_paused(state, page=None):
    _raise_if_send_cancel_requested(
        state,
        message="Campaign canceled by user. All WhatsApp browsers were closed.",
    )
    if not _is_send_pause_requested(state):
        return False

    if state is not None and not bool(state.get("paused")):
        state["paused"] = True
        state["pause_requested_at"] = datetime.now().isoformat()
        _send_log(
            "Processing paused by user. WhatsApp automation windows are still open."
        )

    while _is_send_pause_requested(state):
        _raise_if_send_cancel_requested(
            state,
            message="Campaign canceled by user. All WhatsApp browsers were closed.",
        )
        try:
            if page is not None:
                page.wait_for_timeout(250)
            else:
                time.sleep(0.25)
        except Exception:
            time.sleep(0.25)

    if state is not None and bool(state.get("paused")):
        state["paused"] = False
        state["pause_requested_at"] = ""
        _send_log("Processing resumed by user.")
    _raise_if_send_cancel_requested(
        state,
        message="Campaign canceled by user. All WhatsApp browsers were closed.",
    )
    return True


def _cooperative_wait(page, total_ms, state=None, step_ms=250):
    try:
        remaining = max(0, int(total_ms or 0))
    except Exception:
        remaining = 0
    step = max(40, int(step_ms or 250))
    while remaining > 0:
        _wait_if_send_paused(state, page=page)
        current = min(step, remaining)
        try:
            if page is not None:
                page.wait_for_timeout(current)
            else:
                time.sleep(current / 1000.0)
        except Exception:
            time.sleep(current / 1000.0)
        remaining -= current
    _wait_if_send_paused(state, page=page)


def _set_active_send_contexts(state, context_entries):
    if not isinstance(state, dict):
        return
    normalized = []
    for entry in list(context_entries or []):
        context = None
        profile_dir = ""
        if isinstance(entry, dict):
            context = entry.get("context")
            profile_dir = entry.get("profile_dir")
        elif isinstance(entry, (list, tuple)) and len(entry) >= 2:
            context, profile_dir = entry[0], entry[1]
        if context is None:
            continue
        owner_ident = None
        if isinstance(entry, dict):
            owner_ident = entry.get("thread_ident")
        if owner_ident is None:
            owner_ident = threading.get_ident()
        normalized.append({
            "context": context,
            "profile_dir": str(profile_dir or "").strip(),
            "thread_ident": owner_ident,
        })
    state["_active_send_contexts"] = normalized


def _close_active_send_contexts(state):
    if not isinstance(state, dict):
        return 0
    entries = list(state.get("_active_send_contexts") or [])
    state["_active_send_contexts"] = []
    closed_count = 0
    released_profiles = set()
    for entry in entries:
        context = entry.get("context") if isinstance(entry, dict) else None
        profile_dir = str((entry or {}).get("profile_dir") or "").strip() if isinstance(entry, dict) else ""
        owner_ident = entry.get("thread_ident") if isinstance(entry, dict) else None
        if context is not None and owner_ident not in (None, threading.get_ident()):
            # Owned by a worker thread: it closes itself on cancel. Do not
            # release its profile here either; the worker does that.
            continue
        if context is not None:
            try:
                context.close()
            except Exception:
                pass
            closed_count += 1
        if profile_dir:
            normalized_dir = str(Path(profile_dir).resolve()) if Path(profile_dir).exists() else profile_dir
            if normalized_dir not in released_profiles:
                try:
                    release_playwright_profile(profile_dir)
                except Exception:
                    pass
                released_profiles.add(normalized_dir)
    return closed_count


def _cancel_active_send_job(state, wait_timeout_seconds=8.0):
    if not isinstance(state, dict) or not bool(state.get("running")):
        return {"ok": False, "error": "No send job is running."}
    already_requested = bool(state.get("cancel_requested"))
    state["cancel_requested"] = True
    state["cancel_requested_at"] = datetime.now().isoformat()
    state["pause_requested"] = False
    state["paused"] = False
    state["pause_requested_at"] = ""
    state["batch_pause_active"] = False
    state["batch_pause_remaining_sec"] = 0
    state["batch_next_start_at"] = ""
    state["batch_pauses"] = {}
    if not already_requested:
        _send_log("Campaign cancel requested by user. Closing WhatsApp browsers and stopping all sending.")
    closed_count = _close_active_send_contexts(state)
    deadline = time.time() + max(0.5, float(wait_timeout_seconds or 0))
    while bool(state.get("running")) and time.time() < deadline:
        time.sleep(0.1)
    finished = not bool(state.get("running"))
    return {
        "ok": True,
        "running": bool(state.get("running")),
        "paused": bool(state.get("paused")),
        "pause_requested": bool(state.get("pause_requested")),
        "session_opening": bool(state.get("session_opening")),
        "awaiting_session_start": bool(state.get("awaiting_session_start")),
        "cancel_requested": bool(state.get("cancel_requested")),
        "closed_contexts": int(closed_count),
        "message": (
            "Campaign canceled. WhatsApp browsers closed and sending stopped."
            if finished
            else "Campaign cancel requested. Closing WhatsApp browsers and stopping the active send..."
        ),
    }


def _wait_random_send_delay(
    page,
    min_seconds=0,
    max_seconds=30,
    log_func=None,
    label="Send",
    on_delay=None,
    state=None,
):
    try:
        low = max(0, int(min_seconds))
    except Exception:
        low = 0
    try:
        high = max(low, int(max_seconds))
    except Exception:
        high = max(low, 30)

    delay_seconds = random.randint(low, high)
    planned_click_at = (datetime.now() + timedelta(seconds=delay_seconds)).strftime("%H:%M:%S")
    log_line = (
        f"{label} will click in {delay_seconds}s "
        f"(planned click time: {planned_click_at})."
    )
    print(f"[LOG] {log_line}")
    if callable(log_func):
        try:
            log_func(log_line)
        except Exception:
            pass
    if callable(on_delay):
        try:
            on_delay(delay_seconds, planned_click_at)
        except Exception:
            pass
    if delay_seconds > 0:
        _cooperative_wait(page, delay_seconds * 1000, state=state)
    return delay_seconds


ATTACH_MENU_BUTTON_SELECTORS = (
    "button[aria-label='Add file']",
    "button[aria-label='Attach']",
    "button[title='Attach']",
    "div[title='Attach']",
    "div[role='button'][aria-label='Attach']",
    "button[aria-label*='Attach' i]",
    "div[role='button'][aria-label*='Attach' i]",
    "span[data-icon='plus-rounded']",
    "span[data-icon='plus']",
    "span[data-icon='attach-menu-plus']",
    "span[data-icon='clip']",
    "footer button:has(svg title:has-text('ic-plus'))",
    "footer button:has(svg title:has-text('ic-attach'))",
    "footer span[data-icon*='plus']",
    "footer span[data-icon*='attach']",
)

ATTACH_MENU_OPEN_PROBES = (
    "[role='menu']",
    "[role='menuitem']",
    "li[role='button']",
    "input[type='file'][accept*='video']",
    ":text-matches('Photos\\s*(&|and)\\s*videos', 'i')",
    ":text-matches('^Document$', 'i')",
)


def _attachment_menu_is_open(page, timeout_ms=1200):
    deadline = time.time() + max(0.2, float(timeout_ms) / 1000.0)
    while time.time() < deadline:
        for selector in ATTACH_MENU_OPEN_PROBES:
            try:
                locator = page.locator(selector)
                count = min(max(0, locator.count()), 8)
            except Exception:
                continue
            for i in range(count):
                try:
                    item = locator.nth(i)
                    if selector.startswith("input["):
                        # Hidden file inputs exist only while the menu is mounted.
                        return True
                    if item.is_visible():
                        return True
                except Exception:
                    continue
        page.wait_for_timeout(80)
    return False


def _open_attachment_menu(page, state=None):
    if _attachment_menu_is_open(page, timeout_ms=150):
        return True
    for selector in ATTACH_MENU_BUTTON_SELECTORS:
        try:
            locator = page.locator(selector)
            count = min(max(0, locator.count()), 6)
        except Exception:
            continue
        for i in range(count):
            btn = locator.nth(i)
            try:
                if not btn.is_visible():
                    continue
            except Exception:
                continue
            clicked = False
            try:
                clicked = _delayed_click_fallback(
                    page,
                    btn,
                    label="Open + attachment menu",
                    force=True,
                    state=state,
                )
            except Exception:
                clicked = False
            if not clicked:
                try:
                    parent_btn = btn.locator("xpath=ancestor::*[self::button or self::div[@role='button']][1]")
                    clicked = _delayed_click_fallback(
                        page,
                        parent_btn,
                        label="Open + attachment menu (ancestor)",
                        force=True,
                        state=state,
                    )
                except Exception:
                    clicked = False
            if not clicked:
                continue
            page.wait_for_timeout(180)
            if _attachment_menu_is_open(page, timeout_ms=1500):
                _send_log(f"Attachment menu opened via: {selector}")
                return True
            # Click landed but no menu appeared; try the next candidate.
    _send_log("Attachment menu button not found or menu did not open.")
    return False


def _open_document_picker(page, state=None):
    document_selectors = [
        "button[role='menuitem'][aria-label='Document']",
        "div[role='menuitem'][aria-label='Document']",
        "div[role='button'][aria-label='Document']",
        "button[aria-label='Document']",
        "div[role='button']:has-text('Document')",
        "button:has-text('Document')",
    ]
    for selector in document_selectors:
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                if not item.is_visible():
                    continue
                aria_label = (item.get_attribute("aria-label") or "").strip().lower()
                title_attr = (item.get_attribute("title") or "").strip().lower()
                # Avoid clicking preview thumbnail/remove controls that can trigger downloads.
                if "thumbnail" in aria_label or "remove attachment" in aria_label:
                    continue
                if "thumbnail" in title_attr or "remove attachment" in title_attr:
                    continue
                try:
                    if not _delayed_click_fallback(
                        page,
                        item,
                        label="Pick Document option",
                        force=True,
                        state=state,
                    ):
                        raise RuntimeError("Could not click Document option")
                except Exception:
                    try:
                        parent_btn = item.locator("xpath=ancestor::*[self::button or self::div][1]")
                        if not _delayed_click_fallback(
                            page,
                            parent_btn,
                            label="Pick Document option (ancestor)",
                            force=True,
                            state=state,
                        ):
                            raise RuntimeError("Could not click ancestor Document option")
                    except Exception:
                        continue
                page.wait_for_timeout(160)
                return True
        except Exception:
            continue
    return False


def _open_media_picker(page, state=None):
    media_selectors = [
        "button[role='menuitem'][type='button'][aria-label='Photos &amp; videos']",
        "button[role='menuitem'][type='button'][aria-label*='Photos'][aria-label*='videos']",
        "button[role='menuitem'][type='button'][aria-label='Photos & videos']",
        "button[role='menuitem'][aria-label='Photos & videos']",
        "div[role='menuitem'][aria-label='Photos & videos']",
        "div[role='button'][aria-label='Photos & videos']",
        "button[aria-label='Photos & videos']",
        "div[role='button']:has-text('Photos & videos')",
        "button:has-text('Photos & videos')",
        "div[role='button']:has-text('Photos and videos')",
        "button:has-text('Photos and videos')",
        "[role='menuitem']:text-matches('Photos\\s*(&|and)\\s*videos', 'i')",
        "li:text-matches('Photos\\s*(&|and)\\s*videos', 'i')",
        "span:text-matches('^Photos\\s*(&|and)\\s*videos$', 'i')",
    ]
    for selector in media_selectors:
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                if not item.is_visible():
                    continue
                try:
                    if not _delayed_click_fallback(
                        page,
                        item,
                        label="Pick Photos & videos option",
                        force=True,
                        state=state,
                    ):
                        raise RuntimeError("Could not click Photos & videos option")
                except Exception:
                    try:
                        parent_btn = item.locator("xpath=ancestor::*[self::button or self::div][1]")
                        if not _delayed_click_fallback(
                            page,
                            parent_btn,
                            label="Pick Photos & videos option (ancestor)",
                            force=True,
                            state=state,
                        ):
                            raise RuntimeError("Could not click ancestor Photos & videos option")
                    except Exception:
                        continue
                page.wait_for_timeout(160)
                return True
        except Exception:
            continue
    return False


def _open_audio_picker(page, state=None):
    audio_selectors = [
        "button[role='menuitem'][aria-label='Audio']",
        "div[role='menuitem'][aria-label='Audio']",
        "div[role='button'][aria-label='Audio']",
        "button[aria-label='Audio']",
        "div[role='button']:has-text('Audio')",
        "button:has-text('Audio')",
    ]
    for selector in audio_selectors:
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                if not item.is_visible():
                    continue
                try:
                    if not _delayed_click_fallback(
                        page,
                        item,
                        label="Pick Audio option",
                        force=True,
                        state=state,
                    ):
                        raise RuntimeError("Could not click Audio option")
                except Exception:
                    try:
                        parent_btn = item.locator("xpath=ancestor::*[self::button or self::div][1]")
                        if not _delayed_click_fallback(
                            page,
                            parent_btn,
                            label="Pick Audio option (ancestor)",
                            force=True,
                            state=state,
                        ):
                            raise RuntimeError("Could not click ancestor Audio option")
                    except Exception:
                        continue
                page.wait_for_timeout(160)
                return True
        except Exception:
            continue
    return False


def _attachment_menu_item_selectors(label, extra_texts=None):
    values = [str(label or "").strip()]
    for item in list(extra_texts or []):
        text = str(item or "").strip()
        if text and text not in values:
            values.append(text)
    selectors = []
    for text in values:
        selectors.extend(
            [
                f"button[role='menuitem'][type='button'][aria-label='{text}']",
                f"button[role='menuitem'][aria-label='{text}']",
                f"div[role='menuitem'][aria-label='{text}']",
                f"div[role='button'][aria-label='{text}']",
                f"button[aria-label='{text}']",
                f"div[role='button']:has-text('{text}')",
                f"button:has-text('{text}')",
            ]
        )
    return selectors


def _click_attachment_menu_item(page, selectors, label, state=None):
    for selector in list(selectors or []):
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                if not item.is_visible():
                    continue
                try:
                    if not _delayed_click_fallback(
                        page,
                        item,
                        label=label,
                        force=True,
                        state=state,
                    ):
                        raise RuntimeError(f"Could not click {label}")
                except Exception:
                    try:
                        parent_btn = item.locator("xpath=ancestor::*[self::button or self::div][1]")
                        if not _delayed_click_fallback(
                            page,
                            parent_btn,
                            label=f"{label} (ancestor)",
                            force=True,
                            state=state,
                        ):
                            raise RuntimeError(f"Could not click ancestor {label}")
                    except Exception:
                        continue
                page.wait_for_timeout(160)
                return True
        except Exception:
            continue
    return False


def _peek_attachment_menu_item_with_filechooser(
    page,
    selectors,
    label,
    detail_text,
    state=None,
):
    if not _open_attachment_menu(page, state=state):
        return False, ""
    for selector in list(selectors or []):
        locator = page.locator(selector)
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count):
            item = locator.nth(i)
            try:
                if not item.is_visible():
                    continue
                with page.expect_file_chooser(timeout=2200) as chooser_info:
                    if not _delayed_click_fallback(
                        page,
                        item,
                        label=label,
                        force=True,
                        state=state,
                    ):
                        raise RuntimeError(f"Could not click {label}")
                chooser = chooser_info.value
                try:
                    chooser.set_files([])
                except Exception:
                    pass
                _pause_after_optional_action(page, state=state, min_ms=220)
                _dismiss_optional_side_action_overlays(page, state=state)
                return True, detail_text
            except Exception:
                continue
    _dismiss_optional_side_action_overlays(page, state=state)
    return False, ""


def _peek_attachment_menu_item_click_only(
    page,
    selectors,
    label,
    detail_text,
    state=None,
):
    if not _open_attachment_menu(page, state=state):
        return False, ""
    if not _click_attachment_menu_item(page, selectors, label, state=state):
        _dismiss_optional_side_action_overlays(page, state=state)
        return False, ""
    _pause_after_optional_action(page, state=state, min_ms=240)
    _dismiss_optional_side_action_overlays(page, state=state)
    return True, detail_text


def _upload_via_menu_filechooser(page, selectors, attachment_path, label, state=None):
    for selector in selectors:
        locator = page.locator(selector)
        try:
            count = locator.count()
        except Exception:
            count = 0
        for i in range(count):
            item = locator.nth(i)
            try:
                if not item.is_visible():
                    continue
                with page.expect_file_chooser(timeout=2200) as chooser_info:
                    if not _delayed_click_fallback(
                        page,
                        item,
                        label=label,
                        force=True,
                        state=state,
                    ):
                        raise RuntimeError(f"Could not click {label}")
                chooser = chooser_info.value
                chooser.set_files(attachment_path)
                _send_log(f"Uploaded via menu chooser: {label}")
                return True  # file handed to WhatsApp; the caller waits for the preview
            except Exception:
                continue
    return False


def _accept_matches_kind(accept_attr, attachment_kind, attachment_path):
    accept = str(accept_attr or "").lower().replace(" ", "")
    ext = Path(str(attachment_path or "")).suffix.lower()
    if not accept:
        return True

    if attachment_kind == "image":
        return ("image" in accept) or ("*" in accept) or (ext and ext in accept)
    if attachment_kind == "video":
        return ("video" in accept) or ("*" in accept) or (ext and ext in accept)
    if attachment_kind == "audio":
        # Audio uses dedicated audio option in WhatsApp attachment menu.
        return (
            ("audio" in accept)
            or ("*" in accept)
            or (ext and ext in accept)
        )

    # Document/other file mode: never allow media-only pickers like "image/*".
    if "image" in accept or "video" in accept:
        if "application" not in accept and "text" not in accept:
            if ext and ext in accept:
                return True
            # "image/*" or "video/*" should not be used for docs.
            return False
    if not accept:
        return True
    if ext and ext in accept:
        return True
    if "application" in accept or "text" in accept:
        return True
    if "." in accept:
        return True
    if accept == "*/*" or accept == "*":
        return True
    return False


def _attachment_preview_present(page):
    selectors = [
        "[aria-label='Remove attachment']",
        "button[aria-label='Remove attachment']",
        "div[aria-label='Remove attachment']",
        "[aria-label*='Remove attachment' i]",
        "div[role='tab'][aria-label*='thumbnail']",
        "div[contenteditable='true'][aria-label*='caption' i]",
        "div[contenteditable='true'][aria-placeholder*='caption' i]",
    ]
    for selector in selectors:
        try:
            locator = page.locator(selector)
            if locator.count() > 0 and locator.first.is_visible():
                return True
        except Exception:
            continue
    return False


def _wait_for_attachment_ready(page, timeout_ms=3000, poll_ms=80):
    started = time.time()
    timeout_s = max(0.3, float(timeout_ms) / 1000.0)
    poll = max(40, int(poll_ms))
    while (time.time() - started) < timeout_s:
        if _has_unsupported_attachment_notice(page):
            return False
        if _send_button_ready(page) or _attachment_preview_present(page):
            return True
        page.wait_for_timeout(poll)
    return _send_button_ready(page) or _attachment_preview_present(page)


def _wait_for_attachment_preview_to_close(page, timeout_ms=4500, poll_ms=80):
    started = time.time()
    timeout_s = max(0.3, float(timeout_ms) / 1000.0)
    poll = max(40, int(poll_ms))
    while (time.time() - started) < timeout_s:
        if not _attachment_preview_present(page):
            return True
        page.wait_for_timeout(poll)
    return not _attachment_preview_present(page)


def _click_attachment_send(
    page,
    on_delay=None,
    random_delay=True,
    delay_min_seconds=0,
    delay_max_seconds=30,
    state=None,
):
    send_selectors = [
        "div[role='button'][aria-label='Send'][aria-disabled='false']",
        "button[aria-label='Send'][aria-disabled='false']",
        "div[aria-label='Send'][aria-disabled='false']",
        "div[role='button'][aria-label='Send']",
        "button[aria-label='Send']",
        "button[aria-label='Send message']",
        "span[data-icon='wds-ic-send-filled']",
        "span[data-icon='send-filled']",
        "span[data-icon='send']",
        "button:has(svg title:has-text('ic-send'))",
        "div[role='button']:has(svg title:has-text('ic-send'))",
    ]
    delay_applied = False
    for _ in range(8):
        for selector in send_selectors:
            locator = page.locator(selector)
            try:
                count = locator.count()
            except Exception:
                count = 0
            for i in range(count):
                item = locator.nth(i)
                try:
                    if not item.is_visible():
                        continue
                    if not delay_applied:
                        min_seconds = max(0, int(delay_min_seconds or 0))
                        max_seconds = max(min_seconds, int(delay_max_seconds or min_seconds))
                        if not random_delay:
                            max_seconds = min_seconds
                        _wait_random_send_delay(
                            page,
                            min_seconds,
                            max_seconds,
                            log_func=_send_log,
                            label="Attachment send button",
                            on_delay=on_delay,
                            state=state,
                        )
                        delay_applied = True
                    if not _delayed_click_fallback(
                        page,
                        item,
                        label="Click Send",
                        force=True,
                        state=state,
                    ):
                        continue
                    # If click succeeded on a visible Send control, treat as sent.
                    # WhatsApp preview UI can linger briefly and previously caused false "failed" status.
                    page.wait_for_timeout(120)
                    return True
                except Exception:
                    continue
        page.wait_for_timeout(70)
    return False


def _find_attachment_input(page, attachment_kind, attachment_path):
    candidates = [
        "div[role='dialog'] input[type='file']",
        "input[type='file']",
    ]
    fallback = None
    for selector in candidates:
        locator = page.locator(selector)
        try:
            for i in range(locator.count()):
                item = locator.nth(i)
                accept_attr = ""
                try:
                    accept_attr = item.get_attribute("accept") or ""
                except Exception:
                    accept_attr = ""
                if _accept_matches_kind(accept_attr, attachment_kind, attachment_path):
                    accept_norm = str(accept_attr or "").lower()
                    if attachment_kind == "file":
                        # For documents, prefer explicit document/file inputs and avoid media pickers.
                        if any(token in accept_norm for token in ["application", "text", ".", "*"]):
                            print(f"[LOG] Using file input accept='{accept_attr}' for document upload")
                            return item
                        if not accept_norm and fallback is None:
                            fallback = item
                            continue
                    else:
                        # For media, only accept media-style inputs. Avoid generic/document
                        # pickers and the image-only "New sticker" input.
                        if _file_input_belongs_to_sticker_item(item):
                            continue
                        if attachment_kind == "audio":
                            if "audio" in accept_norm:
                                _send_log(f"Using file input accept='{accept_attr}' for audio upload")
                                return item
                        elif "image" in accept_norm and "video" in accept_norm:
                            _send_log(f"Using file input accept='{accept_attr}' for media upload")
                            return item
                        if not accept_norm and fallback is None:
                            fallback = item
                            continue
                        # Skip wildcard/document-style input for media to force Photos & videos path.
                        continue
        except Exception:
            continue
    if attachment_kind == "file":
        return fallback
    return None


def _has_unsupported_attachment_notice(page):
    notices = [
        "file you tried adding is not supported",
        "files you tried adding are not supported",
        "not supported",
    ]
    for text in notices:
        try:
            if page.locator(f"text={text}").first.is_visible(timeout=300):
                return True
        except Exception:
            continue
    return False


def _file_input_belongs_to_sticker_item(file_input):
    """True when the hidden file input sits under a 'sticker' menu entry."""
    try:
        text_value = str(
            file_input.evaluate(
                "(el) => { let n = el; for (let i = 0; i < 6 && n; i++) { "
                "const t = (n.innerText || n.textContent || '') + ' ' + (n.getAttribute && (n.getAttribute('aria-label') || '')); "
                "if (/sticker/i.test(t)) return 'sticker'; n = n.parentElement; } return ''; }"
            )
            or ""
        )
        return text_value == "sticker"
    except Exception:
        return False


def _await_attachment_preview(page, timeout_ms=20000):
    """After a file was handed to WhatsApp, wait for the preview/send UI.

    Rendering the preview of a multi-megabyte image on a small VM can take
    well over 5 seconds. Returning too early made the caller think the upload
    failed and try to reopen a menu that was already gone."""
    started = time.time()
    if _wait_for_attachment_ready(page, timeout_ms=timeout_ms, poll_ms=150):
        _send_log(f"Attachment preview ready after {time.time() - started:.1f}s.")
        return True
    if _has_unsupported_attachment_notice(page):
        _send_log("WhatsApp rejected the attachment as unsupported.")
        return False
    _send_log(f"Attachment preview did not appear within {timeout_ms / 1000:.0f}s after upload.")
    return False


def _set_files_on_menu_input(page, attachment_kind, attachment_path):
    """WhatsApp mounts hidden <input type=file> elements under the attach menu.
    Setting files directly on the right one is the most reliable upload path."""
    try:
        inputs = page.locator("input[type='file']")
        count = min(max(0, inputs.count()), 12)
    except Exception:
        return False
    chosen = None
    for i in range(count):
        item = inputs.nth(i)
        try:
            accept = str(item.get_attribute("accept") or "").lower().replace(" ", "")
        except Exception:
            accept = ""
        if _file_input_belongs_to_sticker_item(item):
            continue
        if attachment_kind in ("image", "video"):
            # The "Photos & videos" input accepts BOTH image and video types.
            # An image-only input is the "New sticker" picker: never use it,
            # stickers cannot carry captions.
            if "image" in accept and "video" in accept:
                chosen = item
                break
        elif attachment_kind == "audio":
            if "audio" in accept:
                chosen = item
                break
        else:
            # Document: prefer the wildcard / non-media input.
            if accept in ("*", "*/*") or ("image" not in accept and "video" not in accept and "audio" not in accept):
                chosen = item
                break
    if chosen is None:
        return False
    try:
        chosen.set_input_files(attachment_path)
    except Exception as exc:
        _send_log(f"Direct file input upload failed: {exc}")
        return False
    _send_log(f"Attachment uploaded via hidden file input ({attachment_kind}).")
    return _await_attachment_preview(page)


MEDIA_MENU_ITEM_SELECTORS = (
    "button[role='menuitem'][type='button'][aria-label='Photos &amp; videos']",
    "button[role='menuitem'][type='button'][aria-label*='Photos'][aria-label*='videos']",
    "button[role='menuitem'][type='button'][aria-label='Photos & videos']",
    "button[role='menuitem'][aria-label='Photos & videos']",
    "div[role='menuitem'][aria-label='Photos & videos']",
    "[role='menuitem']:text-matches('Photos\\s*(&|and)\\s*videos', 'i')",
    "li:text-matches('Photos\\s*(&|and)\\s*videos', 'i')",
    "div[role='button']:text-matches('Photos\\s*(&|and)\\s*videos', 'i')",
    "button:text-matches('Photos\\s*(&|and)\\s*videos', 'i')",
    "span:text-matches('^Photos\\s*(&|and)\\s*videos$', 'i')",
)

DOCUMENT_MENU_ITEM_SELECTORS = (
    "button[role='menuitem'][type='button'][aria-label='Document']",
    "button[role='menuitem'][aria-label='Document']",
    "div[role='menuitem'][aria-label='Document']",
    "button[aria-label='Document']",
    "[role='menuitem']:text-matches('^Document$', 'i')",
    "li:text-matches('^Document$', 'i')",
    "div[role='button']:text-matches('^Document$', 'i')",
    "span:text-matches('^Document$', 'i')",
)

AUDIO_MENU_ITEM_SELECTORS = (
    "button[role='menuitem'][type='button'][aria-label='Audio']",
    "button[role='menuitem'][aria-label='Audio']",
    "div[role='menuitem'][aria-label='Audio']",
    "button[aria-label='Audio']",
    "[role='menuitem']:text-matches('^Audio$', 'i')",
    "li:text-matches('^Audio$', 'i')",
    "div[role='button']:text-matches('^Audio$', 'i')",
    "span:text-matches('^Audio$', 'i')",
)


def _attach_file_via_plus_button(page, attachment_path, attachment_kind, state=None):
    if not _open_attachment_menu(page, state=state):
        return False

    # Fastest and most robust: feed the hidden input that belongs to the menu.
    if _set_files_on_menu_input(page, attachment_kind, attachment_path):
        return True

    if attachment_kind == "file":
        _send_log("Choosing Document picker for file upload")
        if _upload_via_menu_filechooser(
            page,
            list(DOCUMENT_MENU_ITEM_SELECTORS),
            attachment_path,
            "Pick Document option",
            state=state,
        ):
            return _await_attachment_preview(page)
        _open_document_picker(page, state=state)
    elif attachment_kind == "audio":
        _send_log("Choosing Audio picker for audio upload")
        if _upload_via_menu_filechooser(
            page,
            list(AUDIO_MENU_ITEM_SELECTORS),
            attachment_path,
            "Pick Audio option",
            state=state,
        ):
            return _await_attachment_preview(page)
        if not _open_audio_picker(page, state=state):
            _send_log("Audio picker not found; refusing to use document flow for audio")
            return False
    else:
        _send_log("Choosing Photos & videos picker for image/video upload")
        if _upload_via_menu_filechooser(
            page,
            list(MEDIA_MENU_ITEM_SELECTORS),
            attachment_path,
            "Pick Photos & videos option",
            state=state,
        ):
            return _await_attachment_preview(page)
        if not _open_media_picker(page, state=state):
            _send_log("Photos & videos menu item not found in attachment menu.")
            return False
    file_input = _find_attachment_input(page, attachment_kind, attachment_path)
    if file_input is None:
        if attachment_kind != "file":
            _send_log("No media file input found after opening Photos & videos")
        return False
    try:
        file_input.set_input_files(attachment_path)
    except Exception as exc:
        _send_log(f"set_input_files failed: {exc}")
        return False
    return _await_attachment_preview(page)


def _paste_attachment_into_composer(page, state=None):
    if not _focus_compose_box(page, state=state):
        return False
    for _ in range(2):
        try:
            pyautogui.hotkey("ctrl", "v")
        except Exception:
            page.keyboard.press("Control+V")
        page.wait_for_timeout(240)
        if _has_unsupported_attachment_notice(page):
            return False
        if _wait_for_attachment_ready(page, timeout_ms=2000, poll_ms=80):
            return True
    return False


def send_single_attachment(
    page,
    attachment,
    caption="",
    on_delay=None,
    on_send_delay=None,
    random_delay=True,
    delay_min_seconds=0,
    delay_max_seconds=30,
    send_click_random_delay=True,
    send_click_delay_min_seconds=0,
    send_click_delay_max_seconds=30,
    typing_pause_chars_min=0,
    typing_pause_chars_max=0,
    typing_pause_delay_min_ms=0,
    typing_pause_delay_max_ms=0,
    remove_chars_during_typing_enabled=False,
    remove_chars_times_min=0,
    remove_chars_times_max=0,
    remove_chars_count_min=0,
    remove_chars_count_max=0,
    state=None,
):
    attachment_path = str(attachment.get("path") or "").strip()
    if not attachment_path:
        raise RuntimeError("Attachment path is missing")
    attachment_kind = str(attachment.get("kind") or detect_attachment_kind(attachment_path))
    _send_log(f"Attaching '{Path(attachment_path).name}' as kind='{attachment_kind}'")
    try:
        low = max(0, int(delay_min_seconds or 0))
    except Exception:
        low = 0
    try:
        high = max(low, int(delay_max_seconds or low))
    except Exception:
        high = low
    if not random_delay:
        high = low
    attach_delay_seconds = random.randint(low, high)
    _send_log(f"[media will attach in delay :{attach_delay_seconds} sec]")
    if callable(on_delay):
        try:
            planned_at = (datetime.now() + timedelta(seconds=attach_delay_seconds)).strftime("%H:%M:%S")
            on_delay(attach_delay_seconds, planned_at)
        except Exception:
            pass
    if attach_delay_seconds > 0:
        _cooperative_wait(page, attach_delay_seconds * 1000, state=state)

    _wait_if_send_paused(state, page=page)
    attached = _attach_file_via_plus_button(page, attachment_path, attachment_kind, state=state)
    if not attached:
        if attachment_kind == "file":
            copy_file_to_windows_clipboard(attachment_path)
            attached = _paste_attachment_into_composer(page, state=state)
        else:
            # For media (image/video/audio), never fall back to generic paste,
            # because WhatsApp may treat it as document and skip media preview.
            raise RuntimeError(
                f"Could not attach {attachment_kind} via Photos & videos picker "
                f"(menu/file input not found, or the preview never appeared - see log lines above)."
            )
    if not attached:
        if _has_unsupported_attachment_notice(page):
            raise RuntimeError(
                f"WhatsApp rejected attachment as unsupported: {Path(attachment_path).name}"
            )
        raise RuntimeError("Could not attach file into WhatsApp chat")

    # Wait for preview/send UI so we do not press send before upload is ready.
    preview_selectors = [
        "[aria-label='Remove attachment']",
        "[aria-label*='Remove' i][role='button']",
        "div[role='button'][aria-label='Send']",
        "button[aria-label='Send']",
        "span[data-icon='wds-ic-send-filled']",
        "span[data-icon='send-filled']",
        "span[data-icon='send']",
        "button:has(svg title:has-text('ic-send'))",
    ]
    preview_ready = False
    for selector in preview_selectors:
        try:
            page.wait_for_selector(selector, timeout=1000)
            preview_ready = True
            break
        except Exception:
            continue
    if not preview_ready:
        print("[WARN] File preview not detected before send. Trying to send anyway.")

    _fill_caption_if_available(
        page,
        caption,
        typing_pause_chars_min=typing_pause_chars_min,
        typing_pause_chars_max=typing_pause_chars_max,
        typing_pause_delay_min_ms=typing_pause_delay_min_ms,
        typing_pause_delay_max_ms=typing_pause_delay_max_ms,
        remove_chars_during_typing_enabled=remove_chars_during_typing_enabled,
        remove_chars_times_min=remove_chars_times_min,
        remove_chars_times_max=remove_chars_times_max,
        remove_chars_count_min=remove_chars_count_min,
        remove_chars_count_max=remove_chars_count_max,
        state=state,
    )

    if not _click_attachment_send(
        page,
        on_delay=on_send_delay if callable(on_send_delay) else on_delay,
        random_delay=send_click_random_delay,
        delay_min_seconds=send_click_delay_min_seconds,
        delay_max_seconds=send_click_delay_max_seconds,
        state=state,
    ):
        raise RuntimeError("Send button not found for attachment after waiting for preview.")
    _wait_for_attachment_preview_to_close(page, timeout_ms=4500, poll_ms=90)
    page.wait_for_timeout(160)


def _normalized_text(value):
    return " ".join(str(value or "").strip().lower().split())


def _fill_search_field(page, locator, text_value, state=None):
    """Focus the search field, clear it, and type the target exactly once.

    A second typing pass happens only if the field is verified empty after
    the first pass, so the number can never be appended twice.
    """
    text_value = str(text_value or "")
    focused = False
    try:
        focused = _human_click_item(
            page,
            locator,
            label="Focus search field",
            timeout_ms=1500,
            force=True,
            state=state,
        )
    except Exception:
        focused = False
    if not focused:
        try:
            focused = _delayed_click_fallback(
                page,
                locator,
                label="Focus search field",
                timeout_ms=1500,
                force=True,
                state=state,
            )
        except Exception:
            focused = False
    if not focused:
        try:
            locator.click(force=True, timeout=1500)
            focused = True
        except Exception:
            return False

    _clear_search_field(page, locator)
    _human_pause(page, min_ms=25, max_ms=75, state=state, step_ms=60)
    if not text_value:
        return True

    # Already contains the number (e.g. a previous attempt) -> do not retype.
    if _search_field_has_text(locator, text_value):
        return True

    try:
        if not _human_type_text_basic(page, text_value, state=state, min_delay_ms=26, max_delay_ms=52):
            page.keyboard.type(text_value, delay=8)
    except Exception:
        try:
            page.keyboard.type(text_value, delay=8)
        except Exception:
            pass
    page.wait_for_timeout(180)
    if _search_field_has_text(locator, text_value):
        return True

    current = _search_field_value(locator)
    if current.strip():
        # Something is in the field but it is not our number. Clear it and
        # retry once; if it will not clear, give up rather than append.
        if not _clear_search_field(page, locator):
            _send_log("Search field could not be cleared; skipping retype to avoid duplicates.")
            return False
    _send_log("Search text did not register after typing; retrying once with direct input.")
    try:
        locator.click(force=True, timeout=1200)
        page.keyboard.insert_text(text_value)
    except Exception:
        try:
            locator.fill(text_value, timeout=2000)
        except Exception:
            pass
    page.wait_for_timeout(180)
    if _search_field_has_text(locator, text_value):
        return True
    # Could not read the field back at all; assume the single typing pass
    # landed and let the result matching decide.
    return True


def _clear_search_field(page, locator):
    """Empty the search field and return True only when verified empty."""
    for _ in range(3):
        if not _search_field_value(locator).strip():
            return True
        try:
            locator.click(force=True, timeout=1200)
        except Exception:
            pass
        try:
            page.keyboard.press("Control+A")
            page.keyboard.press("Backspace")
        except Exception:
            pass
        page.wait_for_timeout(80)
        if not _search_field_value(locator).strip():
            return True
        try:
            locator.evaluate(
                "(el) => { if (el.tagName === 'INPUT' || el.tagName === 'TEXTAREA') { el.value = ''; "
                "el.dispatchEvent(new Event('input', {bubbles: true})); } else { el.focus(); "
                "document.execCommand('selectAll', false, null); document.execCommand('delete', false, null); } }"
            )
        except Exception:
            pass
        page.wait_for_timeout(80)
        if not _search_field_value(locator).strip():
            return True
        try:
            remaining = len(_search_field_value(locator))
            page.keyboard.press("End")
            for _i in range(min(remaining + 2, 80)):
                page.keyboard.press("Backspace")
        except Exception:
            pass
        page.wait_for_timeout(80)
    return not _search_field_value(locator).strip()


def _digits_only(value):
    return "".join(ch for ch in str(value or "") if ch.isdigit())


def _read_current_chat_title(page):
    ignored = {
        "menu",
        "search",
        "voice call",
        "video call",
        "attach",
        "search or start new chat",
    }
    selectors = (
        "header [data-testid='conversation-info-header-chat-title']",
        "header span[title]",
        "header h1 span",
        "header div[title]",
    )
    for selector in selectors:
        locator = page.locator(selector)
        try:
            count = min(max(0, locator.count()), 6)
        except Exception:
            count = 0
        for i in range(count):
            item = locator.nth(i)
            try:
                if not item.is_visible():
                    continue
            except Exception:
                continue
            text_value = ""
            try:
                text_value = str(item.get_attribute("title") or "").strip()
            except Exception:
                text_value = ""
            if not text_value:
                try:
                    text_value = str(item.inner_text() or "").strip()
                except Exception:
                    text_value = ""
            if not text_value:
                continue
            if str(text_value).strip().lower() in ignored:
                continue
            return text_value
    return ""


def _row_text(row):
    try:
        title_node = row.locator("span[title]").first
        if title_node.count() > 0:
            title = str(title_node.get_attribute("title") or "").strip()
            if title:
                return title
    except Exception:
        pass
    try:
        return str(row.inner_text() or "").strip()
    except Exception:
        return ""


def _is_left_pane_click_target(page, item):
    try:
        box = item.bounding_box()
    except Exception:
        box = None
    if not box:
        return False
    try:
        viewport_width = None
        try:
            vp = page.viewport_size
            if vp:
                viewport_width = int(vp.get("width") or 0)
        except Exception:
            viewport_width = None
        if not viewport_width:
            viewport_width = int(
                page.evaluate(
                    "() => window.innerWidth || document.documentElement.clientWidth || 1366"
                )
            )
        center_x = float(box.get("x", 0)) + (float(box.get("width", 0)) / 2.0)
        # WhatsApp layout is two-pane; left pane stays on left half.
        return center_x <= (viewport_width * 0.62)
    except Exception:
        return False


# WhatsApp Web keeps renaming/restructuring the search field (plain <input>
# vs. contenteditable <div>, different aria-labels). Try every known variant
# and fall back to opening the search/new-chat pane if nothing is visible yet.
SEARCH_BOX_SELECTORS = (
    "input[aria-label='Search name or number']",
    "input[aria-label*='Search' i]",
    "input[placeholder*='Search' i]",
    "div[role='textbox'][contenteditable='true'][aria-label*='Search' i]",
    "div[contenteditable='true'][aria-label='Search input textbox']",
    "div[contenteditable='true'][aria-placeholder*='Search' i]",
    "div[contenteditable='true'][title*='Search' i]",
    "div[contenteditable='true'][data-tab='3']",
    "#side div[contenteditable='true'][role='textbox']",
    "div[contenteditable='true'][role='textbox'][data-lexical-editor='true']:not([aria-label='Type a message'])",
)

SEARCH_OPENER_SELECTORS = (
    "button[aria-label*='Search' i]",
    "div[role='button'][aria-label*='Search' i]",
    "span[data-icon='search']",
    "span[data-icon='search-refreshed']",
    "button[aria-label*='New chat' i]",
    "div[role='button'][aria-label*='New chat' i]",
    "span[data-icon='new-chat-outline']",
)


def _first_visible_search_box(page):
    for selector in SEARCH_BOX_SELECTORS:
        try:
            locator = page.locator(selector)
            count = min(max(0, locator.count()), 6)
        except Exception:
            continue
        for i in range(count):
            item = locator.nth(i)
            try:
                if not item.is_visible():
                    continue
                # Never confuse the message composer (right pane) with search.
                if item.locator("xpath=ancestor::footer").count() > 0:
                    continue
                if not _is_left_pane_click_target(page, item):
                    continue
            except Exception:
                continue
            return item, selector
    return None, ""


def _locate_search_box(page, state=None, timeout_ms=5000):
    timeout_ms = max(500, int(timeout_ms or 0))
    deadline = time.time() + (timeout_ms / 1000.0)
    opener_tried = False
    while True:
        box, selector = _first_visible_search_box(page)
        if box is not None:
            try:
                _send_log(f"Search box located via: {selector}")
            except Exception:
                pass
            return box
        if not opener_tried and (deadline - time.time()) < (timeout_ms / 2000.0):
            # Half the budget is gone and nothing is visible: the search pane
            # may be closed. Try to open it, then keep polling.
            opener_tried = True
            for opener in SEARCH_OPENER_SELECTORS:
                try:
                    item = page.locator(opener).first
                    if item.count() > 0 and item.is_visible():
                        item.click(timeout=1200)
                        break
                except Exception:
                    continue
        if time.time() >= deadline:
            try:
                _send_log("Search box not found with any known selector.")
            except Exception:
                pass
            return None
        _cooperative_wait(page, 150, state=state)


def _search_field_value(locator):
    try:
        return str(
            locator.evaluate(
                "(el) => {"
                "  const read = (n) => { if (!n) return ''; "
                "    if (n.tagName === 'INPUT' || n.tagName === 'TEXTAREA') return n.value || ''; "
                "    return n.innerText || n.textContent || ''; };"
                "  let v = read(el);"
                "  if (!v.trim()) v = read(document.activeElement);"
                "  if (!v.trim()) {"
                "    const scope = document.querySelector('#side') || el.parentElement || document;"
                "    const cand = scope.querySelector(\"input, [contenteditable='true'], [role='textbox']\");"
                "    v = read(cand);"
                "  }"
                "  return v;"
                "}"
            )
            or ""
        )
    except Exception:
        return ""


def _search_field_has_text(locator, text_value):
    expected = str(text_value or "").strip()
    if not expected:
        return True
    current = _search_field_value(locator)
    expected_digits = _digits_only(expected)
    if expected_digits:
        return expected_digits in _digits_only(current)
    return _normalized_text(expected) in _normalized_text(current)


def _phone_digits_match(target_digits, candidate_digits):
    target_digits = str(target_digits or "")
    candidate_digits = str(candidate_digits or "")
    if not target_digits or not candidate_digits:
        return False
    if target_digits in candidate_digits or candidate_digits in target_digits:
        return True
    # Tolerate country-code / leading-zero differences: compare the tail.
    tail = min(len(target_digits), len(candidate_digits), 10)
    return tail >= 9 and target_digits[-tail:] == candidate_digits[-tail:]


SEARCH_RESULT_TITLE_SELECTORS = (
    "[data-testid='cell-frame-title'] span[title]",
    "div[data-testid='cell-frame-container'] span[title]",
    "div[role='button'][data-tab='4'] span[title]",
    "div[role='listitem'] span[title]",
    "div[role='row'] span[title]",
    "#side span[title]",
)

SEARCH_RESULT_ROW_XPATH = (
    "xpath=ancestor::*[@data-testid='cell-frame-container' or @role='button' "
    "or @role='listitem' or @role='row'][1]"
)


def _find_search_result_by_title(page, target_text):
    """Return (title_span, row_container, title_text) for the visible search
    result whose displayed number matches target_text, or (None, None, [])
    plus the list of visible titles seen (for logging)."""
    target_digits = _digits_only(target_text)
    seen_titles = []
    for selector in SEARCH_RESULT_TITLE_SELECTORS:
        try:
            spans = page.locator(selector)
            count = min(max(0, spans.count()), 40)
        except Exception:
            continue
        for i in range(count):
            span = spans.nth(i)
            try:
                if not span.is_visible():
                    continue
                title = str(span.get_attribute("title") or "").strip()
                if not title:
                    title = str(span.inner_text() or "").strip()
                if not title or title in seen_titles:
                    continue
                if not _is_left_pane_click_target(page, span):
                    continue
                seen_titles.append(title)
                title_digits = _digits_only(title)
                matched = False
                if target_digits and title_digits:
                    matched = _phone_digits_match(target_digits, title_digits)
                elif target_text and title:
                    matched = str(target_text).strip().lower() in title.lower()
                if matched:
                    try:
                        row = span.locator(SEARCH_RESULT_ROW_XPATH).first
                        if row.count() == 0:
                            row = None
                    except Exception:
                        row = None
                    return span, row, seen_titles
            except Exception:
                continue
    return None, None, seen_titles


def _open_contact_chat_from_search(page, target, state=None, result_click_delay_seconds=0):
    target_text = str(target or "").strip()
    if not target_text:
        return False, "Empty target"
    try:
        result_click_delay_seconds = max(0, int(result_click_delay_seconds or 0))
    except Exception:
        result_click_delay_seconds = 0

    # Every number starts from the full chat list: an optional action may have
    # left the Groups / Unread / … filter selected.
    try:
        _click_all_filter_before_search(page, state=state)
    except Exception:
        pass

    search_box = _locate_search_box(page, state=state, timeout_ms=5000)
    if search_box is None:
        return False, "Search box not found"
    previous_chat_title = _read_current_chat_title(page)
    _wait_if_send_paused(state, page=page)
    if not _fill_search_field(page, search_box, target_text, state=state):
        return False, "Could not fill search field"

    # Keep the searched number in the search box and wait until a matching
    # left-pane result actually appears before clicking it.
    _cooperative_wait(page, 260, state=state)
    # WhatsApp renders search results with varying roles depending on the
    # build (listitem / row / gridcell / button). Match broadly and filter
    # by position + text instead of relying on one DOM shape.
    results = page.locator(
        "div[role='listitem'], "
        "div[role='row'], "
        "div[role='gridcell'], "
        "div[data-testid='cell-frame-container'], "
        "div[role='button'][tabindex='0'], "
        "#side span[title]"
    )
    invalid_text = page.locator(
        "span:has-text('not on WhatsApp'), "
        "div:has-text('isn\u2019t on WhatsApp'), "
        "div:has-text(\"isn't on WhatsApp\"), "
        "span:has-text('No results found'), "
        "span:has-text('No chats, contacts or messages found')"
    )
    target_digits = _digits_only(target_text)
    started = time.time()
    timeout_s = 9.0
    result_click_delay_applied = False
    enter_tried = False
    titles_logged = False

    def _is_search_input(item):
        try:
            if item.locator("input, [contenteditable='true'], [role='textbox']").count() > 0:
                return True
            tag = str(item.evaluate("(el) => el.tagName") or "").upper()
            if tag in ("INPUT", "TEXTAREA"):
                return True
            if item.evaluate("(el) => el.isContentEditable === true || el.getAttribute('role') === 'textbox'"):
                return True
        except Exception:
            pass
        return False

    def _row_matches(item):
        row_text = _row_text(item)
        row_digits = _digits_only(row_text)
        if target_digits and row_digits:
            return _phone_digits_match(target_digits, row_digits)
        return bool(target_text) and target_text.lower() in str(row_text or "").lower()

    def _wait_chat_ready(matched_row_text, matched_digits):
        ready_started = time.time()
        composer_ready = False
        search_hidden = False
        current_title = ""
        while (time.time() - ready_started) < 6.5:
            _wait_if_send_paused(state, page=page)
            composer = _locate_active_message_box(page, timeout_ms=800)
            composer_ready = composer is not None
            search_hidden = False
            try:
                search_hidden = (search_box.count() == 0) or (not search_box.is_visible())
            except Exception:
                search_hidden = True
            current_title = _read_current_chat_title(page)
            current_digits = _digits_only(current_title)
            title_match = False
            if target_digits and current_digits:
                title_match = _phone_digits_match(target_digits, current_digits)
            elif matched_digits and current_digits:
                title_match = (matched_digits in current_digits) or (current_digits in matched_digits)
            elif matched_row_text and current_title:
                title_match = matched_row_text.lower() in current_title.lower()
            elif target_text and current_title:
                title_match = target_text.lower() in current_title.lower()
            title_changed = bool(current_title) and current_title != previous_chat_title
            elapsed = time.time() - ready_started
            if composer_ready and search_hidden and (title_match or title_changed or not previous_chat_title):
                return True
            if composer_ready and title_match and elapsed >= 1.2:
                # Search bar stayed visible (main chat-list search) but the
                # header already shows the target number: chat is open.
                return True
            if composer_ready and title_changed and search_hidden and elapsed >= 1.5:
                return True
            _cooperative_wait(page, 160, state=state)
        try:
            _send_log(
                f"Chat not ready for {target_text}: composer={composer_ready} search_hidden={search_hidden} "
                f"title='{current_title}' prev='{previous_chat_title}'"
            )
        except Exception:
            pass
        return False

    while (time.time() - started) < timeout_s:
        _wait_if_send_paused(state, page=page)
        try:
            # "No results" can flash while WhatsApp is still searching; only
            # trust it once the search has settled and no result row exists.
            if (time.time() - started) >= 1.5 and invalid_text.count() > 0 and invalid_text.first.is_visible():
                probe_span, _probe_row, probe_titles = _find_search_result_by_title(page, target_text)
                if probe_span is None and not probe_titles:
                    try:
                        _send_log(f"Not on WhatsApp for {target_text}: '{str(invalid_text.first.inner_text() or '')[:80]}'")
                    except Exception:
                        pass
                    return False, "Not on WhatsApp"
        except Exception:
            pass

        matched_item = None
        saw_visible_rows = False
        visible_row_count = 0
        fallback_item = None

        # Preferred path: WhatsApp renders each result as
        # div[data-testid='cell-frame-container'][role='button'] containing a
        # [data-testid='cell-frame-title'] span[title='+92 3xx xxxxxxx'].
        # Click the title span itself so we never hit "Add to contacts".
        title_span, title_row, seen_titles = _find_search_result_by_title(page, target_text)
        if title_span is not None:
            saw_visible_rows = True
            visible_row_count = 1
            matched_item = title_span
            fallback_item = title_row
        elif seen_titles and not titles_logged:
            titles_logged = True
            _send_log(f"Search results visible for {target_text} (no digit match yet): {seen_titles[:5]}")

        try:
            result_count = 0 if matched_item is not None else min(max(0, results.count()), 60)
        except Exception:
            result_count = 0
        for i in range(result_count):
            item = results.nth(i)
            try:
                if not item.is_visible() or not _is_left_pane_click_target(page, item):
                    continue
                if _is_search_input(item):
                    continue
                saw_visible_rows = True
                visible_row_count += 1
                if fallback_item is None:
                    fallback_item = item
                if _row_matches(item):
                    matched_item = item
                    break
            except Exception:
                continue

        if matched_item is None and fallback_item is not None and visible_row_count == 1:
            if (time.time() - started) >= 1.2:
                matched_item = fallback_item

        if matched_item is not None:
            matched_row_text = str(_row_text(matched_item) or "").strip()
            matched_digits = _digits_only(matched_row_text)
            if result_click_delay_seconds > 0 and not result_click_delay_applied:
                _wait_random_send_delay(
                    page,
                    result_click_delay_seconds,
                    result_click_delay_seconds,
                    log_func=_send_log,
                    label=f"Search result for {target_text}",
                    state=state,
                )
                result_click_delay_applied = True
            _send_log(f"Clicking search result row for {target_text}: {matched_row_text[:60]}")
            if not _human_click_item(
                page,
                matched_item,
                label=f"Open search result {target_text}",
                timeout_ms=1800,
                state=state,
            ):
                _delayed_dom_click_fallback(
                    page,
                    matched_item,
                    label=f"Open search result {target_text}",
                    state=state,
                )
            if _wait_chat_ready(matched_row_text, matched_digits):
                return True, ""
            # Click did not open the chat: DOM-click the row container itself.
            row_container = fallback_item if fallback_item is not None else matched_item
            try:
                _send_log(f"Row click did not open chat for {target_text}; DOM-clicking result container.")
                row_container.evaluate(
                    "(el) => { el.scrollIntoView({block: 'center'}); "
                    "const opts = {bubbles: true, cancelable: true, view: window}; "
                    "el.dispatchEvent(new MouseEvent('mousedown', opts)); "
                    "el.dispatchEvent(new MouseEvent('mouseup', opts)); "
                    "el.click(); }"
                )
            except Exception:
                try:
                    row_container.click(force=True, timeout=1500)
                except Exception:
                    pass
            if _wait_chat_ready(matched_row_text, matched_digits):
                return True, ""
            # Still not open: try pressing Enter in the search box.
            if not enter_tried:
                enter_tried = True
                try:
                    _send_log(f"Row click did not open chat for {target_text}; pressing Enter in search box.")
                    search_box.click(force=True, timeout=1200)
                    page.keyboard.press("Enter")
                except Exception:
                    pass
                if _wait_chat_ready(matched_row_text, matched_digits):
                    return True, ""
        elif not enter_tried and (time.time() - started) >= 2.5:
            # Nothing matched our selectors but the user can see a result:
            # Enter opens the top search result in WhatsApp Web.
            enter_tried = True
            try:
                _send_log(f"No result row matched selectors for {target_text}; pressing Enter in search box.")
                search_box.click(force=True, timeout=1200)
                page.keyboard.press("Enter")
            except Exception:
                pass
            if _wait_chat_ready("", ""):
                return True, ""

        if not saw_visible_rows:
            _cooperative_wait(page, 220, state=state)
        else:
            _cooperative_wait(page, 180, state=state)

    try:
        if invalid_text.count() > 0 and invalid_text.first.is_visible():
            probe_span, _probe_row, probe_titles = _find_search_result_by_title(page, target_text)
            if probe_span is None and not probe_titles:
                return False, "Not on WhatsApp"
    except Exception:
        pass
    return False, "No matching result found"


def _is_whatsapp_profile_logged_in(page):
    try:
        selector = _wait_for_any_selector(
            page,
            (
                "[data-testid='chat-list']",
                "div[aria-label='Chat list']",
                "#pane-side",
            ),
            timeout_ms=1500,
        )
        return bool(selector)
    except Exception:
        return False


def _is_whatsapp_login_required(page):
    try:
        selector = _wait_for_any_selector(
            page,
            (
                "canvas[aria-label*='Scan']",
                "canvas[aria-label*='scan']",
                "[data-testid='qrcode']",
                "[data-testid='qr-code']",
                "text='Log into WhatsApp Web'",
                "text='Keep your phone connected'",
            ),
            timeout_ms=800,
        )
        return bool(selector)
    except Exception:
        return False


def _send_plain_text_single_message(
    page,
    msg_box,
    text_value,
    on_delay=None,
    random_delay=True,
    delay_min_seconds=0,
    delay_max_seconds=30,
    typing_pause_chars_min=0,
    typing_pause_chars_max=0,
    typing_pause_delay_min_ms=0,
    typing_pause_delay_max_ms=0,
    remove_chars_during_typing_enabled=False,
    remove_chars_times_min=0,
    remove_chars_times_max=0,
    remove_chars_count_min=0,
    remove_chars_count_max=0,
    state=None,
):
    text_value = str(text_value or "")
    normalized = text_value.replace("\r\n", "\n").replace("\r", "\n")
    prepared_box = _prepare_active_message_box(page, state=state, timeout_ms=2200)
    if prepared_box is not None:
        msg_box = prepared_box
    try:
        if not _human_click_item(
            page,
            msg_box.first,
            label="Focus message box",
            timeout_ms=1800,
            force=True,
            state=state,
        ):
            _delayed_click_fallback(
                page,
                msg_box.first,
                label="Focus message box",
                timeout_ms=1800,
                force=True,
                state=state,
            )
    except Exception:
        try:
            fallback_box = page.locator("div[contenteditable='true'][role='textbox']").first
            if not _human_click_item(
                page,
                fallback_box,
                label="Focus message box",
                timeout_ms=1800,
                force=True,
                state=state,
            ):
                _delayed_click_fallback(
                    page,
                    fallback_box,
                    label="Focus message box",
                    timeout_ms=1800,
                    force=True,
                    state=state,
                )
        except Exception:
            pass
    try:
        msg_box.first.evaluate(
            "(el) => { if (!el) return; if (el.focus) el.focus(); el.innerHTML = ''; el.textContent = ''; }"
        )
    except Exception:
        try:
            page.keyboard.press("Control+A")
            page.keyboard.press("Backspace")
        except Exception:
            pass

    used_human_typing = _type_text_with_random_pauses(
        page,
        normalized,
        typing_pause_chars_min=typing_pause_chars_min,
        typing_pause_chars_max=typing_pause_chars_max,
        typing_pause_delay_min_ms=typing_pause_delay_min_ms,
        typing_pause_delay_max_ms=typing_pause_delay_max_ms,
        remove_chars_during_typing_enabled=remove_chars_during_typing_enabled,
        remove_chars_times_min=remove_chars_times_min,
        remove_chars_times_max=remove_chars_times_max,
        remove_chars_count_min=remove_chars_count_min,
        remove_chars_count_max=remove_chars_count_max,
        state=state,
    )
    if not used_human_typing:
        typed = _human_type_text_basic(page, normalized, state=state, min_delay_ms=24, max_delay_ms=54)
        if not typed and normalized:
            try:
                msg_box.first.fill(normalized)
            except Exception:
                page.keyboard.type(normalized, delay=12)
    min_seconds = max(0, int(delay_min_seconds or 0))
    max_seconds = max(min_seconds, int(delay_max_seconds or min_seconds))
    if not random_delay:
        max_seconds = min_seconds
    _wait_random_send_delay(
        page,
        min_seconds,
        max_seconds,
        log_func=_send_log,
        label="Text message send",
        on_delay=on_delay,
        state=state,
    )
    page.keyboard.press("Enter")
    return True


def _click_all_filter_before_search(page, timeout_ms=2500, state=None):
    """Select the "All" chat filter (the row of filters under the search bar).

    Optional actions click Groups / Unread / Favourites / Archived and leave that
    filter active; the next number search then runs inside a filtered list and
    may show nothing. Clicking All first restores the full chat list. Best-effort:
    a build without the filter row just continues."""
    candidates = [
        page.locator("#all-filter").first,
        page.locator("#pane-side, #side").get_by_role("tab", name="All", exact=True).first,
        page.get_by_role("tab", name="All", exact=True).first,
        page.locator("#side button:has-text('All'), #side div[role='button']:has-text('All')").first,
    ]
    for item in candidates:
        try:
            if item.count() == 0:
                continue
            item.wait_for(state="visible", timeout=timeout_ms)
            try:
                if (item.get_attribute("aria-selected") or "").lower() == "true":
                    return True                      # already on All
            except Exception:
                pass
            if not _human_click_item(
                page,
                item,
                label="Open All filter",
                timeout_ms=timeout_ms,
                state=state,
            ):
                if not _delayed_click_fallback(
                    page,
                    item,
                    label="Open All filter",
                    timeout_ms=timeout_ms,
                    state=state,
                ):
                    continue
            page.wait_for_timeout(150)
            _send_log("Chat filter reset to All before searching.")
            return True
        except Exception:
            continue
    return False


def _click_groups_filter_for_send(page, timeout_ms=7000, state=None):
    candidates = [
        page.locator("#group-filter").first,
        page.get_by_role("tab", name="Groups").first,
        page.locator("button:has-text('Groups')").first,
    ]
    for item in candidates:
        try:
            if item.count() == 0:
                continue
            item.wait_for(state="visible", timeout=timeout_ms)
            if not _human_click_item(
                page,
                item,
                label="Open Groups filter",
                timeout_ms=timeout_ms,
                state=state,
            ):
                if not _delayed_click_fallback(
                    page,
                    item,
                    label="Open Groups filter",
                    timeout_ms=timeout_ms,
                    state=state,
                ):
                    continue
            page.wait_for_timeout(120)
            return True
        except Exception:
            continue
    return False


def _open_group_chat_like_scrape_flow(page, group_name, state=None):
    target = str(group_name or "").strip()
    if not target:
        return False, "Group name is empty"

    try:
        page.wait_for_selector("#pane-side", timeout=12000)
    except Exception:
        return False, "Chat list pane not ready"

    if not _click_groups_filter_for_send(page, state=state):
        return False, "Could not open Groups filter"

    rows_selector = "div[aria-label='Chat list'][role='grid'] > div[role='row']"
    rows = page.locator(rows_selector)
    if rows.count() == 0:
        rows = page.locator("div[aria-label='Chat list'] div[role='row']")
    if rows.count() == 0:
        return False, "No group rows found in Groups list"

    target_norm = _normalized_text(target)
    matched_row = None
    max_rounds = 18
    no_move_rounds = 0
    last_tail = ""

    for _ in range(max_rounds):
        rows = page.locator(rows_selector)
        if rows.count() == 0:
            rows = page.locator("div[aria-label='Chat list'] div[role='row']")

        for i in range(rows.count()):
            row = rows.nth(i)
            try:
                if not row.is_visible():
                    continue
                title_node = row.locator("span[title]").first
                title_text = (title_node.get_attribute("title") or "").strip() if title_node.count() else ""
                row_text = (row.inner_text() or "").strip()
                title_norm = _normalized_text(title_text)
                row_norm = _normalized_text(row_text)
                if target_norm and (title_norm == target_norm or target_norm in title_norm or target_norm in row_norm):
                    matched_row = row
                    break
            except Exception:
                continue
        if matched_row is not None:
            break

        try:
            current_tail = ""
            if rows.count() > 0:
                current_tail = _normalized_text(rows.nth(rows.count() - 1).inner_text())
            if current_tail and current_tail == last_tail:
                no_move_rounds += 1
            else:
                no_move_rounds = 0
            last_tail = current_tail
        except Exception:
            pass

        if no_move_rounds >= 4:
            break

        try:
            rows.last.scroll_into_view_if_needed(timeout=1800)
            rows.last.hover()
            page.mouse.wheel(0, 1200)
            page.wait_for_timeout(220)
        except Exception:
            break

    if matched_row is None:
        return False, f"Group '{target}' not found in Groups list"

    try:
        matched_row.scroll_into_view_if_needed(timeout=3000)
    except Exception:
        pass
    try:
        if not _human_click_item(
            page,
            matched_row,
            label=f"Open group {target}",
            timeout_ms=2200,
            state=state,
        ):
            if not _delayed_click_fallback(
                page,
                matched_row,
                label=f"Open group {target}",
                timeout_ms=2200,
                state=state,
            ):
                raise RuntimeError("Could not click matched group row")
    except Exception:
        try:
            if not _human_click_item(
                page,
                matched_row,
                label=f"Open group {target}",
                timeout_ms=2200,
                force=True,
                state=state,
            ):
                if not _delayed_click_fallback(
                    page,
                    matched_row,
                    label=f"Open group {target}",
                    timeout_ms=2200,
                    force=True,
                    state=state,
                ):
                    raise RuntimeError("Could not force-click matched group row")
        except Exception:
            return False, f"Could not click group '{target}' row"

    page.wait_for_timeout(180)

    try:
        page.locator("footer div[contenteditable='true'][role='textbox']").first.wait_for(
            state="visible", timeout=5000
        )
    except Exception:
        pass

    return True, ""


def _collect_visible_chat_rows(page, limit=18):
    rows_selector = "div[aria-label='Chat list'][role='grid'] > div[role='row']"
    rows = page.locator(rows_selector)
    try:
        if rows.count() == 0:
            rows = page.locator("div[aria-label='Chat list'] div[role='row']")
    except Exception:
        rows = page.locator("div[aria-label='Chat list'] div[role='row']")
    visible_rows = []
    try:
        total = min(max(0, rows.count()), max(1, int(limit or 18)))
    except Exception:
        total = 0
    for i in range(total):
        row = rows.nth(i)
        try:
            if not row.is_visible():
                continue
            title = _row_text(row)
            if not str(title or "").strip():
                continue
            visible_rows.append((row, str(title or "").strip()))
        except Exception:
            continue
    return visible_rows


def _dismiss_optional_side_action_overlays(page, state=None):
    for _ in range(2):
        try:
            page.keyboard.press("Escape")
            _cooperative_wait(page, 90, state=state, step_ms=60)
        except Exception:
            break


def _find_first_visible_locator(page, selectors, timeout_ms=1400):
    for selector in selectors:
        try:
            node = page.locator(selector).first
            if node.count() == 0:
                continue
            node.wait_for(state="visible", timeout=timeout_ms)
            return node
        except Exception:
            continue
    return None


def _locate_active_message_box(page, timeout_ms=1400):
    return _find_first_visible_locator(
        page,
        (
            "footer div[contenteditable='true'][role='textbox'][data-lexical-editor='true']",
            "footer div[contenteditable='true'][role='textbox']",
            "footer div[contenteditable='true'][data-tab='10']",
            "footer div[contenteditable='true'][data-tab='6']",
            "div[contenteditable='true'][role='textbox'][data-lexical-editor='true']",
            "div[contenteditable='true'][role='textbox']",
        ),
        timeout_ms=timeout_ms,
    )


def _prepare_active_message_box(page, state=None, timeout_ms=5000):
    deadline = time.time() + max(0.5, float(timeout_ms or 0) / 1000.0)
    while time.time() < deadline:
        box = _locate_active_message_box(page, timeout_ms=700)
        if box is not None:
            try:
                box.scroll_into_view_if_needed(timeout=700)
            except Exception:
                pass
            try:
                if not _human_click_item(
                    page,
                    box.first,
                    label="Prepare message box",
                    timeout_ms=1500,
                    force=True,
                    state=state,
                ):
                    _delayed_click_fallback(
                        page,
                        box.first,
                        label="Prepare message box",
                        timeout_ms=1500,
                        force=True,
                        state=state,
                    )
            except Exception:
                pass
            try:
                box.first.evaluate("(el) => { if (el && el.focus) el.focus(); }")
            except Exception:
                pass
            return box
        try:
            if _focus_compose_box(page, state=state):
                box = _locate_active_message_box(page, timeout_ms=400)
                if box is not None:
                    try:
                        box.first.evaluate("(el) => { if (el && el.focus) el.focus(); }")
                    except Exception:
                        pass
                    return box
        except Exception:
            pass
        _cooperative_wait(page, 120, state=state, step_ms=60)
    return None


def _pause_after_optional_action(page, state=None, min_ms=0):
    pause_ms = max(int(min_ms or 0), int(_optional_side_action_pause_ms(state) or 0))
    if pause_ms > 0:
        _cooperative_wait(page, pause_ms, state=state, step_ms=120)
    return pause_ms


def _scroll_locator_region(page, locator, state=None, cycles=1):
    if locator is None:
        return False
    try:
        page.bring_to_front()
    except Exception:
        pass
    try:
        locator.scroll_into_view_if_needed(timeout=1200)
    except Exception:
        pass
    try:
        box = locator.bounding_box()
    except Exception:
        box = None
    if not box:
        return False

    point = _human_move_mouse_to_box(page, box)
    if not point:
        try:
            page.mouse.move(
                float(box.get("x", 0)) + max(10.0, float(box.get("width", 0)) / 2.0),
                float(box.get("y", 0)) + max(10.0, float(box.get("height", 0)) / 2.0),
                steps=random.randint(8, 18),
            )
        except Exception:
            return False

    did_scroll = False
    total_cycles = max(1, int(cycles or 1))
    for _ in range(total_cycles):
        for delta in (
            random.randint(260, 760),
            -random.randint(180, 620),
        ):
            _wait_if_send_paused(state, page=page)
            try:
                page.mouse.wheel(0, delta)
                did_scroll = True
            except Exception:
                continue
            _cooperative_wait(page, random.randint(120, 320), state=state, step_ms=110)
    return did_scroll


def _open_whatsapp_menu_item(page, item_labels, label="Optional menu item", state=None):
    if not _click_first_selector(
        page,
        (
            "button[aria-label='Menu']",
            "div[aria-label='Menu']",
            "button[title='Menu']",
            "div[title='Menu']",
            "span[data-icon='menu']",
        ),
        timeout_ms=1800,
        label="Optional menu open",
        state=state,
    ):
        return False

    selectors = []
    for item_label in list(item_labels or []):
        value = str(item_label or "").strip()
        if not value:
            continue
        selectors.extend(
            [
                f"button[role='menuitem']:has-text('{value}')",
                f"div[role='menuitem']:has-text('{value}')",
                f"button[aria-label='{value}']",
                f"div[aria-label='{value}']",
                f"button:has-text('{value}')",
                f"div[role='button']:has-text('{value}')",
            ]
        )
    if not selectors:
        _dismiss_optional_side_action_overlays(page, state=state)
        return False

    if _click_first_selector(
        page,
        tuple(selectors),
        timeout_ms=1800,
        label=label,
        state=state,
    ):
        return True
    _dismiss_optional_side_action_overlays(page, state=state)
    return False


def _open_current_chat_info_drawer(page, state=None, label="Optional chat info"):
    if not _click_first_selector(
        page,
        (
            "xpath=//header//*[self::span[@title] or @data-testid='conversation-info-header-chat-title']/ancestor::*[@role='button' or self::div or self::button][1]",
            "header [data-testid='conversation-info-header']",
            "header div[role='button']",
        ),
        timeout_ms=2200,
        label=label,
        state=state,
    ):
        return False
    opened = _wait_for_any_selector(
        page,
        (
            "[data-testid='drawer-right']",
            "text='Contact info'",
            "text='Group info'",
            "button[aria-label='Edit Name']",
            "button[aria-label='Edit profile photo']",
        ),
        timeout_ms=1800,
    )
    return bool(opened)


def _run_optional_groups_scroll(page, state=None):
    if not _click_groups_filter_for_send(page, state=state):
        return False, ""
    chat_list = _find_first_visible_locator(
        page,
        (
            "div[aria-label='Chat list'][role='grid']",
            "div[aria-label='Chat list']",
            "#pane-side",
        ),
        timeout_ms=1600,
    )
    if chat_list is None:
        return False, ""
    _scroll_locator_region(page, chat_list, state=state, cycles=1)
    _pause_after_optional_action(page, state=state, min_ms=220)
    return True, "opened Groups and scrolled the list"


def _run_optional_group_members_peek(page, state=None):
    if not _click_groups_filter_for_send(page, state=state):
        return False, ""
    visible_rows = _collect_visible_chat_rows(page, limit=14)
    if not visible_rows:
        return False, ""
    chosen_row, chosen_title = random.choice(visible_rows)
    if not _human_click_item(
        page,
        chosen_row,
        label=f"Optional open group {chosen_title}",
        timeout_ms=2200,
        state=state,
    ):
        return False, ""
    _cooperative_wait(page, random.randint(160, 340), state=state, step_ms=110)
    if not _open_current_chat_info_drawer(
        page,
        state=state,
        label=f"Optional group info {chosen_title}",
    ):
        return True, f"opened group '{chosen_title}' briefly"
    drawer = _find_first_visible_locator(
        page,
        (
            "[data-testid='drawer-right']",
            "div[role='dialog']",
        ),
        timeout_ms=1500,
    )
    _scroll_locator_region(page, drawer, state=state, cycles=1)
    _pause_after_optional_action(page, state=state, min_ms=320)
    _dismiss_optional_side_action_overlays(page, state=state)
    return True, f"checked members/info for group '{chosen_title}'"


def _run_optional_menu_peek(page, state=None):
    if not _click_first_selector(
        page,
        (
            "button[aria-label='Menu']",
            "div[aria-label='Menu']",
            "button[title='Menu']",
            "div[title='Menu']",
            "span[data-icon='menu']",
        ),
        timeout_ms=1800,
        label="Optional menu peek",
        state=state,
    ):
        return False, ""
    action_pause_ms = _optional_side_action_pause_ms(state)
    if action_pause_ms > 0:
        _cooperative_wait(page, action_pause_ms, state=state, step_ms=120)
    _dismiss_optional_side_action_overlays(page, state=state)
    return True, "opened WhatsApp menu briefly"


def _run_optional_profile_peek(page, state=None):
    opened = _open_whatsapp_menu_item(
        page,
        ("Profile", "Your profile"),
        label="Optional profile peek",
        state=state,
    )
    if not opened:
        opened = _click_first_selector(
            page,
            (
                "button[aria-label='Profile']",
                "div[aria-label='Profile']",
                "button[title='Profile']",
                "div[title='Profile']",
            ),
            timeout_ms=1800,
            label="Optional direct profile peek",
            state=state,
        )
    if not opened:
        return False, ""
    try:
        profile_open = _is_profile_panel_open(page)
    except Exception:
        profile_open = False
    if not profile_open:
        _wait_for_any_selector(
            page,
            (
                "[data-testid='drawer-right']",
                "text='Profile'",
                "button[aria-label='Edit profile photo']",
            ),
            timeout_ms=1500,
        )
    drawer = _find_first_visible_locator(
        page,
        (
            "[data-testid='drawer-right']",
            "div[role='dialog']",
        ),
        timeout_ms=1200,
    )
    _scroll_locator_region(page, drawer, state=state, cycles=1)
    _pause_after_optional_action(page, state=state, min_ms=260)
    _dismiss_optional_side_action_overlays(page, state=state)
    return True, "opened WhatsApp profile briefly"


def _run_optional_settings_peek(page, state=None):
    if not _open_whatsapp_menu_item(
        page,
        ("Settings",),
        label="Optional settings peek",
        state=state,
    ):
        return False, ""
    _wait_for_any_selector(
        page,
        (
            "text='Settings'",
            "[data-testid='drawer-right']",
            "div[role='dialog']",
        ),
        timeout_ms=1500,
    )
    drawer = _find_first_visible_locator(
        page,
        (
            "[data-testid='drawer-right']",
            "div[role='dialog']",
        ),
        timeout_ms=1200,
    )
    _scroll_locator_region(page, drawer, state=state, cycles=1)
    _pause_after_optional_action(page, state=state, min_ms=260)
    _dismiss_optional_side_action_overlays(page, state=state)
    return True, "opened WhatsApp settings briefly"


def _run_optional_filter_peek(page, state=None):
    filter_candidates = [
        ("Unread", ("button:has-text('Unread')", "div[role='button']:has-text('Unread')")),
        ("Favorites", ("button:has-text('Favorites')", "div[role='button']:has-text('Favorites')")),
        ("Favourites", ("button:has-text('Favourites')", "div[role='button']:has-text('Favourites')")),
        ("Archived", ("button:has-text('Archived')", "div[role='button']:has-text('Archived')")),
    ]
    random.shuffle(filter_candidates)
    for filter_name, selectors in filter_candidates:
        if _click_first_selector(
            page,
            selectors,
            timeout_ms=1600,
            label=f"Optional {filter_name} peek",
            state=state,
        ):
            action_pause_ms = _optional_side_action_pause_ms(state)
            if action_pause_ms > 0:
                _cooperative_wait(page, action_pause_ms, state=state, step_ms=120)
            _dismiss_optional_side_action_overlays(page, state=state)
            return True, f"peeked {filter_name}"
    return False, ""


def _run_optional_chat_scroll(page, state=None):
    current_chat = _find_first_visible_locator(
        page,
        (
            "div[aria-label='Message list']",
            "[data-testid='conversation-panel-body']",
            "[data-testid='conversation-panel-messages']",
            "#main",
        ),
        timeout_ms=1400,
    )
    if current_chat is None:
        return False, ""
    _scroll_locator_region(page, current_chat, state=state, cycles=1)
    _pause_after_optional_action(page, state=state, min_ms=220)
    return True, "scrolled the opened chat"


def _run_optional_contact_info_peek(page, state=None):
    if not _open_current_chat_info_drawer(page, state=state, label="Optional contact info peek"):
        return False, ""
    drawer = _find_first_visible_locator(
        page,
        (
            "[data-testid='drawer-right']",
            "div[role='dialog']",
        ),
        timeout_ms=1500,
    )
    _scroll_locator_region(page, drawer, state=state, cycles=1)
    _pause_after_optional_action(page, state=state, min_ms=260)
    _dismiss_optional_side_action_overlays(page, state=state)
    return True, "opened the searched number info briefly"


def _run_optional_attachment_document_click(page, state=None):
    return _peek_attachment_menu_item_with_filechooser(
        page,
        _attachment_menu_item_selectors("Document"),
        "Optional Document menu click",
        "clicked Document in the attachment menu",
        state=state,
    )


def _run_optional_attachment_media_click(page, state=None):
    return _peek_attachment_menu_item_with_filechooser(
        page,
        _attachment_menu_item_selectors("Photos & videos", extra_texts=("Photos and videos",)),
        "Optional Photos & videos menu click",
        "clicked Photos & videos in the attachment menu",
        state=state,
    )


def _run_optional_attachment_camera_click(page, state=None):
    return _peek_attachment_menu_item_click_only(
        page,
        _attachment_menu_item_selectors("Camera"),
        "Optional Camera menu click",
        "clicked Camera in the attachment menu",
        state=state,
    )


def _run_optional_attachment_audio_click(page, state=None):
    return _peek_attachment_menu_item_with_filechooser(
        page,
        _attachment_menu_item_selectors("Audio"),
        "Optional Audio menu click",
        "clicked Audio in the attachment menu",
        state=state,
    )


def _run_optional_attachment_contact_click(page, state=None):
    return _peek_attachment_menu_item_click_only(
        page,
        _attachment_menu_item_selectors("Contact"),
        "Optional Contact menu click",
        "clicked Contact in the attachment menu",
        state=state,
    )


def _run_optional_attachment_poll_click(page, state=None):
    return _peek_attachment_menu_item_click_only(
        page,
        _attachment_menu_item_selectors("Poll"),
        "Optional Poll menu click",
        "clicked Poll in the attachment menu",
        state=state,
    )


def _run_optional_attachment_event_click(page, state=None):
    return _peek_attachment_menu_item_click_only(
        page,
        _attachment_menu_item_selectors("Event"),
        "Optional Event menu click",
        "clicked Event in the attachment menu",
        state=state,
    )


def _run_optional_attachment_sticker_click(page, state=None):
    return _peek_attachment_menu_item_click_only(
        page,
        _attachment_menu_item_selectors("New sticker"),
        "Optional New sticker menu click",
        "clicked New sticker in the attachment menu",
        state=state,
    )


def _automation_action_runners_catalog():
    return [
        ("groups_scroll", _run_optional_groups_scroll, "after_send"),
        ("group_members_peek", _run_optional_group_members_peek, "after_send"),
        ("menu_peek", _run_optional_menu_peek, "after_send"),
        ("filter_peek", _run_optional_filter_peek, "after_send"),
        ("profile_peek", _run_optional_profile_peek, "after_send"),
        ("settings_peek", _run_optional_settings_peek, "after_send"),
        ("chat_scroll", _run_optional_chat_scroll, "after_send"),
        ("contact_info_peek", _run_optional_contact_info_peek, "after_send"),
        ("attach_document_click", _run_optional_attachment_document_click, "after_send"),
        ("attach_media_click", _run_optional_attachment_media_click, "after_send"),
        ("attach_camera_click", _run_optional_attachment_camera_click, "after_send"),
        ("attach_audio_click", _run_optional_attachment_audio_click, "after_send"),
        ("attach_contact_click", _run_optional_attachment_contact_click, "after_send"),
        ("attach_poll_click", _run_optional_attachment_poll_click, "after_send"),
        ("attach_event_click", _run_optional_attachment_event_click, "after_send"),
        ("attach_sticker_click", _run_optional_attachment_sticker_click, "after_send"),
    ]


def _set_automation_action_assignment(state, row_index, action_id, phase=""):
    action_meta = _action_definition_by_id(action_id)
    resolved_phase = str(phase or action_meta.get("phase") or "after_send").strip() or "after_send"
    action_label = str(action_meta.get("label") or action_id or "-").strip() or "-"
    _update_automation_action_row(
        state,
        row_index,
        action_id=str(action_id or "").strip(),
        action=action_label,
        phase=resolved_phase,
    )


def _mark_automation_action_performed(
    state,
    row_index,
    action_id,
    phase="",
    detail="",
    target_label="",
):
    action_meta = _action_definition_by_id(action_id)
    action_label = str(action_meta.get("label") or action_id or "-").strip() or "-"
    performed_at = datetime.now().strftime("%H:%M:%S")
    _update_automation_action_row(
        state,
        row_index,
        action_id=str(action_id or "").strip(),
        action=action_label,
        phase=str(phase or action_meta.get("phase") or "after_send").strip() or "after_send",
        checked=True,
        status="performed",
        detail=str(detail or "").strip(),
        performed_at=performed_at,
    )
    event_text = action_label
    if str(target_label or "").strip():
        event_text += f" for {str(target_label or '').strip()}"
    if str(detail or "").strip():
        event_text += f" ({str(detail or '').strip()})"
    state["automation_action_event_text"] = event_text.strip()


def _perform_required_automation_action(
    page,
    state,
    row_index,
    target_label="",
    phase="after_send",
    preferred_action_id="",
):
    phase_key = str(phase or "after_send").strip().lower() or "after_send"
    catalog = [item for item in _automation_action_runners_catalog() if str(item[2] or "").strip().lower() == phase_key]
    if not catalog:
        return False, "", ""

    ordered = []
    preferred_key = str(preferred_action_id or "").strip()
    if preferred_key:
        for item in catalog:
            if item[0] == preferred_key:
                ordered.append(item)
                break
    for item in catalog:
        if item not in ordered:
            ordered.append(item)

    for action_id, runner, runner_phase in ordered:
        try:
            ok, message = runner(page, state=state)
        except SendJobCancelled:
            raise
        except Exception as exc:
            _send_log(f"Required automation action skipped ({action_id}): {exc}")
            ok = False
            message = ""
        if not ok:
            continue
        detail = str(message or "").strip()
        _set_automation_action_assignment(state, row_index, action_id, phase=runner_phase)
        _mark_automation_action_performed(
            state,
            row_index,
            action_id,
            phase=runner_phase,
            detail=detail,
            target_label=target_label,
        )
        _mark_campaign_action_done(state, action_id, target=target_label, detail=detail)
        _send_log(
            f"Automation action after sending to {target_label or ('row ' + str(int(row_index) + 1))}: {detail or action_id}."
        )
        return True, action_id, detail
    return False, "", ""


def _perform_final_safe_automation_fallback(page, state, row_index, target_label=""):
    safe_candidates = [
        ("menu_peek", _run_optional_menu_peek, "after_send"),
        ("settings_peek", _run_optional_settings_peek, "after_send"),
        ("profile_peek", _run_optional_profile_peek, "after_send"),
        ("groups_scroll", _run_optional_groups_scroll, "after_send"),
        ("chat_scroll", _run_optional_chat_scroll, "after_send"),
        ("contact_info_peek", _run_optional_contact_info_peek, "after_send"),
    ]
    for action_id, runner, runner_phase in safe_candidates:
        try:
            ok, message = runner(page, state=state)
        except SendJobCancelled:
            raise
        except Exception as exc:
            _send_log(f"Automation fallback skipped ({action_id}): {exc}")
            ok = False
            message = ""
        if not ok:
            continue
        detail = str(message or "").strip()
        _set_automation_action_assignment(state, row_index, action_id, phase=runner_phase)
        _mark_automation_action_performed(
            state,
            row_index,
            action_id,
            phase=runner_phase,
            detail=detail,
            target_label=target_label,
        )
        _mark_campaign_action_done(state, action_id, target=target_label, detail=detail)
        _send_log(
            f"Automation fallback action after sending to {target_label or ('row ' + str(int(row_index) + 1))}: {detail or action_id}."
        )
        return True, action_id, detail
    return False, "", ""


def _maybe_run_optional_side_action(page, state=None, target_label="", phase="after_send"):
    state = _resolve_interaction_state(state)
    settings = _get_optional_side_action_settings(state)
    if not settings.get("enabled"):
        return False
    phase_key = str(phase or "after_send").strip().lower() or "after_send"
    action_runners = _automation_action_runners_catalog()
    eligible_runners = [item for item in action_runners if str(item[2] or "").strip().lower() == phase_key]
    if not eligible_runners:
        return False

    pending_ids = set(_pending_campaign_action_ids(state=state, phase=phase_key))
    runner_queue = []
    if pending_ids:
        pending_runners = [item for item in eligible_runners if item[0] in pending_ids]
        other_runners = [item for item in eligible_runners if item[0] not in pending_ids]
        random.shuffle(pending_runners)
        random.shuffle(other_runners)
        runner_queue = pending_runners + other_runners
    else:
        chance_percent = max(0, min(100, int(settings.get("chance_percent") or 0)))
        if chance_percent <= 0:
            return False
        if random.randint(1, 100) > chance_percent:
            return False
        runner_queue = list(eligible_runners)
        random.shuffle(runner_queue)

    for action_id, runner, _runner_phase in runner_queue:
        try:
            ok, message = runner(page, state=state)
        except SendJobCancelled:
            raise
        except Exception as exc:
            _send_log(f"Optional side action skipped ({action_id}): {exc}")
            ok = False
            message = ""
        if ok:
            detail = str(message or "completed a side action").strip()
            _mark_campaign_action_done(state, action_id, target=target_label, detail=detail)
            if target_label:
                _send_log(f"Optional side action ({phase_key}) after {target_label}: {detail}.")
            else:
                _send_log(f"Optional side action ({phase_key}): {detail}.")
            return True
    return False


def send_messages(
    targets,
    payload,
    delay_seconds,
    allow_resend,
    state,
    send_target="contacts",
    campaign_id=None,
    campaign_name="",
    batch_min_size=20,
    batch_max_size=30,
    range_config=None,
):
    session_dir = get_playwright_user_data_dir()
    db_conn = get_db_connection()
    db_lock = threading.Lock()
    worker_contexts = []
    worker_specs = []
    state["cancel_requested"] = False
    state["cancel_requested_at"] = ""
    _set_active_send_contexts(state, [])
    ensure_message_log_tables(db_conn)
    _send_log("Connected to PostgreSQL.")
    if campaign_id:
        _send_log(f"Campaign ID: {campaign_id}")
    if campaign_name:
        _send_log(f"Campaign Name: {campaign_name}")

    def set_contact_status(target, status, reason=""):
        with _SEND_STATE_LOCK:
            if "contact_statuses" not in state:
                state["contact_statuses"] = {}
            state["contact_statuses"][target] = {"status": status, "reason": reason or ""}

    def record_failed(target, reason):
        current_status = (
            ((state.get("contact_statuses") or {}).get(target) or {}).get("status", "")
        )
        if current_status == "sent":
            _send_log(
                f"Post-send warning for {target}: {reason}. Keeping final status as sent."
            )
            return
        with _SEND_STATE_LOCK:
            state["failed"].append({"number": target, "reason": reason})
        set_contact_status(target, "failed", reason)
        try:
            with db_lock:
                save_not_sent_sms(db_conn, campaign_id, campaign_name, target, reason)
        except Exception as db_exc:
            _send_log(f"DB warning: failed to save sms_not_sent for {target}: {db_exc}")

    try:
        if True:  # browsers are launched inside per-profile worker threads (see _worker_thread_main)
            launch_kwargs = get_playwright_launch_kwargs(default_headless=False)
            _send_log(
                f"Playwright config -> headless={launch_kwargs.get('headless')} "
                f"start_maximized={'args' in launch_kwargs} "
                f"no_viewport={launch_kwargs.get('no_viewport', False)} "
                f"slow_mo={launch_kwargs.get('slow_mo', 0)}ms"
            )
            config = dict(range_config or {})
            def _config_int(key, default):
                raw_value = config.get(key, default)
                if raw_value in (None, ""):
                    raw_value = default
                return int(raw_value)

            profile_numbers_map = dict(config.get("profile_numbers_map") or {})
            skipped_profiles = set()
            for raw in (config.get("skipped_profiles") or []):
                try:
                    parsed = int(raw)
                    if parsed > 0:
                        skipped_profiles.add(parsed)
                except Exception:
                    continue
            active_profile_indexes = []
            for raw in (config.get("active_profile_indexes") or []):
                try:
                    parsed = int(raw)
                except Exception:
                    continue
                if parsed > 0 and parsed not in active_profile_indexes:
                    active_profile_indexes.append(parsed)
            active_profile_indexes.sort()
            random_message_enabled = bool(config.get("random_message_enabled", True))
            message_panel_count = max(1, _config_int("message_panel_count", 1))
            message_panel_count = min(100, message_panel_count)
            random_attachment_enabled = bool(config.get("random_attachment_enabled", True))
            media_count_min = max(1, _config_int("media_count_min", 1))
            media_count_max = max(media_count_min, _config_int("media_count_max", media_count_min))
            random_attachment_delay_enabled = bool(config.get("random_attachment_delay_enabled", True))
            attachment_delay_min = max(0, _config_int("attachment_delay_min", 0))
            attachment_delay_max = max(attachment_delay_min, _config_int("attachment_delay_max", 30))
            random_search_result_delay_enabled = bool(
                config.get("random_search_result_delay_enabled", True)
            )
            search_result_delay_min = max(0, _config_int("search_result_delay_min", 1))
            search_result_delay_max = max(
                search_result_delay_min,
                _config_int("search_result_delay_max", 4),
            )
            random_click_delay_enabled = bool(config.get("random_click_delay_enabled", True))
            click_delay_min_ms = max(0, _config_int("click_delay_min_ms", 150))
            click_delay_max_ms = max(
                click_delay_min_ms,
                _config_int("click_delay_max_ms", 900),
            )
            random_send_delay_enabled = bool(config.get("random_send_delay_enabled", True))
            send_delay_min = max(0, _config_int("send_delay_min", 0))
            send_delay_max = max(send_delay_min, _config_int("send_delay_max", 30))
            typing_pause_chars_min = max(0, _config_int("typing_pause_chars_min", 10))
            typing_pause_chars_max = max(
                typing_pause_chars_min,
                _config_int("typing_pause_chars_max", 20),
            )
            typing_pause_delay_min_ms = max(0, _config_int("typing_pause_delay_min", 2000))
            typing_pause_delay_max_ms = max(
                typing_pause_delay_min_ms,
                _config_int("typing_pause_delay_max", 5000),
            )
            remove_chars_during_typing_enabled = bool(
                config.get("remove_chars_during_typing_enabled", True)
            )
            remove_chars_times_min = max(0, _config_int("remove_chars_times_min", 3))
            remove_chars_times_max = max(
                remove_chars_times_min,
                _config_int("remove_chars_times_max", 7),
            )
            remove_chars_count_min = max(0, _config_int("remove_chars_count_min", 1))
            remove_chars_count_max = max(
                remove_chars_count_min,
                _config_int("remove_chars_count_max", 3),
            )
            random_batch_size_enabled = bool(config.get("random_batch_size_enabled", True))
            random_batch_start_enabled = bool(config.get("random_batch_start_enabled", True))
            batch_start_min_minutes = max(0, _config_int("batch_start_min_minutes", 3))
            batch_start_max_minutes = max(batch_start_min_minutes, _config_int("batch_start_max_minutes", 6))
            profiles_numbers = max(1, min(100, _config_int("profiles_numbers", 1)))
            optional_side_actions_enabled = bool(config.get("optional_side_actions_enabled", True))
            side_action_chance_percent = max(0, min(100, _config_int("side_action_chance_percent", 20)))
            side_action_pause_min_ms = max(0, _config_int("side_action_pause_min_ms", 700))
            side_action_pause_max_ms = max(
                side_action_pause_min_ms,
                _config_int("side_action_pause_max_ms", 2000),
            )
            profile_slot_indexes = [
                idx for idx in active_profile_indexes
                if idx > 0 and idx not in skipped_profiles
            ]
            if not profile_slot_indexes:
                raise RuntimeError(
                    "No logged-in WhatsApp sessions selected. Open WhatsApp Sessions and log in at least one profile first."
                )
            highest_active_idx = max(profile_slot_indexes) if profile_slot_indexes else 0
            session_profile_count = max(profiles_numbers, highest_active_idx)
            stored_numbers_map = _load_profile_numbers_from_session_dirs(
                session_profile_count,
                create_missing=True,
            )
            random_profile_pick_enabled = bool(config.get("random_profile_pick_enabled", True))
            state["_click_delay_config"] = {
                "enabled": random_click_delay_enabled,
                "min_ms": click_delay_min_ms,
                "max_ms": click_delay_max_ms,
            }
            state["_optional_side_action_config"] = {
                "enabled": optional_side_actions_enabled,
                "chance_percent": side_action_chance_percent,
                "pause_min_ms": side_action_pause_min_ms,
                "pause_max_ms": side_action_pause_max_ms,
            }
            _send_log(f"Playwright base user_data_dir: {session_dir}")
            _send_log(
                "Logged-in WhatsApp sessions selected for sending: "
                + ", ".join(f"Profile {idx}" for idx in profile_slot_indexes)
            )

            base_session_dir = Path(session_dir).resolve()
            worker_pages = []
            state["profile_states"] = []
            state["session_opening"] = True
            state["awaiting_session_start"] = False
            state["session_start_requested"] = False
            state["session_logged_in_count"] = 0
            state["session_logged_out_count"] = 0

            def _upsert_profile_state(profile_name, **fields):
                with _SEND_STATE_LOCK:
                    rows = list(state.get("profile_states") or [])
                    idx = -1
                    for i, row in enumerate(rows):
                        if str((row or {}).get("name") or "") == str(profile_name or ""):
                            idx = i
                            break
                    if idx >= 0:
                        row = dict(rows[idx] or {})
                    else:
                        row = {
                            "name": str(profile_name or "").strip() or "Profile",
                            "profile_number": "",
                            "status": "unknown",
                            "reason": "",
                            "targets_assigned": 0,
                            "sent": 0,
                            "failed": 0,
                            "last_target": "",
                        }
                    for k, v in fields.items():
                        row[k] = v
                    if idx >= 0:
                        rows[idx] = row
                    else:
                        rows.append(row)
                    state["profile_states"] = rows

            def _profile_index_from_label(profile_name):
                match = re.search(r"(\d+)$", str(profile_name or ""))
                if not match:
                    return 0
                try:
                    return int(match.group(1))
                except Exception:
                    return 0

            worker_pages = []
            closed_logged_out_count = 0
            workers_lock = threading.Lock()
            cancel_message = "Campaign canceled by user. All WhatsApp browsers were closed."

            def _worker_profile_number_refresh(worker):
                profile_dir = worker.get("profile_dir")
                profile_idx = int(worker.get("profile_idx") or 0)
                profile_number = _ensure_and_read_profile_number(
                    profile_dir,
                    seed_value=str(worker.get("profile_number") or "").strip(),
                )
                worker["profile_number"] = profile_number
                if profile_idx > 0:
                    _set_profile_number_in_sequence(
                        profile_idx,
                        profile_number,
                        min_size=profiles_numbers,
                    )
                return profile_number

            def _recount_session_states():
                checked = [w for w in worker_specs if w["ready_event"].is_set()]
                logged_in_count = sum(1 for w in checked if w.get("logged_in"))
                state["session_logged_in_count"] = logged_in_count
                state["session_logged_out_count"] = max(0, len(checked) - logged_in_count)

            def _batch_no_for_target(target_index):
                try:
                    if 0 <= int(target_index) < len(batch_numbers):
                        return int(batch_numbers[int(target_index)])
                except Exception:
                    pass
                return 1

            # ---- Per-session batches ---------------------------------------
            # Every WhatsApp session splits ITS OWN assigned numbers into batches
            # of the configured size range and takes its own break between
            # them. Sessions never wait for each other.
            def _assign_per_worker_batches(active_workers):
                max_batches = 0
                for worker in active_workers:
                    assigned = list(worker.get("assigned") or [])
                    if not assigned:
                        continue
                    if random_batch_size_enabled:
                        sizes = _build_random_batch_sizes(
                            len(assigned),
                            min_size=batch_min_size,
                            max_size=batch_max_size,
                        )
                    else:
                        sizes = _build_fixed_batch_sizes(len(assigned), batch_size=batch_min_size)
                    local_numbers = _expand_batch_numbers(sizes)
                    for pos, target_index in enumerate(assigned):
                        local_no = int(local_numbers[pos]) if pos < len(local_numbers) else max(1, len(sizes))
                        if 0 <= int(target_index) < len(batch_numbers):
                            batch_numbers[int(target_index)] = local_no
                        plan = preassigned_plan.get(target_index)
                        if isinstance(plan, dict):
                            plan["batch_no"] = local_no
                        _update_live_row(state, target_index, batch=local_no)
                    worker["batch_sizes"] = list(sizes)
                    max_batches = max(max_batches, len(sizes))
                    _send_log(
                        f"[{worker['name']}] Batch plan for {len(assigned)} target(s): "
                        + ", ".join(f"B{i + 1}={size}" for i, size in enumerate(sizes))
                    )
                with _SEND_STATE_LOCK:
                    state["total_batches"] = max_batches
                    state["current_batch_no"] = 1 if max_batches else 0

            def _set_worker_batch_pause(worker, remaining, prev_batch_no, next_batch_no):
                name = str(worker.get("name") or "")
                display = str(worker.get("profile_number") or "").strip() or name
                with _SEND_STATE_LOCK:
                    pauses = dict(state.get("batch_pauses") or {})
                    if remaining > 0:
                        pauses[name] = {
                            "name": name,
                            "profile": display,
                            "remaining_sec": int(remaining),
                            "prev_batch": int(prev_batch_no),
                            "next_batch": int(next_batch_no),
                            "next_start_at": (
                                datetime.now() + timedelta(seconds=int(remaining))
                            ).strftime("%H:%M:%S"),
                        }
                    else:
                        pauses.pop(name, None)
                    state["batch_pauses"] = pauses
                    active = [item for item in pauses.values() if int(item.get("remaining_sec") or 0) > 0]
                    state["batch_pause_active"] = bool(active)
                    if active:
                        soonest = min(active, key=lambda item: int(item.get("remaining_sec") or 0))
                        state["batch_pause_remaining_sec"] = int(soonest.get("remaining_sec") or 0)
                        state["batch_next_start_at"] = str(soonest.get("next_start_at") or "")
                    else:
                        state["batch_pause_remaining_sec"] = 0
                        state["batch_next_start_at"] = ""

            def _run_batch_pause(worker, page, prev_batch_no, next_batch_no):
                name = str(worker.get("name") or "")
                pause_min_seconds = batch_start_min_minutes * 60
                pause_max_seconds = batch_start_max_minutes * 60
                if random_batch_start_enabled:
                    pause_seconds = random.randint(pause_min_seconds, pause_max_seconds)
                else:
                    pause_seconds = pause_min_seconds
                _send_log(
                    f"[{name}] Batch {prev_batch_no} complete. Next batch {next_batch_no} starts after "
                    f"{pause_seconds // 60}m {pause_seconds % 60:02d}s."
                )
                try:
                    for remaining in range(pause_seconds, 0, -1):
                        _set_worker_batch_pause(worker, remaining, prev_batch_no, next_batch_no)
                        _cooperative_wait(page, 1000, state=state)
                finally:
                    _set_worker_batch_pause(worker, 0, prev_batch_no, next_batch_no)
                with _SEND_STATE_LOCK:
                    state["current_batch_no"] = max(int(state.get("current_batch_no") or 0), int(next_batch_no))
                _send_log(f"[{name}] Batch {next_batch_no} started.")

            def _worker_session_body(worker, context, page, owns_browser):
                """Runs inside the thread that owns `page`: either this worker's
                own browser or an already-open logged-in session window."""
                idx = int(worker.get("profile_idx") or 0)
                profile_label = str(worker.get("name") or f"Profile {idx}")
                profile_dir = worker.get("profile_dir")
                worker["page"] = page
                worker["context"] = context
                worker["owns_browser"] = bool(owns_browser)
                if owns_browser:
                    with workers_lock:
                        worker_contexts.append((context, profile_dir))
                        _set_active_send_contexts(
                            state,
                            [
                                {"context": c, "profile_dir": d, "thread_ident": None}
                                for c, d in worker_contexts
                            ],
                        )
                    worker["registered"] = True
                    _send_log(f"[Profile {idx}] Opening WhatsApp Web... ({profile_dir})")
                    page.goto(
                        "https://web.whatsapp.com/",
                        wait_until="commit",
                        timeout=15000,
                    )
                    _send_log(f"[Profile {idx}] WhatsApp session window opened. Scan QR if needed.")
                else:
                    try:
                        page.bring_to_front()
                    except Exception:
                        pass
                    try:
                        current_url = str(page.url or "")
                    except Exception:
                        current_url = ""
                    if "web.whatsapp.com" not in current_url.lower():
                        page.goto(
                            "https://web.whatsapp.com/",
                            wait_until="commit",
                            timeout=15000,
                        )
                    _send_log(f"[Profile {idx}] Using the already opened logged-in WhatsApp window.")
                _upsert_profile_state(
                    profile_label,
                    profile_number=str(worker.get("profile_number") or ""),
                    status="opening",
                    reason=(
                        "Using already opened WhatsApp session window..."
                        if not owns_browser
                        else "Opening WhatsApp session window..."
                    ),
                )

                logged_in = False
                probe_until = time.time() + 8.0
                while True:
                    _raise_if_send_cancel_requested(state, message=cancel_message)
                    profile_number = _worker_profile_number_refresh(worker)
                    try:
                        logged_in = _is_whatsapp_profile_logged_in(page)
                    except Exception:
                        logged_in = False
                    if logged_in:
                        _upsert_profile_state(
                            profile_label,
                            profile_number=profile_number,
                            status="logged_in",
                            reason="",
                        )
                        break
                    _upsert_profile_state(
                        profile_label,
                        profile_number=profile_number,
                        status="logged_out",
                        reason="Scan QR in this browser window to log in.",
                    )
                    if time.time() >= probe_until:
                        break
                    _cooperative_wait(page, 350, state=state)
                worker["logged_in"] = bool(logged_in)
                worker["ready_event"].set()
                _recount_session_states()

                # Wait until the coordinator has assigned targets to this
                # session and released all sessions together.
                while not worker["start_event"].is_set() and not worker["abort_event"].is_set():
                    _raise_if_send_cancel_requested(state, message=cancel_message)
                    try:
                        page.wait_for_timeout(150)
                    except Exception:
                        time.sleep(0.15)
                if worker["abort_event"].is_set():
                    return

                profile_number = str(worker.get("profile_number") or "").strip()
                if profile_number:
                    _write_profile_number(profile_dir, profile_number)
                    _set_profile_number_in_sequence(idx, profile_number, min_size=profiles_numbers)
                try:
                    _, session_line = save_whatsapp_session_snapshot(
                        page=page,
                        flow_name=f"send_messages_profile_{idx}",
                        user_data_dir=str(profile_dir),
                        login_state="reused_logged_in",
                    )
                    _send_log(f"[Profile {idx}] Session snapshot -> {session_line}")
                except Exception as session_exc:
                    _send_log(f"[Profile {idx}] Session snapshot warning: {session_exc}")

                assigned = list(worker.get("assigned") or [])
                _send_log(f"[{profile_label}] Starting {len(assigned)} assigned target(s).")
                for position, target_index in enumerate(assigned):
                    _raise_if_send_cancel_requested(state, message=cancel_message)
                    batch_no = _batch_no_for_target(target_index)
                    if position > 0:
                        prev_batch_no = _batch_no_for_target(assigned[position - 1])
                        if batch_no != prev_batch_no:
                            _run_batch_pause(worker, page, prev_batch_no, batch_no)
                    target_value = str(targets[target_index] or "").strip()
                    _process_target(
                        worker,
                        target_index,
                        target_value,
                        preassigned_plan.get(target_index) or {},
                    )
                _send_log(f"[{profile_label}] Finished all assigned targets.")
                try:
                    page.keyboard.press("Escape")
                except Exception:
                    pass

            def _worker_thread_main(worker):
                """One thread per WhatsApp session. If the session already has an
                open logged-in window (from the Sessions panel), the work runs
                inside that window's own thread so the window is kept and reused.
                Otherwise this thread launches its own browser. Either way all
                selected sessions send in parallel."""
                idx = int(worker.get("profile_idx") or 0)
                profile_label = str(worker.get("name") or f"Profile {idx}")
                profile_dir = worker.get("profile_dir")
                context = None
                owns_browser = False
                try:
                    _raise_if_send_cancel_requested(state, message=cancel_message)
                    if _is_open_session_preview_active(idx):
                        request = _run_in_open_session_preview_window(
                            idx,
                            lambda ctx, pg: _worker_session_body(worker, ctx, pg, owns_browser=False),
                            wait_for_start_seconds=15.0,
                        )
                        if request is not None:
                            while not request["done_event"].wait(timeout=0.5):
                                pass
                            if request["error"] is not None:
                                raise request["error"]
                            return
                        _send_log(
                            f"[Profile {idx}] The open session window did not respond; opening a new browser for this session."
                        )
                        close_result = _request_close_open_session_preview_windows(
                            [idx],
                            wait_timeout_seconds=12.0,
                        )
                        if list(close_result.get("remaining") or []):
                            raise RuntimeError(
                                "Could not use or close the already opened WhatsApp window. "
                                "Please use Close All in WhatsApp Sessions and try again."
                            )
                    with sync_playwright() as p:
                        _raise_if_send_cancel_requested(state, message=cancel_message)
                        context = launch_persistent_context_safe(
                            browser_type=p.chromium,
                            profile_dir=profile_dir,
                            launch_kwargs=launch_kwargs,
                            owner_label=f"message sending ({profile_label})",
                            log_fn=_send_log,
                        )
                        owns_browser = True
                        page = context.pages[0] if context.pages else context.new_page()
                        _worker_session_body(worker, context, page, owns_browser=True)
                except SendJobCancelled:
                    worker["error"] = "cancelled"
                except Exception as exc:
                    worker["error"] = str(exc) or exc.__class__.__name__
                    if not _is_send_cancel_requested(state):
                        _send_log(f"[{profile_label}] Worker stopped: {exc}")
                        _upsert_profile_state(
                            profile_label,
                            status="logged_out",
                            reason=f"Worker stopped: {exc}"[:280],
                        )
                        for target_index in list(worker.get("assigned") or []):
                            rows = list(state.get("live_rows") or [])
                            row = rows[target_index] if 0 <= target_index < len(rows) else {}
                            if str((row or {}).get("status") or "pending") in ("pending", "processing"):
                                target_value = str(targets[target_index] or "").strip()
                                record_failed(target_value, f"Session stopped: {exc}"[:200])
                                _update_live_row(state, target_index, status="failed")
                finally:
                    worker["ready_event"].set()
                    if owns_browser and context is not None:
                        try:
                            context.close()
                        except Exception:
                            pass
                    if worker.get("registered"):
                        with workers_lock:
                            worker_contexts[:] = [entry for entry in worker_contexts if entry[0] is not context]
                            _set_active_send_contexts(
                                state,
                                [
                                    {"context": c, "profile_dir": d, "thread_ident": None}
                                    for c, d in worker_contexts
                                ],
                            )
                    if owns_browser:
                        release_playwright_profile(profile_dir)
                    try:
                        _set_worker_batch_pause(worker, 0, 0, 0)
                    except Exception:
                        pass
                    worker["done_event"].set()

            for idx in profile_slot_indexes:
                _raise_if_send_cancel_requested(state, message=cancel_message)
                profile_label = f"Profile {idx}"
                profile_dir = _get_send_profile_dir(base_session_dir, idx, create=True)
                provided_profile_number = str(
                    profile_numbers_map.get(str(idx)) or profile_numbers_map.get(idx) or ""
                ).strip()
                stored_profile_number = str(stored_numbers_map.get(idx) or "").strip()
                profile_number = _ensure_and_read_profile_number(
                    profile_dir,
                    seed_value=stored_profile_number,
                )
                if provided_profile_number:
                    _write_profile_number(profile_dir, provided_profile_number)
                    profile_number = provided_profile_number
                _set_profile_number_in_sequence(idx, profile_number, min_size=profiles_numbers)
                if idx in skipped_profiles:
                    _upsert_profile_state(
                        profile_label,
                        profile_number=profile_number,
                        status="logged_out",
                        reason="Skipped by user for this run.",
                    )
                    continue
                worker_specs.append(
                    {
                        "name": profile_label,
                        "page": None,
                        "context": None,
                        "profile_dir": profile_dir,
                        "profile_idx": idx,
                        "profile_number": profile_number,
                        "logged_in": False,
                        "error": "",
                        "assigned": [],
                        "last_batch_no": None,
                        "processed_count": 0,
                        "owns_browser": False,
                        "registered": False,
                        "ready_event": threading.Event(),
                        "start_event": threading.Event(),
                        "abort_event": threading.Event(),
                        "done_event": threading.Event(),
                        "thread": None,
                    }
                )

            if not worker_specs:
                raise RuntimeError("No browser profile window could be initialized for sending.")

            for worker in worker_specs:
                thread = threading.Thread(
                    target=_worker_thread_main,
                    args=(worker,),
                    daemon=True,
                    name=f"send-{worker['name']}",
                )
                worker["thread"] = thread
                thread.start()
            _send_log(f"Opening {len(worker_specs)} WhatsApp session window(s) in parallel...")

            ready_deadline = time.time() + 120.0
            while time.time() < ready_deadline:
                _raise_if_send_cancel_requested(state, message=cancel_message)
                _recount_session_states()
                if all(w["ready_event"].is_set() for w in worker_specs):
                    break
                time.sleep(0.25)
            _recount_session_states()

            for worker in worker_specs:
                if worker.get("logged_in") and not worker.get("error") and worker.get("page") is not None:
                    worker_pages.append(worker)
                else:
                    closed_logged_out_count += 1
                    worker["abort_event"].set()
                    if not worker.get("error"):
                        _upsert_profile_state(
                            worker["name"],
                            profile_number=str(worker.get("profile_number") or ""),
                            status="logged_out",
                            reason="Closed before sending because QR/login was still pending.",
                        )

            state["session_logged_in_count"] = len(worker_pages)
            state["session_logged_out_count"] = closed_logged_out_count
            state["awaiting_session_start"] = False
            state["session_start_requested"] = False
            state["session_opening"] = False
            if not worker_pages:
                for worker in worker_specs:
                    worker["abort_event"].set()
                raise RuntimeError(
                    "No logged in WhatsApp sessions available. Open WhatsApp Sessions and log in at least one profile first."
                )

            _send_log(f"Cursor view: {'ON' if _cursor_view_enabled() else 'OFF'}")
            if closed_logged_out_count > 0:
                _send_log(
                    f"Skipped {closed_logged_out_count} selected session(s) because they were no longer logged in."
                )
            _send_log(f"Starting sending with {len(worker_pages)} logged in WhatsApp session(s).")
            _send_log(f"Starting messaging loop for {len(targets)} target(s).")
            _send_log(
                f"Profile pick mode: {'random' if random_profile_pick_enabled else 'sequential'} | "
                f"active profiles={len(worker_pages)}"
            )

            if random_batch_size_enabled:
                batch_sizes = _build_random_batch_sizes(
                    len(targets),
                    min_size=batch_min_size,
                    max_size=batch_max_size,
                )
            else:
                batch_sizes = _build_fixed_batch_sizes(
                    len(targets),
                    batch_size=batch_min_size,
                )
            batch_numbers = _expand_batch_numbers(batch_sizes)
            if len(batch_numbers) < len(targets):
                batch_numbers.extend([len(batch_sizes) or 1] * (len(targets) - len(batch_numbers)))
            state["total_batches"] = len(batch_sizes)
            state["current_batch_no"] = int(batch_numbers[0]) if batch_numbers else 0
            state["batch_pause_active"] = False
            state["batch_pause_remaining_sec"] = 0
            state["batch_next_start_at"] = ""
            state["batch_pauses"] = {}
            state["live_rows"] = _init_live_rows(targets, batch_numbers=batch_numbers)
            if batch_sizes:
                _send_log(
                    f"Batch size mode: {'random' if random_batch_size_enabled else 'fixed'} | "
                    f"range: {int(batch_min_size)}-{int(batch_max_size)}"
                )
                _send_log(
                    "Batch plan: "
                    + ", ".join(f"B{i+1}={size}" for i, size in enumerate(batch_sizes))
                )
            active_message_variants = list(payload.get("message_variants") or [])
            if active_message_variants:
                slots = ", ".join(str(int(item.get("slot") or 0)) for item in active_message_variants)
                _send_log(
                    f"Message mode: {'random' if random_message_enabled else 'fixed'} | "
                    f"message windows: {message_panel_count} | "
                    f"active panels: {len(active_message_variants)} (Message {slots})."
                )
            total_attachments = len(payload.get("attachments") or [])
            if total_attachments > 0:
                _send_log(
                    f"Attachment mode: {'random' if random_attachment_enabled else 'fixed'} | "
                    f"count range: {media_count_min}-{media_count_max} | uploaded: {total_attachments}."
                )
            _send_log(
                f"Search result click delay: "
                f"{'random' if random_search_result_delay_enabled else 'fixed'} "
                f"({search_result_delay_min}-{search_result_delay_max}s)."
            )
            _send_log(
                f"Automation click delay: "
                f"{'random' if random_click_delay_enabled else 'fixed'} "
                f"({click_delay_min_ms}-{click_delay_max_ms} ms)."
            )
            _send_log(
                f"Text send delay: {'random' if random_send_delay_enabled else 'fixed'} "
                f"({send_delay_min}-{send_delay_max}s). "
                f"Attachment delay: {'random' if random_attachment_delay_enabled else 'fixed'} "
                f"({attachment_delay_min}-{attachment_delay_max}s)."
            )
            _send_log(
                f"Remove characters while typing: "
                f"{'enabled' if remove_chars_during_typing_enabled else 'disabled'}"
                + (
                    f" | times={remove_chars_times_min}-{remove_chars_times_max} per SMS"
                    f" | chars={remove_chars_count_min}-{remove_chars_count_max}."
                    if remove_chars_during_typing_enabled
                    else "."
                )
            )
            _send_log(
                f"Optional WhatsApp side actions: "
                f"{'enabled' if optional_side_actions_enabled else 'disabled'}"
                + (
                    f" | chance={side_action_chance_percent}% | pause={side_action_pause_min_ms}-{side_action_pause_max_ms} ms."
                    if optional_side_actions_enabled
                    else "."
                )
            )
            _send_log(
                f"Batch start delay: {'random' if random_batch_start_enabled else 'fixed'} "
                f"({batch_start_min_minutes}-{batch_start_max_minutes} min)."
            )
            preassigned_plan = {}
            for target_index, target in enumerate(targets):
                _raise_if_send_cancel_requested(
                    state,
                    message="Campaign canceled by user. All WhatsApp browsers were closed.",
                )
                target = str(target or "").strip()
                if not target:
                    continue
                if random_profile_pick_enabled and worker_pages:
                    worker = random.choice(worker_pages)
                else:
                    worker = worker_pages[target_index % max(1, len(worker_pages))]
                worker_name = str(worker.get("name") or "Profile 1")
                worker_profile_display = str(worker.get("profile_number") or "").strip() or worker_name
                worker_page = worker.get("page")
                batch_no = int(batch_numbers[target_index]) if target_index < len(batch_numbers) else 1
                selected_slot, selected_message, active_count = _pick_random_message_for_target(
                    payload,
                    random_enabled=random_message_enabled,
                    slot_min=1,
                    slot_max=message_panel_count,
                )
                selected_attachments, selected_attachment_count, selected_attachment_total, selected_attachment_indexes = _pick_random_attachment_for_target(
                    payload,
                    random_enabled=random_attachment_enabled,
                    min_count=media_count_min,
                    max_count=media_count_max,
                )
                assigned_send_delay = random.randint(send_delay_min, send_delay_max) if random_send_delay_enabled else send_delay_min
                assigned_search_result_delay = (
                    random.randint(search_result_delay_min, search_result_delay_max)
                    if random_search_result_delay_enabled
                    else search_result_delay_min
                )
                assigned_media_delay = None
                if selected_attachment_count > 0:
                    assigned_media_delay = random.randint(attachment_delay_min, attachment_delay_max) if random_attachment_delay_enabled else attachment_delay_min
                message_label = f"txt{int(selected_slot)}" if active_count > 0 else "txt1"
                media_label = "-"
                planned_action = _planned_automation_action_for_target(target_index)
                if selected_attachment_count > 0:
                    attachment_names = [
                        str((item or {}).get("name") or "").strip() or "unnamed"
                        for item in selected_attachments
                    ]
                    media_label = ", ".join(attachment_names)
                preassigned_plan[target_index] = {
                    "target": target,
                    "worker_name": worker_name,
                    "worker_profile_display": worker_profile_display,
                    "worker_page": worker_page,
                    "batch_no": batch_no,
                    "selected_slot": selected_slot,
                    "selected_message": selected_message,
                    "selected_attachments": selected_attachments,
                    "selected_attachment_count": selected_attachment_count,
                    "selected_attachment_total": selected_attachment_total,
                    "selected_attachment_indexes": selected_attachment_indexes,
                    "active_count": active_count,
                    "message_label": message_label,
                    "media_label": media_label,
                    "assigned_search_result_delay": int(assigned_search_result_delay),
                    "assigned_send_delay": int(assigned_send_delay),
                    "assigned_media_delay": (int(assigned_media_delay) if assigned_media_delay is not None else None),
                    "automation_action_id": str(planned_action.get("id") or "").strip(),
                    "automation_action_label": str(planned_action.get("label") or "").strip(),
                    "automation_action_phase": str(planned_action.get("phase") or "after_send").strip() or "after_send",
                }
                _set_automation_action_assignment(
                    state,
                    target_index,
                    str(planned_action.get("id") or "").strip(),
                    phase=str(planned_action.get("phase") or "after_send").strip() or "after_send",
                )
                worker["assigned"].append(target_index)
                _upsert_profile_state(
                    worker_name,
                    targets_assigned=int(((next((r for r in (state.get("profile_states") or []) if str((r or {}).get("name") or "") == worker_name), {}) or {}).get("targets_assigned") or 0)) + 1,
                )
                _update_live_row(
                    state,
                    target_index,
                    profile=worker_profile_display,
                    batch=batch_no,
                    message=message_label,
                    media=media_label,
                    media_delay=(f"{int(assigned_media_delay)}s" if assigned_media_delay is not None else "-"),
                    send_delay=f"{int(assigned_send_delay)}s",
                    status="pending",
                )
            _send_log("Pre-assignment complete for all targets (Profile, Message, Media, Media_delay, Send_delay).")
            def _process_target(worker, target_index, target, plan_item):
                _raise_if_send_cancel_requested(
                    state,
                    message="Campaign canceled by user. All WhatsApp browsers were closed.",
                )
                target = str(target or "").strip()
                if not target:
                    return
                plan_item = dict(plan_item or {})
                worker_name = str(plan_item.get("worker_name") or "Profile 1")
                worker_profile_display = str(plan_item.get("worker_profile_display") or "").strip() or worker_name
                page = worker.get("page")
                if page is None:
                    raise RuntimeError("WhatsApp browser page is not available for this worker.")
                _wait_if_send_paused(state, page=page)
                batch_no = int(plan_item.get("batch_no") or 1)
                worker["last_batch_no"] = batch_no
                worker["processed_count"] = int(worker.get("processed_count") or 0) + 1
                with _SEND_STATE_LOCK:
                    state["current_batch_no"] = max(int(state.get("current_batch_no") or 0), batch_no)
                _send_log(f"Processing: {target} via {worker_name}")
                if not _is_whatsapp_profile_logged_in(page):
                    worker_idx = _profile_index_from_label(worker_name)
                    worker_profile_number = ""
                    try:
                        worker_profile_number = _ensure_and_read_profile_number(
                            worker.get("profile_dir"),
                            seed_value=str(worker.get("profile_number") or "").strip(),
                        )
                    except Exception:
                        worker_profile_number = str(worker.get("profile_number") or "").strip()
                    if worker_idx > 0:
                        _set_profile_number_in_sequence(
                            worker_idx,
                            worker_profile_number,
                            min_size=profiles_numbers,
                        )
                    _upsert_profile_state(
                        worker_name,
                        profile_number=worker_profile_number,
                        status="logged_out",
                        reason="WhatsApp session logged out",
                        last_target=target,
                    )
                    worker_profile_display = str(worker_profile_number or "").strip() or worker_name
                else:
                    _upsert_profile_state(
                        worker_name,
                        status="logged_in",
                        reason="",
                        last_target=target,
                    )
                _update_live_row(state, target_index, profile=worker_profile_display, status="processing")
                selected_slot = int(plan_item.get("selected_slot") or 1)
                selected_message = str(plan_item.get("selected_message") or "")
                active_count = int(plan_item.get("active_count") or 0)
                if active_count > 0:
                    preview = " ".join(selected_message.split())
                    if len(preview) > 70:
                        preview = preview[:67] + "..."
                    _send_log(
                        f"Message selected for {target}: Message {selected_slot} "
                        f"from {active_count} panel(s)"
                        + (f" | text='{preview}'" if preview else " | text=(empty)")
                    )
                selected_attachments = list(plan_item.get("selected_attachments") or [])
                selected_attachment_count = int(plan_item.get("selected_attachment_count") or 0)
                selected_attachment_total = int(plan_item.get("selected_attachment_total") or 0)
                selected_attachment_indexes = list(plan_item.get("selected_attachment_indexes") or [])
                message_label = str(plan_item.get("message_label") or (f"txt{int(selected_slot)}" if active_count > 0 else "txt1"))
                media_label = str(plan_item.get("media_label") or "-")
                assigned_action_id = str(plan_item.get("automation_action_id") or "").strip()
                assigned_action_label = str(plan_item.get("automation_action_label") or "").strip()
                assigned_action_phase = str(plan_item.get("automation_action_phase") or "after_send").strip() or "after_send"
                assigned_search_result_delay = int(
                    plan_item.get("assigned_search_result_delay") or search_result_delay_min
                )
                assigned_send_delay = int(plan_item.get("assigned_send_delay") or send_delay_min)
                assigned_media_delay = plan_item.get("assigned_media_delay")
                _update_automation_action_row(
                    state,
                    target_index,
                    action_id=assigned_action_id,
                    action=assigned_action_label or "-",
                    phase=assigned_action_phase,
                    checked=False,
                    status="assigned",
                    detail="",
                    performed_at="",
                )
                if selected_attachment_count > 0:
                    _send_log(
                        f"Attachments selected for {target}: "
                        f"{selected_attachment_count}/{selected_attachment_total} "
                        f"(indexes: {', '.join(str(i) for i in selected_attachment_indexes)})"
                    )
                def _on_media_delay_selected(delay_seconds, planned_click_at):
                    _update_live_row(
                        state,
                        target_index,
                        media_delay=f"{int(delay_seconds)}s -> {planned_click_at}",
                    )

                def _on_send_delay_selected(delay_seconds, planned_click_at):
                    _update_live_row(
                        state,
                        target_index,
                        send_delay=f"{int(delay_seconds)}s -> {planned_click_at}",
                    )
                target_payload = build_payload(
                    selected_message,
                    selected_attachments,
                    message_variants=payload.get("message_variants") or [],
                )

                if not allow_resend and send_target != "groups":
                    if not is_valid_phone(target):
                        _send_log(f"Invalid format: {target}")
                        record_failed(target, "Invalid format")
                        _update_live_row(state, target_index, status="failed")
                        return
                    try:
                        with db_lock:
                            already_sent = is_already_sent(db_conn, target, target_payload)
                        if already_sent:
                            _send_log(f"Already exists in table. Skipping before open: {target}")
                            set_contact_status(target, "skipped", "Already sent")
                            _update_live_row(state, target_index, status="skipped")
                            return
                    except Exception as db_exc:
                        _send_log(f"DB warning: duplicate check failed for {target}: {db_exc}")

                sent_committed = False
                try:
                    page.keyboard.press("Escape")
                    page.wait_for_timeout(20)

                    if send_target == "groups":
                        open_started = time.time()
                        opened, open_err = _open_group_chat_like_scrape_flow(page, target, state=state)
                        if not opened:
                            _send_log(f"Group open failed: {target} ({open_err})")
                            record_failed(target, open_err or "Could not open group chat")
                            _update_live_row(state, target_index, status="failed")
                            return
                        elapsed_ms = int((time.time() - open_started) * 1000)
                        _send_log(f"Opened group chat via Groups tab: {target} ({elapsed_ms} ms) [{worker_name}]")
                    else:
                        open_started = time.time()
                        if not _click_first_selector(
                            page,
                            (
                                "button[aria-label*='New']",
                                "div[role='button'][aria-label*='New chat']",
                                "[title='New chat']",
                            ),
                            timeout_ms=3200,
                            label="Open New Chat",
                            state=state,
                        ):
                            _delayed_click_fallback(
                                page,
                                page.locator("button[aria-label*='New']").first,
                                label="Open New Chat",
                                timeout_ms=3200,
                                state=state,
                            )
                        _wait_if_send_paused(state, page=page)
                        opened, open_err = _open_contact_chat_from_search(
                            page,
                            target,
                            state=state,
                            result_click_delay_seconds=assigned_search_result_delay,
                        )
                        if not opened:
                            _send_log(f"Search/open failed: {target} ({open_err})")
                            record_failed(target, open_err or "No result found")
                            _update_live_row(state, target_index, status="failed")
                            page.keyboard.press("Escape")
                            return
                        elapsed_ms = int((time.time() - open_started) * 1000)
                        _send_log(f"Opened chat from search: {target} ({elapsed_ms} ms) [{worker_name}]")

                    if target_payload.get("attachments"):
                        first_caption = target_payload.get("message") or ""
                        for index, attachment in enumerate(target_payload.get("attachments") or []):
                            send_single_attachment(
                                page,
                                attachment,
                                first_caption if index == 0 else "",
                                on_delay=_on_media_delay_selected,
                                on_send_delay=_on_send_delay_selected,
                                random_delay=False,
                                delay_min_seconds=int(assigned_media_delay),
                                delay_max_seconds=int(assigned_media_delay),
                                send_click_random_delay=False,
                                send_click_delay_min_seconds=assigned_send_delay,
                                send_click_delay_max_seconds=assigned_send_delay,
                                typing_pause_chars_min=typing_pause_chars_min,
                                typing_pause_chars_max=typing_pause_chars_max,
                                typing_pause_delay_min_ms=typing_pause_delay_min_ms,
                                typing_pause_delay_max_ms=typing_pause_delay_max_ms,
                                remove_chars_during_typing_enabled=remove_chars_during_typing_enabled,
                                remove_chars_times_min=remove_chars_times_min,
                                remove_chars_times_max=remove_chars_times_max,
                                remove_chars_count_min=remove_chars_count_min,
                                remove_chars_count_max=remove_chars_count_max,
                                state=state,
                            )
                    else:
                        msg_box = _prepare_active_message_box(page, state=state, timeout_ms=5000)
                        if msg_box is None:
                            raise RuntimeError("Message box not ready after opening selected chat")
                        _send_plain_text_single_message(
                            page,
                            msg_box,
                            target_payload.get("message") or "",
                            on_delay=_on_send_delay_selected,
                            random_delay=False,
                            delay_min_seconds=assigned_send_delay,
                            delay_max_seconds=assigned_send_delay,
                            typing_pause_chars_min=typing_pause_chars_min,
                            typing_pause_chars_max=typing_pause_chars_max,
                            typing_pause_delay_min_ms=typing_pause_delay_min_ms,
                            typing_pause_delay_max_ms=typing_pause_delay_max_ms,
                            remove_chars_during_typing_enabled=remove_chars_during_typing_enabled,
                            remove_chars_times_min=remove_chars_times_min,
                            remove_chars_times_max=remove_chars_times_max,
                            remove_chars_count_min=remove_chars_count_min,
                            remove_chars_count_max=remove_chars_count_max,
                            state=state,
                        )

                    _send_log(f"Message sent to: {target} [{worker_name}]")
                    with _SEND_STATE_LOCK:
                        state["sent"] += 1
                    _upsert_profile_state(
                        worker_name,
                        sent=int(((next((r for r in (state.get("profile_states") or []) if str((r or {}).get("name") or "") == worker_name), {}) or {}).get("sent") or 0)) + 1,
                        last_target=target,
                    )
                    set_contact_status(target, "sent", "")
                    _update_live_row(state, target_index, profile=worker_profile_display, status="sent")
                    sent_committed = True
                    try:
                        with db_lock:
                            save_sent_sms(
                                db_conn,
                                target,
                                target_payload,
                                campaign_id=campaign_id,
                                campaign_name=campaign_name,
                            )
                    except Exception as db_exc:
                        _send_log(f"DB warning: failed to save sms_sent for {target}: {db_exc}")

                    action_done_for_target = False
                    actual_action_id = ""
                    action_detail = ""
                    _update_automation_action_row(
                        state,
                        target_index,
                        status="performing",
                        detail="",
                    )
                    action_done_for_target, actual_action_id, action_detail = _perform_required_automation_action(
                        page,
                        state=state,
                        row_index=target_index,
                        target_label=target,
                        phase=assigned_action_phase,
                        preferred_action_id=assigned_action_id,
                    )
                    if not action_done_for_target:
                        action_done_for_target, actual_action_id, action_detail = _perform_final_safe_automation_fallback(
                            page,
                            state=state,
                            row_index=target_index,
                            target_label=target,
                        )
                    if not action_done_for_target:
                        _update_automation_action_row(
                            state,
                            target_index,
                            status="not_performed",
                            detail="No automation action could be completed after send.",
                        )
                        _send_log(
                            f"Automation action warning for {target}: no action could be completed after send."
                        )

                    page.keyboard.press("Escape")
                    page.wait_for_timeout(60)

                except SendJobCancelled:
                    raise
                except Exception as e:
                    if _is_send_cancel_requested(state):
                        raise SendJobCancelled(
                            "Campaign canceled by user. All WhatsApp browsers were closed."
                        )
                    if sent_committed:
                        _send_log(
                            f"Post-send warning for {target}: {e}. Message already sent; status kept as sent."
                        )
                        try:
                            page.keyboard.press("Escape")
                            page.wait_for_timeout(120)
                        except Exception:
                            pass
                        return
                    _send_log(f"Error for {target}: {e}")
                    err_text = str(e or "")
                    if any(
                        token in err_text.lower()
                        for token in ("qr", "log in", "login", "logged out", "not logged", "pane-side")
                    ):
                        worker_idx = _profile_index_from_label(worker_name)
                        worker_profile_number = ""
                        try:
                            worker_profile_number = _ensure_and_read_profile_number(
                                worker.get("profile_dir"),
                                seed_value=str(worker.get("profile_number") or "").strip(),
                            )
                        except Exception:
                            worker_profile_number = str(worker.get("profile_number") or "").strip()
                        if worker_idx > 0:
                            _set_profile_number_in_sequence(
                                worker_idx,
                                worker_profile_number,
                                min_size=profiles_numbers,
                            )
                        _upsert_profile_state(
                            worker_name,
                            profile_number=worker_profile_number,
                            status="logged_out",
                            reason=f"Send error: {err_text}"[:280],
                            last_target=target,
                        )
                        worker_profile_display = str(worker_profile_number or "").strip() or worker_name
                    _upsert_profile_state(
                        worker_name,
                        failed=int(((next((r for r in (state.get("profile_states") or []) if str((r or {}).get("name") or "") == worker_name), {}) or {}).get("failed") or 0)) + 1,
                        last_target=target,
                    )
                    record_failed(target, f"Error: {e}")
                    _update_live_row(state, target_index, profile=worker_profile_display, status="failed")
                    try:
                        page.keyboard.press("Escape")
                        page.wait_for_timeout(120)
                    except Exception:
                        pass
                    return


            # Release every logged-in session at the same time; each one works
            # through its own assigned targets in its own thread.
            _assign_per_worker_batches(worker_pages)
            for worker in worker_pages:
                _send_log(
                    f"[{worker['name']}] {len(worker.get('assigned') or [])} target(s) assigned; starting in parallel."
                )
                worker["start_event"].set()
            while True:
                _raise_if_send_cancel_requested(
                    state,
                    message="Campaign canceled by user. All WhatsApp browsers were closed.",
                )
                if all(w["done_event"].is_set() for w in worker_pages):
                    break
                time.sleep(0.3)
            stopped_workers = [
                w for w in worker_pages if w.get("error") and w.get("error") != "cancelled"
            ]
            if stopped_workers:
                _send_log(
                    "Some sessions stopped early: "
                    + "; ".join(f"{w['name']}: {w['error']}" for w in stopped_workers)
                )
            _send_log("Done. You can close the browser or keep it open.")
    finally:
        for worker in list(worker_specs):
            try:
                worker["abort_event"].set()
            except Exception:
                pass
        for worker in list(worker_specs):
            thread = worker.get("thread")
            try:
                if thread is not None and thread.is_alive() and thread is not threading.current_thread():
                    thread.join(timeout=20.0)
            except Exception:
                pass
        _set_active_send_contexts(state, [])
        # Browsers are closed by their own worker threads; release any profile
        # lock that a thread could not release in time.
        for ctx, profile_dir in list(worker_contexts):
            release_playwright_profile(profile_dir)
        try:
            db_conn.close()
        except Exception:
            pass


app = Flask(__name__)
app.secret_key = "whatsapp-bot-secret"

send_state = {
    "running": False,
    "last_error": None,
    "last_result": None,
    "last_failed": [],
    "failed": [],
    "sent": 0,
    "total": 0,
    "last_message": "",
    "last_messages": [""] * 100,
    "last_delay": "0",
    "last_country_filter": "",
    "last_filename": "",
    "last_attachment_names": [],
    "last_resend": False,
    "last_source_type": "file",
    "last_send_target": "contacts",
    "last_selected_tables": [],
    "last_group_filter_name": "",
    "last_group_filter_names": [],
    "last_selected_contacts": [],
    "contacts": [],
    "last_contacts": [],
    "contact_statuses": {},
    "last_contact_statuses": {},
    "campaign_id": "",
    "last_campaign_id": "",
    "campaign_name": "",
    "last_campaign_name": "",
    "live_rows": [],
    "last_live_rows": [],
    "automation_action_rows": [],
    "last_automation_action_rows": [],
    "automation_action_event_text": "",
    "last_automation_action_event_text": "",
    "campaign_actions": [],
    "last_campaign_actions": [],
    "campaign_action_event_text": "",
    "last_campaign_action_event_text": "",
    "batch_pause_active": False,
    "batch_pause_remaining_sec": 0,
    "batch_next_start_at": "",
    "batch_pauses": {},
    "pause_requested": False,
    "paused": False,
    "pause_requested_at": "",
    "cancel_requested": False,
    "cancel_requested_at": "",
    "session_opening": False,
    "awaiting_session_start": False,
    "session_start_requested": False,
    "session_logged_in_count": 0,
    "session_logged_out_count": 0,
    "current_batch_no": 0,
    "total_batches": 0,
    "last_batch_min_size": 20,
    "last_batch_max_size": 30,
    "last_random_message_enabled": True,
    "last_message_panel_count": 1,
    "last_random_attachment_enabled": True,
    "last_media_count_min": 1,
    "last_media_count_max": 1,
    "last_random_attachment_delay_enabled": True,
    "last_attachment_delay_min": 0,
    "last_attachment_delay_max": 30,
    "last_random_search_result_delay_enabled": True,
    "last_search_result_delay_min": 1,
    "last_search_result_delay_max": 4,
    "last_random_click_delay_enabled": True,
    "last_click_delay_min_ms": 150,
    "last_click_delay_max_ms": 900,
    "last_random_send_delay_enabled": True,
    "last_send_delay_min": 0,
    "last_send_delay_max": 30,
    "last_typing_pause_chars_min": 10,
    "last_typing_pause_chars_max": 20,
    "last_typing_pause_delay_min": 2000,
    "last_typing_pause_delay_max": 5000,
    "last_typing_pause_delay_unit": "ms",
    "last_remove_chars_during_typing_enabled": True,
    "last_remove_chars_times_min": 3,
    "last_remove_chars_times_max": 7,
    "last_remove_chars_count_min": 1,
    "last_remove_chars_count_max": 3,
    "last_optional_side_actions_enabled": True,
    "last_side_action_chance_percent": 20,
    "last_side_action_pause_min_ms": 700,
    "last_side_action_pause_max_ms": 2000,
    "last_random_batch_size_enabled": True,
    "last_random_batch_start_enabled": True,
    "last_batch_start_min_minutes": 3,
    "last_batch_start_max_minutes": 6,
    "last_profiles_numbers": 1,
    "last_random_profile_pick_enabled": True,
    "profile_states": [],
    "last_profile_states": [],
    "_active_send_contexts": [],
}

app_log_state = {
    "logs": [],
}
app_log_lock = threading.Lock()

scrape_state = {
    "running": False,
    "last_error": None,
    "last_result": None,
    "logs": [],
    "groups_total": 0,
    "groups_done": 0,
    "groups_skipped": 0,
    "groups": [],
    "started_at": None,
    "finished_at": None,
    "scrap_id": "",
}


def _append_app_log(message, source="APP"):
    ts = datetime.now().strftime("%H:%M:%S")
    line = f"[{ts}] [{str(source or 'APP').upper()}] {message}"
    with app_log_lock:
        app_log_state["logs"].append(line)
        if len(app_log_state["logs"]) > 1000:
            app_log_state["logs"] = app_log_state["logs"][-1000:]
    print(f"[{str(source or 'APP').upper()}] {message}")
    return line


def _send_log(message):
    return _append_app_log(message, "SEND")


def _scrape_log(message):
    scrape_state["logs"].append(_append_app_log(message, "SCRAPE"))
    # Keep memory bounded
    if len(scrape_state["logs"]) > 400:
        scrape_state["logs"] = scrape_state["logs"][-400:]


def _reset_scrape_state():
    scrape_state["last_error"] = None
    scrape_state["last_result"] = None
    scrape_state["logs"] = []
    scrape_state["groups_total"] = 0
    scrape_state["groups_done"] = 0
    scrape_state["groups_skipped"] = 0
    scrape_state["groups"] = []
    scrape_state["started_at"] = None
    scrape_state["finished_at"] = None
    scrape_state["scrap_id"] = ""


def ensure_scrape_runs_table(db_conn):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                CREATE TABLE IF NOT EXISTS {schema_name}.scrape_group_runs (
                    id BIGSERIAL PRIMARY KEY,
                    scrap_id TEXT NOT NULL,
                    group_index INTEGER NOT NULL,
                    groups_total INTEGER NOT NULL DEFAULT 0,
                    group_name TEXT NOT NULL,
                    scrapped_contacts_count INTEGER NOT NULL DEFAULT 0,
                    total_contacts_in_loop INTEGER NOT NULL DEFAULT 0,
                    imported_groups_in_loop INTEGER NOT NULL DEFAULT 0,
                    loop_started_at TIMESTAMPTZ NOT NULL,
                    event_time TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        cur.execute(
            sql.SQL(
                """
                CREATE INDEX IF NOT EXISTS idx_scrape_group_runs_scrap_id
                ON {schema_name}.scrape_group_runs (scrap_id, event_time)
                """
            ).format(schema_name=sql.Identifier(schema))
        )


def save_scrape_group_row(
    db_conn,
    scrap_id,
    group_index,
    groups_total,
    group_name,
    scrapped_contacts_count,
    total_contacts_in_loop,
    imported_groups_in_loop,
    loop_started_at,
):
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                INSERT INTO {schema_name}.scrape_group_runs (
                    scrap_id,
                    group_index,
                    groups_total,
                    group_name,
                    scrapped_contacts_count,  
                    total_contacts_in_loop,
                    imported_groups_in_loop,
                    loop_started_at,
                    event_time 
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, NOW())
                """
            ).format(schema_name=sql.Identifier(schema)),
            (
                scrap_id,
                int(group_index),
                int(groups_total),
                group_name,
                int(scrapped_contacts_count),
                int(total_contacts_in_loop),
                int(imported_groups_in_loop),
                loop_started_at,
            ),
        )


def update_scrape_groups_total_for_run(db_conn, scrap_id, groups_total):
    if not scrap_id:
        return
    schema = get_db_schema()
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                UPDATE {schema_name}.scrape_group_runs
                SET groups_total = %s
                WHERE scrap_id = %s
                """
            ).format(schema_name=sql.Identifier(schema)),
            (int(groups_total), scrap_id),
        )


def load_scrape_rows(scrap_id):
    if not scrap_id:
        return []
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT
                        group_index,
                        groups_total,
                        group_name,
                        scrapped_contacts_count,
                        total_contacts_in_loop,
                        imported_groups_in_loop,
                        loop_started_at,
                        event_time
                    FROM {schema_name}.scrape_group_runs
                    WHERE scrap_id = %s
                    ORDER BY group_index ASC, event_time ASC
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (scrap_id,),
            )
            rows = cur.fetchall()
            return [
                {
                    "group_index": r[0],
                    "groups_total": r[1],
                    "group_name": r[2],
                    "scrapped_contacts_count": r[3],
                    "total_contacts_in_loop": r[4],
                    "imported_groups_in_loop": r[5],
                    "loop_started_at": r[6].isoformat() if r[6] else "",
                    "event_time": r[7].isoformat() if r[7] else "",
                }
                for r in rows
            ]
    except Exception as exc:
        print(f"[SCRAPE DB WARN] Could not load scrape rows: {exc}")
        return []
    finally:
        db_conn.close()


def list_scrap_ids(limit=100):
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT scrap_id, MAX(event_time) AS last_time
                    FROM {schema_name}.scrape_group_runs
                    GROUP BY scrap_id
                    ORDER BY last_time DESC
                    LIMIT %s
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (int(limit),),
            )
            return [row[0] for row in cur.fetchall() if row and row[0]]
    except Exception as exc:
        print(f"[SCRAPE DB WARN] Could not list scrap IDs: {exc}")
        return []
    finally:
        db_conn.close()


def generate_unique_numeric_scrap_id(db_conn, digits=6, max_attempts=300):
    schema = get_db_schema()
    d = max(6, int(digits))
    with db_conn.cursor() as cur:
        cur.execute(
            sql.SQL(
                """
                SELECT COALESCE(
                    MAX(
                        CASE
                            WHEN scrap_id ~ '^[0-9]+$' THEN scrap_id::BIGINT
                            ELSE NULL
                        END
                    ),
                    0
                )
                FROM {schema_name}.scrape_group_runs
                """
            ).format(schema_name=sql.Identifier(schema))
        )
        row = cur.fetchone()
        last_value = int(row[0]) if row and row[0] is not None else 0
        next_value = last_value + 1
        return str(next_value).zfill(d)


def search_groups_across_scrap_runs(query_text, limit=300):
    db_conn = get_db_connection()
    schema = get_db_schema()
    q = (query_text or "").strip().lower()
    try:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT
                        scrap_id,
                        group_name,
                        MAX(event_time) AS last_event_time
                    FROM {schema_name}.scrape_group_runs
                    WHERE group_name IS NOT NULL
                      AND group_name <> ''
                      AND scrap_id NOT LIKE '%-GROUP%'
                      AND scrap_id NOT LIKE '%-RECOVERY-%'
                      AND LOWER(group_name) LIKE %s
                    GROUP BY scrap_id, group_name
                    ORDER BY last_event_time DESC, scrap_id ASC, group_name ASC
                    LIMIT %s
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (f"%{q}%", int(limit)),
            )
            rows = cur.fetchall()
            results = []
            for row in rows:
                if not row:
                    continue
                scrap_id = row[0] or ""
                group_name = row[1] or ""
                last_event_time = row[2].isoformat() if row[2] else ""
                if not scrap_id or not group_name:
                    continue
                results.append(
                    {
                        "scrap_id": scrap_id,
                        "group_name": group_name,
                        "last_event_time": last_event_time,
                    }
                )
            return results
    finally:
        db_conn.close()


def list_unique_group_names(limit=2000):
    db_conn = get_db_connection()
    schema = get_db_schema()
    try:
        with db_conn.cursor() as cur:
            cur.execute(
                sql.SQL(
                    """
                    SELECT MIN(group_name) AS group_name
                    FROM {schema_name}.scrape_group_runs
                    WHERE group_name IS NOT NULL
                      AND TRIM(group_name) <> ''
                    GROUP BY LOWER(TRIM(group_name))
                    ORDER BY group_name ASC
                    LIMIT %s
                    """
                ).format(schema_name=sql.Identifier(schema)),
                (int(limit),),
            )
            return [row[0] for row in cur.fetchall() if row and row[0]]
    finally:
        db_conn.close()


def build_scrape_summary_from_rows(scrap_id, rows):
    if not rows:
        return {
            "scrap_id": scrap_id,
            "groups_total": 0,
            "groups_done": 0,
            "groups_skipped": 0,
            "total_contacts": 0,
            "last_result": None,
            "started_at": None,
            "finished_at": None,
        }
    groups_total = max(int(r.get("groups_total", 0)) for r in rows)
    groups_done = max(int(r.get("imported_groups_in_loop", 0)) for r in rows)
    groups_skipped = max(int(r.get("groups_skipped_in_loop", 0)) for r in rows)
    total_contacts = max(int(r.get("total_contacts_in_loop", 0)) for r in rows)
    started_at = rows[0].get("loop_started_at")
    finished_at = rows[-1].get("event_time")
    return {
        "scrap_id": scrap_id,
        "groups_total": groups_total,
        "groups_done": groups_done,
        "groups_skipped": groups_skipped,
        "total_contacts": total_contacts,
        "last_result": f"Scrapping complete. Groups imported: {groups_done}. Total contacts: {total_contacts}.",
        "started_at": started_at,
        "finished_at": finished_at,
    }


def run_scrape_job(profile_number):
    scrape_state["running"] = True
    _reset_scrape_state()
    scrape_state["started_at"] = datetime.now().isoformat()
    _scrape_log("Starting WhatsApp contacts scrapping...")

    browser = None
    db_conn = None
    user_data_dir = None
    existing_group_names = set()
    logged_group_checks = set()
    try:
        db_conn = get_db_connection()
        ensure_scrape_runs_table(db_conn)
        existing_group_names = load_existing_group_names_from_all_contacts(db_conn)
        _scrape_log(
            f"Existing groups in DB: {len(existing_group_names)}. These groups will be skipped."
        )
        scrape_state["scrap_id"] = generate_unique_numeric_scrap_id(db_conn, digits=6)
        _scrape_log(f"Scrap ID: {scrape_state['scrap_id']}")

        user_data_dir = get_playwright_user_data_dir()
        profile_number_clean = str(profile_number or "").strip()
        if not profile_number_clean:
            raise RuntimeError("Profile number is required for launch.")

        with sync_playwright() as p:
            launch_kwargs = get_playwright_launch_kwargs(default_headless=False)
            _scrape_log(
                "Playwright config -> "
                f"headless={launch_kwargs.get('headless')} "
                f"start_maximized={'args' in launch_kwargs} "
                f"no_viewport={launch_kwargs.get('no_viewport', False)} "
                f"slow_mo={launch_kwargs.get('slow_mo', 0)}ms"
            )
            _scrape_log(f"Playwright user_data_dir: {user_data_dir}")
            _write_profile_number(user_data_dir, profile_number_clean)
            _set_profile_number_in_sequence(1, profile_number_clean, min_size=1)
            browser = launch_persistent_context_safe(
                browser_type=p.chromium,
                profile_dir=user_data_dir,
                launch_kwargs=launch_kwargs,
                owner_label="scrapping",
                log_fn=_scrape_log,
            )
            page = browser.pages[0] if browser.pages else browser.new_page()
            _scrape_log("Opening WhatsApp Web...")
            page.goto("https://web.whatsapp.com")
            login_selector = _wait_for_any_selector(
                page,
                (
                    "[data-testid='chat-list']",
                    "div[aria-label='Chat list']",
                    "#pane-side",
                ),
                timeout_ms=5000,
            )
            login_state = "reused_logged_in" if login_selector else "logged_in_after_wait"
            if not login_selector:
                _wait_for_any_selector(
                    page,
                    (
                        "[data-testid='chat-list']",
                        "div[aria-label='Chat list']",
                        "#pane-side",
                    ),
                    timeout_ms=120000,
                )
            _write_profile_number(user_data_dir, profile_number_clean)
            _set_profile_number_in_sequence(1, profile_number_clean, min_size=1)
            _scrape_log("WhatsApp loaded. Logged in.")
            try:
                _, session_line = save_whatsapp_session_snapshot(
                    page=page,
                    flow_name="run_scrape_job",
                    user_data_dir=str(user_data_dir),
                    login_state=login_state,
                )
                _scrape_log(f"Session snapshot -> {session_line}")
            except Exception as session_exc:
                _scrape_log(f"Session snapshot warning: {session_exc}")

            if not open_group_and_members(page):
                raise RuntimeError("Could not open Groups filter.")

            rows_selector = "div[aria-label='Chat list'][role='grid'] > div[role='row']"
            scrape_state["groups_total"] = page.locator(rows_selector).count()
            _scrape_log(f"Total groups found: {scrape_state['groups_total']}")

            def _on_group_complete(item):
                group_name = item.get("group", "group")
                contacts = item.get("contacts", []) or []
                metadata = dict(item.get("metadata") or {})
                metadata["group_name"] = group_name
                metadata["source_scrap_id"] = scrape_state["scrap_id"]
                metadata["last_scraped_at"] = scrape_state["started_at"]

                group_key = _normalize_group_key(group_name)
                if group_key:
                    existing_group_names.add(group_key)

                scrape_state["groups_done"] += 1
                loop_total_contacts = sum(
                    g.get("contacts_count", 0) for g in scrape_state["groups"]
                ) + len(contacts)
                group_row = {
                    "name": group_name,
                    "contacts_count": len(contacts),
                    "db_table": "",
                    "group_index": scrape_state["groups_done"],
                    "groups_total": scrape_state["groups_total"],
                    "loop_total_contacts": loop_total_contacts,
                    "imported_groups_in_loop": scrape_state["groups_done"],
                    "metadata": metadata,
                }
                scrape_state["groups"].append(group_row)

                if db_conn:
                    try:
                        save_scrape_group_row(
                            db_conn=db_conn,
                            scrap_id=scrape_state["scrap_id"],
                            group_index=scrape_state["groups_done"],
                            groups_total=scrape_state["groups_total"],
                            group_name=group_name,
                            scrapped_contacts_count=len(contacts),
                            total_contacts_in_loop=loop_total_contacts,
                            imported_groups_in_loop=scrape_state["groups_done"],
                            loop_started_at=scrape_state["started_at"],
                        )
                    except Exception as exc:
                        _scrape_log(f"DB log warning for group {group_name}: {exc}")

                table_name = ""
                try:
                    if not db_conn:
                        raise RuntimeError("Database connection is not available.")
                    table_name = save_scraped_group_contacts(
                        db_conn=db_conn,
                        group_name=group_name,
                        contacts=contacts,
                        metadata=metadata,
                    )
                    group_row["db_table"] = table_name
                except Exception as exc:
                    _scrape_log(f"Contacts DB save warning for group {group_name}: {exc}")

                _scrape_log(
                    f"{group_name}: {len(contacts)} contacts imported "
                    f"(DB: {group_row['db_table'] or 'warning'})."
                )

            def _should_skip_group(group_name):
                group_key = _normalize_group_key(group_name)
                if not group_key:
                    return False
                if group_key in existing_group_names:
                    if group_key not in logged_group_checks:
                        scrape_state["groups_skipped"] += 1
                        _scrape_log(
                            f"Skipped:[ {group_name} ] already exists in list."
                        )
                        logged_group_checks.add(group_key)
                    return True
                return False

            def _on_group_skipped(event):
                reason = str((event or {}).get("reason") or "").strip().lower()
                group_name = str((event or {}).get("group") or "").strip()
                if reason == "community":
                    scrape_state["groups_total"] = max(
                        0, int(scrape_state.get("groups_total", 0)) - 1
                    )
                    if group_name:
                        _scrape_log(
                            f"Community skipped (not counted in groups): {group_name}"
                        )

            extract_all_groups(
                page,
                on_group_complete=_on_group_complete,
                should_skip_group=_should_skip_group,
                on_group_skipped=_on_group_skipped,
            )

            if db_conn and scrape_state.get("scrap_id"):
                try:
                    update_scrape_groups_total_for_run(
                        db_conn=db_conn,
                        scrap_id=scrape_state["scrap_id"],
                        groups_total=scrape_state["groups_total"],
                    )
                except Exception as exc:
                    _scrape_log(f"DB totals update warning: {exc}")

            total_contacts = sum(g["contacts_count"] for g in scrape_state["groups"])
            scrape_state["last_result"] = (
                f"Scrapping complete. Groups imported: {scrape_state['groups_done']}. "
                f"Total contacts: {total_contacts}."
            )
            _scrape_log(scrape_state["last_result"])

    except Exception as exc:
        scrape_state["last_error"] = str(exc)
        _scrape_log(f"Error: {exc}")
        traceback.print_exc()
    finally:
        try:
            if browser:
                browser.close()
        except Exception:
            pass
        release_playwright_profile(user_data_dir)
        try:
            if db_conn:
                db_conn.close()
        except Exception:
            pass
        scrape_state["running"] = False
        scrape_state["finished_at"] = datetime.now().isoformat()


def start_scrape_job(profile_number):
    if scrape_state["running"]:
        _append_app_log("Scrapping job request ignored: another scrapping job is already running.", "SCRAPE")
        return False

    profile_number_clean = str(profile_number or "").strip() or resolve_scrape_profile_number()
    if not profile_number_clean:
        _append_app_log("Scrapping job request ignored: profile number missing.", "SCRAPE")
        return False

    _append_app_log("Scrapping job started.", "SCRAPE")
    threading.Thread(target=run_scrape_job, args=(profile_number_clean,), daemon=True).start()
    return True


def _is_scrape_running():
    return scrape_state["running"]


recovery_manager = RecoveryManager(
    get_db_connection=get_db_connection,
    get_db_schema=get_db_schema,
    sql_module=sql,
    ensure_scrape_runs_table=ensure_scrape_runs_table,
    save_scrape_group_row=save_scrape_group_row,
    is_scrape_running=_is_scrape_running,
)


def start_send_job(
    targets,
    payload,
    delay_seconds,
    allow_resend,
    send_target="contacts",
    campaign_id=None,
    campaign_name="",
    batch_min_size=20,
    batch_max_size=30,
    range_config=None,
):
    if send_state["running"]:
        _send_log("Send job request ignored: another send job is already running.")
        return False

    def _job():
        active_campaign_id = str(campaign_id or generate_campaign_id()).zfill(6)
        active_campaign_name = str(campaign_name or "").strip()
        send_state["running"] = True
        send_state["last_error"] = None
        send_state["last_result"] = None
        send_state["campaign_id"] = active_campaign_id
        send_state["last_campaign_id"] = active_campaign_id
        send_state["campaign_name"] = active_campaign_name
        send_state["last_campaign_name"] = active_campaign_name
        send_state["last_live_rows"] = []
        send_state["automation_action_rows"] = _init_automation_action_rows(targets)
        send_state["last_automation_action_rows"] = []
        send_state["automation_action_event_text"] = ""
        send_state["last_automation_action_event_text"] = ""
        send_state["campaign_actions"] = _build_campaign_action_state()
        send_state["last_campaign_actions"] = []
        send_state["campaign_action_event_text"] = ""
        send_state["last_campaign_action_event_text"] = ""
        send_state["profile_states"] = []
        send_state["batch_pause_active"] = False
        send_state["batch_pause_remaining_sec"] = 0
        send_state["batch_next_start_at"] = ""
        send_state["pause_requested"] = False
        send_state["paused"] = False
        send_state["pause_requested_at"] = ""
        send_state["cancel_requested"] = False
        send_state["cancel_requested_at"] = ""
        send_state["session_opening"] = False
        send_state["awaiting_session_start"] = False
        send_state["session_start_requested"] = False
        send_state["session_logged_in_count"] = 0
        send_state["session_logged_out_count"] = 0
        send_state["current_batch_no"] = 0
        send_state["total_batches"] = 0
        send_state["failed"] = []
        send_state["sent"] = 0
        send_state["total"] = len(targets)
        send_state["contacts"] = list(targets)
        send_state["live_rows"] = _init_live_rows(targets)
        send_state["_active_send_contexts"] = []
        send_state["contact_statuses"] = {
            phone: {"status": "pending", "reason": ""}
            for phone in targets
        }
        _send_log(
            f"Send job started. Targets={len(targets)} | "
            f"target_type={send_target} | campaign={active_campaign_id}_{active_campaign_name or '-'}"
        )
        try:
            send_messages(
                targets,
                payload,
                delay_seconds,
                allow_resend,
                send_state,
                send_target=send_target,
                campaign_id=active_campaign_id,
                campaign_name=active_campaign_name,
                batch_min_size=batch_min_size,
                batch_max_size=batch_max_size,
                range_config=range_config,
            )
            send_state["last_failed"] = list(send_state["failed"])
            send_state["last_contacts"] = list(send_state["contacts"])
            send_state["last_contact_statuses"] = {
                k: {"status": v.get("status", "pending"), "reason": v.get("reason", "")}
                for k, v in send_state["contact_statuses"].items()
            }
            send_state["last_live_rows"] = list(send_state.get("live_rows") or [])
            send_state["last_automation_action_rows"] = list(
                send_state.get("automation_action_rows") or []
            )
            send_state["last_automation_action_event_text"] = str(
                send_state.get("automation_action_event_text") or ""
            ).strip()
            send_state["last_campaign_actions"] = _build_campaign_action_state(
                send_state.get("campaign_actions") or []
            )
            send_state["last_campaign_action_event_text"] = str(
                send_state.get("campaign_action_event_text") or ""
            ).strip()
            send_state["last_profile_states"] = list(send_state.get("profile_states") or [])
            send_state["last_result"] = f"Sent to {send_state['sent']} of {send_state['total']} contacts."
            _send_log(send_state["last_result"])
        except SendJobCancelled as exc:
            send_state["last_error"] = None
            send_state["last_result"] = str(exc or "Campaign canceled by user.")
            _send_log(send_state["last_result"])
        except Exception as exc:
            send_state["last_error"] = str(exc)
            _send_log(f"Send job error: {exc}")
        finally:
            if send_state["contacts"] and not send_state["last_contacts"]:
                send_state["last_contacts"] = list(send_state["contacts"])
            if send_state["contact_statuses"] and not send_state["last_contact_statuses"]:
                send_state["last_contact_statuses"] = {
                    k: {"status": v.get("status", "pending"), "reason": v.get("reason", "")}
                    for k, v in send_state["contact_statuses"].items()
                }
            if send_state["live_rows"] and not send_state["last_live_rows"]:
                send_state["last_live_rows"] = list(send_state.get("live_rows") or [])
            if send_state["automation_action_rows"] and not send_state["last_automation_action_rows"]:
                send_state["last_automation_action_rows"] = list(
                    send_state.get("automation_action_rows") or []
                )
            if (
                str(send_state.get("automation_action_event_text") or "").strip()
                and not str(send_state.get("last_automation_action_event_text") or "").strip()
            ):
                send_state["last_automation_action_event_text"] = str(
                    send_state.get("automation_action_event_text") or ""
                ).strip()
            if send_state["campaign_actions"] and not send_state["last_campaign_actions"]:
                send_state["last_campaign_actions"] = _build_campaign_action_state(
                    send_state.get("campaign_actions") or []
                )
            if (
                str(send_state.get("campaign_action_event_text") or "").strip()
                and not str(send_state.get("last_campaign_action_event_text") or "").strip()
            ):
                send_state["last_campaign_action_event_text"] = str(
                    send_state.get("campaign_action_event_text") or ""
                ).strip()
            if send_state["profile_states"] and not send_state["last_profile_states"]:
                send_state["last_profile_states"] = list(send_state.get("profile_states") or [])
            send_state["campaign_id"] = ""
            send_state["campaign_name"] = ""
            send_state["automation_action_rows"] = []
            send_state["automation_action_event_text"] = ""
            send_state["campaign_actions"] = []
            send_state["campaign_action_event_text"] = ""
            send_state["batch_pause_active"] = False
            send_state["batch_pause_remaining_sec"] = 0
            send_state["batch_next_start_at"] = ""
            send_state["pause_requested"] = False
            send_state["paused"] = False
            send_state["pause_requested_at"] = ""
            send_state["cancel_requested"] = False
            send_state["cancel_requested_at"] = ""
            send_state["session_opening"] = False
            send_state["awaiting_session_start"] = False
            send_state["session_start_requested"] = False
            send_state["session_logged_in_count"] = 0
            send_state["session_logged_out_count"] = 0
            send_state["_active_send_contexts"] = []
            send_state["running"] = False
            _send_log("Send job finished.")

    threading.Thread(target=_job, daemon=True).start()
    return True


@app.route("/", methods=["GET"])
def index():
    return messaging_panel.index(globals())


@app.route("/reset", methods=["GET"])
def reset():
    return messaging_panel.reset(globals())


@app.route("/send", methods=["POST"])
def send():
    return messaging_panel.send(globals())


@app.route("/api/send/pause", methods=["POST"])
def api_send_pause():
    return messaging_panel.api_send_pause(globals())


@app.route("/api/send/resume", methods=["POST"])
def api_send_resume():
    return messaging_panel.api_send_resume(globals())


@app.route("/api/send/cancel", methods=["POST"])
def api_send_cancel():
    return messaging_panel.api_send_cancel(globals())


@app.route("/api/send/session-start", methods=["POST"])
def api_send_session_start():
    return messaging_panel.api_send_session_start(globals())


@app.route("/status", methods=["GET"])
def status():
    return campaign_monitoring_panel.status(globals())


@app.route("/logs/status", methods=["GET"])
def logs_status():
    return logs_panel.logs_status(globals())


@app.route("/assets/pic/<path:filename>", methods=["GET"])
def template_pic(filename):
    return messaging_panel.template_pic(globals(), filename)


@app.route("/assets/send_uploads/<path:filename>", methods=["GET"])
def send_upload_asset(filename):
    return messaging_panel.send_upload_asset(globals(), filename)


@app.route("/api/recent-send-media", methods=["GET"])
def api_recent_send_media():
    return messaging_panel.api_recent_send_media(globals())


@app.route("/api/recent-send-media/delete", methods=["POST"])
def api_recent_send_media_delete():
    return messaging_panel.api_recent_send_media_delete(globals())


@app.route("/api/sms-messages/add", methods=["POST"])
def api_sms_messages_add():
    return messaging_panel.api_add_sms_message(globals())


@app.route("/api/sms-messages/delete", methods=["POST"])
def api_sms_messages_delete():
    return messaging_panel.api_delete_sms_message(globals())


@app.route("/api/range-config/save", methods=["POST"])
def api_range_config_save():
    return messaging_panel.api_range_config_save(globals())


@app.route("/api/campaign-name/suggest", methods=["GET", "POST"])
def api_campaign_name_suggest():
    return messaging_panel.api_campaign_name_suggest(globals())


@app.route("/scrape/start", methods=["POST"])
def scrape_start():
    return new_scrapping_panel.scrape_start(globals())


@app.route("/scrape/status", methods=["GET"])
def scrape_status():
    return previous_scrapping_panel.scrape_status(globals())


@app.route("/scrape/recover/start", methods=["POST"])
def scrape_recover_start():
    return previous_scrapping_panel.scrape_recover_start(globals())


@app.route("/scrape/recover/status", methods=["GET"])
def scrape_recover_status():
    return previous_scrapping_panel.scrape_recover_status(globals())


@app.route("/api/groups-joining/preview", methods=["POST"])
def api_groups_joining_preview():
    return groups_joining_panel.api_groups_joining_preview(globals())


@app.route("/api/groups-joining/start", methods=["POST"])
def api_groups_joining_start():
    return groups_joining_panel.api_groups_joining_start(globals())


@app.route("/api/groups-joining/stop", methods=["POST"])
def api_groups_joining_stop():
    return groups_joining_panel.api_groups_joining_stop(globals())


@app.route("/api/groups-joining/status", methods=["GET"])
def api_groups_joining_status():
    return groups_joining_panel.api_groups_joining_status(globals())


@app.route("/api/group-tables", methods=["GET"])
def api_group_tables():
    return messaging_panel.api_group_tables(globals())


@app.route("/api/table-info", methods=["POST"])
def api_table_info():
    return messaging_panel.api_table_info(globals())


@app.route("/api/preview-table-contacts", methods=["POST"])
def api_preview_table_contacts():
    return messaging_panel.api_preview_table_contacts(globals())


@app.route("/api/preview-file-contacts", methods=["POST"])
def api_preview_file_contacts():
    return messaging_panel.api_preview_file_contacts(globals())


@app.route("/api/preview-manual-all-contacts", methods=["POST"])
def api_preview_manual_all_contacts():
    return messaging_panel.api_preview_manual_all_contacts(globals())


@app.route("/api/import-all-contacts-manual", methods=["POST"])
def api_import_all_contacts_manual():
    return messaging_panel.api_import_all_contacts_manual(globals())


@app.route("/api/groups-panel", methods=["GET"])
def api_groups_panel():
    return groups_panel.api_groups_panel(globals())


@app.route("/api/campaigns-panel", methods=["GET"])
def api_campaigns_panel():
    return campaigns_panel.api_campaigns_panel(globals())


@app.route("/api/campaign-details", methods=["GET"])
def api_campaign_details():
    return campaigns_panel.api_campaign_details(globals())


@app.route("/api/campaign-reuse", methods=["GET"])
def api_campaign_reuse():
    return campaigns_panel.api_campaign_reuse(globals())


@app.route("/api/group-details", methods=["GET"])
def api_group_details(): 
    return groups_panel.api_group_details(globals())


@app.route("/api/group-details-bulk", methods=["POST"])
def api_group_details_bulk():
    return groups_panel.api_group_details_bulk(globals())


@app.route("/api/profile-sessions/status", methods=["POST"])
def api_profile_sessions_status():
    try:
        payload = request.get_json(silent=True) or {}
    except Exception:
        payload = {}
    validate = str(payload.get("validate") or "").strip().lower() in ("1", "true", "yes", "on")
    try:
        count = int(
            payload.get("count")
            or send_state.get("last_profiles_numbers")
            or _discover_profile_session_count()
            or 1
        )
    except Exception:
        count = 1
    count = max(1, min(100, count))
    scan_result = _scan_profile_sessions_status(count, send_state=send_state, validate=validate)
    scan_count = scan_result.get("count")
    if scan_count is None:
        scan_count = count
    rows = list(scan_result.get("rows") or [])
    active_rows = list(scan_result.get("active_rows") or [])
    active_profile_indexes = [
        int(row.get("profile_idx") or 1)
        for row in active_rows
        if int(row.get("profile_idx") or 1) >= 1
    ]
    logged_in_count = sum(1 for row in rows if str((row or {}).get("status") or "").lower() == "logged_in")
    in_use_count = sum(1 for row in rows if str((row or {}).get("status") or "").lower() == "in_use")
    logged_out_count = sum(1 for row in rows if str((row or {}).get("status") or "").lower() == "logged_out")
    error_count = sum(1 for row in rows if str((row or {}).get("status") or "").lower() == "error")
    session_jobs = _get_profile_session_jobs()
    preview_rows = _get_open_session_preview_rows()
    return jsonify(
        {
            "ok": True,
            "count": int(scan_count),
            "sequence": dict(scan_result.get("sequence") or {}),
            "rows": rows,
            "active_rows": active_rows,
            "active_profile_indexes": active_profile_indexes,
            "source": str(scan_result.get("source") or "scan"),
            "logged_in_count": logged_in_count,
            "in_use_count": in_use_count,
            "logged_out_count": logged_out_count,
            "error_count": error_count,
            "usable_count": len(active_rows),
            "session_jobs": session_jobs,
            "preview_open_count": len(preview_rows),
            "preview_rows": preview_rows,
        }
    )


@app.route("/api/profile-numbers/status", methods=["POST"])
def api_profile_numbers_status():
    try:
        payload = request.get_json(silent=True) or {}
    except Exception:
        payload = {}
    try:
        count = int(payload.get("count") or _discover_profile_session_count() or 1)
    except Exception:
        count = 1
    count = max(1, min(100, count))

    stored = _load_profile_numbers_from_session_dirs(count, create_missing=True)
    existing = {}
    missing = []
    for idx in range(1, count + 1):
        value = str(stored.get(idx) or stored.get(str(idx)) or "").strip()
        if value:
            existing[str(idx)] = value
        else:
            missing.append(idx)
    return jsonify({"ok": True, "count": count, "existing": existing, "missing": missing})


@app.route("/api/profile-numbers/save", methods=["POST"])
def api_profile_numbers_save():
    try:
        payload = request.get_json(silent=True) or {}
    except Exception:
        payload = {}
    try:
        count = int(
            payload.get("count")
            or send_state.get("last_profiles_numbers")
            or _discover_profile_session_count()
            or 1
        )
    except Exception:
        count = 1
    count = max(1, min(100, count))
    try:
        profile_idx = int(payload.get("index") or 0)
    except Exception:
        profile_idx = 0
    if profile_idx < 1 or profile_idx > count:
        return jsonify({"ok": False, "error": "Invalid profile row."}), 400

    value = str(payload.get("value") or "").strip()
    try:
        saved_value, profile_dir = _save_profile_number_for_session(profile_idx, value, count)
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400

    send_state["last_profiles_numbers"] = max(
        int(send_state.get("last_profiles_numbers") or 1),
        count,
    )
    _sync_profile_number_into_send_state(send_state, profile_idx, saved_value)
    existing = _load_profile_numbers_from_session_dirs(count, create_missing=True)
    return jsonify(
        {
            "ok": True,
            "count": count,
            "index": profile_idx,
            "value": saved_value,
            "profile_dir": str(profile_dir),
            "existing": {str(k): v for k, v in existing.items()},
        }
    )


@app.route("/api/profile-sessions/add", methods=["POST"])
def api_profile_sessions_add():
    try:
        next_idx, profile_dir, profile_number = _create_next_profile_session_slot(
            send_state=send_state,
            limit=100,
        )
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    _set_profile_session_job(
        next_idx,
        "opening",
        f"Opening WhatsApp Web for session slot #{next_idx}...",
        profile_number=str(profile_number or "").strip(),
    )
    _run_profile_session_login_window(next_idx)
    return jsonify(
        {
            "ok": True,
            "index": int(next_idx),
            "count": int(max(int(send_state.get("last_profiles_numbers") or 1), next_idx)),
            "profile_dir": str(profile_dir),
            "value": str(profile_number or "").strip(),
            "message": f"WhatsApp Web opened for session slot #{next_idx}. Scan QR and complete login.",
        }
    )


@app.route("/api/profile-sessions/login", methods=["POST"])
def api_profile_sessions_login():
    """Open the WhatsApp window for one session slot so the user can log in
    (or keep it open, logged in, ready for the next campaign)."""
    if bool(send_state.get("running")) or bool(send_state.get("session_opening")):
        return jsonify({"ok": False, "error": "Cannot open a session window while sending is active."}), 400
    try:
        payload = request.get_json(silent=True) or {}
    except Exception:
        payload = {}
    idx = _parse_profile_index(payload.get("index"))
    if idx < 1:
        return jsonify({"ok": False, "error": "Invalid session slot."}), 400
    if _is_open_session_preview_active(idx):
        return jsonify(
            {
                "ok": True,
                "index": idx,
                "already_open": True,
                "message": f"WhatsApp window for session slot #{idx} is already open. Scan the QR code there if login is required.",
            }
        )
    profile_dir = _get_send_profile_dir(get_playwright_user_data_dir(), idx, create=True)
    profile_number = _ensure_and_read_profile_number(profile_dir, seed_value="")
    try:
        started = _start_open_session_preview_window(idx, profile_number=profile_number)
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    return jsonify(
        {
            "ok": True,
            "index": idx,
            "already_open": not bool(started),
            "message": (
                f"Opening WhatsApp window for session slot #{idx}. Scan the QR code to log in; "
                "the window stays open and the campaign will use it."
            ),
        }
    )


@app.route("/api/profile-sessions/open-logged-in", methods=["POST"])
def api_profile_sessions_open_logged_in():
    if bool(send_state.get("running")) or bool(send_state.get("session_opening")):
        return jsonify(
            {
                "ok": False,
                "error": "Cannot open preview sessions while sending is active.",
            }
        ), 400
    try:
        payload = request.get_json(silent=True) or {}
    except Exception:
        payload = {}
    try:
        count = int(
            payload.get("count")
            or send_state.get("last_profiles_numbers")
            or _discover_profile_session_count()
            or 1
        )
    except Exception:
        count = 1
    count = max(1, min(100, count))

    scan_result = _scan_profile_sessions_status(count, send_state=send_state, validate=True)
    active_rows = list(scan_result.get("active_rows") or [])
    if not active_rows:
        return jsonify(
            {
                "ok": False,
                "error": "No logged-in WhatsApp sessions found to open.",
            }
        ), 400

    started_indexes = []
    already_open_indexes = []
    start_errors = []
    for row in active_rows:
        profile_idx = _parse_profile_index((row or {}).get("profile_idx") or (row or {}).get("name"))
        if profile_idx < 1:
            continue
        if _is_open_session_preview_active(profile_idx):
            already_open_indexes.append(profile_idx)
            continue
        try:
            started = _start_open_session_preview_window(
                profile_idx,
                profile_number=str((row or {}).get("profile_number") or "").strip(),
            )
            if started:
                started_indexes.append(profile_idx)
            else:
                already_open_indexes.append(profile_idx)
        except Exception as exc:
            start_errors.append(f"#{profile_idx}: {exc}")

    snapshot = _scan_profile_sessions_status(count, send_state=send_state, validate=False)
    preview_rows = _get_open_session_preview_rows()
    message_parts = []
    if started_indexes:
        message_parts.append(
            "Opening preview window(s) for session slot(s): "
            + ", ".join(f"#{idx}" for idx in started_indexes)
            + "."
        )
    if already_open_indexes:
        message_parts.append(
            "Already open: "
            + ", ".join(f"#{idx}" for idx in already_open_indexes)
            + "."
        )
    if start_errors:
        message_parts.append(
            "Some sessions could not be opened: " + "; ".join(start_errors)
        )
    if not message_parts:
        message_parts.append("No logged-in session preview window was opened.")
    if not started_indexes and not already_open_indexes and start_errors:
        return jsonify(
            {
                "ok": False,
                "error": " ".join(message_parts).strip(),
                "preview_open_count": len(preview_rows),
                "preview_rows": preview_rows,
            }
        ), 400
    return jsonify(
        {
            "ok": True,
            "count": int(snapshot.get("count") or count),
            "sequence": snapshot.get("sequence") or {},
            "rows": snapshot.get("rows") or [],
            "active_rows": snapshot.get("active_rows") or [],
            "active_profile_indexes": snapshot.get("active_profile_indexes") or [],
            "session_jobs": _get_profile_session_jobs(),
            "preview_open_count": len(preview_rows),
            "preview_rows": preview_rows,
            "opened_indexes": started_indexes,
            "already_open_indexes": already_open_indexes,
            "message": " ".join(message_parts).strip(),
        }
    )


@app.route("/api/profile-sessions/close-all", methods=["POST"])
def api_profile_sessions_close_all():
    try:
        payload = request.get_json(silent=True) or {}
    except Exception:
        payload = {}
    try:
        count = int(
            payload.get("count")
            or send_state.get("last_profiles_numbers")
            or _discover_profile_session_count()
            or 1
        )
    except Exception:
        count = 1
    count = max(1, min(100, count))

    close_result = _request_close_open_session_preview_windows(wait_timeout_seconds=15.0)
    snapshot = _scan_profile_sessions_status(count, send_state=send_state, validate=False)
    preview_rows = _get_open_session_preview_rows()
    requested = int(close_result.get("requested") or 0)
    closed = int(close_result.get("closed") or 0)
    remaining = list(close_result.get("remaining") or [])
    if requested <= 0:
        message = "No open WhatsApp preview windows were found."
    elif remaining:
        message = (
            "Requested close for preview windows, but some are still open: "
            + ", ".join(f"#{idx}" for idx in remaining)
            + "."
        )
    else:
        message = f"Closed {closed} WhatsApp preview window(s)."
    return jsonify(
        {
            "ok": True,
            "count": int(snapshot.get("count") or count),
            "sequence": snapshot.get("sequence") or {},
            "rows": snapshot.get("rows") or [],
            "active_rows": snapshot.get("active_rows") or [],
            "active_profile_indexes": snapshot.get("active_profile_indexes") or [],
            "session_jobs": _get_profile_session_jobs(),
            "preview_open_count": len(preview_rows),
            "preview_rows": preview_rows,
            "closed_count": closed,
            "remaining_preview_indexes": remaining,
            "message": message,
        }
    )


@app.route("/api/profile-sessions/remove", methods=["POST"])
def api_profile_sessions_remove():
    try:
        payload = request.get_json(silent=True) or {}
    except Exception:
        payload = {}
    try:
        profile_idx = int(payload.get("index") or 0)
    except Exception:
        profile_idx = 0
    try:
        result = _remove_profile_session_slot(profile_idx, send_state=send_state, limit=100)
    except Exception as exc:
        print(
            f"[profile-sessions/remove] failed to remove #{profile_idx}: "
            f"{type(exc).__name__}: {exc}"
        )
        traceback.print_exc()
        return jsonify({"ok": False, "error": str(exc)}), 400
    remaining_count = int(result.get("remaining_count") or 0)
    snapshot = _scan_profile_sessions_status(
        remaining_count,
        send_state=send_state,
        validate=False,
    )
    preview_rows = _get_open_session_preview_rows()
    return jsonify(
        {
            "ok": True,
            "removed_index": int(result.get("removed_index") or profile_idx),
            "remaining_count": remaining_count,
            "count": int(snapshot.get("count") or remaining_count),
            "sequence": snapshot.get("sequence") or {},
            "rows": snapshot.get("rows") or [],
            "active_rows": snapshot.get("active_rows") or [],
            "active_profile_indexes": snapshot.get("active_profile_indexes") or [],
            "session_jobs": _get_profile_session_jobs(),
            "preview_open_count": len(preview_rows),
            "preview_rows": preview_rows,
            "message": f"Profile session #{int(result.get('removed_index') or profile_idx)} removed permanently.",
        }
    )


@app.route("/api/search-scrape-groups", methods=["GET"])
def api_search_scrape_groups():
    return previous_scrapping_panel.api_search_scrape_groups(globals())


@app.route("/api/unique-scrape-groups", methods=["GET"])
def api_unique_scrape_groups():
    return previous_scrapping_panel.api_unique_scrape_groups(globals())

 
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
